#!/usr/bin/env python3
"""Build ประกาศผู้ชนะ Word template matching แบบฟอร์ม_ประกาศชื่อผู้ชนะในการเสนอราคา."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips

# Same font as docs/แบบฟอร์ม_ประกาศชื่อผู้ชนะในการเสนอราคา.docx
FONT = "TH SarabunIT๙"
OUT = Path(__file__).with_name("winner_announcement.docx")
GARUDA = Path(__file__).resolve().parents[1] / "img" / "garuda.jpeg"
# Form inline shape ≈ 2.72 × 3.01 cm
GARUDA_WIDTH = Cm(2.72)


def set_run_font(run, size_pt=16, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), FONT)


def add_para(
    doc,
    text,
    *,
    size=16,
    bold=False,
    align="left",
    before=0,
    after=0,
    first_line=None,
):
    para = doc.add_paragraph()
    para.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "thai_justify": WD_ALIGN_PARAGRAPH.THAI_JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    pf = para.paragraph_format
    pf.space_before = Twips(before)
    pf.space_after = Twips(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    run = para.add_run(text)
    set_run_font(run, size_pt=size, bold=bold)
    return para


def add_garuda(doc):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = para.paragraph_format
    pf.space_before = Twips(0)
    pf.space_after = Twips(60)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = para.add_run()
    if not GARUDA.is_file():
        raise SystemExit("missing garuda image: %s" % GARUDA)
    run.add_picture(str(GARUDA), width=GARUDA_WIDTH)
    return para


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(16)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), FONT)

    # Layout matches docs/แบบฟอร์ม_ประกาศชื่อผู้ชนะในการเสนอราคา.docx
    add_garuda(doc)
    add_para(doc, "ประกาศ{{province}}", size=16, bold=True, align="center")
    add_para(doc, "เรื่อง  {{subject}}", size=16, align="center")
    add_para(
        doc,
        "---------------------------------------------------------------------",
        size=16,
        align="center",
    )
    add_para(
        doc,
        "{{body_intro}}",
        size=16,
        align="thai_justify",
        before=120,
        first_line=2.5,
    )
    add_para(
        doc,
        "{{body_winner}}",
        size=16,
        align="thai_justify",
        before=120,
        first_line=2.5,
    )
    add_para(
        doc,
        "ประกาศ  ณ  วันที่  {{announce_date}}",
        size=16,
        align="thai_justify",
        before=240,
        first_line=6.25,
    )
    add_para(doc, "", size=16)
    add_para(doc, "", size=16)
    add_para(doc, "", size=16)
    add_para(doc, "({{signer_full}})", size=16, first_line=8.0)
    add_para(doc, "{{signer_title}}", size=16, first_line=7.5)
    add_para(doc, "{{signer_acting}}", size=16, first_line=6.5)

    doc.save(OUT)
    print("wrote", OUT, "font=", FONT, "garuda=", GARUDA.name)


if __name__ == "__main__":
    main()
