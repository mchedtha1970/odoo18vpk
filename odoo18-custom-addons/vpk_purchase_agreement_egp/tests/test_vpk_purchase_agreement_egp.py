# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

import base64
from io import BytesIO

from odoo.exceptions import UserError
from odoo.tests import tagged
from odoo.tests.common import TransactionCase


@tagged("post_install", "-at_install")
class TestVpkPurchaseAgreementEgp(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.partner = cls.env["res.partner"].create({"name": "Vendor A", "supplier_rank": 1})
        cls.partner_b = cls.env["res.partner"].create({"name": "Vendor B", "supplier_rank": 1})
        cls.product = cls.env["product.product"].create(
            {
                "name": "Medical Supply",
                "purchase_ok": True,
                "type": "consu",
            }
        )
        cls.requisition = cls.env["purchase.requisition"].create(
            {
                "requisition_type": "egp_procurement",
                "egp_reference": "EGP-PRJ-001",
                "line_ids": [
                    (
                        0,
                        0,
                        {
                            "product_id": cls.product.id,
                            "product_qty": 5.0,
                        },
                    )
                ],
            }
        )

    def test_egp_agreement_sequence_and_confirm(self):
        self.assertTrue(self.requisition.name.startswith("EGP/"))
        self.requisition.action_confirm()
        self.assertEqual(self.requisition.state, "confirmed")

    def test_create_rfq_propagates_origin(self):
        self.requisition.action_confirm()
        wizard = self.env["purchase.requisition.create.rfq"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "egp_bid_reference": "BID-001",
            }
        )
        action = wizard.action_create_rfq()
        purchase_order = self.env["purchase.order"].browse(action["res_id"])
        self.assertEqual(purchase_order.requisition_id, self.requisition)
        self.assertIn("EGP-PRJ-001", purchase_order.origin)
        self.assertIn(self.requisition.name, purchase_order.origin)
        self.assertIn("BID-001", purchase_order.origin)
        self.assertEqual(purchase_order.order_line[0].product_qty, 5.0)
        self.assertEqual(purchase_order.order_line[0].price_unit, 0.0)

    def test_alternative_keeps_requisition(self):
        self.requisition.action_confirm()
        first_rfq = self.env["purchase.order"].create(
            self.requisition._prepare_rfq_vals(self.partner, "BID-A")
        )
        alternative_wizard = self.env["purchase.requisition.create.alternative"].create(
            {
                "origin_po_id": first_rfq.id,
                "partner_id": self.partner_b.id,
                "copy_products": True,
            }
        )
        action = alternative_wizard.action_create_alternative()
        second_rfq = self.env["purchase.order"].browse(action["res_id"])
        self.assertEqual(second_rfq.requisition_id, self.requisition)
        self.assertEqual(second_rfq.origin, first_rfq.origin)

    def test_egp_document_auto_reference(self):
        self.requisition.egp_reference = "EGP-PROJECT-999"
        invitation_type = self.env.ref(
            "vpk_purchase_agreement_egp.egp_document_type_invitation"
        )
        document = self.env["purchase.requisition.egp.document"].new(
            {
                "requisition_id": self.requisition.id,
                "document_type_id": invitation_type.id,
            }
        )
        document._onchange_document_type_id()
        self.assertEqual(document.egp_reference, "EGP-PROJECT-999")

    def test_egp_document_preview(self):
        invitation_type = self.env.ref(
            "vpk_purchase_agreement_egp.egp_document_type_invitation"
        )
        document = self.env["purchase.requisition.egp.document"].create(
            {
                "requisition_id": self.requisition.id,
                "document_type_id": invitation_type.id,
                "egp_reference": "EGP-INV-001",
                "document_file": "dGVzdA==",
                "document_filename": "invitation.pdf",
            }
        )
        self.assertTrue(document.attachment_id)
        action = document.action_preview_document()
        self.assertEqual(action["type"], "ir.actions.act_url")
        self.assertIn(str(document.attachment_id.id), action["url"])
        self.requisition._update_egp_reference_from_documents()
        self.assertEqual(self.requisition.egp_reference, "EGP-INV-001")

    def test_generate_winner_announcement_document(self):
        winner_type = self.env.ref(
            "vpk_purchase_agreement_egp.egp_document_type_winner_announcement"
        )
        self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "amount_total": 97370.0,
                "is_winner": True,
            }
        )
        document = self.env["purchase.requisition.egp.document"].create(
            {
                "requisition_id": self.requisition.id,
                "document_type_id": winner_type.id,
                "egp_reference": "6904571",
            }
        )
        values = document._winner_announcement_values()
        self.assertIn("Vendor A", values["body_winner"])
        self.assertIn("ประกาศผู้ชนะการเสนอราคา", values["subject"])
        self.assertIn("ตามที่", values["body_intro"])
        document.action_generate_winner_document()
        self.assertTrue(document.document_file)
        self.assertTrue(document.document_filename)
        self.assertTrue(document.attachment_id)
        content = base64.b64decode(document.document_file)
        if document.document_filename.endswith(".docx"):
            from docx import Document

            text = "\n".join(p.text for p in Document(BytesIO(content)).paragraphs)
            self.assertIn("Vendor A", text)
            self.assertIn("ประกาศผู้ชนะการเสนอราคา", text)
            self.assertIn("ผู้ได้รับการคัดเลือก", text)
            self.assertNotIn("{{", text)
            self.assertNotIn("สำเนา", text)

    def test_generate_winner_announcement_requires_type_and_winner(self):
        invitation_type = self.env.ref(
            "vpk_purchase_agreement_egp.egp_document_type_invitation"
        )
        winner_type = self.env.ref(
            "vpk_purchase_agreement_egp.egp_document_type_winner_announcement"
        )
        invitation = self.env["purchase.requisition.egp.document"].create(
            {
                "requisition_id": self.requisition.id,
                "document_type_id": invitation_type.id,
            }
        )
        with self.assertRaises(UserError):
            invitation.action_generate_winner_document()
        winner_doc = self.env["purchase.requisition.egp.document"].create(
            {
                "requisition_id": self.requisition.id,
                "document_type_id": winner_type.id,
            }
        )
        with self.assertRaises(UserError):
            winner_doc.action_generate_winner_document()

    def test_egp_flow_stage_progression(self):
        self.assertEqual(self.requisition.egp_flow_stage, "draft")
        self.requisition.action_confirm()
        self.assertEqual(self.requisition.egp_flow_stage, "confirmed")
        invitation_type = self.env.ref(
            "vpk_purchase_agreement_egp.egp_document_type_invitation"
        )
        self.env["purchase.requisition.egp.document"].create(
            {
                "requisition_id": self.requisition.id,
                "document_type_id": invitation_type.id,
            }
        )
        self.assertEqual(self.requisition.egp_flow_stage, "egp_docs")
        self.env["purchase.order"].create(
            self.requisition._prepare_rfq_vals(self.partner, "BID-A")
        )
        self.assertEqual(self.requisition.egp_flow_stage, "rfq")
        self.env["purchase.order"].create(
            self.requisition._prepare_rfq_vals(self.partner_b, "BID-B")
        )
        self.assertEqual(self.requisition.egp_flow_stage, "compare")
        self.assertIn("vpk_egp_flow_current", self.requisition.egp_flow_html)

    def test_one_bid_sets_compare_until_statusbar_is_clicked(self):
        self.requisition.action_confirm()
        self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "egp_bid_reference": "Q-1",
            }
        )
        self.assertEqual(self.requisition.egp_flow_stage, "compare")
        self.assertFalse(self.requisition.egp_flow_stage_manual)
        self.requisition.egp_flow_stage = "rfq"
        self.assertTrue(self.requisition.egp_flow_stage_manual)
        self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner_b.id,
                "egp_bid_reference": "Q-2",
            }
        )
        self.assertEqual(self.requisition.egp_flow_stage, "rfq")

    def test_winner_tick_keeps_compare_stage_and_confirmed_state(self):
        self.requisition.action_confirm()
        bid = self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "egp_bid_reference": "Q-1",
            }
        )
        self.assertEqual(self.requisition.egp_flow_stage, "compare")
        self.assertEqual(self.requisition.state, "confirmed")
        bid.write({"is_winner": True})
        self.assertTrue(bid.is_winner)
        self.assertEqual(self.requisition.state, "confirmed")
        self.assertEqual(self.requisition.egp_flow_stage, "compare")
        self.requisition.egp_flow_stage = "compare"
        onchange = self.requisition.onchange(
            {"egp_bid_ids": [(1, bid.id, {"is_winner": True})]},
            ["egp_bid_ids"],
            {"egp_flow_stage": {}, "state": {}, "egp_bid_ids": {"fields": {"is_winner": {}}}},
        )
        self.assertNotIn(onchange["value"].get("egp_flow_stage"), ("done", "awarded"))
        self.assertNotEqual(onchange["value"].get("state"), "done")

    def test_pr_create_rfq_from_linked_agreement(self):
        purchase_type = self.env.ref("l10n_th_gov_purchase_request.purchase_type_001")
        procurement_type = self.env.ref("l10n_th_gov_purchase_request.procurement_type_001")
        procurement_method = self.env.ref("l10n_th_gov_purchase_request.procurement_specific")
        request = self.env["purchase.request"].create(
            {
                "purchase_type_id": purchase_type.id,
                "procurement_type_id": procurement_type.id,
                "procurement_method_id": procurement_method.id,
                "line_ids": [
                    (
                        0,
                        0,
                        {
                            "product_id": self.product.id,
                            "product_qty": 5.0,
                        },
                    )
                ],
            }
        )
        request.write({"state": "approved"})
        self.requisition.action_confirm()
        self.requisition.line_ids[0].write(
            {"purchase_request_lines": [(4, request.line_ids[0].id)]}
        )
        request.invalidate_recordset()
        self.assertEqual(request.egp_requisition_id, self.requisition)
        self.assertTrue(request.can_create_egp_rfq)
        action = request.action_open_create_rfq_wizard()
        self.assertEqual(action["res_model"], "purchase.requisition.create.rfq")
        wizard = self.env["purchase.requisition.create.rfq"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "egp_bid_reference": "BID-PR-001",
            }
        )
        po_action = wizard.action_create_rfq()
        purchase_order = self.env["purchase.order"].browse(po_action["res_id"])
        self.assertEqual(purchase_order.requisition_id, self.requisition)

    def test_create_rfq_from_pr_keeps_qty_uom_and_winner_price(self):
        purchase_type = self.env.ref("l10n_th_gov_purchase_request.purchase_type_001")
        procurement_type = self.env.ref("l10n_th_gov_purchase_request.procurement_type_001")
        procurement_method = self.env.ref("l10n_th_gov_purchase_request.procurement_specific")
        request = self.env["purchase.request"].create(
            {
                "purchase_type_id": purchase_type.id,
                "procurement_type_id": procurement_type.id,
                "procurement_method_id": procurement_method.id,
                "line_ids": [
                    (
                        0,
                        0,
                        {
                            "product_id": self.product.id,
                            "product_qty": 300.0,
                            "product_uom_id": self.product.uom_id.id,
                            "estimated_cost": 26643.0,
                        },
                    )
                ],
            }
        )
        request.write({"state": "approved"})
        self.requisition.write({"purchase_request_id": request.id})
        self.requisition.action_confirm()
        self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "amount_total": 26643.0,
                "is_winner": True,
            }
        )
        vals = self.requisition._prepare_rfq_vals(self.partner, "BID-WIN")
        purchase_order = self.env["purchase.order"].create(vals)
        line = purchase_order.order_line
        self.assertEqual(len(line), 1)
        self.assertEqual(line.product_qty, 300.0)
        self.assertEqual(line.product_uom, self.product.uom_id)
        self.assertAlmostEqual(line.price_unit, 88.81, places=2)
        self.assertAlmostEqual(line.price_subtotal, 26643.0, places=2)

    def _vpk_make_egp_purchase_request(self):
        purchase_type = self.env.ref("l10n_th_gov_purchase_request.purchase_type_001")
        procurement_type = self.env.ref("l10n_th_gov_purchase_request.procurement_type_001")
        procurement_method = self.env.ref(
            "l10n_th_gov_purchase_request.procurement_specific"
        )
        return self.env["purchase.request"].create(
            {
                "purchase_type_id": purchase_type.id,
                "procurement_type_id": procurement_type.id,
                "procurement_method_id": procurement_method.id,
                "line_ids": [
                    (
                        0,
                        0,
                        {
                            "product_id": self.product.id,
                            "product_qty": 5.0,
                        },
                    )
                ],
            }
        )

    def test_send_to_egp_button_requires_related_documents(self):
        request = self._vpk_make_egp_purchase_request()
        request.write({"state": "sent_to_procurement"})
        request.invalidate_recordset()
        if "procurement_verified" in request._fields:
            self.assertFalse(request.can_send_to_egp)
            request.write({"procurement_verified": True})
            request.invalidate_recordset()
        self.assertTrue(request.can_send_to_egp)
        action = request.action_send_to_egp()
        self.assertEqual(
            action["res_model"],
            "purchase.request.line.make.purchase.requisition",
        )
        self.assertEqual(request.state, "approved")

    def test_invitation_tab_pulls_products_from_referenced_pr(self):
        request = self._vpk_make_egp_purchase_request()
        request.line_ids.product_qty = 4.0
        self.requisition.write({"purchase_request_id": request.id})
        invitation = self.env["purchase.requisition.egp.invitation"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "submission_deadline": "2026-09-30 10:00:00",
            }
        )
        self.assertEqual(invitation.line_ids.product_qty, 4.0)
        request.line_ids.product_qty = 12.0
        self.requisition.action_fill_invitation_lines_from_pr()
        self.assertEqual(len(invitation.line_ids), 1)
        self.assertEqual(invitation.line_ids.product_id, self.product)
        self.assertEqual(invitation.line_ids.product_qty, 12.0)
        self.assertEqual(invitation.line_ids.product_uom_id, self.product.uom_id)

    def test_award_approval_docx_template_fills_items(self):
        import importlib.util
        from pathlib import Path

        module_path = Path(__file__).resolve().parents[1] / "models" / "award_approval_docx.py"
        spec = importlib.util.spec_from_file_location("award_approval_docx", module_path)
        renderer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(renderer)
        content = renderer.render_award_approval_docx(
            renderer.load_template_bytes(),
            {
                "agency": "กลุ่มงานพัสดุ",
                "memo_number": "AR-EGP/๒๕๖๙/๐๐๐๑",
                "memo_date": "๑๘ กันยายน ๒๕๖๙",
                "subject": "รายงานผลการพิจารณาและขออนุมัติสั่งซื้อสั่งจ้าง",
                "recipient": "ผู้ว่าราชการจังหวัดภูเก็ต",
                "intro": "ขอรายงานผลการพิจารณา จำนวน ๑ รายการ โดยวิธีเฉพาะเจาะจง ดังนี้",
                "items": [{
                    "item_desc": "๑. โฟมปั๊มเท้า\nจำนวน ๔ แพ็ค",
                    "vendor_name": "บริษัททดสอบ",
                    "offer_price": "๑๐๐.๐๐",
                    "agreed_price": "๑๐๐.๐๐",
                }],
                "total_amount": "๑๐๐.๐๐",
                "criteria": "โดยใช้หลักเกณฑ์ราคา",
                "hospital_opinion": "เห็นสมควรจัดซื้อ",
                "request_text": "จึงเรียนมาเพื่อโปรดพิจารณา",
                "officer_name": "นางทดสอบ",
                "officer_position": "เจ้าหน้าที่",
                "head_officer_name": "นางสาวทดสอบ",
                "head_officer_position": "หัวหน้าเจ้าหน้าที่",
                "approver_authority": "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต",
                "signer_name": "นายทดสอบ",
                "signer_position": "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต",
                "signer_acting": "ปฏิบัติราชการแทนผู้ว่าราชการจังหวัดภูเก็ต",
            },
        )
        from docx import Document
        import io

        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n" + " ".join(cell.text for cell in row.cells)
        self.assertIn("บันทึกข้อความ", text)
        self.assertIn("โฟมปั๊มเท้า", text)
        self.assertIn("๑๐๐.๐๐", text)
        self.assertNotIn("{{", text)

    def test_file_winner_announcement_from_compare(self):
        self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "amount_total": 2320.0,
                "is_winner": True,
            }
        )
        self.requisition.action_file_winner_announcement()
        documents = self.requisition.egp_document_ids.filtered(
            lambda document: document.document_type_id.code == "winner_announcement"
        )
        self.assertEqual(len(documents), 1)
        self.assertEqual(documents.document_type_id.name, "ประกาศผู้ชนะ")
        self.assertTrue(documents.document_file)
        self.assertTrue(documents.document_filename)
        self.requisition.action_file_winner_announcement()
        documents = self.requisition.egp_document_ids.filtered(
            lambda document: document.document_type_id.code == "winner_announcement"
        )
        self.assertEqual(len(documents), 1)

    def test_award_print_files_egp_document(self):
        report = self.env["purchase.requisition.award.report"].create({
            "requisition_id": self.requisition.id,
        })
        report._file_award_egp_document(b"%PDF-1.4 award")
        documents = self.requisition.egp_document_ids.filtered(
            lambda document: document.document_type_id.code == "award_approval"
        )
        self.assertEqual(len(documents), 1)
        self.assertEqual(documents.award_report_id, report)
        self.assertEqual(documents.document_type_id.name, "รายงานผลการพิจารณา")
        self.assertTrue(documents.document_file)
        report._file_award_egp_document(b"%PDF-1.4 award-2")
        documents = self.requisition.egp_document_ids.filtered(
            lambda document: document.document_type_id.code == "award_approval"
        )
        self.assertEqual(len(documents), 1)

    def _egp_document_type(self, code, name):
        DocumentType = self.env["purchase.requisition.egp.document.type"]
        doc_type = DocumentType.search([("code", "=", code)], limit=1)
        if not doc_type:
            doc_type = DocumentType.create({"name": name, "code": code})
        return doc_type

    def _attach_signed_egp_document(self, requisition, code, name, filename):
        doc_type = self._egp_document_type(code, name)
        document = self.env["purchase.requisition.egp.document"].create(
            {
                "requisition_id": requisition.id,
                "document_type_id": doc_type.id,
                "document_file": "dGVzdA==",
                "document_filename": filename,
            }
        )
        signing = document._ensure_attached_egp_studio_document()
        signing.with_context(skip_validation_check=True).write({"state": "approved"})
        return document

    def test_create_po_ready_when_consideration_and_winner_signed(self):
        if "studio_document_id" not in self.env["purchase.requisition.egp.document"]._fields:
            self.skipTest("vpk_document_studio_procurement is required for signed e-GP documents")
        self.requisition.action_confirm()
        self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
                "amount_total": 1000,
                "is_winner": True,
            }
        )
        self.assertFalse(self.requisition.egp_award_approved)
        self._attach_signed_egp_document(
            self.requisition,
            "consider_procurement",
            "รายงานผลการพิจารณาจัดซื้อจัดจ้าง",
            "consider.pdf",
        )
        self.assertFalse(self.requisition.egp_award_approved)
        self._attach_signed_egp_document(
            self.requisition,
            "winner_announcement",
            "ประกาศผู้ชนะ",
            "winner.pdf",
        )
        self.requisition.invalidate_recordset(["egp_award_approved"])
        self.assertTrue(self.requisition.egp_has_winner)
        self.assertTrue(self.requisition.egp_award_approved)
        action = self.requisition.action_open_create_rfq_wizard()
        self.assertEqual(action["res_model"], "purchase.requisition.create.rfq")
        wizard = self.env["purchase.requisition.create.rfq"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
            }
        )
        po_action = wizard.action_create_rfq()
        purchase_order = self.env["purchase.order"].browse(po_action["res_id"])
        self.assertEqual(purchase_order.partner_id, self.partner)

    def test_egp_committee_shows_pr_confirm_state(self):
        employee = self.env["hr.employee"].create({"name": "กรรมการ e-GP ยินยอม"})
        request = self._vpk_make_egp_purchase_request()
        member = self.env["procurement.committee"].create(
            {
                "request_id": request.id,
                "employee_id": employee.id,
                "name": employee.name,
                "committee_type": "work_acceptance",
                "approve_role": "chairman",
                "confirm_state": "accepted",
                "confirm_note": False,
            }
        )
        self.requisition.write({"purchase_request_id": request.id})
        self.requisition._copy_committees_from_purchase_request()
        egp_member = self.requisition.work_acceptance_committee_ids.filtered(
            lambda line: line.employee_id == employee
        )
        self.assertTrue(egp_member)
        self.assertEqual(egp_member.source_confirm_state, "accepted")
        member.write({"confirm_state": "declined", "confirm_note": "ติดราชการ"})
        egp_member.invalidate_recordset(
            ["source_confirm_state", "source_confirm_note", "confirm_state"]
        )
        self.assertEqual(egp_member.source_confirm_state, "declined")
        self.assertEqual(egp_member.source_confirm_note, "ติดราชการ")

    @tagged("post_install", "-at_install", "egp_bid_line_entry")
    def test_line_consideration_adds_vendor_product(self):
        self.requisition.egp_consideration_type = "line"
        empty_bid = self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": self.requisition.id,
                "partner_id": self.partner.id,
            }
        )
        self.assertFalse(empty_bid.line_ids)
        BidLine = self.env["purchase.requisition.egp.bid.line"].with_context(
            default_requisition_id=self.requisition.id
        )
        line = BidLine.create(
            {
                "partner_id": self.partner.id,
                "product_id": self.product.id,
                "product_qty": 2.0,
                "price_unit": 150.0,
            }
        )
        self.assertEqual(line.requisition_id, self.requisition)
        self.assertEqual(line.bid_id, empty_bid)
        self.assertEqual(line.bid_id.partner_id, self.partner)
        self.assertEqual(line.bid_id.amount_total, 300.0)
        self.assertEqual(self.requisition.egp_bid_ids, line.bid_id)
        self.assertEqual(self.requisition.egp_bid_line_ids, line)

        product_b = self.env["product.product"].create(
            {
                "name": "Second Supply",
                "purchase_ok": True,
                "type": "consu",
            }
        )
        line_b = BidLine.create(
            {
                "partner_id": self.partner.id,
                "product_id": product_b.id,
                "product_qty": 1.0,
                "price_unit": 50.0,
            }
        )
        self.assertEqual(line_b.bid_id, line.bid_id)
        self.assertEqual(line.bid_id.amount_total, 350.0)

        line_c = BidLine.create(
            {
                "partner_id": self.partner_b.id,
                "product_id": self.product.id,
                "product_qty": 2.0,
                "price_unit": 140.0,
            }
        )
        self.assertEqual(len(self.requisition.egp_bid_ids), 2)
        self.assertEqual(line_c.bid_id.amount_total, 280.0)

        total_req = self.env["purchase.requisition"].create(
            {
                "requisition_type": "egp_procurement",
                "egp_consideration_type": "total",
                "line_ids": [
                    (
                        0,
                        0,
                        {
                            "product_id": self.product.id,
                            "product_qty": 5.0,
                        },
                    )
                ],
            }
        )
        total_bid = self.env["purchase.requisition.egp.bid"].create(
            {
                "requisition_id": total_req.id,
                "partner_id": self.partner.id,
                "amount_total": 1000.0,
            }
        )
        self.assertEqual(len(total_bid.line_ids), 1)
        self.assertEqual(total_bid.amount_total, 1000.0)

        from_total = BidLine.create(
            {
                "partner_id": self.partner_b.id,
                "product_id": product_b.id,
                "product_qty": 5.0,
                "price_subtotal": 1000.0,
            }
        )
        self.assertAlmostEqual(from_total.price_unit, 200.0)
        self.assertAlmostEqual(from_total.price_subtotal, 1000.0)
        self.assertAlmostEqual(from_total.bid_id.amount_total, 1280.0)
        from_total.write({"price_subtotal": 250.0})
        self.assertAlmostEqual(from_total.price_unit, 50.0)
        self.assertAlmostEqual(from_total.price_subtotal, 250.0)

    @tagged("post_install", "-at_install", "egp_bid_line_entry")
    def test_bid_line_pulls_budget_from_purchase_request(self):
        request = self._vpk_make_egp_purchase_request()
        request.line_ids.write({"estimated_cost": 1000.0, "product_qty": 5.0})
        if "budget_unit_price" in request.line_ids._fields:
            request.line_ids.write({"budget_unit_price": 200.0})
        self.requisition.write(
            {
                "egp_consideration_type": "line",
                "purchase_request_id": request.id,
            }
        )
        self.requisition.line_ids.write(
            {
                "product_qty": 5.0,
                "purchase_request_lines": [(6, 0, request.line_ids.ids)],
            }
        )
        BidLine = self.env["purchase.requisition.egp.bid.line"].with_context(
            default_requisition_id=self.requisition.id
        )
        line = BidLine.create(
            {
                "partner_id": self.partner.id,
                "product_id": self.product.id,
            }
        )
        self.assertEqual(line.requisition_line_id, self.requisition.line_ids[:1])
        self.assertAlmostEqual(line.budget_amount, 1000.0)
        self.assertAlmostEqual(line.budget_unit_price, 200.0)
        self.assertEqual(line.product_qty, 5.0)

    @tagged("post_install", "-at_install", "egp_bid_line_entry")
    def test_summarize_offered_totals_from_product_lines(self):
        self.requisition.egp_consideration_type = "line"
        BidLine = self.env["purchase.requisition.egp.bid.line"].with_context(
            default_requisition_id=self.requisition.id
        )
        line_a = BidLine.create(
            {
                "partner_id": self.partner.id,
                "product_id": self.product.id,
                "product_qty": 2.0,
                "price_unit": 150.0,
            }
        )
        product_b = self.env["product.product"].create(
            {
                "name": "Budget Item B",
                "purchase_ok": True,
                "type": "consu",
            }
        )
        BidLine.create(
            {
                "partner_id": self.partner.id,
                "product_id": product_b.id,
                "product_qty": 1.0,
                "price_unit": 50.0,
            }
        )
        BidLine.create(
            {
                "partner_id": self.partner_b.id,
                "product_id": self.product.id,
                "product_qty": 2.0,
                "price_unit": 140.0,
            }
        )
        line_a.bid_id.amount_total = 9999.0
        action = self.requisition.action_summarize_offered_totals()
        self.assertEqual(action["tag"], "display_notification")
        self.assertAlmostEqual(line_a.bid_id.amount_total, 350.0)
        other_bid = self.requisition.egp_bid_ids.filtered(
            lambda bid: bid.partner_id == self.partner_b
        )
        self.assertAlmostEqual(other_bid.amount_total, 280.0)

    @tagged("post_install", "-at_install", "egp_bid_line_entry")
    def test_line_winner_is_exclusive_per_product(self):
        self.requisition.egp_consideration_type = "line"
        BidLine = self.env["purchase.requisition.egp.bid.line"].with_context(
            default_requisition_id=self.requisition.id
        )
        line_a = BidLine.create(
            {
                "partner_id": self.partner.id,
                "product_id": self.product.id,
                "product_qty": 2.0,
                "price_unit": 150.0,
            }
        )
        line_b = BidLine.create(
            {
                "partner_id": self.partner_b.id,
                "product_id": self.product.id,
                "product_qty": 2.0,
                "price_unit": 140.0,
            }
        )
        line_a.write({"is_winner": True})
        self.assertTrue(line_a.is_winner)
        self.assertFalse(line_b.is_winner)
        line_b.write({"is_winner": True})
        self.assertTrue(line_b.is_winner)
        self.assertFalse(line_a.is_winner)
        self.assertAlmostEqual(line_a.bid_id.budget_amount_total, line_a.budget_amount)

    @tagged("post_install", "-at_install", "egp_bid_line_entry")
    def test_contract_essence_values_from_line_winners(self):
        self.requisition.write(
            {
                "egp_consideration_type": "line",
                "egp_reference": "68124983449",
                "egp_project_name": "ประกวดราคาซื้อยา",
            }
        )
        BidLine = self.env["purchase.requisition.egp.bid.line"].with_context(
            default_requisition_id=self.requisition.id
        )
        BidLine.create(
            {
                "partner_id": self.partner.id,
                "product_id": self.product.id,
                "name": "Medical Supply",
                "product_qty": 2.0,
                "price_unit": 150.0,
                "is_winner": True,
            }
        )
        values = self.requisition._get_contract_essence_values()
        self.assertEqual(values["project_no"], "68124983449")
        self.assertEqual(values["consideration"], "ราคาต่อรายการ")
        self.assertEqual(len(values["offer_rows"]), 1)
        self.assertEqual(values["offer_rows"][0]["vendor"], "Vendor A")
        self.assertEqual(values["offer_rows"][0]["amount"], "300.00")
        self.assertEqual(len(values["winner_rows"]), 1)
        self.assertEqual(values["winner_rows"][0]["vendor"], "Vendor A")
        self.assertEqual(values["winner_rows"][0]["egp_control"], "-")
        action = self.requisition.action_print_contract_essence()
        self.assertEqual(action["type"], "ir.actions.report")
        self.assertIn("contract_essence", action["report_name"])
