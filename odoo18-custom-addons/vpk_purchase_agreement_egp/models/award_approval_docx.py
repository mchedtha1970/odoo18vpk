# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

"""Fill รายงานผลการพิจารณาจัดซื้อจัดจ้าง.docx from {{placeholders}}."""

import io
import os
from copy import deepcopy

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    Document = None
    qn = None


ITEM_KEYS = ("item_desc", "vendor_name", "offer_price", "agreed_price")


class AwardApprovalRenderError(ValueError):
    """Raised when the award-report Word template cannot be filled."""


def _require_docx():
    if Document is None:
        raise AwardApprovalRenderError(
            "ไม่พบไลบรารี python-docx กรุณาติดตั้งในสภาพแวดล้อม Odoo "
            "ก่อนพิมพ์รายงานผลการพิจารณา"
        )


def _paragraph_text(paragraph):
    return "".join(run.text or "" for run in paragraph.runs)


def _set_paragraph_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_placeholders(document, mapping):
    for paragraph in _iter_paragraphs(document):
        original = _paragraph_text(paragraph) or paragraph.text
        if "{{" not in original:
            continue
        updated = original
        for key, value in mapping.items():
            updated = updated.replace("{{%s}}" % key, value if value is not None else "")
        if updated != original:
            _set_paragraph_text(paragraph, updated)


def _cell_plain(cell):
    return "\n".join(
        _paragraph_text(paragraph) or paragraph.text for paragraph in cell.paragraphs
    )


def _set_cell_lines(cell, text):
    lines = (text or "").split("\n") or [""]
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        return
    _set_paragraph_text(paragraphs[0], lines[0])
    template = paragraphs[0]._element
    for paragraph in paragraphs[1:]:
        parent = paragraph._element.getparent()
        if parent is not None:
            parent.remove(paragraph._element)
    parent = paragraphs[0]._element.getparent()
    current = paragraphs[0]._element
    for line in lines[1:]:
        clone = deepcopy(template)
        current.addnext(clone)
        current = clone
        for node in clone.findall(".//" + qn("w:t")):
            node.text = line
            break


def _item_table(document):
    for table in document.tables:
        for row in table.rows:
            if "{{item_desc}}" in _cell_plain(row.cells[0]):
                return table, row
    return None, None


def _fill_items(document, items):
    table, template_row = _item_table(document)
    if template_row is None:
        raise AwardApprovalRenderError(
            "ไฟล์แบบฟอร์มไม่มีตารางรายการพิจารณา กรุณาใช้เทมเพลตต้นฉบับ"
        )
    rows_needed = items or [
        {
            "item_desc": "-",
            "vendor_name": "",
            "offer_price": "",
            "agreed_price": "",
        }
    ]
    current = template_row._tr
    for _extra in rows_needed[1:]:
        clone = deepcopy(template_row._tr)
        current.addnext(clone)
        current = clone
    data_rows = [
        row
        for row in table.rows
        if "{{item_desc}}" in _cell_plain(row.cells[0])
    ]
    if len(data_rows) < len(rows_needed):
        raise AwardApprovalRenderError("เพิ่มแถวรายการในรายงานผลไม่สำเร็จ")
    for row, item in zip(data_rows, rows_needed):
        values = [item.get(key) or "" for key in ITEM_KEYS]
        for cell, value in zip(row.cells, values):
            _set_cell_lines(cell, value)


def template_path():
    """Filesystem copy owned by vpk_official_document. Installed prints use the template record."""
    return os.path.normpath(
        os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "..",
            "vpk_official_document",
            "static",
            "src",
            "templates",
            "award_approval.docx",
        )
    )


def load_template_bytes(path=None):
    path = path or template_path()
    if not os.path.isfile(path):
        raise AwardApprovalRenderError(
            "ไม่พบไฟล์เทมเพลตรายงานผลการพิจารณา (%s)" % path
        )
    with open(path, "rb") as handle:
        return handle.read()


def render_award_approval_docx(template_bytes, values):
    """Return filled .docx bytes for a procurement consideration report."""
    _require_docx()
    if not template_bytes:
        raise AwardApprovalRenderError("เทมเพลตรายงานผลการพิจารณาว่างเปล่า")
    document = Document(io.BytesIO(template_bytes))
    items = list(values.get("items") or [])
    _fill_items(document, items)
    mapping = {
        key: "" if value is None else str(value)
        for key, value in values.items()
        if key != "items"
    }
    _replace_placeholders(document, mapping)
    leftover = []
    for paragraph in _iter_paragraphs(document):
        text = _paragraph_text(paragraph) or paragraph.text
        if "{{" in text:
            leftover.append(text)
    if leftover:
        raise AwardApprovalRenderError(
            "เทมเพลตยังมีช่องที่ไม่ได้แทนค่า: %s" % leftover[0]
        )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
