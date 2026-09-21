# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

"""Fill ประกาศผู้ชนะการเสนอราคา.docx from {{placeholders}}."""

import io
import os

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    Document = None
    qn = None


class WinnerAnnouncementRenderError(ValueError):
    """Raised when the winner announcement Word template cannot be filled."""


def _require_docx():
    if Document is None:
        raise WinnerAnnouncementRenderError(
            "ไม่พบไลบรารี python-docx กรุณาติดตั้งในสภาพแวดล้อม Odoo ก่อนสร้างประกาศผู้ชนะ"
        )


def _set_paragraph_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_placeholders(document, mapping):
    for paragraph in _iter_paragraphs(document):
        original = "".join(run.text or "" for run in paragraph.runs) or paragraph.text
        if "{{" not in original:
            continue
        updated = original
        for key, value in mapping.items():
            updated = updated.replace("{{%s}}" % key, value if value is not None else "")
        if updated != original:
            _set_paragraph_text(paragraph, updated)


def template_path():
    return os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "static",
        "src",
        "templates",
        "winner_announcement.docx",
    )


def load_template_bytes(path=None):
    path = path or template_path()
    if not os.path.isfile(path):
        raise WinnerAnnouncementRenderError(
            "ไม่พบไฟล์เทมเพลตประกาศผู้ชนะ (%s)" % path
        )
    with open(path, "rb") as handle:
        return handle.read()


def render_winner_announcement_docx(template_bytes, values):
    """Return filled .docx bytes for a winner announcement."""
    _require_docx()
    if not template_bytes:
        raise WinnerAnnouncementRenderError("เทมเพลตประกาศผู้ชนะว่างเปล่า")
    document = Document(io.BytesIO(template_bytes))
    mapping = {key: "" if value is None else str(value) for key, value in values.items()}
    _replace_placeholders(document, mapping)
    leftover = []
    for paragraph in _iter_paragraphs(document):
        text = "".join(run.text or "" for run in paragraph.runs) or paragraph.text
        if "{{" in text:
            leftover.append(text)
    if leftover:
        raise WinnerAnnouncementRenderError(
            "เทมเพลตยังมีช่องที่ไม่ได้แทนค่า: %s" % leftover[0]
        )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
