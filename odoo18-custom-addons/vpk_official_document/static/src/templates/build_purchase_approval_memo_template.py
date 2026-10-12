#!/usr/bin/env python3
"""บันทึกข้อความอนุมัติขอซื้อ — layout from the hospital memo (ครุฑกึ่งกลาง)."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips

FONT = "TH Sarabun New"
OUT = Path(__file__).with_name("purchase_approval_memo.docx")
GARUDA = (
    Path(__file__).resolve().parents[4]
    / "vpk_purchase_agreement_egp"
    / "static"
    / "src"
    / "img"
    / "garuda.jpeg"
)
CONTENT_WIDTH = Cm(16.5)


def _set_half_point(r_pr, tag, half):
    node = r_pr.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        r_pr.append(node)
    node.set(qn("w:val"), str(half))


def set_run_font(run, size_pt=16, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), FONT)
    half = int(size_pt * 2)
    _set_half_point(r_pr, "w:sz", half)
    _set_half_point(r_pr, "w:szCs", half)


def add_para(doc, text, *, size=16, bold=False, align="left", before=0, after=0, first_line=None):
    para = doc.add_paragraph()
    para.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "thai_justify": WD_ALIGN_PARAGRAPH.THAI_JUSTIFY,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = para.paragraph_format
    pf.space_before = Twips(before)
    pf.space_after = Twips(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    run = para.add_run(text)
    set_run_font(run, size_pt=size, bold=bold)
    return para


def _set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement("w:%s" % edge)
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tc_pr.append(borders)


def _set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width)))
    tc_w.set(qn("w:type"), "dxa")


def _cell_text(cell, text, *, size=16, bold=False, align="left", before=0):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = paragraph.paragraph_format
    pf.space_before = Twips(before)
    pf.space_after = Twips(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if paragraph.runs:
        paragraph.runs[0].text = text
        run = paragraph.runs[0]
        for extra in paragraph.runs[1:]:
            extra.text = ""
    else:
        run = paragraph.add_run(text)
    set_run_font(run, size_pt=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    total = sum(widths)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl_pr.addnext(grid)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            _set_cell_width(cell, width)
            _set_cell_borders(cell)


def add_sign_cell(cell, name_key, position_key, date_key=None):
    _cell_text(cell, "", before=0)
    lines = [
        "({{%s}})" % name_key,
        "{{%s}}" % position_key,
    ]
    if date_key:
        lines.append("{{%s}}" % date_key)
    _cell_text(cell, lines[0], align="center", before=1100)
    for line in lines[1:]:
        extra = cell.add_paragraph()
        extra.alignment = WD_ALIGN_PARAGRAPH.CENTER
        extra.paragraph_format.space_before = Twips(0)
        extra.paragraph_format.space_after = Twips(0)
        extra.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = extra.add_run(line)
        set_run_font(run, size_pt=16)


def main():
    if not GARUDA.is_file():
        raise SystemExit("missing garuda image: %s" % GARUDA)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.27)
    section.bottom_margin = Cm(1.27)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(16)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), FONT)
    _set_half_point(r_pr, "w:sz", 32)
    _set_half_point(r_pr, "w:szCs", 32)

    emblem = add_para(doc, "", align="center", before=0, after=0)
    emblem.runs[0].add_picture(str(GARUDA), width=Cm(1.8))
    add_para(doc, "บันทึกข้อความ", size=20, bold=True, align="center", before=40, after=80)
    add_para(doc, "ส่วนราชการ  {{agency}}", before=40, after=0)
    ref = doc.add_table(rows=1, cols=2)
    widths = [int(Cm(8.2).twips), int(Cm(8.3).twips)]
    _set_table_widths(ref, widths)
    _cell_text(ref.rows[0].cells[0], "ที่  {{memo_number}}")
    _cell_text(ref.rows[0].cells[1], "วันที่  {{memo_date}}")
    add_para(doc, "เรื่อง  {{subject}}", before=40, after=0)
    add_para(doc, "เรียน  {{recipient}}", before=40, after=80)
    add_para(doc, "{{intro}}", align="thai_justify", before=60, after=40, first_line=2.5)
    for index in range(1, 6):
        add_para(doc, "{{item_%s}}" % index, align="justify", before=0, after=0, first_line=2.5)
    add_para(doc, "{{body_2}}", align="thai_justify", before=60, after=40, first_line=2.5)
    add_para(doc, "จึงเรียนมาเพื่อโปรดพิจารณาอนุมัติ", align="thai_justify", before=80, after=60, first_line=2.5)

    proposer = doc.add_table(rows=1, cols=1)
    _set_table_widths(proposer, [int(CONTENT_WIDTH.twips)])
    add_sign_cell(proposer.rows[0].cells[0], "proposer_name", "proposer_position")

    signs = doc.add_table(rows=1, cols=2)
    half = int(Cm(8.25).twips)
    _set_table_widths(signs, [half, half])
    add_sign_cell(signs.rows[0].cells[0], "reviewer_name", "reviewer_position", "reviewer_date")
    add_sign_cell(signs.rows[0].cells[1], "approver_name", "approver_position", "approver_date")

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
