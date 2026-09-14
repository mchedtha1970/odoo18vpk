# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

"""Fill the official Word form for คำสั่งแต่งตั้งคณะกรรมการตรวจรับพัสดุ."""

import io
import re
from copy import deepcopy

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover
    Document = None
    WD_ALIGN_PARAGRAPH = None
    OxmlElement = None
    qn = None
    Paragraph = None

# Match the TTF family name used by THSarabunNew-*.ttf so LibreOffice can embed it.
THAI_PDF_FONT = "TH Sarabun New"


MEMBER_LINE_RE = re.compile(r"^\s*\d+[\.．)]")
DUTY_DEFAULT = (
    "โดยให้มีอำนาจหน้าที่  ทำการตรวจรับพัสดุให้เป็นไปตามเงื่อนไขของสัญญาหรือข้อตกลงนั้น  "
    "และปฏิบัติตามระเบียบกระทรวงการคลังว่าด้วยการจัดซื้อจัดจ้างและการบริหารพัสดุภาครัฐ  "
    "พ.ศ. ๒๕๖๐  หมวด  ๖  การบริหารสัญญา  และการตรวจรับพัสดุ  ข้อ  ๑๗๕"
)
EFFECTIVE_DEFAULT = (
    "ทั้งนี้  ตั้งแต่บัดนี้เป็นต้นไป  หรือจนกว่าจะดำเนินการตรวจรับพัสดุแล้วเสร็จ"
)


class OfficialDocumentRenderError(ValueError):
    """Raised when the Word template cannot be filled."""


def _require_docx():
    if Document is None:
        raise OfficialDocumentRenderError(
            "ไม่พบไลบรารี python-docx กรุณาติดตั้งในสภาพแวดล้อม Odoo ก่อนพิมพ์หนังสือราชการ"
        )


def _paragraph_text(paragraph):
    return "".join(run.text or "" for run in paragraph.runs)


def _set_paragraph_text(paragraph, text):
    """Replace visible text while keeping the first run's font/style."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _clone_paragraph_after(paragraph):
    new_elm = deepcopy(paragraph._element)
    paragraph._element.addnext(new_elm)
    return Paragraph(new_elm, paragraph._parent)


def _insert_paragraph_after(paragraph, text, style_from=None):
    source = style_from or paragraph
    new_elm = deepcopy(source._element)
    paragraph._element.addnext(new_elm)
    new_para = Paragraph(new_elm, paragraph._parent)
    _set_paragraph_text(new_para, text)
    return new_para


def _remove_paragraph(paragraph):
    elm = paragraph._element
    parent = elm.getparent()
    if parent is not None:
        parent.remove(elm)


def _set_run_font(run, font_name=THAI_PDF_FONT):
    if OxmlElement is None:
        return
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), font_name)


def _iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _apply_thai_font(document, font_name=THAI_PDF_FONT):
    """Force a system-available Thai font so PDF conversion can embed glyphs."""
    for paragraph in _iter_paragraphs(document):
        for run in paragraph.runs:
            _set_run_font(run, font_name)


def _is_member_line(text):
    return bool(MEMBER_LINE_RE.match(text or ""))


def _classify_paragraphs(doc):
    """Map template paragraphs of the official order form."""
    classified = {
        "title": None,
        "order_no": None,
        "subject": None,
        "body": None,
        "members": [],
        "duty": None,
        "effective": None,
        "order_date": None,
        "signer_name": None,
        "signer_position": None,
    }
    text_paras = []
    for paragraph in doc.paragraphs:
        text = _paragraph_text(paragraph).strip()
        if text:
            text_paras.append((paragraph, text))

    for paragraph, text in text_paras:
        if classified["title"] is None and text.startswith("คำสั่ง"):
            classified["title"] = paragraph
        elif classified["order_no"] is None and text.startswith("ที่"):
            classified["order_no"] = paragraph
        elif classified["subject"] is None and text.startswith("เรื่อง"):
            classified["subject"] = paragraph
        elif classified["body"] is None and text.startswith("ด้วย"):
            classified["body"] = paragraph
        elif _is_member_line(text):
            classified["members"].append(paragraph)
        elif classified["duty"] is None and text.startswith("โดยให้มีอำนาจ"):
            classified["duty"] = paragraph
        elif classified["effective"] is None and text.startswith("ทั้งนี้"):
            classified["effective"] = paragraph
        elif classified["order_date"] is None and text.startswith("สั่ง"):
            classified["order_date"] = paragraph
        elif classified["signer_name"] is None and text.startswith("("):
            classified["signer_name"] = paragraph
        elif classified["signer_name"] is not None and classified["signer_position"] is None:
            classified["signer_position"] = paragraph

    return classified


def format_member_line(index, member):
    name = (member.get("name") or "").strip() or "................................"
    position = (member.get("position") or "").strip()
    role = (member.get("role") or "กรรมการ").strip()
    if position:
        return f"{index}. {name}  ตำแหน่ง {position}  {role}"
    return f"{index}. {name}  {role}"


def render_wa_committee_order_docx(template_bytes, values):
    """Return filled .docx bytes from the official appointment-order template.

    ``values`` keys:
        agency, order_no, year_be, subject, body, members (list of dict),
        duty, effective, order_date, signer_name, signer_position
    """
    _require_docx()
    members = list(values.get("members") or [])
    if not members:
        raise OfficialDocumentRenderError(
            "กรุณาระบุรายชื่อคณะกรรมการตรวจรับอย่างน้อย 1 คน"
        )

    document = Document(io.BytesIO(template_bytes))
    parts = _classify_paragraphs(document)
    if not parts["title"] or not parts["members"]:
        raise OfficialDocumentRenderError(
            "ไฟล์แบบฟอร์มไม่ตรงกับคำสั่งแต่งตั้งคณะกรรมการตรวจรับพัสดุ "
            "กรุณาใช้ไฟล์ต้นฉบับที่กำหนด"
        )

    agency = (values.get("agency") or "").strip()
    order_no = (values.get("order_no") or "").strip() or "...."
    year_be = (values.get("year_be") or "").strip() or "........"
    order_no_line = (values.get("order_no_line") or "").strip()
    if not order_no_line:
        order_no_line = f"ที่  {order_no} / {year_be}"
    elif not order_no_line.startswith("ที่"):
        order_no_line = f"ที่  {order_no_line}"
    subject = (values.get("subject") or "").strip()
    body = (values.get("body") or "").strip()
    duty = (values.get("duty") or "").strip() or DUTY_DEFAULT
    effective = (values.get("effective") or "").strip() or EFFECTIVE_DEFAULT
    order_date = (values.get("order_date") or "").strip()
    signer_name = (values.get("signer_name") or "").strip()
    signer_position = (values.get("signer_position") or "").strip()

    _set_paragraph_text(parts["title"], f"คำสั่ง{agency}" if agency else "คำสั่ง")
    if parts["order_no"] is not None:
        _set_paragraph_text(parts["order_no"], order_no_line)
    if parts["subject"] is not None:
        prefix = "" if subject.startswith("เรื่อง") else "เรื่อง  "
        _set_paragraph_text(parts["subject"], f"{prefix}{subject}")
    if parts["body"] is not None and body:
        _set_paragraph_text(parts["body"], body)

    member_paras = parts["members"]
    template_member = member_paras[0]
    for extra in member_paras[1:]:
        _remove_paragraph(extra)

    _set_paragraph_text(template_member, format_member_line(1, members[0]))
    current = template_member
    for index, member in enumerate(members[1:], start=2):
        current = _clone_paragraph_after(current)
        _set_paragraph_text(current, format_member_line(index, member))

    if parts["duty"] is not None:
        _set_paragraph_text(parts["duty"], duty)
    if parts["effective"] is not None:
        _set_paragraph_text(parts["effective"], effective)
    if parts["order_date"] is not None and order_date:
        prefix = "" if order_date.startswith("สั่ง") else "สั่ง  ณ  วันที่  "
        _set_paragraph_text(parts["order_date"], f"{prefix}{order_date}")
    if parts["signer_name"] is not None:
        name = signer_name or "................................"
        if not name.startswith("("):
            name = f"({name})"
        _set_paragraph_text(parts["signer_name"], name)
    if parts["signer_position"] is not None and signer_position:
        _set_paragraph_text(parts["signer_position"], signer_position)

    _apply_thai_font(document)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


INTEGRITY_ROLE_LABELS = {
    "head_officer": "หัวหน้าเจ้าหน้าที่",
    "officer": "เจ้าหน้าที่",
    "chairman": "ประธานกรรมการตรวจรับพัสดุ",
    "committee": "กรรมการตรวจรับพัสดุ",
    "secretary": "กรรมการและเลขานุการ",
}


def _clone_table_row(row):
    new_tr = deepcopy(row._tr)
    row._tr.addnext(new_tr)
    return new_tr


def _is_item_list_table(table):
    if not table.rows:
        return False
    header = " ".join((cell.text or "") for cell in table.rows[0].cells)
    if "ราคาต่อหน่วย" in header:
        return True
    return "ลำดับ" in header and "รายการ" in header and "จำนวนเงิน" in header


def _remove_item_list_tables(document):
    """Drop leftover PR item grids so section ๒ stays text-only in Word and PDF."""
    for table in list(document.tables):
        if not _is_item_list_table(table):
            continue
        elm = table._tbl
        parent = elm.getparent()
        if parent is not None:
            parent.remove(elm)


def _ensure_child(parent, tag):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_table_column_widths(table, widths_twips):
    """Force fixed column widths so Thai labels like ข้าพเจ้า stay on one line."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    layout = _ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths_twips):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        cells = _unique_row_cells(row)
        for cell, width in zip(cells, widths_twips):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = _ensure_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def _keep_cell_on_one_line(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    _ensure_child(tc_pr, "w:noWrap")
    for paragraph in cell.paragraphs:
        if WD_ALIGN_PARAGRAPH is not None:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_pr = paragraph._p.get_or_add_pPr()
        jc = _ensure_child(p_pr, "w:jc")
        jc.set(qn("w:val"), "left")


def _layout_integrity_name_table(table):
    # Original col0 is 910 twips (~1.6cm), too narrow for ข้าพเจ้า at 16pt.
    _set_table_column_widths(table, (1800, 4012, 3214))
    for row in table.rows:
        cells = _unique_row_cells(row)
        if cells:
            _keep_cell_on_one_line(cells[0])


def _fill_cell_text(cell, text):
    if not cell.paragraphs:
        cell.add_paragraph(text or "")
        return
    _set_paragraph_text(cell.paragraphs[0], text or "")
    for paragraph in cell.paragraphs[1:]:
        _set_paragraph_text(paragraph, "")


def _person_display_name(person):
    return (person.get("name") or "").strip()


def _person_position(person):
    return (person.get("position") or "").strip()


def _fill_integrity_name_row(row, person, role_key):
    cells = row.cells
    if len(cells) < 3:
        return
    _fill_cell_text(cells[0], "ข้าพเจ้า")
    _fill_cell_text(cells[1], _person_display_name(person))
    label = INTEGRITY_ROLE_LABELS.get(role_key, role_key or "กรรมการตรวจรับพัสดุ")
    _fill_cell_text(cells[2], f"({label})")


def _fill_integrity_sign_cell(cell, person, role_key):
    name = _person_display_name(person) or "................................"
    position = _person_position(person) or "..."
    label = INTEGRITY_ROLE_LABELS.get(role_key, role_key or "กรรมการตรวจรับพัสดุ")
    for paragraph in cell.paragraphs:
        text = _paragraph_text(paragraph)
        if "นาย/นาง/นางสาว" in text:
            _set_paragraph_text(paragraph, f"({name})")
        elif text.startswith("ตำแหน่ง"):
            _set_paragraph_text(
                paragraph,
                f"ตำแหน่ง {position}" if position != "..." else "ตำแหน่ง...",
            )
        elif any(
            token in text
            for token in ("หัวหน้าเจ้าหน้าที่", "เจ้าหน้าที่", "ประธานกรรมการ", "กรรมการ")
        ):
            _set_paragraph_text(paragraph, f"({label})")


def _unique_row_cells(row):
    seen = set()
    cells = []
    for cell in row.cells:
        ident = id(cell._tc)
        if ident in seen:
            continue
        seen.add(ident)
        cells.append(cell)
    return cells


def _set_cell_text(cell, text):
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.text = text
        return
    _set_paragraph_text(paragraphs[0], text)
    for extra in paragraphs[1:]:
        _set_paragraph_text(extra, "")


def _thai_index(number):
    return str(number).translate(str.maketrans("0123456789", "๐๑๒๓๔๕๖๗๘๙"))


def _paren_name(name):
    name = (name or "").strip()
    if not name:
        return "(................................)"
    if name.startswith("("):
        return name
    return f"({name})"


def _is_spec_member_line(text):
    value = (text or "").strip()
    return bool(MEMBER_LINE_RE.match(value)) and "ตำแหน่ง" in value


def format_spec_member_line(index, member):
    name = (member.get("name") or "").strip() or "..............."
    position = (member.get("position") or "").strip() or "..."
    role = (member.get("role") or "กรรมการ").strip()
    return f"{_thai_index(index)}. {name}\tตำแหน่ง {position}\t{role}"


def render_spec_price_committee_docx(template_bytes, values):
    """Fill บันทึกข้อความขออนุมัติแต่งตั้งคณะกรรมการกำหนดรายละเอียดคุณลักษณะ."""
    _require_docx()
    members = list(values.get("members") or [])
    if not members:
        raise OfficialDocumentRenderError(
            "กรุณาระบุรายชื่อคณะกรรมการกำหนดคุณลักษณะอย่างน้อย 1 คน"
        )

    document = Document(io.BytesIO(template_bytes))
    if not document.tables:
        raise OfficialDocumentRenderError(
            "ไฟล์แบบฟอร์มไม่ตรงกับบันทึกข้อความขออนุมัติแต่งตั้งคณะกรรมการ "
            "กรุณาใช้ไฟล์ต้นฉบับที่กำหนด"
        )

    agency = (values.get("agency") or "").strip()
    memo_number = (values.get("memo_number") or "").strip()
    order_date = (values.get("order_date") or "").strip()
    subject = (values.get("subject") or "").strip()
    recipient = (values.get("recipient") or "").strip()
    body = (values.get("body") or "").strip()
    signer_name = (values.get("signer_name") or "").strip()
    signer_position = (values.get("signer_position") or "").strip()
    signer_org = (values.get("signer_org") or "").strip()
    if not signer_org:
        if "ผู้อำนวยการ" in signer_position:
            signer_org = signer_position
        else:
            signer_org = "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต"

    table = document.tables[0]
    if len(table.rows) < 5:
        raise OfficialDocumentRenderError(
            "ไฟล์แบบฟอร์มไม่ตรงกับบันทึกข้อความขออนุมัติแต่งตั้งคณะกรรมการ "
            "กรุณาใช้ไฟล์ต้นฉบับที่กำหนด"
        )

    agency_cells = _unique_row_cells(table.rows[1])
    number_cells = _unique_row_cells(table.rows[2])
    subject_cells = _unique_row_cells(table.rows[3])
    recipient_cells = _unique_row_cells(table.rows[4])
    if len(agency_cells) > 1 and agency:
        _set_cell_text(agency_cells[1], agency)
    if len(number_cells) > 1 and memo_number:
        _set_cell_text(number_cells[1], memo_number)
    if len(number_cells) > 3 and order_date:
        _set_cell_text(number_cells[3], order_date)
    if len(subject_cells) > 1 and subject:
        _set_cell_text(subject_cells[1], subject)
    if len(recipient_cells) > 1 and recipient:
        _set_cell_text(recipient_cells[1], recipient)

    member_paras = []
    body_para = None
    signer_name_para = None
    signer_position_para = None
    signer_org_para = None
    for paragraph in document.paragraphs:
        text = _paragraph_text(paragraph).strip()
        if body_para is None and text.startswith("ด้วย"):
            body_para = paragraph
        elif _is_spec_member_line(text):
            member_paras.append(paragraph)
        elif signer_name_para is None and text.startswith("("):
            signer_name_para = paragraph
        elif signer_name_para is not None and signer_position_para is None and (
            text.startswith("ตำแหน่ง") or not text
        ):
            signer_position_para = paragraph
        elif (
            signer_position_para is not None
            and signer_org_para is None
            and (text.startswith("ผู้อำนวยการ") or text.startswith("ปฏิบัติ"))
        ):
            signer_org_para = paragraph

    if body_para is not None and body:
        _set_paragraph_text(body_para, body)
    if not member_paras:
        raise OfficialDocumentRenderError(
            "ไฟล์แบบฟอร์มไม่ตรงกับบันทึกข้อความขออนุมัติแต่งตั้งคณะกรรมการ "
            "กรุณาใช้ไฟล์ต้นฉบับที่กำหนด"
        )

    template_member = member_paras[0]
    for extra in member_paras[1:]:
        _remove_paragraph(extra)
    _set_paragraph_text(template_member, format_spec_member_line(1, members[0]))
    current = template_member
    for index, member in enumerate(members[1:], start=2):
        current = _clone_paragraph_after(current)
        _set_paragraph_text(current, format_spec_member_line(index, member))

    if signer_name_para is not None and signer_name:
        _set_paragraph_text(signer_name_para, _paren_name(signer_name))
    if signer_position_para is not None and signer_position:
        prefix = "" if signer_position.startswith("ตำแหน่ง") else "ตำแหน่ง "
        _set_paragraph_text(signer_position_para, f"{prefix}{signer_position}")
    if signer_org_para is not None and signer_org:
        _set_paragraph_text(signer_org_para, signer_org)

    _apply_thai_font(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_integrity_over_100k_docx(template_bytes, values):
    """Return filled .docx bytes for แบบแสดงความบริสุทธิ์ใจ (วงเงินมากกว่า 100,000 บาท)."""
    _require_docx()
    document = Document(io.BytesIO(template_bytes))
    if len(document.tables) < 2:
        raise OfficialDocumentRenderError(
            "ไฟล์แบบฟอร์มไม่ตรงกับแบบแสดงความบริสุทธิ์ใจ กรุณาใช้ไฟล์ต้นฉบับที่กำหนด"
        )

    head = values.get("head_officer") or {}
    officer = values.get("officer") or {}
    members = list(values.get("members") or [])
    chairman = next((item for item in members if item.get("role_key") == "chairman"), None)
    others = [item for item in members if item is not chairman]
    if chairman is None and members:
        chairman = members[0]
        others = members[1:]

    name_table = document.tables[0]
    while len(name_table.rows) < 3 + max(2, len(others)):
        _clone_table_row(name_table.rows[-1])

    _fill_integrity_name_row(name_table.rows[0], head, "head_officer")
    _fill_integrity_name_row(name_table.rows[1], officer, "officer")
    _fill_integrity_name_row(name_table.rows[2], chairman or {}, "chairman")
    for index, person in enumerate(others):
        row = name_table.rows[3 + index]
        role_key = person.get("role_key") or "committee"
        if role_key not in INTEGRITY_ROLE_LABELS:
            role_key = "committee"
        _fill_integrity_name_row(row, person, role_key)
    for extra_index in range(len(others), max(2, len(name_table.rows) - 3)):
        row_idx = 3 + extra_index
        if row_idx >= len(name_table.rows):
            break
        _fill_integrity_name_row(name_table.rows[row_idx], {}, "committee")

    sign_table = document.tables[1]
    needed_sign_rows = 3
    if len(others) > 3:
        needed_sign_rows = 3 + (len(others) - 3 + 1) // 2
    while len(sign_table.rows) < needed_sign_rows:
        _clone_table_row(sign_table.rows[-1])

    sign_slots = [
        (0, 0, head, "head_officer"),
        (0, 1, officer, "officer"),
        (1, 0, chairman or {}, "chairman"),
    ]
    for index, person in enumerate(others):
        role_key = person.get("role_key") or "committee"
        if role_key not in INTEGRITY_ROLE_LABELS:
            role_key = "committee"
        if index == 0:
            row_idx, col_idx = 1, 1
        elif index == 1:
            row_idx, col_idx = 2, 0
        elif index == 2:
            row_idx, col_idx = 2, 1
        else:
            offset = index - 3
            row_idx = 3 + offset // 2
            col_idx = offset % 2
        sign_slots.append((row_idx, col_idx, person, role_key))

    for row_idx, col_idx, person, role_key in sign_slots:
        if row_idx >= len(sign_table.rows):
            continue
        cells = _unique_row_cells(sign_table.rows[row_idx])
        if col_idx >= len(cells):
            continue
        _fill_integrity_sign_cell(cells[col_idx], person, role_key)

    _layout_integrity_name_table(name_table)
    _apply_thai_font(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


SPECIFIC_MEMBER_RE = re.compile(r"^\s*[๘8]\.\s*[๐-๙0-9]+")


def _fill_memo_header_table(table, values):
    agency = (values.get("agency") or "").strip()
    memo_number = (values.get("memo_number") or "").strip()
    order_date = (values.get("order_date") or "").strip()
    subject = (values.get("subject") or "").strip()
    recipient = (values.get("recipient") or "").strip()
    agency_cells = _unique_row_cells(table.rows[1])
    number_cells = _unique_row_cells(table.rows[2])
    subject_cells = _unique_row_cells(table.rows[3])
    recipient_cells = _unique_row_cells(table.rows[4])
    if len(agency_cells) > 1 and agency:
        _set_cell_text(agency_cells[1], agency)
    if len(number_cells) > 1 and memo_number:
        _set_cell_text(number_cells[1], memo_number)
    if len(number_cells) > 3 and order_date:
        _set_cell_text(number_cells[3], order_date)
    if len(subject_cells) > 1 and subject:
        _set_cell_text(subject_cells[1], subject)
    if len(recipient_cells) > 1 and recipient:
        _set_cell_text(recipient_cells[1], recipient)


def format_specific_member_line(index, member):
    name = (member.get("name") or "").strip() or "..............."
    position = (member.get("position") or "").strip() or "..."
    role = (member.get("role") or "กรรมการ").strip()
    return f"๑๐.{_thai_index(index)} {name}\tตำแหน่ง {position}\t{role}"


def _prefixed_section(text, prefix):
    value = (text or "").strip()
    if not value:
        return ""
    stripped = value.lstrip()
    if stripped.startswith(prefix) or stripped[:1] in "๓๔๕๖3456":
        return value
    return f"{prefix}  {value}"


def _find_sign_table(document):
    for table in reversed(document.tables):
        text = "\n".join(
            cell.text for row in table.rows for cell in _unique_row_cells(row)
        )
        if "หัวหน้าเจ้าหน้าที่" in text or "อนุมัติ" in text:
            return table
    if len(document.tables) > 1:
        return document.tables[-1]
    return None


def _fill_named_paragraphs(paragraphs, name, position):
    for paragraph in paragraphs:
        text = _paragraph_text(paragraph).strip()
        if text.startswith("(") and "นาย/นาง" in text:
            _set_paragraph_text(paragraph, _paren_name(name))
        elif text.startswith("ตำแหน่ง"):
            prefix = "" if position.startswith("ตำแหน่ง") else "ตำแหน่ง "
            _set_paragraph_text(paragraph, f"{prefix}{position}" if position else text)


def render_specific_method_approval_docx(template_bytes, values):
    """Fill รายงานขออนุมัติจัดซื้อจัดจ้างโดยวิธีเฉพาะเจาะจง."""
    _require_docx()
    document = Document(io.BytesIO(template_bytes))
    if not document.tables:
        raise OfficialDocumentRenderError(
            "ไฟล์แบบฟอร์มไม่ตรงกับรายงานขออนุมัติจัดซื้อจัดจ้างโดยวิธีเฉพาะเจาะจง "
            "กรุณาใช้ไฟล์ต้นฉบับที่กำหนด"
        )

    members = list(values.get("members") or [])
    body = (values.get("body") or "").strip()
    reason = (values.get("reason") or "").strip()
    work_detail = (values.get("work_detail") or "").strip()
    price_mid_text = (values.get("price_mid_text") or "").strip()
    budget_text = (values.get("budget_text") or "").strip()
    delivery_text = (values.get("delivery_text") or "").strip()
    method_text = (values.get("method_text") or "").strip()
    criteria_text = (values.get("criteria_text") or "").strip()
    announcement_text = (values.get("announcement_text") or "").strip()
    legal_text = (values.get("legal_text") or "").strip()
    other_proposal_text = (values.get("other_proposal_text") or "").strip()
    amount_text = (values.get("amount_text") or "").strip()
    baht_text = (values.get("baht_text") or "").strip()
    price_mid_source = (values.get("price_mid_source") or "").strip()
    budget_source = (values.get("budget_source") or "").strip()
    delivery_days = (values.get("delivery_days") or "").strip()
    recipient = (values.get("recipient") or "").strip()
    officer_name = (values.get("officer_name") or "").strip()
    officer_position = (values.get("officer_position") or "").strip()
    head_name = (values.get("head_officer_name") or "").strip()
    head_position = (values.get("head_officer_position") or "").strip()
    signer_name = (values.get("signer_name") or "").strip()
    signer_position = (values.get("signer_position") or "").strip()
    approver_acting = (values.get("approver_acting") or "").strip()
    inspector = members[0] if members else {}
    if not price_mid_text and amount_text:
        source = price_mid_source or "ราคากลางตามท้องตลาด"
        price_mid_text = (
            f"จำนวน  {amount_text}  บาท  ({baht_text})  จาก{source}"
        )
    if not budget_text and amount_text:
        budget_text = f"ภายในวงเงินงบประมาณ  {amount_text}  บาท  ({baht_text})"
        if budget_source:
            budget_text = f"{budget_text}  โดยเบิกจ่ายจาก{budget_source}"
    if not delivery_text and delivery_days:
        delivery_text = (
            f"กำหนดเวลาส่งมอบงานหรือให้งานแล้วเสร็จภายใน  {delivery_days}  วัน  "
            "นับถัดจากวันที่ลงนามในสัญญา"
        )

    header = document.tables[0]
    if len(header.rows) < 5:
        raise OfficialDocumentRenderError(
            "ไฟล์แบบฟอร์มไม่ตรงกับรายงานขออนุมัติจัดซื้อจัดจ้างโดยวิธีเฉพาะเจาะจง "
            "กรุณาใช้ไฟล์ต้นฉบับที่กำหนด"
        )
    _fill_memo_header_table(header, values)

    member_paras = []
    criteria_heading_para = None
    criteria_content_para = None
    other_heading_para = None
    junk_paras = []
    for paragraph in document.paragraphs:
        text = _paragraph_text(paragraph).strip()
        if text.startswith("ด้วย") and body:
            _set_paragraph_text(paragraph, body)
        elif text.startswith("เพื่อ") and "ระบุเหตุผล" in text and reason:
            _set_paragraph_text(paragraph, reason)
        elif text.startswith("๑. เหตุผล"):
            _set_paragraph_text(paragraph, "๑. เหตุผลและความจำเป็น")
        elif text.startswith("๒. รายละเอียดงาน"):
            _set_paragraph_text(paragraph, "๒. รายละเอียดงานที่จัดซื้อ จัดจ้าง")
        elif text.startswith("รายละเอียดตามเอกสารแนบท้าย"):
            if work_detail:
                lines = [
                    line.strip() for line in work_detail.splitlines() if line.strip()
                ] or [work_detail]
                _set_paragraph_text(paragraph, lines[0])
                current = paragraph
                for extra in lines[1:]:
                    current = _insert_paragraph_after(
                        current, extra, style_from=paragraph
                    )
        elif text.startswith("๓. ราคากลาง") and price_mid_text:
            _set_paragraph_text(
                paragraph,
                _prefixed_section(
                    price_mid_text,
                    "๓. ราคากลางและรายละเอียดของราคากลาง",
                ),
            )
        elif text.startswith("ภายในวงเงินงบประมาณ") and budget_text:
            _set_paragraph_text(paragraph, budget_text)
        elif text.startswith("กำหนดเวลาส่งมอบ") and delivery_text:
            _set_paragraph_text(paragraph, delivery_text)
        elif text.startswith("๖.๑") and method_text:
            _set_paragraph_text(
                paragraph,
                _prefixed_section(method_text, "๖.๑"),
            )
        elif text.startswith("๗. หลักเกณฑ์"):
            criteria_heading_para = paragraph
        elif text.startswith("การพิจารณาคัดเลือกข้อเสนอ"):
            criteria_content_para = paragraph
            if criteria_text:
                _set_paragraph_text(paragraph, criteria_text)
        elif text.startswith("/8.") or text.startswith("/๘"):
            junk_paras.append(paragraph)
        elif text.startswith("๘. ข้อเสนอ") or text.startswith("8. ข้อเสนอ"):
            other_heading_para = paragraph
            _set_paragraph_text(paragraph, "๑๐. ข้อเสนออื่นๆ")
        elif "จึงขอแต่งตั้ง  นาย/นาง" in text or "จึงขอแต่งตั้ง นาย/นาง" in text:
            name = (inspector.get("name") or "").strip()
            position = (inspector.get("position") or "").strip()
            if name:
                replaced = re.sub(
                    r"นาย/นาง/นางสาว\.\.\.\s+ตำแหน่ง\.\.\.",
                    f"{name}  ตำแหน่ง {position or '...'}",
                    text,
                    count=1,
                )
                _set_paragraph_text(paragraph, replaced)
        elif SPECIFIC_MEMBER_RE.match(text) and "ตำแหน่ง" in text:
            member_paras.append(paragraph)
        elif text.startswith("๑. อนุมัติให้ความเห็นชอบ") and amount_text:
            pay = f"โดยเบิกจ่ายจาก{budget_source}" if budget_source else "โดยเบิกจ่ายจากงบประมาณของหน่วยงาน"
            _set_paragraph_text(
                paragraph,
                "๑. อนุมัติให้ความเห็นชอบรายงานขออนุมัติจัดซื้อ/จัดจ้าง  "
                "และอนุมัติให้ดำเนินการตามรายละเอียดข้างต้น  "
                f"ภายในวงเงิน  {amount_text}  บาท  ({baht_text})  {pay}",
            )
        elif text.startswith("เรียน") and recipient:
            _set_paragraph_text(paragraph, f"เรียน  {recipient}")
        elif officer_name and text.startswith("(") and "นาย/นาง" in text:
            _set_paragraph_text(paragraph, _paren_name(officer_name))
        elif officer_position and text.startswith("ตำแหน่ง"):
            prefix = "" if officer_position.startswith("ตำแหน่ง") else "ตำแหน่ง "
            _set_paragraph_text(paragraph, f"{prefix}{officer_position}")

    if member_paras:
        template_member = member_paras[0]
        for extra in member_paras[1:]:
            _remove_paragraph(extra)
        if members:
            _set_paragraph_text(
                template_member, format_specific_member_line(1, members[0])
            )
            current = template_member
            for index, member in enumerate(members[1:], start=2):
                current = _clone_paragraph_after(current)
                _set_paragraph_text(current, format_specific_member_line(index, member))

    for junk in junk_paras:
        _remove_paragraph(junk)

    insert_after = criteria_content_para
    heading_style = criteria_heading_para or criteria_content_para
    content_style = criteria_content_para or criteria_heading_para
    if insert_after is not None:
        if announcement_text:
            insert_after = _insert_paragraph_after(
                insert_after,
                "๘. ร่างประกาศ และร่างเอกสารประกวดราคา",
                style_from=heading_style,
            )
            insert_after = _insert_paragraph_after(
                insert_after,
                announcement_text,
                style_from=content_style,
            )
        if legal_text:
            insert_after = _insert_paragraph_after(
                insert_after,
                "๙. ข้อระเบียบและกฎหมาย",
                style_from=heading_style,
            )
            insert_after = _insert_paragraph_after(
                insert_after,
                legal_text,
                style_from=content_style,
            )
    if other_heading_para is not None and other_proposal_text:
        _insert_paragraph_after(
            other_heading_para,
            other_proposal_text,
            style_from=content_style or other_heading_para,
        )

    sign_table = _find_sign_table(document)
    if sign_table:
        sign_cells = _unique_row_cells(sign_table.rows[0])
        if sign_cells:
            _fill_named_paragraphs(
                sign_cells[0].paragraphs, head_name, head_position
            )
        if len(sign_cells) > 1:
            right = sign_cells[1]
            for paragraph in right.paragraphs:
                text = _paragraph_text(paragraph).strip()
                if text.startswith("(") and signer_name:
                    _set_paragraph_text(paragraph, _paren_name(signer_name))
                elif text.startswith("สาธารณสุข") and signer_position:
                    _set_paragraph_text(paragraph, signer_position)
                elif text.startswith("ปฏิบัติ") and approver_acting:
                    _set_paragraph_text(paragraph, approver_acting)
                elif text.startswith("ปฏิบัติ") and not approver_acting and signer_position:
                    _set_paragraph_text(paragraph, signer_position)

    _remove_item_list_tables(document)
    _apply_thai_font(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
