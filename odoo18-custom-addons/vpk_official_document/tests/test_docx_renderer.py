# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

import io

from docx import Document
from odoo.tests import TransactionCase, tagged
from odoo.tools.misc import file_path

from odoo.addons.vpk_official_document.models.docx_renderer import (
    render_integrity_over_100k_docx,
    render_spec_price_committee_docx,
    render_specific_method_approval_docx,
    render_wa_committee_order_docx,
)
from odoo.addons.vpk_official_document.models.pdf_converter import (
    convert_docx_bytes_to_pdf,
    find_soffice,
)


def _docx_text(content):
    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


@tagged("post_install", "-at_install")
class TestWaCommitteeOrderDocx(TransactionCase):
    def test_renderer_fills_committee_names(self):
        path = file_path(
            "vpk_official_document/static/src/templates/wa_committee_order.docx"
        )
        with open(path, "rb") as handle:
            template = handle.read()
        content = render_wa_committee_order_docx(
            template,
            {
                "agency": "โรงพยาบาลวชิระภูเก็ต",
                "order_no": "๑๒๓",
                "year_be": "๒๕๖๙",
                "subject": "แต่งตั้งคณะกรรมการตรวจรับพัสดุ  สำหรับการซื้อ/จ้างวัสดุสำนักงาน  โดยวิธีเฉพาะเจาะจง",
                "body": "ด้วย  กลุ่มงานพัสดุ  มีความประสงค์จะดำเนินการจัดซื้อจัดจ้างวัสดุสำนักงาน  โดยวิธีเฉพาะเจาะจง  ประกอบด้วย",
                "members": [
                    {
                        "name": "นายทดสอบ ประธาน",
                        "position": "หัวหน้ากลุ่มงาน",
                        "role": "ประธานกรรมการ",
                    },
                    {
                        "name": "นางทดสอบ กรรมการ",
                        "position": "นักวิชาการพัสดุ",
                        "role": "กรรมการ",
                    },
                    {
                        "name": "นางสาวทดสอบ คนที่สาม",
                        "position": "เจ้าพนักงานพัสดุ",
                        "role": "กรรมการ",
                    },
                    {
                        "name": "นายทดสอบ คนที่สี่",
                        "position": "เจ้าหน้าที่",
                        "role": "กรรมการและเลขานุการ",
                    },
                ],
                "order_date": "๗ กันยายน พ.ศ. ๒๕๖๙",
                "signer_name": "นายวีระศักดิ์ หล่อทองคำ",
                "signer_position": "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต",
            },
        )
        text = _docx_text(content)
        self.assertIn("คำสั่งโรงพยาบาลวชิระภูเก็ต", text)
        self.assertIn("ที่  ๑๒๓ / ๒๕๖๙", text)
        self.assertIn("นายทดสอบ ประธาน", text)
        self.assertIn("นางทดสอบ กรรมการ", text)
        self.assertIn("นายทดสอบ คนที่สี่", text)
        self.assertIn("นายวีระศักดิ์ หล่อทองคำ", text)
        self.assertNotIn("สำนักงานสาธารณสุขอำเภอดำเนินสะดวก", text)
        self.assertNotIn("นาย/นาง/นางสาว...", text)

    def test_committee_order_prints_saraban_number(self):
        path = file_path(
            "vpk_official_document/static/src/templates/wa_committee_order.docx"
        )
        with open(path, "rb") as handle:
            template = handle.read()
        content = render_wa_committee_order_docx(
            template,
            {
                "agency": "โรงพยาบาลวชิระภูเก็ต",
                "order_no": "๑๒๓",
                "year_be": "๒๕๖๙",
                "order_no_line": "ที่  ภก ๐๐๓๓.๒๐๑/๒๕๖๙/๐๐๐๓",
                "subject": "แต่งตั้งคณะกรรมการตรวจรับพัสดุ",
                "body": "ด้วย  กลุ่มงานพัสดุ  มีความประสงค์จะดำเนินการจัดซื้อจัดจ้างวัสดุสำนักงาน",
                "members": [
                    {
                        "name": "นายทดสอบ ประธาน",
                        "position": "หัวหน้ากลุ่มงาน",
                        "role": "ประธานกรรมการ",
                    }
                ],
                "order_date": "๗ กันยายน พ.ศ. ๒๕๖๙",
                "signer_name": "นายวีระศักดิ์ หล่อทองคำ",
                "signer_position": "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต",
            },
        )
        text = _docx_text(content)
        self.assertIn("ที่  ภก ๐๐๓๓.๒๐๑/๒๕๖๙/๐๐๐๓", text)
        self.assertNotIn("ที่  ๑๒๓ / ๒๕๖๙", text)

    def test_filled_docx_converts_to_pdf_with_thai_text(self):
        if not find_soffice():
            self.skipTest("LibreOffice is required to convert the Word template to PDF")
        path = file_path(
            "vpk_official_document/static/src/templates/wa_committee_order.docx"
        )
        with open(path, "rb") as handle:
            template = handle.read()
        docx_content = render_wa_committee_order_docx(
            template,
            {
                "agency": "โรงพยาบาลวชิระภูเก็ต",
                "order_no": "๑๒๓",
                "year_be": "๒๕๖๙",
                "subject": "แต่งตั้งคณะกรรมการตรวจรับพัสดุ",
                "body": "ด้วย  กลุ่มงานพัสดุ  มีความประสงค์จะดำเนินการจัดซื้อจัดจ้างวัสดุสำนักงาน",
                "members": [
                    {
                        "name": "นายทดสอบ ประธาน",
                        "position": "หัวหน้ากลุ่มงาน",
                        "role": "ประธานกรรมการ",
                    }
                ],
                "order_date": "๗ กันยายน พ.ศ. ๒๕๖๙",
                "signer_name": "นายวีระศักดิ์ หล่อทองคำ",
                "signer_position": "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต",
            },
        )
        pdf_content = convert_docx_bytes_to_pdf(docx_content)
        self.assertTrue(pdf_content.startswith(b"%PDF"))
        self.assertGreater(len(pdf_content), 5000)

    def test_integrity_form_fills_names(self):
        path = file_path(
            "vpk_official_document/static/src/templates/integrity_over_100k.docx"
        )
        with open(path, "rb") as handle:
            template = handle.read()
        content = render_integrity_over_100k_docx(
            template,
            {
                "head_officer": {
                    "name": "นางสาวปรีดา สุดขาว",
                    "position": "หัวหน้าเจ้าหน้าที่",
                },
                "officer": {
                    "name": "นางทดสอบ เจ้าหน้าที่",
                    "position": "นักวิชาการพัสดุ",
                },
                "members": [
                    {
                        "name": "นายทดสอบ ประธาน",
                        "position": "หัวหน้ากลุ่มงาน",
                        "role_key": "chairman",
                    },
                    {
                        "name": "นางทดสอบ กรรมการ",
                        "position": "นักวิชาการพัสดุ",
                        "role_key": "committee",
                    },
                    {
                        "name": "นางสาวทดสอบ เลขา",
                        "position": "เจ้าพนักงานพัสดุ",
                        "role_key": "secretary",
                    },
                ],
            },
        )
        text = _docx_text(content)
        self.assertIn("นางสาวปรีดา สุดขาว", text)
        self.assertIn("นางทดสอบ เจ้าหน้าที่", text)
        self.assertIn("นายทดสอบ ประธาน", text)
        self.assertIn("นางทดสอบ กรรมการ", text)
        self.assertIn("นางสาวทดสอบ เลขา", text)
        self.assertIn("กรรมการและเลขานุการ", text)
        self.assertNotIn("นาย/นาง/นางสาว...", text)
        from docx import Document
        import io
        from docx.oxml.ns import qn

        filled = Document(io.BytesIO(content))
        first_col = filled.tables[0]._tbl.tblGrid.gridCol_lst[0].get(qn("w:w"))
        self.assertGreaterEqual(int(first_col), 1600)

    def test_spec_price_committee_form_fills_names(self):
        path = file_path(
            "vpk_official_document/static/src/templates/spec_price_committee.docx"
        )
        with open(path, "rb") as handle:
            template = handle.read()
        content = render_spec_price_committee_docx(
            template,
            {
                "agency": "กลุ่มงานพัสดุ โรงพยาบาลวชิระภูเก็ต",
                "memo_number": "กค/๒๕๖๙/๐๐๑๒",
                "order_date": "๗ กันยายน ๒๕๖๙",
                "subject": "ขออนุมัติแต่งตั้งคณะกรรมการกำหนดรายละเอียดคุณลักษณะเฉพาะและกำหนดราคากลางของวัสดุสำนักงาน",
                "recipient": "ผู้ว่าราชการจังหวัดภูเก็ต",
                "body": "ด้วย  กลุ่มงานพัสดุ  มีความประสงค์จะดำเนินการจัดซื้อ/จัดจ้างวัสดุสำนักงาน  จำนวน  ๑  รายการ  โดยวิธีเฉพาะเจาะจง  จำนวนเงิน  ๑๒๓,๔๕๖.๐๐  บาท  (หนึ่งแสนสองหมื่นสามพันสี่ร้อยห้าสิบหกบาทถ้วน)  เพื่อใช้ในราชการ  โดยเตรียมความพร้อมสำหรับการดำเนินงานจัดหาพัสดุตามแนวทางปฏิบัติและระเบียบที่เกี่ยวข้อง",
                "members": [
                    {
                        "name": "นายทดสอบ ประธาน",
                        "position": "หัวหน้ากลุ่มงาน",
                        "role": "ประธานกรรมการ",
                    },
                    {
                        "name": "นางทดสอบ กรรมการ",
                        "position": "นักวิชาการพัสดุ",
                        "role": "กรรมการ",
                    },
                    {
                        "name": "นางสาวทดสอบ คนที่สาม",
                        "position": "เจ้าพนักงานพัสดุ",
                        "role": "กรรมการ",
                    },
                    {
                        "name": "นายทดสอบ คนที่สี่",
                        "position": "เจ้าหน้าที่",
                        "role": "กรรมการและเลขานุการ",
                    },
                ],
                "signer_name": "นายวีระศักดิ์ หล่อทองคำ",
                "signer_position": "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต",
                "signer_org": "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต",
            },
        )
        text = _docx_text(content)
        self.assertIn("บันทึกข้อความ", text)
        self.assertIn("กลุ่มงานพัสดุ โรงพยาบาลวชิระภูเก็ต", text)
        self.assertIn("กค/๒๕๖๙/๐๐๑๒", text)
        self.assertIn("ผู้ว่าราชการจังหวัดภูเก็ต", text)
        self.assertIn("นายทดสอบ ประธาน", text)
        self.assertIn("นางทดสอบ กรรมการ", text)
        self.assertIn("นายทดสอบ คนที่สี่", text)
        self.assertIn("นายวีระศักดิ์ หล่อทองคำ", text)
        self.assertNotIn("สาธารณสุขอำเภอดำเนินสะดวก", text)
        self.assertNotIn("โรงพยาบาลส่งเสริมสุขภาพตำบล", text)
        self.assertNotIn("นาย/นาง/นางสาว...", text)
        self.assertNotIn("...............", text)

    def test_specific_method_approval_form_fills_values(self):
        path = file_path(
            "vpk_official_document/static/src/templates/specific_method_approval.docx"
        )
        with open(path, "rb") as handle:
            template = handle.read()
        content = render_specific_method_approval_docx(
            template,
            {
                "agency": "กลุ่มงานพัสดุ โรงพยาบาลวชิระภูเก็ต",
                "memo_number": "กค/๒๕๖๙/๐๐๒๐",
                "order_date": "๗ กันยายน ๒๕๖๙",
                "subject": "รายงานขออนุมัติซื้อ/จ้างวัสดุสำนักงาน  โดยวิธีเฉพาะเจาะจง",
                "recipient": "ผู้ว่าราชการจังหวัดภูเก็ต",
                "body": "ด้วยกลุ่มงานพัสดุ  มีความประสงค์ในการจัดซื้อจัดจ้างวัสดุสำนักงาน  เป็นจำนวนเงินทั้งสิ้น  ๑๒๓,๔๕๖.๐๐  บาท  (หนึ่งแสนสองหมื่นสามพันสี่ร้อยห้าสิบหกบาทถ้วน)  เพื่อใช้ในราชการ  โดยใช้งบประมาณจากงบประมาณของหน่วยงาน  ซึ่งขอจัดซื้อจัดจ้างตามระเบียบกระทรวงการคลังว่าด้วยการจัดซื้อจัดจ้างและการบริหารพัสดุภาครัฐ  พ.ศ. ๒๕๖๐  ข้อ  ๒๒  โดยมีรายละเอียด  ดังต่อไปนี้",
                "reason": "เพื่อใช้ในราชการ",
                "work_detail": (
                    "๑. กระดาษถ่ายเอกสาร A4  จำนวน  ๑๐  รีม  เป็นเงิน  ๑๒๓,๔๕๖.๐๐  บาท"
                ),
                "price_mid_text": "จำนวน  ๑๒๓,๔๕๖.๐๐  บาท  (หนึ่งแสนสองหมื่นสามพันสี่ร้อยห้าสิบหกบาทถ้วน)  จากราคากลางตามท้องตลาด",
                "budget_text": "ภายในวงเงินงบประมาณ  ๑๒๓,๔๕๖.๐๐  บาท  (หนึ่งแสนสองหมื่นสามพันสี่ร้อยห้าสิบหกบาทถ้วน)  โดยเบิกจ่ายจากงบประมาณของหน่วยงาน",
                "delivery_text": "กำหนดเวลาส่งมอบงานหรือให้งานแล้วเสร็จภายใน  ๓๐  วัน  นับถัดจากวันที่ลงนามในสัญญา",
                "method_text": "โดยวิธีเฉพาะเจาะจง  ตามพระราชบัญญัติการจัดซื้อจัดจ้างและการบริหารพัสดุภาครัฐ  พ.ศ. ๒๕๖๐  มาตรา  ๕๖  (๒)  (ข)",
                "criteria_text": "การพิจารณาคัดเลือกข้อเสนอโดยใช้เกณฑ์ราคา",
                "announcement_text": "ไม่ต้องจัดทำร่างประกาศและร่างเอกสารประกวดราคา  เนื่องจากเป็นการจัดซื้อจัดจ้างโดยวิธีเฉพาะเจาะจง",
                "legal_text": "ตามพระราชบัญญัติการจัดซื้อจัดจ้างและการบริหารพัสดุภาครัฐ  พ.ศ. ๒๕๖๐",
                "other_proposal_text": "แต่งตั้งคณะกรรมการตรวจรับพัสดุ  ตามรายชื่อที่เสนอในหนังสือฉบับนี้",
                "amount_text": "๑๒๓,๔๕๖.๐๐",
                "baht_text": "หนึ่งแสนสองหมื่นสามพันสี่ร้อยห้าสิบหกบาทถ้วน",
                "price_mid_source": "ราคากลางตามท้องตลาด",
                "budget_source": "งบประมาณของหน่วยงาน",
                "delivery_days": "๓๐",
                "members": [
                    {
                        "name": "นายทดสอบ ประธาน",
                        "position": "หัวหน้ากลุ่มงาน",
                        "role": "ประธานกรรมการ",
                    },
                    {
                        "name": "นางทดสอบ กรรมการ",
                        "position": "นักวิชาการพัสดุ",
                        "role": "กรรมการ",
                    },
                ],
                "officer_name": "นางทดสอบ เจ้าหน้าที่",
                "officer_position": "นักวิชาการพัสดุ",
                "head_officer_name": "นางสาวปรีดา สุดขาว",
                "head_officer_position": "หัวหน้าเจ้าหน้าที่",
                "signer_name": "นายวีระศักดิ์ หล่อทองคำ",
                "signer_position": "ผู้อำนวยการโรงพยาบาลวชิระภูเก็ต",
                "approver_acting": "ปฏิบัติราชการแทนผู้ว่าราชการจังหวัดภูเก็ต",
            },
        )
        text = _docx_text(content)
        self.assertIn("บันทึกข้อความ", text)
        self.assertIn("กลุ่มงานพัสดุ โรงพยาบาลวชิระภูเก็ต", text)
        self.assertIn("ผู้ว่าราชการจังหวัดภูเก็ต", text)
        self.assertIn("นายทดสอบ ประธาน", text)
        self.assertIn("นางทดสอบ กรรมการ", text)
        self.assertIn("นางทดสอบ เจ้าหน้าที่", text)
        self.assertIn("นางสาวปรีดา สุดขาว", text)
        self.assertIn("นายวีระศักดิ์ หล่อทองคำ", text)
        self.assertIn("๑๒๓,๔๕๖.๐๐", text)
        self.assertIn("กระดาษถ่ายเอกสาร A4", text)
        self.assertIn("จำนวน  ๑๐  รีม", text)
        self.assertNotIn("ราคาต่อหน่วย (บาท)", text)
        self.assertNotIn("จำนวนเงิน (บาท)", text)
        document = Document(io.BytesIO(content))
        item_headers = [
            " ".join(cell.text for cell in table.rows[0].cells)
            for table in document.tables
            if table.rows
        ]
        self.assertFalse(
            any("ราคาต่อหน่วย" in header for header in item_headers),
            item_headers,
        )
        self.assertIn("มาตรา  ๕๖", text)
        self.assertIn("นับถัดจากวันที่ลงนามในสัญญา", text)
        self.assertIn("๑. เหตุผลและความจำเป็น", text)
        self.assertIn("๒. รายละเอียดงานที่จัดซื้อ จัดจ้าง", text)
        self.assertIn("๗. หลักเกณฑ์การพิจารณาคัดเลือกข้อเสนอ", text)
        self.assertIn("การพิจารณาคัดเลือกข้อเสนอโดยใช้เกณฑ์ราคา", text)
        self.assertIn("๘. ร่างประกาศ และร่างเอกสารประกวดราคา", text)
        self.assertIn("๙. ข้อระเบียบและกฎหมาย", text)
        self.assertIn("๑๐. ข้อเสนออื่นๆ", text)
        self.assertIn("แต่งตั้งคณะกรรมการตรวจรับพัสดุ", text)
        self.assertNotIn("สาธารณสุขอำเภอดำเนินสะดวก", text)
        self.assertNotIn("โรงพยาบาลส่งเสริมสุขภาพตำบล", text)
        self.assertNotIn("นาย/นาง/นางสาว...", text)
