# -*- coding: utf-8 -*-
import base64
import io

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError


class GfmisDocument(models.Model):
    _name = "gfmis.document"
    _description = "เอกสารกระบวนการ GFMIS"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "document_date desc, id desc"
    _check_company_auto = True

    name = fields.Char(
        string="เลขที่ ERP",
        required=True,
        copy=False,
        default="/",
        tracking=True,
    )
    company_id = fields.Many2one(
        "res.company",
        required=True,
        default=lambda self: self.env.company,
        index=True,
    )
    company_partner_id = fields.Many2one(
        related="company_id.partner_id",
    )
    currency_id = fields.Many2one(
        related="company_id.currency_id",
        store=True,
    )
    doc_class = fields.Selection(
        [
            ("request", "ขอเบิกเงิน"),
            ("agency_pay", "ขอจ่ายโดยส่วนราชการ"),
            ("return", "เบิกเกินส่งคืน"),
        ],
        string="ประเภทเอกสาร",
        required=True,
        default="request",
        tracking=True,
        index=True,
    )
    form_type_id = fields.Many2one(
        "gfmis.form.type",
        string="แบบฟอร์ม GFMIS",
        required=True,
        ondelete="restrict",
        tracking=True,
        domain="[('doc_class', '=', doc_class), ('active', '=', True)]",
    )
    form_code = fields.Char(related="form_type_id.code", store=True)
    requires_po = fields.Boolean(related="form_type_id.requires_po")

    state = fields.Selection(
        [
            ("draft", "ร่าง"),
            ("recorded", "บันทึกใน GFMIS"),
            ("wait_approve_request", "รออนุมัติขอเบิก (อม.01)"),
            ("wait_approve_pay", "รออนุมัติสั่งจ่าย (อม.02)"),
            ("sent_treasury", "ส่งกรมบัญชีกลาง/คลังจังหวัด"),
            ("paid", "จ่ายตรงแล้ว"),
            ("wait_agency_pay", "รอจ่ายผู้มีสิทธิ (จ่ายผ่าน)"),
            ("agency_paid", "จ่ายให้ผู้มีสิทธิแล้ว"),
            ("returned", "เบิกเกินส่งคืน"),
            ("cancelled", "ยกเลิก"),
        ],
        string="สถานะ",
        default="draft",
        required=True,
        tracking=True,
        index=True,
        copy=False,
    )

    # Header — คู่มือบทที่ 3 ข้อมูลส่วนหัว
    agency_code = fields.Char(string="รหัสหน่วยงาน", size=5)
    area_code = fields.Char(string="รหัสพื้นที่", size=4)
    disbursing_unit_code = fields.Char(string="รหัสหน่วยเบิกจ่าย", size=10)
    reference = fields.Char(
        string="การอ้างอิง",
        help="เลขที่ใบแจ้งหนี้ หรือเอกสารหลักฐานการขอเบิก",
        tracking=True,
    )
    document_date = fields.Date(
        string="วันที่เอกสาร",
        required=True,
        default=fields.Date.context_today,
        tracking=True,
    )
    posting_date = fields.Date(
        string="วันที่ผ่านรายการ",
        required=True,
        default=fields.Date.context_today,
        tracking=True,
    )
    period = fields.Char(string="งวด", compute="_compute_period", store=True)
    gfmis_po_number = fields.Char(
        string="เลขที่ใบสั่งซื้อสั่งจ้างระบบ GFMIS",
        help="เลขที่ PO ใน New GFMIS Thai เช่น 4001004668",
        tracking=True,
    )
    gfmis_doc_number = fields.Char(
        string="เลขที่เอกสาร GFMIS",
        copy=False,
        tracking=True,
        help="เลขที่เอกสารที่ระบบ GFMIS ออกให้หลังบันทึก เช่น 31XXXXXXXX",
        index=True,
    )

    # ข้อมูลทั่วไป
    fund_type = fields.Selection(
        [
            ("budget", "เงินงบประมาณ"),
            ("reserved", "เงินกันไว้เบิกเหลื่อมปี"),
            ("extra", "เงินนอกงบประมาณ"),
            ("loan", "เงินกู้"),
            ("allocated", "เงินรายได้จัดสรร"),
        ],
        string="ประเภทรายการขอเบิก / แหล่งเงิน",
        default="budget",
        required=True,
        tracking=True,
    )
    payment_method = fields.Selection(
        [
            ("direct", "จ่ายตรงผู้ขาย"),
            ("through", "จ่ายผ่านส่วนราชการ"),
        ],
        string="วิธีการชำระเงิน",
        default="direct",
        required=True,
        tracking=True,
    )
    partner_id = fields.Many2one(
        "res.partner",
        string="เจ้าหนี้ / ผู้มีสิทธิรับเงิน",
        required=True,
        tracking=True,
        check_company=True,
    )
    vat = fields.Char(related="partner_id.vat", string="เลขประจำตัวผู้เสียภาษี")
    partner_bank_id = fields.Many2one(
        "res.partner.bank",
        string="บัญชีเงินฝากธนาคารผู้รับเงิน",
        domain="[('partner_id', '=', partner_id)]",
        check_company=True,
    )
    company_bank_id = fields.Many2one(
        "res.partner.bank",
        string="บัญชีธนาคารโรงพยาบาล",
        help="ใช้เมื่อจ่ายผ่านส่วนราชการ เงินเข้าบัญชีของหน่วยงานก่อนจ่ายต่อ",
        domain="[('partner_id', '=', company_partner_id)]",
        check_company=True,
    )
    description = fields.Char(string="คำอธิบายเอกสาร", size=99)
    dika_number = fields.Char(string="ฎีกาที่", tracking=True, index=True)
    transfer_date = fields.Date(string="วันที่โอน", tracking=True, index=True)

    # ERP links
    purchase_id = fields.Many2one(
        "purchase.order",
        string="ใบสั่งซื้อ ERP",
        check_company=True,
        tracking=True,
    )
    invoice_id = fields.Many2one(
        "account.move",
        string="ใบแจ้งหนี้ผู้ขาย ERP",
        domain="[('move_type', 'in', ('in_invoice', 'in_refund')), ('partner_id', '=', partner_id)]",
        check_company=True,
        tracking=True,
        index=True,
    )
    payment_id = fields.Many2one(
        "account.payment",
        string="การจ่ายชำระ ERP",
        check_company=True,
        tracking=True,
    )
    source_document_id = fields.Many2one(
        "gfmis.document",
        string="เอกสาร GFMIS อ้างอิง",
        help="ใช้กับใบลดหนี้ ขจ.05 หรือเบิกเกินส่งคืน ที่อ้างเอกสารขอเบิกเดิม",
        index=True,
    )

    line_ids = fields.One2many(
        "gfmis.document.line",
        "document_id",
        string="รายการขอเบิก",
        copy=True,
    )
    wht_ids = fields.One2many(
        "gfmis.document.wht",
        "document_id",
        string="ภาษีหัก ณ ที่จ่าย / ค่าปรับ",
        copy=True,
    )

    amount_untaxed = fields.Monetary(
        string="ยอดรายการ",
        compute="_compute_amounts",
        store=True,
        currency_field="currency_id",
    )
    amount_wht = fields.Monetary(
        string="ภาษีหัก ณ ที่จ่าย",
        compute="_compute_amounts",
        store=True,
        currency_field="currency_id",
    )
    amount_penalty = fields.Monetary(
        string="ค่าปรับ",
        compute="_compute_amounts",
        store=True,
        currency_field="currency_id",
    )
    amount_net = fields.Monetary(
        string="ยอดสุทธิจ่าย",
        compute="_compute_amounts",
        store=True,
        currency_field="currency_id",
    )

    note = fields.Text(string="หมายเหตุ")
    gfmis_message = fields.Text(
        string="ผลการตรวจสอบ GFMIS",
        help="รหัสข้อความ / คำอธิบายจากการจำลองบันทึกใน GFMIS",
    )

    @api.depends("document_date")
    def _compute_period(self):
        for rec in self:
            rec.period = rec.document_date.strftime("%m/%Y") if rec.document_date else False

    @api.depends("line_ids.amount", "wht_ids.wht_amount", "wht_ids.penalty_amount")
    def _compute_amounts(self):
        for rec in self:
            rec.amount_untaxed = sum(rec.line_ids.mapped("amount"))
            rec.amount_wht = sum(rec.wht_ids.mapped("wht_amount"))
            rec.amount_penalty = sum(rec.wht_ids.mapped("penalty_amount"))
            rec.amount_net = rec.amount_untaxed - rec.amount_wht - rec.amount_penalty

    @api.onchange("company_id")
    def _onchange_company_id(self):
        company = self.company_id
        if company:
            self.agency_code = company.gfmis_agency_code
            self.area_code = company.gfmis_area_code
            self.disbursing_unit_code = company.gfmis_disbursing_unit_code

    @api.onchange("form_type_id")
    def _onchange_form_type_id(self):
        form = self.form_type_id
        if not form:
            return
        self.doc_class = form.doc_class
        if form.fund_type and form.fund_type != "any":
            mapping = {"budget": "budget", "extra": "extra", "loan": "loan", "allocated": "allocated"}
            self.fund_type = mapping.get(form.fund_type, self.fund_type)
        if form.payment_method == "direct":
            self.payment_method = "direct"
        elif form.payment_method == "through":
            self.payment_method = "through"

    @api.onchange("partner_id")
    def _onchange_partner_id(self):
        if self.partner_id:
            banks = self.partner_id.bank_ids.filtered("active")
            self.partner_bank_id = banks[:1]
        else:
            self.partner_bank_id = False

    @api.onchange("invoice_id")
    def _onchange_invoice_id(self):
        invoice = self.invoice_id
        if not invoice:
            return
        self.partner_id = invoice.partner_id
        self.reference = invoice.ref or invoice.name
        self.purchase_id = invoice.invoice_line_ids.purchase_line_id.order_id[:1]
        if invoice.invoice_date:
            self.document_date = invoice.invoice_date
        if not self.line_ids:
            self.line_ids = self._prepare_lines_from_invoice(invoice)

    @api.model
    def _receipt_doc_number_from_invoice(self, invoice):
        """เลขที่ใบตรวจรับที่ผูกกับใบแจ้งหนี้ผู้ขาย."""
        wa = invoice.wa_id if "wa_id" in invoice._fields else invoice.env["work.acceptance"]
        if not wa and "wa_ids" in invoice.invoice_line_ids.purchase_line_id.order_id._fields:
            wa = invoice.invoice_line_ids.purchase_line_id.order_id.wa_ids.filtered(
                lambda rec: rec.state == "accept"
            )[:1]
        name = (wa.name or "").strip() if wa else ""
        if not name or name in ("New", "/"):
            return False
        return name

    @api.model
    def _prepare_lines_from_invoice(self, invoice):
        receipt_doc_number = self._receipt_doc_number_from_invoice(invoice)
        lines = []
        for line in invoice.invoice_line_ids.filtered(lambda item: item.display_type == "product"):
            lines.append(
                (
                    0,
                    0,
                    {
                        "name": line.name,
                        "quantity": line.quantity,
                        "price_unit": line.price_unit,
                        "amount": line.price_subtotal,
                        "receipt_doc_number": receipt_doc_number,
                        "bill_number": invoice.ref or invoice.name or False,
                        "invoice_line_id": line.id,
                        "purchase_line_id": line.purchase_line_id.id,
                    },
                )
            )
        return lines

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            company = self.env["res.company"].browse(vals.get("company_id") or self.env.company.id)
            vals.setdefault("agency_code", company.gfmis_agency_code)
            vals.setdefault("area_code", company.gfmis_area_code)
            vals.setdefault("disbursing_unit_code", company.gfmis_disbursing_unit_code)
            if vals.get("name", "/") == "/":
                vals["name"] = self.env["ir.sequence"].next_by_code("gfmis.document") or "/"
        return super().create(vals_list)

    @api.constrains("form_type_id", "purchase_id", "gfmis_po_number")
    def _check_requires_po(self):
        for rec in self:
            if rec.form_type_id.requires_po and rec.state != "cancelled":
                if not rec.purchase_id and not rec.gfmis_po_number:
                    raise ValidationError(
                        _("แบบฟอร์ม %s ต้องอ้างใบสั่งซื้อสั่งจ้าง (PO ใน ERP หรือเลขที่ PO ของ GFMIS)")
                        % rec.form_type_id.code
                    )

    def _ensure_draftish(self, allowed):
        for rec in self:
            if rec.state not in allowed:
                raise UserError(_("ไม่สามารถทำรายการนี้ในสถานะ %s") % rec.state)

    def action_record_gfmis(self):
        self._ensure_draftish(("draft",))
        for rec in self:
            if not rec.line_ids and rec.doc_class == "request":
                raise UserError(_("กรุณาบันทึกรายการขอเบิกอย่างน้อย 1 บรรทัด"))
            rec.state = "recorded"
            rec.message_post(
                body=_("บันทึกเอกสารในกระบวนการ GFMIS แล้ว กรอกเลขที่เอกสาร GFMIS เมื่อระบบออกเลขให้")
            )

    def action_submit_approve(self):
        self._ensure_draftish(("recorded",))
        self.write({"state": "wait_approve_request"})
        self.message_post(body=_("ส่งรออนุมัติขอเบิก (อม.01)"))

    def action_approve_request(self):
        """อม.01 อนุมัติเอกสารขอเบิก"""
        self._ensure_draftish(("wait_approve_request",))
        self.write({"state": "wait_approve_pay"})
        self.message_post(body=_("อนุมัติขอเบิก (อม.01) แล้ว"))

    def action_approve_pay(self):
        """อม.02 อนุมัติเอกสารเพื่อสั่งจ่าย"""
        self._ensure_draftish(("wait_approve_pay",))
        self.write({"state": "sent_treasury"})
        self.message_post(body=_("อนุมัติสั่งจ่าย (อม.02) แล้ว ส่งกรมบัญชีกลาง/คลังจังหวัด"))

    def action_mark_paid(self):
        self._ensure_draftish(("sent_treasury", "wait_agency_pay"))
        for rec in self:
            if rec.payment_method == "through" and rec.state == "sent_treasury":
                rec.state = "wait_agency_pay"
                rec.message_post(body=_("เงินเข้าบัญชีส่วนราชการแล้ว รอจ่ายให้ผู้มีสิทธิ (ขจ.05)"))
            elif rec.payment_method == "through":
                rec.state = "agency_paid"
                rec.message_post(body=_("จ่ายให้ผู้มีสิทธิแล้ว"))
            else:
                rec.state = "paid"
                rec.message_post(body=_("กรมบัญชีกลางประมวลผลจ่ายตรงผู้ขายแล้ว"))

    def action_mark_returned(self):
        self._ensure_draftish(("paid", "agency_paid", "wait_agency_pay", "sent_treasury"))
        self.write({"state": "returned"})
        self.message_post(body=_("บันทึกเบิกเกินส่งคืน"))

    def action_cancel(self):
        self._ensure_draftish(("draft", "recorded", "wait_approve_request", "wait_approve_pay"))
        self.write({"state": "cancelled"})

    def action_draft(self):
        self._ensure_draftish(("cancelled", "recorded"))
        self.write({"state": "draft"})

    def action_open_invoice(self):
        self.ensure_one()
        if not self.invoice_id:
            return False
        return {
            "type": "ir.actions.act_window",
            "res_model": "account.move",
            "res_id": self.invoice_id.id,
            "view_mode": "form",
            "target": "current",
        }

    @api.model
    def action_create_from_invoice(self, invoice):
        invoice.ensure_one()
        if invoice.move_type not in ("in_invoice", "in_refund"):
            raise UserError(_("สร้างเอกสาร GFMIS ได้เฉพาะใบแจ้งหนี้ / ใบลดหนี้ผู้ขาย"))
        existing = self.search([("invoice_id", "=", invoice.id)], limit=1)
        if existing:
            return existing
        form_code = "ขบ.01" if invoice.invoice_line_ids.purchase_line_id else "ขบ.02"
        if invoice.move_type == "in_refund":
            form_code = "ขบ.23" if invoice.invoice_line_ids.purchase_line_id else "ขบ.24"
        form = self.env["gfmis.form.type"].search([("code", "=", form_code)], limit=1)
        if not form:
            form = self.env["gfmis.form.type"].search([("doc_class", "=", "request")], limit=1)
        lines = self._prepare_lines_from_invoice(invoice)
        return self.create(
            {
                "doc_class": "request" if invoice.move_type == "in_invoice" else "request",
                "form_type_id": form.id,
                "partner_id": invoice.partner_id.id,
                "invoice_id": invoice.id,
                "purchase_id": invoice.invoice_line_ids.purchase_line_id.order_id[:1].id,
                "reference": invoice.ref or invoice.name,
                "document_date": invoice.invoice_date or fields.Date.context_today(self),
                "posting_date": invoice.date or fields.Date.context_today(self),
                "description": (invoice.payment_reference or invoice.narration or invoice.name or "")[:99],
                "line_ids": lines,
                "payment_method": "direct",
            }
        )


    def action_export_gf_excel(self):
        if not self.line_ids:
            raise UserError(_("ยังไม่มีรายการสำหรับออกรายงาน GF"))
        return self.line_ids.action_export_gf_excel()


class GfmisDocumentLine(models.Model):
    _name = "gfmis.document.line"
    _description = "บรรทัดรายการขอเบิก GFMIS"
    _order = "transfer_date, dika_number, id"

    document_id = fields.Many2one(
        "gfmis.document",
        required=True,
        ondelete="cascade",
        index=True,
    )
    company_id = fields.Many2one(related="document_id.company_id", store=True)
    currency_id = fields.Many2one(related="document_id.currency_id")
    name = fields.Char(string="รายการ", required=True)
    delivery_date = fields.Date(string="วันที่ส่งมอบ")
    receipt_doc_number = fields.Char(string="เลขที่เอกสารตรวจรับ")
    gpsc_code = fields.Char(string="รหัส GPSC")
    bill_number = fields.Char(string="เลขที่บิล")
    tax_amount = fields.Monetary(string="ภาษี", currency_field="currency_id")
    note = fields.Char(string="หมายเหตุ")
    transfer_date = fields.Date(
        related="document_id.transfer_date",
        store=True,
        index=True,
    )
    partner_id = fields.Many2one(
        related="document_id.partner_id",
        store=True,
        index=True,
    )
    form_code = fields.Char(related="document_id.form_code", store=True)
    dika_number = fields.Char(related="document_id.dika_number", store=True, index=True)
    quantity = fields.Float(string="จำนวน", default=1.0)
    price_unit = fields.Monetary(string="ราคาต่อหน่วย", currency_field="currency_id")
    amount = fields.Monetary(string="จำนวนเงิน", currency_field="currency_id")
    fund_source_code = fields.Char(string="รหัสแหล่งของเงิน")
    activity_code = fields.Char(string="รหัสกิจกรรมหลัก")
    budget_code = fields.Char(string="รหัสงบประมาณ")
    purchase_line_id = fields.Many2one("purchase.order.line", string="บรรทัด PO")
    invoice_line_id = fields.Many2one("account.move.line", string="บรรทัดใบแจ้งหนี้")

    group_amount = fields.Monetary(
        string="เป็นเงิน",
        compute="_compute_group_amount",
        currency_field="currency_id",
    )
    penalty_amount = fields.Monetary(
        related="document_id.amount_penalty",
        string="ค่าปรับ",
    )

    @api.depends("document_id.line_ids.amount")
    def _compute_group_amount(self):
        for line in self:
            line.group_amount = sum(line.document_id.line_ids.mapped("amount"))

    @api.onchange("quantity", "price_unit")
    def _onchange_amount(self):
        self.amount = (self.quantity or 0.0) * (self.price_unit or 0.0)

    def action_export_gf_excel(self):
        if self:
            lines = self.search([("document_id", "in", self.document_id.ids)])
        else:
            lines = self.search([])
        if not lines:
            raise UserError(_("ยังไม่มีรายการสำหรับออกรายงาน GF"))
        content = lines._gf_register_xlsx()
        attachment = self.env["ir.attachment"].create({
            "name": "รายงาน-GF.xlsx",
            "datas": base64.b64encode(content),
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "res_model": "gfmis.document.line",
        })
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content/%s?download=true" % attachment.id,
            "target": "self",
        }

    def _gf_register_xlsx(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        months = [
            "", "ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
            "ก.ค.", "ส.ค.", "กย.", "ต.ค.", "พ.ย.", "ธ.ค.",
        ]

        def thai_short(value):
            if not value:
                return ""
            return "%s %s%s" % (value.day, months[value.month], str(value.year + 543)[-2:])

        lines = self.sorted(
            key=lambda line: (
                line.transfer_date or fields.Date.from_string("1900-01-01"),
                line.dika_number or "",
                line.document_id.id,
                line.id,
            )
        )
        groups = []
        bucket = []
        current = None
        for line in lines:
            if current != line.document_id:
                if bucket:
                    groups.append(bucket)
                bucket = [line]
                current = line.document_id
            else:
                bucket.append(line)
        if bucket:
            groups.append(bucket)

        headers = [
            "ลำดับที่", "วันที่ โอน", "ชื่อบริษัท", "ขบ.", "ฎีกาที่",
            "เลขที่บิล", "จำนวนเงิน", "ภาษี", "เป็นเงิน", "ค่าปรับ", "หมายเหตุ",
        ]
        book = Workbook()
        sheet = book.active
        sheet.title = "รายงาน GF"
        thin = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )
        header_fill = PatternFill("solid", fgColor="D6E3F0")
        header_font = Font(name="Arial", bold=True, size=11, color="1F4E79")
        title_font = Font(name="Arial", bold=True, size=16, color="1F4E79")
        body_font = Font(name="Arial", size=11)
        money = '#,##0.00'
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center")
        right = Alignment(horizontal="right", vertical="center")

        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title = sheet.cell(1, 1, "รายงาน GF")
        title.font = title_font
        title.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 28
        for col, label in enumerate(headers, start=1):
            cell = sheet.cell(2, col, label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin
        sheet.row_dimensions[2].height = 32
        sheet.auto_filter.ref = "A2:K2"
        sheet.freeze_panes = "A3"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.paperSize = 9
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.horizontalCentered = True
        sheet.oddHeader.center.text = "รายงาน GF"
        sheet.print_title_rows = "1:2"
        sheet.page_setup.horizontalDpi = 300
        sheet.page_setup.verticalDpi = 300
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 110
        sheet.sheet_view.view = "normal"
        sheet.oddFooter.left.text = "กระบวนการ GFMIS"
        sheet.oddFooter.right.text = "หน้า &P / &N"

        row = 3
        for index, group in enumerate(groups, start=1):
            start = row
            document = group[0].document_id
            line_tax = sum(line.tax_amount or 0.0 for line in group)
            tax_value = line_tax or document.amount_wht or 0.0
            penalty_value = document.amount_penalty or 0.0
            for offset, line in enumerate(group):
                first = offset == 0
                last = offset == len(group) - 1
                remark = line.note or ""
                values = [
                    index if first else "..",
                    thai_short(line.transfer_date) if first else "..",
                    (line.partner_id.name or "") if first else "..",
                    (line.form_code or "") if first else "..",
                    (line.dika_number or "") if first else "..",
                    line.bill_number or "",
                    line.amount or 0.0,
                    tax_value if last and tax_value else None,
                    None,
                    penalty_value if last and penalty_value else None,
                    remark,
                ]
                for col, value in enumerate(values, start=1):
                    cell = sheet.cell(row, col, value)
                    cell.font = body_font
                    cell.border = thin
                    cell.alignment = center if col in (1, 2, 4, 5) else left
                    if col in (7, 8, 9, 10) and isinstance(value, (int, float)):
                        cell.number_format = money
                        cell.alignment = right
                end = row
                row += 1
            total_cell = sheet.cell(end, 9, "=SUM(G%s:G%s)" % (start, end))
            total_cell.font = Font(name="Arial", size=11, bold=True)
            total_cell.number_format = money
            total_cell.alignment = right
            total_cell.border = thin

        if groups:
            total_row = row
            label = sheet.cell(total_row, 6, "รวม")
            label.font = Font(name="Arial", size=11, bold=True)
            label.alignment = right
            for col, letter in ((7, "G"), (8, "H"), (9, "I"), (10, "J")):
                cell = sheet.cell(total_row, col, "=SUM(%s3:%s%s)" % (letter, letter, total_row - 1))
                cell.font = Font(name="Arial", size=11, bold=True)
                cell.number_format = money
                cell.alignment = right
                cell.border = Border(top=Side(style="thin", color="1F4E79"), bottom=Side(style="double", color="1F4E79"))
            label.border = Border(top=Side(style="thin", color="1F4E79"), bottom=Side(style="double", color="1F4E79"))
            sheet.auto_filter.ref = "A2:K%s" % (total_row - 1)

        widths = [12, 14, 28, 12, 16, 28, 16, 14, 16, 14, 24]
        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.row_dimensions[2].height = 32
        sheet.page_margins.left = 0.5
        sheet.page_margins.right = 0.5
        sheet.page_margins.top = 0.6
        sheet.page_margins.bottom = 0.6
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()


class GfmisDocumentWht(models.Model):
    _name = "gfmis.document.wht"
    _description = "ภาษีหัก ณ ที่จ่าย / ค่าปรับ GFMIS"

    document_id = fields.Many2one(
        "gfmis.document",
        required=True,
        ondelete="cascade",
        index=True,
    )
    currency_id = fields.Many2one(related="document_id.currency_id")
    person_type = fields.Selection(
        [
            ("individual", "บุคคลธรรมดา"),
            ("company", "นิติบุคคล"),
        ],
        string="ภาษีเงินได้",
        default="company",
        required=True,
    )
    wht_base = fields.Monetary(string="ฐานคำนวณภาษี", currency_field="currency_id")
    wht_amount = fields.Monetary(string="จำนวนเงินที่หักไว้ (ภาษี)", currency_field="currency_id")
    penalty_type = fields.Selection(
        [
            ("land", "รายได้ของแผ่นดิน"),
            ("agency", "รายได้ของหน่วยงาน"),
        ],
        string="ค่าปรับ",
    )
    penalty_base = fields.Monetary(string="ฐานคำนวณค่าปรับ", currency_field="currency_id")
    penalty_amount = fields.Monetary(string="จำนวนเงินค่าปรับ", currency_field="currency_id")
    note = fields.Char(string="หมายเหตุ")
