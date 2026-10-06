import base64
import io

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlsxwriter
except ImportError:
    xlsxwriter = None


HEADER_FIELDS = {
    "รหัสสินค้า": "code",
    "default_code": "code",
    "code": "code",
    "product_code": "code",
    "รายการ": "name",
    "ชื่อสินค้า": "name",
    "name": "name",
    "หน่วยนับ": "uom",
    "uom": "uom",
    "ปริมาณที่ขอ": "quantity",
    "ปริมาณ": "quantity",
    "quantity": "quantity",
    "qty": "quantity",
    "ราคาต่อหน่วย": "price",
    "ราคา/หน่วยนับ": "price",
    "unit_price": "price",
    "price": "price",
    "หมายเหตุ": "note",
    "note": "note",
}


class VpkBudgetMaterialDetailImportWizard(models.TransientModel):
    _name = "vpk.budget.material.detail.import.wizard"
    _description = "Import Material Budget Lines from Excel"

    request_id = fields.Many2one(
        comodel_name="departmental.budget.request",
        required=True,
        readonly=True,
    )
    data_file = fields.Binary(string="ไฟล์ Excel")
    filename = fields.Char(string="ชื่อไฟล์")
    state = fields.Selection(
        [("draft", "Draft"), ("done", "Done")],
        default="draft",
    )
    log_text = fields.Text(string="ผลการนำเข้า", readonly=True)
    created_count = fields.Integer(readonly=True)
    updated_count = fields.Integer(readonly=True)
    skipped_count = fields.Integer(readonly=True)
    error_count = fields.Integer(readonly=True)

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        request_id = self.env.context.get("default_request_id") or self.env.context.get(
            "active_id"
        )
        if request_id and self.env.context.get("active_model") in (
            None,
            "departmental.budget.request",
        ):
            res["request_id"] = request_id
        return res

    def action_download_template(self):
        self.ensure_one()
        if not xlsxwriter:
            raise UserError(
                _("ไม่พบไลบรารี xlsxwriter กรุณาติดตั้งแพ็กเกจ python xlsxwriter")
            )
        request = self.request_id
        subtype = request._require_material_load_subtype()
        categs = request._product_categories_for_material_load(subtype)
        products = self.env["product.product"].search(
            [
                ("purchase_ok", "=", True),
                ("active", "=", True),
                ("categ_id", "child_of", categs.ids),
            ],
            order="default_code, id",
        )

        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        header_fmt = workbook.add_format(
            {"bold": True, "bg_color": "#0B6848", "font_color": "white", "border": 1}
        )
        text_fmt = workbook.add_format({"border": 1})
        code_fmt = workbook.add_format({"border": 1, "num_format": "@"})
        note = workbook.add_worksheet("วิธีใช้")
        note.set_column(0, 0, 88)
        note.write(0, 0, "กรอกปริมาณที่ขอเฉพาะรายการที่ต้องการ แถวที่ปริมาณว่างจะไม่ถูกนำเข้า")
        note.write(1, 0, "จับคู่สินค้าด้วยรหัสสินค้า ถ้ารหัสว่างจะใช้ชื่อรายการ")
        note.write(2, 0, "ราคาต่อหน่วยถ้าว่าง ระบบใช้ราคามาตรฐานของสินค้า")
        note.write(
            3,
            0,
            "หมวดที่เลือก: %s" % (request.load_budget_post_id.display_name or ""),
        )

        sheet = workbook.add_worksheet("รายการวัสดุ")
        headers = [
            "รหัสสินค้า",
            "รายการ",
            "หน่วยนับ",
            "ราคาต่อหน่วย",
            "ปริมาณที่ขอ",
            "หมายเหตุ",
        ]
        widths = [18, 48, 14, 16, 14, 28]
        for col, (title, width) in enumerate(zip(headers, widths)):
            sheet.write(0, col, title, header_fmt)
            sheet.set_column(col, col, width)
        for row, product in enumerate(products, start=1):
            price = product.standard_price or product.lst_price or 0.0
            sheet.write_string(row, 0, product.default_code or "", code_fmt)
            sheet.write(row, 1, product.display_name or "", text_fmt)
            sheet.write(row, 2, product.uom_id.name or "", text_fmt)
            sheet.write(row, 3, price, text_fmt)
        workbook.close()

        attachment = self.env["ir.attachment"].create(
            {
                "name": "template-รายการวัสดุ.xlsx",
                "type": "binary",
                "datas": base64.b64encode(output.getvalue()),
                "mimetype": (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
                "res_model": self._name,
                "res_id": self.id,
            }
        )
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content/%s?download=true" % attachment.id,
            "target": "new",
        }

    def action_import(self):
        self.ensure_one()
        if not openpyxl:
            raise UserError(
                _("ไม่พบไลบรารี openpyxl กรุณาติดตั้งแพ็กเกจ python openpyxl")
            )
        if not self.data_file:
            raise UserError(_("กรุณาเลือกไฟล์ Excel"))
        request = self.request_id
        subtype = request._require_material_load_subtype()
        rows = self._read_excel_rows()
        if not rows:
            raise UserError(_("ไม่พบแถวข้อมูลในชีตรายการวัสดุ"))

        products = self._products_for_subtype(request, subtype)
        by_code, by_name = self._product_indexes(products)
        existing = {
            detail.product_id.id: detail
            for detail in request.material_detail_ids
            if detail.product_id
        }
        Detail = self.env["departmental.budget.request.material.detail"]
        to_create = []
        created = updated = skipped = errors = 0
        logs = []
        base_seq = max(request.material_detail_ids.mapped("sequence") or [0])

        for row_no, row in rows:
            product, reason = self._match_product(row, by_code, by_name)
            if reason == "skip":
                skipped += 1
                continue
            if not product:
                errors += 1
                logs.append(_("แถว %s: %s") % (row_no, reason))
                continue
            quantity = row.get("quantity")
            if quantity is None:
                skipped += 1
                continue
            if quantity is False:
                errors += 1
                logs.append(_("แถว %s: ปริมาณไม่ใช่ตัวเลข") % row_no)
                continue
            if quantity < 0:
                errors += 1
                logs.append(_("แถว %s: ปริมาณติดลบ") % row_no)
                continue
            if quantity == 0:
                skipped += 1
                continue
            price = row.get("price")
            if price is False:
                errors += 1
                logs.append(_("แถว %s: ราคาไม่ใช่ตัวเลข") % row_no)
                continue
            if price is None:
                price = product.standard_price or product.lst_price or 0.0
            if price < 0:
                errors += 1
                logs.append(_("แถว %s: ราคาติดลบ") % row_no)
                continue
            note = row.get("note") or False
            current = existing.get(product.id)
            if current:
                current.with_context(skip_material_amount_sync=True).write(
                    {
                        "quantity": quantity,
                        "unit_price": price,
                        "note": note or current.note,
                        "material_sub_type_id": subtype.id,
                        "budget_post_id": request.load_budget_post_id.id,
                    }
                )
                updated += 1
                continue
            base_seq += 10
            to_create.append(
                {
                    "request_id": request.id,
                    "sequence": base_seq,
                    "product_id": product.id,
                    "name": product.display_name,
                    "material_sub_type_id": subtype.id,
                    "budget_post_id": request.load_budget_post_id.id,
                    "product_uom_id": product.uom_id.id,
                    "quantity": quantity,
                    "unit_price": price,
                    "note": note,
                }
            )
            created += 1

        if to_create:
            Detail.with_context(skip_material_amount_sync=True).create(to_create)
        request._ensure_material_overview_line(subtype, request.load_budget_post_id)
        request._sync_material_amounts_from_details()

        summary = _(
            "เพิ่ม %s รายการ อัปเดต %s รายการ ข้าม %s แถว พบปัญหา %s แถว"
        ) % (created, updated, skipped, errors)
        log_text = summary
        if logs:
            log_text = summary + "\n" + "\n".join(logs[:80])
        self.write(
            {
                "state": "done",
                "created_count": created,
                "updated_count": updated,
                "skipped_count": skipped,
                "error_count": errors,
                "log_text": log_text,
            }
        )
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    def _products_for_subtype(self, request, subtype):
        categs = request._product_categories_for_material_load(subtype)
        if not categs:
            raise ValidationError(
                _("ไม่พบหมวดสินค้าสำหรับกลุ่มวัสดุ \"%s\"")
                % (subtype.budget_group_id.display_name,)
            )
        return self.env["product.product"].search(
            [
                ("purchase_ok", "=", True),
                ("active", "=", True),
                ("categ_id", "child_of", categs.ids),
            ]
        )

    @staticmethod
    def _product_indexes(products):
        by_code = {}
        by_name = {}
        for product in products:
            code = (product.default_code or "").strip().casefold()
            if code and code not in by_code:
                by_code[code] = product
            for raw_name in (product.display_name, product.name):
                name = (raw_name or "").strip().casefold()
                if name and name not in by_name:
                    by_name[name] = product
                if "]" in name:
                    plain = name.split("]", 1)[-1].strip()
                    if plain and plain not in by_name:
                        by_name[plain] = product
        return by_code, by_name

    @staticmethod
    def _match_product(row, by_code, by_name):
        code = (row.get("code") or "").strip()
        name = (row.get("name") or "").strip()
        if not code and not name:
            return None, "skip"
        if code:
            product = by_code.get(code.casefold())
            if product:
                return product, None
            return None, _("ไม่พบรหัสสินค้า %s ในหมวดที่เลือก") % code
        product = by_name.get(name.casefold())
        if product:
            return product, None
        return None, _("ไม่พบรายการ %s ในหมวดที่เลือก") % name

    def _read_excel_rows(self):
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.data_file)),
                read_only=True,
                data_only=True,
            )
        except Exception as error:
            raise UserError(_("อ่านไฟล์ Excel ไม่ได้: %s") % error) from error
        sheet = workbook["รายการวัสดุ"] if "รายการวัสดุ" in workbook.sheetnames else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        header = rows[0] if rows else None
        if not header:
            return []
        rows = rows[1:]
        columns = []
        for cell in header:
            key = str(cell or "").strip().casefold()
            columns.append(HEADER_FIELDS.get(key) or HEADER_FIELDS.get(str(cell or "").strip()))
        if "code" not in columns and "name" not in columns:
            raise UserError(
                _("ไม่พบคอลัมน์ รหัสสินค้า หรือ รายการ ในแถวแรกของไฟล์")
            )
        parsed = []
        for index, values in enumerate(rows, start=2):
            if not values or not any(value not in (None, "") for value in values):
                continue
            row = {}
            for col, value in zip(columns, values):
                if not col or col in row:
                    continue
                row[col] = value
            row["code"] = self._normalize_code(row.get("code"))
            row["name"] = "" if row.get("name") is None else str(row.get("name")).strip()
            row["quantity"] = self._to_number(row.get("quantity"), blank_ok=True)
            row["price"] = self._to_number(row.get("price"), blank_ok=True)
            note = row.get("note")
            row["note"] = "" if note is None else str(note).strip()
            parsed.append((index, row))
        return parsed

    @staticmethod
    def _normalize_code(value):
        if value in (None, ""):
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        text = str(value).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            return text[:-2]
        return text

    @staticmethod
    def _to_number(value, blank_ok=False):
        if value in (None, ""):
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return False if blank_ok else None
