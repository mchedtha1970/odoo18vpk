# -*- coding: utf-8 -*-
import base64
import io
import logging

from odoo import _, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

try:
    import xlsxwriter
except ImportError:  # pragma: no cover
    xlsxwriter = None


BOOL_TRUE = {"1", "true", "yes", "y", "t", "x", "ใช่", "จริง"}
BOOL_FALSE = {"0", "false", "no", "n", "f", "ไม่", "เท็จ", ""}

SECTION_KEY_MAP = {
    "material": "material",
    "พัสดุ": "material",
    "วัสดุ": "material",
    "asset": "asset",
    "ครุภัณฑ์": "asset",
    "construction": "construction",
    "ก่อสร้าง": "construction",
    "project": "project",
    "โครงการ": "project",
}

# Sheet definitions: technical headers + Thai aliases for import matching.
SHEET_SPECS = {
    "FundSource": {
        "model": "vpk.budget.fund.source",
        "label": "แหล่งเงิน",
        "key": "code",
        "columns": [
            ("code", ["code", "รหัส"]),
            ("name", ["name", "ชื่อ"]),
            ("sequence", ["sequence", "ลำดับ"]),
            ("active", ["active", "ใช้งาน"]),
            ("description", ["description", "รายละเอียด"]),
        ],
        "widths": [18, 40, 10, 10, 40],
        "samples": [
            ["STATE_BUDGET", "เงินงบประมาณ", 10, 1, ""],
            ["HOSP_MAINT", "เงินบำรุง", 20, 1, ""],
        ],
        "note": "จับคู่ด้วย code (รหัสต้องไม่ซ้ำ)",
    },
    "BudgetGroup": {
        "model": "vpk.budget.group",
        "label": "กลุ่มงบประมาณ-วัสดุ",
        "key": "code",
        "columns": [
            ("code", ["code", "รหัส"]),
            ("name", ["name", "ชื่อ"]),
            ("sequence", ["sequence", "ลำดับ"]),
            ("active", ["active", "ใช้งาน"]),
            ("description", ["description", "รายละเอียด"]),
        ],
        "widths": [20, 36, 10, 10, 40],
        "samples": [
            ["MEDICINE", "เวชภัณฑ์ ยา", 10, 1, ""],
            ["GENERAL_SUPPLY", "วัสดุทั่วไป", 50, 1, ""],
        ],
        "note": "จับคู่ด้วย code (รหัสต้องไม่ซ้ำ)",
    },
    "MaterialSubType": {
        "model": "vpk.budget.material.sub.type",
        "label": "ประเภทวัสดุย่อย",
        "key": "code",
        "key_with": "budget_group_code",
        "columns": [
            ("code", ["code", "รหัส"]),
            ("name", ["name", "ชื่อ", "ชื่อประเภทวัสดุย่อย"]),
            ("budget_group_code", ["budget_group_code", "รหัสกลุ่มงบประมาณ", "กลุ่มงบประมาณ"]),
            ("sequence", ["sequence", "ลำดับ"]),
            ("active", ["active", "ใช้งาน"]),
            ("description", ["description", "รายละเอียด"]),
        ],
        "widths": [16, 36, 22, 10, 10, 36],
        "samples": [
            ["ORAL", "ยาใช้ภายใน", "MEDICINE", 10, 1, ""],
            ["OFFICE", "วัสดุสำนักงาน", "GENERAL_SUPPLY", 10, 1, ""],
        ],
        "note": (
            "จับคู่ด้วย code ภายในกลุ่มงบประมาณเดียวกัน | "
            "budget_group_code = รหัสจากชีต BudgetGroup"
        ),
    },
    "BudgetType": {
        "model": "vpk.budget.type",
        "label": "ประเภทงบประมาณ",
        "key": "code",
        "columns": [
            ("code", ["code", "รหัส"]),
            ("name", ["name", "ชื่อ", "ชื่อประเภทงบประมาณ"]),
            ("section_key", ["section_key", "กลุ่มแบบฟอร์ม"]),
            ("sequence", ["sequence", "ลำดับ"]),
            ("active", ["active", "ใช้งาน"]),
            ("description", ["description", "รายละเอียด"]),
        ],
        "widths": [22, 28, 16, 10, 10, 36],
        "samples": [
            ["supplies_budget", "วัสดุ", "material", 10, 1, ""],
            ["asset_budget", "ครุภัณฑ์", "asset", 20, 1, ""],
        ],
        "note": (
            "section_key = material/asset/construction/project "
            "(หรือ พัสดุ/ครุภัณฑ์/ก่อสร้าง/โครงการ)"
        ),
    },
    "AssetCategory": {
        "model": "vpk.budget.asset.category",
        "label": "กลุ่มครุภัณฑ์",
        "key": "code",
        "columns": [
            ("code", ["code", "รหัส"]),
            ("name", ["name", "ชื่อ"]),
            ("product_categ_id", ["product_categ_id", "หมวดสินค้า", "product category", "product_category"]),
            ("sequence", ["sequence", "ลำดับ"]),
            ("active", ["active", "ใช้งาน"]),
        ],
        "widths": [14, 40, 36, 10, 10],
        "samples": [
            ["medical", "ครุภัณฑ์การแพทย์", "ครุภัณฑ์ / 01. ครุภัณฑ์การแพทย์", 50, 1],
            ["office", "ครุภัณฑ์สำนักงาน", "ครุภัณฑ์ / 12. ครุภัณฑ์สำนักงาน", 10, 1],
        ],
        "note": (
            "จับคู่ด้วย code (รหัสต้องไม่ซ้ำ) | "
            "product_categ_id = ชื่อหรือ complete name ของหมวดสินค้าที่ต้องการผูก"
        ),
    },
    "FormType": {
        "model": "vpk.budget.request.form.type",
        "label": "แบบฟอร์มคำของบ",
        "key": "code",
        "columns": [
            ("code", ["code", "รหัส"]),
            ("name", ["name", "ชื่อ"]),
            ("sequence", ["sequence", "ลำดับ"]),
            ("active", ["active", "ใช้งาน"]),
        ],
        "widths": [18, 40, 10, 10],
        "samples": [
            ["material_budget", "แบบฟอร์มตั้งงบวัสดุ", 10, 1],
            ["asset", "แบบฟอร์มครุภัณฑ์", 20, 1],
            ["construction", "แบบฟอร์ม ก่อสร้าง", 30, 1],
            ["project", "แบบฟอร์มโครงการ", 40, 1],
        ],
        "note": "จับคู่ด้วย code | กำหนดประเภทงบในชีต FormTypeLine",
    },
    "FormTypeLine": {
        "model": "vpk.budget.request.form.type.line",
        "label": "ประเภทงบในแบบฟอร์ม",
        "key": "budget_type_code",
        "key_with": "form_type_code",
        "columns": [
            ("form_type_code", ["form_type_code", "รหัสแบบฟอร์ม"]),
            ("budget_type_code", ["budget_type_code", "รหัสประเภทงบประมาณ"]),
            ("sequence", ["sequence", "ลำดับ"]),
            ("name", ["name", "ชื่อแท็บ"]),
        ],
        "widths": [18, 22, 10, 28],
        "samples": [
            ["material_budget", "supplies_budget", 10, ""],
            ["asset", "asset_budget", 10, ""],
            ["construction", "construction_budget", 10, ""],
            ["project", "project_budget", 10, ""],
        ],
        "note": (
            "form_type_code จากชีต FormType | "
            "budget_type_code จากชีต BudgetType | "
            "ชื่อแท็บว่างได้"
        ),
    },
}

# Import order respects foreign-key dependencies.
IMPORT_SHEET_ORDER = [
    "FundSource",
    "BudgetGroup",
    "BudgetType",
    "AssetCategory",
    "MaterialSubType",
    "FormType",
    "FormTypeLine",
]


class BudgetMasterImportWizard(models.TransientModel):
    _name = "vpk.budget.master.import.wizard"
    _description = "Import Budget Master Data"
    _transient_max_hours = 72.0
    _transient_max_count = 200

    data_file = fields.Binary(string="Excel File", attachment=False)
    filename = fields.Char(string="Filename")
    import_mode = fields.Selection(
        [
            ("upsert", "สร้างใหม่ + อัปเดตที่มีอยู่แล้ว (ตามรหัส)"),
            ("create", "สร้างใหม่เท่านั้น (ข้ามถ้ารหัสมีอยู่)"),
            ("update", "อัปเดตเท่านั้น (ข้ามถ้ารหัสไม่พบ)"),
        ],
        string="Import Mode",
        default="upsert",
        required=True,
    )
    include_existing = fields.Boolean(
        string="ใส่ข้อมูลที่มีอยู่ในระบบลง Template",
        default=True,
        help="เมื่อดาวน์โหลด Template จะเติมแถวจาก master ที่ตั้งค่าไว้แล้วในระบบ",
    )
    state = fields.Selection(
        [("draft", "Draft"), ("done", "Done")],
        default="draft",
    )
    log_text = fields.Text(string="Import Log", readonly=True)
    created_count = fields.Integer(readonly=True)
    updated_count = fields.Integer(readonly=True)
    skipped_count = fields.Integer(readonly=True)
    error_count = fields.Integer(readonly=True)

    # -------------------------------------------------------------------------
    # Template download
    # -------------------------------------------------------------------------
    def action_download_template(self):
        self.ensure_one()
        if not xlsxwriter:
            raise UserError(
                _("ไม่พบไลบรารี xlsxwriter กรุณาติดตั้งแพ็กเกจ python xlsxwriter")
            )

        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        header_fmt = workbook.add_format(
            {
                "bold": True,
                "bg_color": "#1F4E79",
                "font_color": "white",
                "border": 1,
            }
        )
        sample_fmt = workbook.add_format({"border": 1})
        note_fmt = workbook.add_format({"italic": True, "font_color": "#666666"})
        title_fmt = workbook.add_format({"bold": True, "font_size": 14})
        wrap_fmt = workbook.add_format({"text_wrap": True, "valign": "top"})

        self._write_instructions_sheet(workbook, title_fmt, wrap_fmt, header_fmt)

        for sheet_name in IMPORT_SHEET_ORDER:
            spec = SHEET_SPECS[sheet_name]
            ws = workbook.add_worksheet(sheet_name)
            headers = [col[0] for col in spec["columns"]]
            for col, (header, width) in enumerate(zip(headers, spec["widths"])):
                ws.write(0, col, header, header_fmt)
                ws.set_column(col, col, width)

            rows = []
            if self.include_existing:
                rows = self._export_existing_rows(sheet_name, spec)
            if not rows:
                rows = list(spec.get("samples") or [])

            for r_idx, row in enumerate(rows, start=1):
                for c_idx, value in enumerate(row):
                    ws.write(r_idx, c_idx, value, sample_fmt)

            note_row = max(len(rows) + 2, 4)
            ws.write(
                note_row,
                0,
                "หมายเหตุ: %s | active = 1 หรือ 0 (ใช่/ไม่)" % spec["note"],
                note_fmt,
            )

        workbook.close()
        content = output.getvalue()
        attachment = self.env["ir.attachment"].create(
            {
                "name": "vpk_budget_master_template.xlsx",
                "type": "binary",
                "datas": base64.b64encode(content),
                "mimetype": (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
                "res_model": self._name,
                "res_id": self.id,
            }
        )
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content/%s?download=true" % attachment.id,
            "target": "new",
        }

    def _write_instructions_sheet(self, workbook, title_fmt, wrap_fmt, header_fmt):
        ws = workbook.add_worksheet("Instructions")
        ws.set_column(0, 0, 28)
        ws.set_column(1, 1, 70)
        ws.write(0, 0, "Template นำเข้า Master ระบบงบประมาณ (VPK Budget)", title_fmt)
        rows = [
            ("วิธีใช้", ""),
            ("1", "ดาวน์โหลดไฟล์นี้ แล้วกรอกข้อมูลในแต่ละชีต"),
            ("2", "อัปโหลดไฟล์กลับในเมนู Import Budget Master แล้วกด Import"),
            ("3", "ระบบจับคู่ด้วยรหัส (code) ตามชีต — ดูหมายเหตุท้ายแต่ละชีต"),
            ("", ""),
            ("ลำดับชีตที่แนะนำ", "ควรเตรียมตามลำดับนี้เพราะมีการอ้างอิงรหัสข้ามชีต"),
            ("FundSource", "แหล่งเงิน"),
            ("BudgetGroup", "กลุ่มงบประมาณ-วัสดุ"),
            ("BudgetType", "ประเภทงบประมาณ (section_key)"),
            ("AssetCategory", "กลุ่มครุภัณฑ์ (ผูกหมวดสินค้าได้ที่คอลัมน์ product_categ_id)"),
            ("MaterialSubType", "ประเภทวัสดุย่อย (อ้างอิง BudgetGroup)"),
            ("FormType", "แบบฟอร์มคำของบ"),
            ("FormTypeLine", "ผูกประเภทงบกับแบบฟอร์ม"),
            ("", ""),
            ("ค่า active", "1 / ใช่ = ใช้งาน, 0 / ไม่ = ปิดใช้งาน (ว่าง = 1)"),
            ("section_key", "material, asset, construction, project"),
            ("Import Mode", "upsert = สร้าง+อัปเดต, create = สร้างอย่างเดียว, update = อัปเดตอย่างเดียว"),
        ]
        ws.write(2, 0, "หัวข้อ", header_fmt)
        ws.write(2, 1, "รายละเอียด", header_fmt)
        for idx, (left, right) in enumerate(rows, start=3):
            ws.write(idx, 0, left)
            ws.write(idx, 1, right, wrap_fmt)

    def _export_existing_rows(self, sheet_name, spec):
        Model = self.env[spec["model"]].with_context(active_test=False)
        records = Model.search([], order="sequence, id" if "sequence" in Model._fields else "id")
        rows = []
        for rec in records:
            row = []
            for field_name, _aliases in spec["columns"]:
                row.append(self._export_field_value(sheet_name, rec, field_name))
            rows.append(row)
        return rows

    def _export_field_value(self, sheet_name, rec, field_name):
        if sheet_name == "MaterialSubType" and field_name == "budget_group_code":
            return rec.budget_group_id.code or ""
        if sheet_name == "FormTypeLine" and field_name == "form_type_code":
            return rec.form_type_id.code or ""
        if sheet_name == "FormTypeLine" and field_name == "budget_type_code":
            return rec.budget_type_id.code or ""
        if sheet_name == "AssetCategory" and field_name == "product_categ_id":
            return rec.product_categ_id.complete_name or rec.product_categ_id.name or ""
        if field_name not in rec._fields:
            return ""
        value = rec[field_name]
        if isinstance(value, bool):
            return 1 if value else 0
        return value if value is not None else ""

    # -------------------------------------------------------------------------
    # Import
    # -------------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if not openpyxl:
            raise UserError(
                _("ไม่พบไลบรารี openpyxl กรุณาติดตั้งแพ็กเกจ python openpyxl")
            )
        if not self.data_file:
            raise UserError(_("กรุณาอัปโหลดไฟล์ Excel"))

        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.data_file)),
                data_only=True,
                read_only=True,
            )
        except Exception as exc:
            raise UserError(_("ไม่สามารถเปิดไฟล์ Excel ได้: %s") % exc) from exc

        created = updated = skipped = errors = 0
        logs = []

        for sheet_name in IMPORT_SHEET_ORDER:
            if sheet_name not in workbook.sheetnames:
                logs.append(_("ข้ามชีต %s (ไม่พบในไฟล์)") % sheet_name)
                continue
            c, u, s, e, sheet_logs = self._import_sheet(
                workbook[sheet_name], sheet_name
            )
            created += c
            updated += u
            skipped += s
            errors += e
            logs.extend(sheet_logs)

        self.write(
            {
                "state": "done",
                "created_count": created,
                "updated_count": updated,
                "skipped_count": skipped,
                "error_count": errors,
                "log_text": "\n".join(logs) or _("ไม่มีแถวข้อมูลให้นำเข้า"),
            }
        )
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    def _import_sheet(self, worksheet, sheet_name):
        spec = SHEET_SPECS[sheet_name]
        header_map = self._build_header_map(worksheet, spec)
        if not header_map:
            return 0, 0, 0, 1, [
                _("[%s] ไม่พบหัวคอลัมน์ที่รู้จักในแถวแรก") % sheet_name
            ]

        created = updated = skipped = errors = 0
        logs = [_("--- ชีต %s (%s) ---") % (sheet_name, spec["label"])]
        row_idx = 1
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            row_idx += 1
            if self._row_empty(row):
                continue
            # Skip note / instruction rows
            first = row[0] if row else None
            if isinstance(first, str) and first.strip().startswith("หมายเหตุ"):
                continue
            try:
                vals = self._row_to_vals(row, header_map, spec)
                if not vals:
                    skipped += 1
                    continue
                action, message = self._upsert_row(sheet_name, spec, vals)
                if action == "created":
                    created += 1
                elif action == "updated":
                    updated += 1
                else:
                    skipped += 1
                logs.append(_("แถว %s: %s") % (row_idx, message))
            except Exception as exc:  # noqa: BLE001 - collect per-row errors
                errors += 1
                logs.append(_("แถว %s: ERROR %s") % (row_idx, exc))
                _logger.exception("Budget master import error %s row %s", sheet_name, row_idx)
        return created, updated, skipped, errors, logs

    def _build_header_map(self, worksheet, spec):
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            return {}
        alias_to_field = {}
        for field_name, aliases in spec["columns"]:
            for alias in aliases:
                alias_to_field[self._normalize_header(alias)] = field_name
        header_map = {}
        for idx, cell in enumerate(header_row):
            key = self._normalize_header(cell)
            if key in alias_to_field:
                header_map[idx] = alias_to_field[key]
        return header_map

    @staticmethod
    def _normalize_header(value):
        if value is None:
            return ""
        return str(value).strip().lower().replace(" ", "")

    @staticmethod
    def _row_empty(row):
        return not row or all(cell is None or str(cell).strip() == "" for cell in row)

    def _row_to_vals(self, row, header_map, spec):
        raw = {}
        for idx, field_name in header_map.items():
            if idx >= len(row):
                continue
            raw[field_name] = row[idx]
        if "code" in raw and not self._as_str(raw.get("code")):
            return {}
        if "form_type_code" in raw and not self._as_str(raw.get("form_type_code")):
            return {}
        if "name" in dict(spec["columns"]) and "name" in raw:
            # name required for most masters except FormTypeLine (optional tab name)
            if spec["model"] != "vpk.budget.request.form.type.line" and not self._as_str(
                raw.get("name")
            ):
                raise UserError(_("ต้องระบุ name"))
        return raw

    def _upsert_row(self, sheet_name, spec, raw):
        if sheet_name == "MaterialSubType":
            return self._upsert_material_sub_type(raw)
        if sheet_name == "FormTypeLine":
            return self._upsert_form_type_line(raw)
        if sheet_name == "BudgetType":
            return self._upsert_budget_type(raw)
        return self._upsert_simple_master(spec, raw)

    def _upsert_simple_master(self, spec, raw):
        Model = self.env[spec["model"]].with_context(active_test=False)
        code = self._as_str(raw.get("code"))
        if not code:
            raise UserError(_("ต้องระบุ code"))
        vals = {
            "code": code,
            "name": self._as_str(raw.get("name")),
        }
        if "sequence" in Model._fields and raw.get("sequence") not in (None, ""):
            vals["sequence"] = self._as_int(raw.get("sequence"), default=10)
        if "active" in Model._fields:
            vals["active"] = self._as_bool(raw.get("active"), default=True)
        if "description" in Model._fields and "description" in raw:
            vals["description"] = self._as_str(raw.get("description")) or False
        if spec["model"] == "vpk.budget.asset.category" and "product_categ_id" in raw:
            vals["product_categ_id"] = self._resolve_product_category(
                raw.get("product_categ_id")
            )

        record = Model.search([("code", "=", code)], limit=1)
        return self._apply_mode(Model, record, vals, code)

    def _upsert_budget_type(self, raw):
        Model = self.env["vpk.budget.type"].with_context(active_test=False)
        code = self._as_str(raw.get("code"))
        if not code:
            raise UserError(_("ต้องระบุ code"))
        section_key = self._as_section_key(raw.get("section_key"))
        if not section_key:
            raise UserError(
                _("ต้องระบุ section_key (material/asset/construction/project)")
            )
        vals = {
            "code": code,
            "name": self._as_str(raw.get("name")),
            "section_key": section_key,
            "sequence": self._as_int(raw.get("sequence"), default=10),
            "active": self._as_bool(raw.get("active"), default=True),
            "description": self._as_str(raw.get("description")) or False,
        }
        record = Model.search([("code", "=", code)], limit=1)
        return self._apply_mode(Model, record, vals, code)

    def _upsert_material_sub_type(self, raw):
        Model = self.env["vpk.budget.material.sub.type"].with_context(active_test=False)
        code = self._as_str(raw.get("code"))
        group_code = self._as_str(raw.get("budget_group_code"))
        if not code or not group_code:
            raise UserError(_("ต้องระบุ code และ budget_group_code"))
        group = self.env["vpk.budget.group"].with_context(active_test=False).search(
            [("code", "=", group_code)], limit=1
        )
        if not group:
            raise UserError(_("ไม่พบกลุ่มงบประมาณรหัส %s") % group_code)
        vals = {
            "code": code,
            "name": self._as_str(raw.get("name")),
            "budget_group_id": group.id,
            "sequence": self._as_int(raw.get("sequence"), default=10),
            "active": self._as_bool(raw.get("active"), default=True),
            "description": self._as_str(raw.get("description")) or False,
        }
        record = Model.search(
            [("code", "=", code), ("budget_group_id", "=", group.id)], limit=1
        )
        label = "%s / %s" % (group_code, code)
        return self._apply_mode(Model, record, vals, label)

    def _upsert_form_type_line(self, raw):
        Model = self.env["vpk.budget.request.form.type.line"]
        form_code = self._as_str(raw.get("form_type_code"))
        budget_type_code = self._as_str(raw.get("budget_type_code"))
        if not form_code or not budget_type_code:
            raise UserError(_("ต้องระบุ form_type_code และ budget_type_code"))
        form_type = self.env["vpk.budget.request.form.type"].with_context(
            active_test=False
        ).search([("code", "=", form_code)], limit=1)
        if not form_type:
            raise UserError(_("ไม่พบแบบฟอร์มรหัส %s") % form_code)
        budget_type = self.env["vpk.budget.type"].with_context(active_test=False).search(
            [("code", "=", budget_type_code)], limit=1
        )
        if not budget_type:
            raise UserError(_("ไม่พบประเภทงบประมาณรหัส %s") % budget_type_code)
        vals = {
            "form_type_id": form_type.id,
            "budget_type_id": budget_type.id,
            "sequence": self._as_int(raw.get("sequence"), default=10),
            "name": self._as_str(raw.get("name")) or False,
        }
        record = Model.search(
            [
                ("form_type_id", "=", form_type.id),
                ("budget_type_id", "=", budget_type.id),
            ],
            limit=1,
        )
        label = "%s / %s" % (form_code, budget_type_code)
        return self._apply_mode(Model, record, vals, label)

    def _apply_mode(self, Model, record, vals, label):
        mode = self.import_mode
        if record:
            if mode == "create":
                return "skipped", _("ข้าม %s (มีอยู่แล้ว)") % label
            record.write(vals)
            return "updated", _("อัปเดต %s") % label
        if mode == "update":
            return "skipped", _("ข้าม %s (ไม่พบในระบบ)") % label
        Model.create(vals)
        return "created", _("สร้าง %s") % label

    @staticmethod
    def _as_str(value):
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @staticmethod
    def _as_int(value, default=0):
        if value in (None, ""):
            return default
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _as_bool(value, default=True):
        if value in (None, ""):
            return default
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return bool(value)
        text = str(value).strip().lower()
        if text in BOOL_TRUE:
            return True
        if text in BOOL_FALSE:
            return False
        return default

    @staticmethod
    def _as_section_key(value):
        if value in (None, ""):
            return False
        text = str(value).strip().lower()
        # keep original Thai keys via map
        mapped = SECTION_KEY_MAP.get(text) or SECTION_KEY_MAP.get(str(value).strip())
        return mapped or False

    def _resolve_product_category(self, value):
        raw = self._as_str(value)
        if not raw:
            return False
        Categ = self.env["product.category"].with_context(active_test=False)
        rec = Categ.search(
            ["|", ("complete_name", "=", raw), ("name", "=", raw)],
            limit=1,
        )
        if not rec:
            rec = Categ.search([("complete_name", "ilike", raw)], limit=1)
        if not rec:
            rec = Categ.search([("name", "ilike", raw)], limit=1)
        if not rec and raw.isdigit():
            rec = Categ.browse(int(raw)).exists()
        if not rec:
            raise UserError(_("ไม่พบหมวดสินค้า %s") % raw)
        return rec.id
