#!/usr/bin/env python3
# Copyright 2026 VPK
"""Generate PowerPoint overview of VPK HIS ↔ ERP APIs for the HIS dev team."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parent
    / "VPK-HIS-API-Presentation-สำหรับทีมพัฒนา-HIS.pptx"
)

# Brand colours (match Word specs)
NAVY = RGBColor(0x0B, 0x3D, 0x5C)
TEAL = RGBColor(0x0E, 0x6B, 0x6B)
LIGHT = RGBColor(0xF4, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GREEN = RGBColor(0x1A, 0x7A, 0x4C)
RED = RGBColor(0xC0, 0x39, 0x2B)
PURPLE = RGBColor(0x6C, 0x34, 0x8B)

FONT = "Prompt"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


class HisApiDeck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self._blank = self.prs.slide_layouts[6]

    # ── low-level helpers ──────────────────────────────────────────────

    def _slide(self):
        return self.prs.slides.add_slide(self._blank)

    def _fill_bg(self, slide, color=WHITE):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color

    def _bar(self, slide, color=NAVY, height=Inches(0.12)):
        s = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, height
        )
        s.fill.solid()
        s.fill.fore_color.rgb = color
        s.line.fill.background()

    def _footer(self, slide, text):
        box = slide.shapes.add_textbox(
            Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.35)
        )
        p = box.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(9)
        p.font.color.rgb = GRAY
        p.font.name = FONT

    def _text(
        self,
        slide,
        left,
        top,
        width,
        height,
        text,
        size=18,
        bold=False,
        color=DARK,
        align=PP_ALIGN.LEFT,
    ):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.font.name = FONT
        p.alignment = align
        return box

    def _bullets(self, slide, left, top, width, height, items, size=16, color=DARK):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = item
            p.level = 0
            p.font.size = Pt(size)
            p.font.color.rgb = color
            p.font.name = FONT
            p.space_after = Pt(6)
        return box

    def _code(self, slide, left, top, width, height, text, size=11):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xF3, 0xF4, 0xF6)
        shape.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_top = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.name = "Consolas"
        p.font.color.rgb = DARK
        return shape

    def _box(
        self,
        slide,
        left,
        top,
        width,
        height,
        text,
        fill=NAVY,
        text_color=WHITE,
        size=13,
        bold=True,
    ):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = text_color
        p.font.name = FONT
        p.alignment = PP_ALIGN.CENTER
        return shape

    def _arrow(self, slide, x1, y1, x2, y2, color=NAVY):
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2
        )
        conn.line.color.rgb = color
        conn.line.width = Pt(2.5)
        return conn

    def _table(self, slide, left, top, width, rows, col_widths=None, header=True):
        n_rows = len(rows)
        n_cols = len(rows[0])
        height = Inches(0.38 * n_rows + 0.2)
        tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
        tbl = tbl_shape.table
        if col_widths:
            for i, w in enumerate(col_widths):
                tbl.columns[i].width = w
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.text = val
                for p in cell.text_frame.paragraphs:
                    p.font.name = FONT
                    p.font.size = Pt(11 if r > 0 or not header else 12)
                    p.font.bold = header and r == 0
                    p.font.color.rgb = WHITE if header and r == 0 else DARK
                if header and r == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = NAVY
                elif r % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = LIGHT
        return tbl_shape

    # ── slide templates ──────────────────────────────────────────────

    def title_slide(self, title, subtitle, footer):
        slide = self._slide()
        self._fill_bg(slide, NAVY)
        self._text(
            slide,
            Inches(0.8),
            Inches(1.8),
            Inches(11.7),
            Inches(1.2),
            "โรงพยาบาลวชิระภูเก็ต",
            size=22,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
        self._text(
            slide,
            Inches(0.8),
            Inches(2.8),
            Inches(11.7),
            Inches(1.5),
            title,
            size=36,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
        self._text(
            slide,
            Inches(0.8),
            Inches(4.5),
            Inches(11.7),
            Inches(1.0),
            subtitle,
            size=20,
            color=RGBColor(0xB0, 0xD4, 0xE8),
            align=PP_ALIGN.CENTER,
        )
        self._text(
            slide,
            Inches(0.8),
            Inches(6.2),
            Inches(11.7),
            Inches(0.5),
            footer,
            size=14,
            color=RGBColor(0x90, 0xB0, 0xC8),
            align=PP_ALIGN.CENTER,
        )

    def section_slide(self, num, title, subtitle=""):
        slide = self._slide()
        self._fill_bg(slide, TEAL)
        self._text(
            slide,
            Inches(0.8),
            Inches(2.5),
            Inches(2),
            Inches(1),
            f"{num:02d}",
            size=72,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.LEFT,
        )
        self._text(
            slide,
            Inches(2.5),
            Inches(2.8),
            Inches(10),
            Inches(1.2),
            title,
            size=36,
            bold=True,
            color=WHITE,
        )
        if subtitle:
            self._text(
                slide,
                Inches(2.5),
                Inches(4.0),
                Inches(10),
                Inches(0.8),
                subtitle,
                size=18,
                color=WHITE,
            )

    def content_slide(self, title, footer, bullets=None, note=None):
        slide = self._slide()
        self._bar(slide)
        self._text(
            slide, Inches(0.5), Inches(0.35), Inches(12), Inches(0.7), title, size=28, bold=True, color=NAVY
        )
        y = Inches(1.2)
        if bullets:
            self._bullets(slide, Inches(0.6), y, Inches(12), Inches(5.5), bullets, size=17)
        if note:
            self._text(
                slide,
                Inches(0.6),
                Inches(6.3),
                Inches(12),
                Inches(0.6),
                note,
                size=13,
                color=TEAL,
                bold=True,
            )
        self._footer(slide, footer)

    def two_col_slide(self, title, footer, left_title, left_items, right_title, right_items):
        slide = self._slide()
        self._bar(slide)
        self._text(
            slide, Inches(0.5), Inches(0.35), Inches(12), Inches(0.7), title, size=28, bold=True, color=NAVY
        )
        self._text(slide, Inches(0.6), Inches(1.15), Inches(5.8), Inches(0.4), left_title, size=18, bold=True, color=TEAL)
        self._bullets(slide, Inches(0.6), Inches(1.55), Inches(5.8), Inches(5.2), left_items, size=15)
        self._text(slide, Inches(6.8), Inches(1.15), Inches(5.8), Inches(0.4), right_title, size=18, bold=True, color=TEAL)
        self._bullets(slide, Inches(6.8), Inches(1.55), Inches(5.8), Inches(5.2), right_items, size=15)
        self._footer(slide, footer)

    def flow_slide(self, title, footer, steps, colors=None):
        """Horizontal pipeline: list of (label, sublabel)."""
        slide = self._slide()
        self._bar(slide)
        self._text(
            slide, Inches(0.5), Inches(0.35), Inches(12), Inches(0.7), title, size=28, bold=True, color=NAVY
        )
        n = len(steps)
        box_w = Inches(10.5 / n - 0.15)
        box_h = Inches(1.1)
        y = Inches(2.8)
        x0 = Inches(0.65)
        default_colors = [NAVY, TEAL, ORANGE, GREEN, PURPLE, RED]
        for i, (label, sub) in enumerate(steps):
            x = x0 + i * (box_w + Inches(0.18))
            fill = colors[i] if colors and i < len(colors) else default_colors[i % len(default_colors)]
            self._box(slide, x, y, box_w, box_h, label, fill=fill, size=12)
            if sub:
                self._text(slide, x, y + box_h + Inches(0.08), box_w, Inches(0.55), sub, size=10, color=GRAY, align=PP_ALIGN.CENTER)
            if i < n - 1:
                self._arrow(
                    slide,
                    x + box_w,
                    y + box_h / 2,
                    x + box_w + Inches(0.18),
                    y + box_h / 2,
                )
        self._footer(slide, footer)

    def architecture_slide(self, footer):
        slide = self._slide()
        self._bar(slide)
        self._text(
            slide,
            Inches(0.5),
            Inches(0.35),
            Inches(12),
            Inches(0.7),
            "สถาปัตยกรรมภาพรวม",
            size=28,
            bold=True,
            color=NAVY,
        )
        # HIS zone
        self._box(
            slide,
            Inches(0.5),
            Inches(1.3),
            Inches(3.5),
            Inches(4.8),
            "",
            fill=LIGHT,
            text_color=DARK,
        )
        self._text(
            slide, Inches(0.7), Inches(1.45), Inches(3.1), Inches(0.4), "ระบบ HIS", size=16, bold=True, color=NAVY
        )
        self._bullets(
            slide,
            Inches(0.7),
            Inches(1.9),
            Inches(3.1),
            Inches(4.0),
            [
                "Front HIS / ห้องยา",
                "POS / Cashier",
                "Pharmacy module",
                "",
                "ส่ง JSON ผ่าน HTTPS",
                "Header: X-Api-Key",
            ],
            size=13,
        )
        # API layer
        self._box(
            slide,
            Inches(4.4),
            Inches(2.0),
            Inches(4.2),
            Inches(2.8),
            "REST API\n/vpk/api/v1/his",
            fill=TEAL,
            size=14,
        )
        self._bullets(
            slide,
            Inches(4.55),
            Inches(3.5),
            Inches(3.9),
            Inches(1.2),
            ["/revenue", "/stock-issues", "/stock-requisitions"],
            size=11,
            color=WHITE,
        )
        # ERP zone
        self._box(
            slide,
            Inches(9.0),
            Inches(1.3),
            Inches(3.8),
            Inches(4.8),
            "",
            fill=LIGHT,
            text_color=DARK,
        )
        self._text(
            slide, Inches(9.2), Inches(1.45), Inches(3.4), Inches(0.4), "Odoo 18 ERP", size=16, bold=True, color=NAVY
        )
        self._bullets(
            slide,
            Inches(9.2),
            Inches(1.9),
            Inches(3.4),
            Inches(4.0),
            [
                "Staging Batch (คิว)",
                "Validate → Ready",
                "เจ้าหน้าที่ Post",
                "",
                "→ บัญชีรายได้",
                "→ ตัดสต็อก / โอนคลัง",
            ],
            size=13,
        )
        self._arrow(slide, Inches(4.0), Inches(3.4), Inches(4.4), Inches(3.4))
        self._arrow(slide, Inches(8.6), Inches(3.4), Inches(9.0), Inches(3.4))
        self._text(
            slide,
            Inches(0.5),
            Inches(6.35),
            Inches(12),
            Inches(0.5),
            "หลักการ: Staging ก่อนลงบัญชี/คลัง · Idempotency · แยก endpoint ตามประเภทธุรกิจ",
            size=13,
            color=TEAL,
            bold=True,
        )
        self._footer(slide, footer)

    def data_flow_vertical(self, title, footer, nodes, side_notes=None):
        """Vertical flow with optional side annotations."""
        slide = self._slide()
        self._bar(slide)
        self._text(
            slide, Inches(0.5), Inches(0.35), Inches(12), Inches(0.7), title, size=26, bold=True, color=NAVY
        )
        bx = Inches(1.2)
        bw = Inches(5.5)
        bh = Inches(0.75)
        y = Inches(1.15)
        gap = Inches(0.35)
        colors = [NAVY, TEAL, ORANGE, GREEN, PURPLE]
        for i, (label, sub) in enumerate(nodes):
            self._box(slide, bx, y, bw, bh, label, fill=colors[i % len(colors)], size=12)
            if sub:
                self._text(slide, bx + bw + Inches(0.2), y + Inches(0.08), Inches(5.5), Inches(0.6), sub, size=11, color=GRAY)
            if side_notes and i < len(side_notes):
                self._text(slide, Inches(7.5), y + Inches(0.05), Inches(5.2), Inches(0.65), side_notes[i], size=11, color=TEAL, bold=True)
            y_next = y + bh + gap
            if i < len(nodes) - 1:
                self._arrow(slide, bx + bw / 2, y + bh, bx + bw / 2, y_next)
            y = y_next
        self._footer(slide, footer)

    def comparison_slide(self, footer):
        slide = self._slide()
        self._bar(slide)
        self._text(
            slide,
            Inches(0.5),
            Inches(0.35),
            Inches(12),
            Inches(0.7),
            "เปรียบเทียบ API ทั้ง 3 เส้น",
            size=28,
            bold=True,
            color=NAVY,
        )
        self._table(
            slide,
            Inches(0.4),
            Inches(1.15),
            Inches(12.5),
            [
                ["", "Revenue", "Stock Issues", "Stock Requisitions"],
                ["Endpoint", "POST /revenue", "POST /stock-issues", "POST /stock-requisitions"],
                ["วัตถุประสงค์", "สรุปรายได้รายวัน", "ตัดจ่ายให้คนไข้", "เบิกเติมคลัง front"],
                ["HN / VN", "ห้ามส่ง (ตัดทิ้ง)", "จำเป็น (patient_use)", "ไม่เกี่ยวข้อง"],
                ["เมื่อ Post", "Invoice + Payment", "HISOUT → Consumption", "HISINT Internal Transfer"],
                ["ทิศทางสต็อก", "—", "UNIT → Consumption", "PHAR → UNIT"],
                ["คิว ERP", "HIS Revenue Queue", "HIS ตัดจ่ายสินค้า", "HIS เบิกเติมคลัง"],
            ],
            col_widths=[Inches(2.2), Inches(3.4), Inches(3.4), Inches(3.5)],
        )
        self._footer(slide, footer)

    def save(self, path):
        self.prs.save(path)
        print(path)


def build():
    d = HisApiDeck()
    footer = "VPK HIS ↔ ERP Integration  ·  Odoo 18  ·  vpk_his_api  ·  สำหรับทีมพัฒนา HIS"

    # ── Opening ──────────────────────────────────────────────────────
    d.title_slide(
        "HIS → ERP Integration",
        "REST API Overview  ·  Revenue · Stock Issue · Stock Requisition",
        "ร่างนำเสนอ  ·  26 สิงหาคม 2569  ·  โรงพยาบาลวชิระภูเก็ต",
    )

    d.content_slide(
        "Agenda",
        footer,
        [
            "1. ภาพรวมและหลักการออกแบบ",
            "2. สถาปัตยกรรมและเส้นทางข้อมูล (Data Flow)",
            "3. รูปแบบร่วม: Auth · Batch · Idempotency · Reversal",
            "4. API รายได้ — POST /revenue",
            "5. API ตัดจ่ายสินค้า — POST /stock-issues",
            "6. API เบิกเติมคลัง — POST /stock-requisitions",
            "7. Lookups · ข้อผิดพลาด · Checklist สำหรับทีม HIS",
        ],
    )

    # ── Design principles ────────────────────────────────────────────
    d.section_slide(1, "หลักการออกแบบ", "Design Principles")

    d.two_col_slide(
        "ทำไมต้องแยก 3 API?",
        footer,
        "ปัญหาที่แก้",
        [
            "รายได้ ตัดสต็อก และเบิกเติมคลัง เป็นคนละธุรกิจ",
            "ข้อมูล PII ต่างกัน — รายได้ห้าม HN/VN",
            "เอกสาร Odoo ต่างกัน — Invoice vs Picking",
            "ทีมตรวจสอบต่างกัน — บัญชี vs เภสัช",
        ],
        "หลักการหลัก",
        [
            "Staging ก่อน Post — ไม่ลงบัญชี/คลังทันที",
            "Idempotency: source_system + external_id",
            "Validate อัตโนมัติ → state ready | error",
            "Reversal สำหรับแก้ชุดที่ Post แล้ว",
            "ไม่สร้าง master data ใหม่ — รหัสต้องมีใน ERP",
        ],
    )

    d.architecture_slide(footer)

    d.flow_slide(
        "Lifecycle ชุดข้อมูล (Batch) — ใช้ร่วมทุก API",
        footer,
        [
            ("1. HIS ส่ง JSON", "POST + API Key"),
            ("2. Staging", "draft → validate"),
            ("3. Ready", "รอเจ้าหน้าที่"),
            ("4. Post", "ลงบัญชี/คลัง"),
            ("5. Posted", "เสร็จสิ้น"),
        ],
        colors=[TEAL, NAVY, ORANGE, GREEN, PURPLE],
    )

    d.content_slide(
        "รูปแบบร่วมที่ทีม HIS ต้องรู้",
        footer,
        [
            "Base URL: /vpk/api/v1/his  ·  Auth: X-Api-Key หรือ Authorization: Bearer",
            "ฟิลด์บังคับทุก endpoint: external_id, source_system, business_date",
            "Idempotency: คีย์ = source_system + external_id — ส่งซ้ำก่อน Post แทนที่ได้, หลัง Post ได้ HTTP 200 unchanged",
            "Reversal: batch_type=reversal + original_external_id ชี้ไปชุดที่ Post แล้ว",
            "สถานะ: draft → ready | error → posted (หรือ cancelled ก่อน Post)",
            "ตรวจสถานะ: GET /batches/{external_id}?source_system=...",
        ],
        note="Health check (ไม่ต้อง auth): GET /vpk/api/v1/his/health",
    )

    # ── Revenue ──────────────────────────────────────────────────────
    d.section_slide(2, "API รายได้", "POST /vpk/api/v1/his/revenue")

    d.data_flow_vertical(
        "Data Flow — Revenue (สรุปรายได้รายวัน)",
        footer,
        [
            ("HIS ปิดกะ / ปิดวัน", "รวมยอดขายและรับชำระ"),
            ("POST /revenue", "sales[] + payments[] + control_totals"),
            ("ERP Validate", "แมปสิทธิ์ · วิธีชำระ · บัญชี"),
            ("คิว HIS Revenue", "state = ready"),
            ("เจ้าหน้าที่ Post", "บัญชี/การเงิน"),
            ("Invoice + Payment", "ลูกหนี้สิทธิ์ / เงินสด"),
        ],
        side_notes=[
            "",
            "HTTP 202",
            "",
            "",
            "",
            "HTTP 200 ถ้าซ้ำ",
        ],
    )

    d.two_col_slide(
        "Revenue — ข้อมูลสำคัญ",
        footer,
        "ส่งอะไร",
        [
            "sales[] — ยอดขายตามสิทธิ (entitlement_code)",
            "payments[] — วิธีรับชำระ (cash, sso, uc…)",
            "ticket_external_id — ผูกหลายสิทธิในใบเดียว",
            "control_totals — คุมยอด sales/payments",
            "service_type: op | ip",
        ],
        "ห้ามส่ง (PII)",
        [
            "hn, vn, cid, patient_name",
            "national_id, birthdate",
            "→ ระบบตัดทิ้งทั้งหมด",
            "",
            "เอกสาร spec: VPK-HIS-SPEC-REV-001",
        ],
    )

    slide = d._slide()
    d._bar(slide)
    d._text(slide, Inches(0.5), Inches(0.35), Inches(12), Inches(0.7), "Revenue — ตัวอย่าง Payload (ย่อ)", size=28, bold=True, color=NAVY)
    d._code(
        slide,
        Inches(0.5),
        Inches(1.15),
        Inches(12.3),
        Inches(5.5),
        (
            'POST /vpk/api/v1/his/revenue\n\n'
            '{\n'
            '  "external_id": "HIS-REV-2026-08-26-SHIFT1",\n'
            '  "source_system": "front_his",\n'
            '  "business_date": "2026-08-26",\n'
            '  "shift": "1",\n'
            '  "sales": [\n'
            '    {"entitlement_code": "SELF_PAY", "service_type": "op",\n'
            '     "amount_total": 150.00, "ticket_external_id": "POS-001"},\n'
            '    {"entitlement_code": "UC", "service_type": "op",\n'
            '     "amount_total": 80.00, "ticket_external_id": "POS-001"}\n'
            '  ],\n'
            '  "payments": [\n'
            '    {"payment_method_code": "cash", "amount": 150.00},\n'
            '    {"payment_method_code": "ar_claim", "entitlement_code": "UC", "amount": 80.00}\n'
            '  ],\n'
            '  "control_totals": {"sales_total": 230.00, "payments_total": 230.00}\n'
            '}'
        ),
        size=12,
    )
    d._footer(slide, footer)

    # ── Stock Issues ─────────────────────────────────────────────────
    d.section_slide(3, "API ตัดจ่ายสินค้า", "POST /vpk/api/v1/his/stock-issues")

    d.data_flow_vertical(
        "Data Flow — Stock Issue (จ่ายยา/เวชภัณฑ์ให้คนไข้)",
        footer,
        [
            ("HIS บันทึกการจ่าย", "ตาม HN + VN + รหัสยา"),
            ("POST /stock-issues", "issues[] + warehouse + lot"),
            ("ERP Validate", "สินค้า · Lot · คลัง · HN/VN"),
            ("คิว HIS ตัดจ่าย", "state = ready"),
            ("เภสัช Post", "ตัดสต็อก"),
            ("Picking HISOUT", "UNIT → Consumption"),
        ],
    )

    d.two_col_slide(
        "Stock Issues — กฎสำคัญ",
        footer,
        "บังคับ",
        [
            "hn + vn เมื่อ reason = patient_use",
            "product_code ต้องมีใน ERP",
            "lot_name / lots[] เมื่อสินค้า tracked",
            "warehouse_code หรือ department_code",
            "control_totals.qty_total (แนะนำ)",
        ],
        "ไม่เกี่ยว",
        [
            "ไม่ใช่การเบิกเติมคลัง",
            "ชื่อ/CID ตัดทิ้ง (เก็บแค่ HN/VN)",
            "reason: patient_use | ward_use | expired",
            "",
            "Spec: VPK-HIS-SPEC-STK-001",
        ],
    )

    slide = d._slide()
    d._bar(slide)
    d._text(slide, Inches(0.5), Inches(0.35), Inches(12), Inches(0.7), "Stock Issues — ตัวอย่าง Payload (ย่อ)", size=28, bold=True, color=NAVY)
    d._code(
        slide,
        Inches(0.5),
        Inches(1.15),
        Inches(12.3),
        Inches(5.5),
        (
            'POST /vpk/api/v1/his/stock-issues\n\n'
            '{\n'
            '  "external_id": "HIS-STK-2026-08-26-OPD-1",\n'
            '  "source_system": "front_his",\n'
            '  "business_date": "2026-08-26",\n'
            '  "issues": [\n'
            '    {\n'
            '      "hn": "6500123", "vn": "6808260001",\n'
            '      "product_code": "00001", "item_type": "drug",\n'
            '      "warehouse_code": "UNIT", "qty": 10,\n'
            '      "reason": "patient_use",\n'
            '      "lots": [{"lot_name": "LOT-A", "qty": 6},\n'
            '               {"lot_name": "LOT-B", "qty": 4}]\n'
            '    }\n'
            '  ],\n'
            '  "control_totals": {"qty_total": 10}\n'
            '}'
        ),
        size=12,
    )
    d._footer(slide, footer)

    # ── Stock Requisitions ───────────────────────────────────────────
    d.section_slide(4, "API เบิกเติมคลัง", "POST /vpk/api/v1/his/stock-requisitions")

    d.data_flow_vertical(
        "Data Flow — Stock Requisition (เติมสต็อก front จากคลังหลัก)",
        footer,
        [
            ("หน่วยบริการขอเบิก", "OPD / ห้องยาชั้น 1"),
            ("POST /stock-requisitions", "lines[] + source/dest warehouse"),
            ("ERP Validate", "สินค้า · คลังต้นทาง/ปลายทาง"),
            ("คิว HIS เบิกเติม", "state = ready"),
            ("เภสัช Post", "โอนสต็อก"),
            ("Picking HISINT", "PHAR → UNIT"),
        ],
    )

    d.two_col_slide(
        "Stock Requisitions — กฎสำคัญ",
        footer,
        "ส่งอะไร",
        [
            "source_warehouse_code: PHAR (คลังหลัก)",
            "dest_warehouse_code: UNIT (คลัง front)",
            "dest_location_code / department_code",
            "lines[] — product_code + qty",
            "lot ไม่บังคับ — FEFO ตอน Post",
        ],
        "ต่างจาก stock-issues",
        [
            "ไม่มี HN/VN",
            "ไม่มี reason",
            "Internal Transfer ไม่ใช่ Consumption",
            "ทิศทาง PHAR → UNIT",
            "Spec: VPK-HIS-SPEC-STK-002",
        ],
    )

    slide = d._slide()
    d._bar(slide)
    d._text(slide, Inches(0.5), Inches(0.35), Inches(12), Inches(0.7), "Stock Requisitions — ตัวอย่าง Payload (ย่อ)", size=28, bold=True, color=NAVY)
    d._code(
        slide,
        Inches(0.5),
        Inches(1.15),
        Inches(12.3),
        Inches(5.5),
        (
            'POST /vpk/api/v1/his/stock-requisitions\n\n'
            '{\n'
            '  "external_id": "HIS-REQ-2026-08-26-OPD-1",\n'
            '  "source_system": "front_his",\n'
            '  "business_date": "2026-08-26",\n'
            '  "source_warehouse_code": "PHAR",\n'
            '  "dest_warehouse_code": "UNIT",\n'
            '  "department_code": "OPD",\n'
            '  "lines": [\n'
            '    {"product_code": "00001", "item_type": "drug", "qty": 100},\n'
            '    {"product_code": "00012", "item_type": "drug", "qty": 50,\n'
            '     "lot_name": "AMX-2508"}\n'
            '  ],\n'
            '  "control_totals": {"qty_total": 150}\n'
            '}'
        ),
        size=12,
    )
    d._footer(slide, footer)

    # ── Combined view ────────────────────────────────────────────────
    d.section_slide(5, "ภาพรวมทั้งระบบ", "Data Flow Summary")

    slide = d._slide()
    d._bar(slide)
    d._text(
        slide,
        Inches(0.5),
        Inches(0.35),
        Inches(12),
        Inches(0.7),
        "เส้นทางข้อมูลทั้ง 3 API ในมุมมองเดียว",
        size=28,
        bold=True,
        color=NAVY,
    )
    # HIS box
    d._box(slide, Inches(0.4), Inches(1.2), Inches(2.8), Inches(1.0), "ระบบ HIS", fill=NAVY, size=14)
    # Three API boxes
    apis = [
        (Inches(3.6), "POST /revenue", "รายได้\nInvoice", TEAL),
        (Inches(6.8), "POST /stock-issues", "ตัดจ่าย\nHISOUT", ORANGE),
        (Inches(10.0), "POST /stock-requisitions", "เบิกเติม\nHISINT", GREEN),
    ]
    for x, endpoint, desc, color in apis:
        d._arrow(slide, Inches(3.2), Inches(1.7), x, Inches(1.7))
        d._box(slide, x, Inches(1.2), Inches(2.8), Inches(1.0), endpoint, fill=color, size=11)
        d._box(slide, x, Inches(2.5), Inches(2.8), Inches(1.2), desc, fill=LIGHT, text_color=NAVY, size=11)
        d._arrow(slide, x + Inches(1.4), Inches(3.7), x + Inches(1.4), Inches(4.2))
    # Staging
    d._box(slide, Inches(3.6), Inches(4.2), Inches(9.2), Inches(0.9), "Staging Batch (Odoo) — Validate → Ready → Post", fill=NAVY, size=13)
    d._arrow(slide, Inches(6.2), Inches(5.1), Inches(6.2), Inches(5.6))
    # Outputs
    outs = [
        (Inches(1.0), "บัญชีรายได้\nInvoice / Payment", TEAL),
        (Inches(5.0), "Consumption\nลดสต็อก front", ORANGE),
        (Inches(9.0), "Internal Transfer\nPHAR → UNIT", GREEN),
    ]
    for x, label, color in outs:
        d._box(slide, x, Inches(5.6), Inches(3.5), Inches(1.0), label, fill=color, size=11)
    d._footer(slide, footer)

    d.comparison_slide(footer)

    d.content_slide(
        "Lookups & Monitoring",
        footer,
        [
            "GET /lookups/entitlements — รหัสสิทธิการรักษา (revenue)",
            "GET /lookups/payment-methods — วิธีรับชำระ (revenue)",
            "GET /lookups/warehouses — รหัสคลัง Odoo (stock)",
            "GET /lookups/products?codes=00001 — ตรวจ lot_required (stock)",
            "GET /batches/{external_id}?source_system=... — ติดตามสถานะชุดข้อมูล",
            "API Logs — ฝั่ง ERP (Settings → HIS Interface → API Logs)",
        ],
    )

    d.content_slide(
        "HTTP Response & ข้อผิดพลาดที่พบบ่อย",
        footer,
        [
            "202 — รับชุดข้อมูลแล้ว (draft/ready/error) · อ่าน state และ errors[]",
            "200 — คีย์ซ้ำของชุดที่ posted แล้ว (unchanged)",
            "400 — JSON ผิด / ขาดฟิลด์ / validation fail",
            "401 — API key ผิด · 503 — API ปิดหรือยังไม่ตั้งคีย์",
            "",
            "Revenue: sales/payments ไม่ balance · entitlement ไม่รู้จัก",
            "Stock: Unknown product · ขาด HN/VN · ขาด Lot · Unknown warehouse",
            "Requisition: source=dest · ไม่มี lines · Unknown product",
        ],
    )

    d.content_slide(
        "Checklist สำหรับทีมพัฒนา HIS",
        footer,
        [
            "□ ขอ URL + API Key จาก IT โรงพยาบาล (อย่า hard-code ใน source)",
            "□ ใช้ external_id ไม่ซ้ำต่อ source_system + business_date/shift",
            "□ Revenue: ส่ง control_totals ทุกครั้ง · ห้ามส่ง PII",
            "□ Stock Issues: ส่ง HN+VN ทุกบรรทัดจ่ายคนไข้ · ตรวจ lot_required ก่อนส่ง",
            "□ Requisition: แยก endpoint จากตัดจ่าย · ระบุ PHAR→UNIT",
            "□ หลัง POST เก็บ response batch_id/state · poll GET /batches ถ้าต้องการ",
            "□ Reversal: ใช้ external_id ใหม่ + original_external_id",
            "□ อ่าน spec เต็ม: REV-001 · STK-001 · STK-002 · payload-examples.md",
        ],
        note="เอกสาร spec Word อยู่ใน odoo18-custom-addons/vpk_his_api/doc/",
    )

    d.title_slide(
        "Thank You",
        "Q & A  ·  ติดต่อทีม ERP / IT โรงพยาบาลวชิระภูเก็ต",
        "VPK-HIS-API-Presentation  ·  26 สิงหาคม 2569",
    )

    d.save(OUT)


if __name__ == "__main__":
    build()
