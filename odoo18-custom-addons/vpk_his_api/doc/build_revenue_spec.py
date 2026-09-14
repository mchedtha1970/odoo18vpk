#!/usr/bin/env python3
# Copyright 2026 VPK
"""Generate HIS revenue ingest specification as a Word document."""

import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Prompt"
FONT_DIR = Path("/opt/odoo18vpk/docs/fonts/Prompt")
OUT = Path(__file__).resolve().parent / "VPK-HIS-SPEC-REV-001-การรับข้อมูลรายได้.docx"

NAVY = RGBColor(0x0B, 0x3D, 0x5C)
TEAL = RGBColor(0x0E, 0x6B, 0x6B)
HEADER_BG = "0B3D5C"
ALT_ROW = "F4F8FA"
CODE_BG = "F3F4F6"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
FONT_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
ODTTF_CT = "application/vnd.openxmlformats-officedocument.obfuscatedFont"


def _obfuscate_font(ttf_bytes, guid):
    key = guid.bytes_le
    out = bytearray(ttf_bytes)
    for i in range(min(32, len(out))):
        out[i] ^= key[i % 16]
    return bytes(out)


def embed_prompt_fonts(docx_path):
    """Embed Prompt Regular/Bold so Word shows the typeface without a local install."""
    faces = (
        ("embedRegular", FONT_DIR / "Prompt-Regular.ttf", "promptRegular.odttf"),
        ("embedBold", FONT_DIR / "Prompt-Bold.ttf", "promptBold.odttf"),
    )
    for _tag, ttf, _name in faces:
        if not ttf.is_file():
            return

    ET.register_namespace("w", W_NS)
    ET.register_namespace("r", R_NS)
    ET.register_namespace("", CT_NS)

    src = zipfile.ZipFile(docx_path, "r")
    parts = {info.filename: src.read(info.filename) for info in src.infolist()}
    src.close()

    font_files = {}
    embed_meta = []
    for i, (tag, ttf, name) in enumerate(faces, 1):
        guid = uuid.uuid4()
        font_key = "{" + str(guid).upper() + "}"
        rid = f"rIdFont{i}"
        font_files[f"word/fonts/{name}"] = _obfuscate_font(ttf.read_bytes(), guid)
        embed_meta.append((tag, rid, font_key, name))

    font_table = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:fonts xmlns:w="{W_NS}" xmlns:r="{R_NS}">'
        f'<w:font w:name="{FONT}">'
        f'<w:charset w:val="00"/>'
        f'<w:family w:val="swiss"/>'
        f'<w:pitch w:val="variable"/>'
    )
    for tag, rid, font_key, _name in embed_meta:
        font_table += f'<w:{tag} r:id="{rid}" w:fontKey="{font_key}"/>'
    font_table += "</w:font></w:fonts>"
    parts["word/fontTable.xml"] = font_table.encode("utf-8")

    rels = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{PKG_REL_NS}">'
    )
    for _tag, rid, _key, name in embed_meta:
        rels += (
            f'<Relationship Id="{rid}" Type="{FONT_REL_TYPE}" '
            f'Target="fonts/{name}"/>'
        )
    rels += "</Relationships>"
    parts["word/_rels/fontTable.xml.rels"] = rels.encode("utf-8")

    ct_root = ET.fromstring(parts["[Content_Types].xml"])
    if not any(
        el.get("Extension") == "odttf" for el in ct_root.findall(f"{{{CT_NS}}}Default")
    ):
        default = ET.SubElement(ct_root, f"{{{CT_NS}}}Default")
        default.set("Extension", "odttf")
        default.set("ContentType", ODTTF_CT)
    parts["[Content_Types].xml"] = ET.tostring(
        ct_root, encoding="utf-8", xml_declaration=True
    )

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in parts.items():
            dst.writestr(name, data)
        for name, data in font_files.items():
            dst.writestr(name, data)
    Path(docx_path).write_bytes(buf.getvalue())


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run_font(run, size=16, bold=False, italic=False, color=None, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:cs"), font)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), kwargs[edge].get("val", "single"))
            element.set(qn("w:sz"), kwargs[edge].get("sz", "4"))
            element.set(qn("w:color"), kwargs[edge].get("color", "BFCFD6"))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=12, color=NAVY)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


class SpecDoc:
    def __init__(self):
        self.doc = Document()
        self._setup_page()
        self._setup_styles()
        self.sec_no = 0

    def _setup_page(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("โรงพยาบาลวชิระภูเก็ต  ·  HIS Revenue Ingest Specification")
        set_run_font(r, size=12, color=NAVY)
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("VPK-HIS-SPEC-REV-001  ·  ร่าง  ·  หน้า ")
        set_run_font(r, size=12, color=NAVY)
        add_page_number(fp)

    def _setup_styles(self):
        styles_el = self.doc.styles.element
        doc_defaults = styles_el.find(qn("w:docDefaults"))
        if doc_defaults is None:
            doc_defaults = OxmlElement("w:docDefaults")
            styles_el.insert(0, doc_defaults)
        rpr_default = doc_defaults.find(qn("w:rPrDefault"))
        if rpr_default is None:
            rpr_default = OxmlElement("w:rPrDefault")
            doc_defaults.append(rpr_default)
        rPr = rpr_default.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            rpr_default.append(rPr)
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), FONT)
        styles = self.doc.styles
        for name, size, bold, space_before, space_after in (
            ("Normal", 16, False, 0, 6),
            ("Heading 1", 22, True, 16, 8),
            ("Heading 2", 18, True, 12, 6),
            ("Heading 3", 16, True, 10, 4),
        ):
            st = styles[name]
            st.font.name = FONT
            st.font.size = Pt(size)
            st.font.bold = bold
            st.font.color.rgb = NAVY if name != "Normal" else RGBColor(0x22, 0x22, 0x22)
            rPr = st.element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rFonts.set(qn(attr), FONT)
            pf = st.paragraph_format
            pf.space_before = Pt(space_before)
            pf.space_after = Pt(space_after)
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    def p(self, text, size=16, bold=False, align="left", space_after=6, color=None):
        para = self.doc.add_paragraph()
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "justify":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(0)
        run = para.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
        return para

    def h1(self, title):
        self.sec_no += 1
        self._h2_no = 0
        para = self.doc.add_heading(f"{self.sec_no}.  {title}", level=1)
        for run in para.runs:
            set_run_font(run, size=22, bold=True, color=NAVY)
        return self.sec_no

    def h2(self, title):
        self._h2_no = getattr(self, "_h2_no", 0) + 1
        self._h3_no = 0
        para = self.doc.add_heading(f"{self.sec_no}.{self._h2_no}  {title}", level=2)
        for run in para.runs:
            set_run_font(run, size=18, bold=True, color=TEAL)
        return f"{self.sec_no}.{self._h2_no}"

    def h3(self, title):
        self._h3_no = getattr(self, "_h3_no", 0) + 1
        para = self.doc.add_heading(
            f"{self.sec_no}.{self._h2_no}.{self._h3_no}  {title}", level=3
        )
        for run in para.runs:
            set_run_font(run, size=16, bold=True, color=NAVY)
        return f"{self.sec_no}.{self._h2_no}.{self._h3_no}"

    def bullets(self, items, numbered=False):
        for i, item in enumerate(items, 1):
            para = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            para.clear()
            prefix = f"{i}.  " if numbered else "•  "
            run = para.add_run(prefix + item)
            set_run_font(run, size=16)
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.left_indent = Cm(0.75)

    def table(self, headers, rows, col_widths=None):
        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        for i, h in enumerate(headers):
            cell = tbl.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(h)
            set_run_font(run, size=14, bold=True, color=RGBColor(255, 255, 255))
            shade_cell(cell, HEADER_BG)
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = tbl.rows[r_i + 1].cells[c_i]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(str(val) if val is not None else "")
                set_run_font(run, size=14)
                if r_i % 2 == 1:
                    shade_cell(cell, ALT_ROW)
        if col_widths:
            for row in tbl.rows:
                for i, w in enumerate(col_widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)
        return tbl

    def code(self, text):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.left_indent = Cm(0.3)
        run = para.add_run(text.strip("\n"))
        set_run_font(run, size=11, font="Consolas")
        pPr = para._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), CODE_BG)
        shd.set(qn("w:val"), "clear")
        pPr.append(shd)
        return para

    def note(self, text, label="หมายเหตุ"):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(8)
        r1 = para.add_run(f"{label}: ")
        set_run_font(r1, size=16, bold=True, color=TEAL)
        r2 = para.add_run(text)
        set_run_font(r2, size=16, italic=True)

    def save(self, path):
        self.doc.save(str(path))
        embed_prompt_fonts(path)


def build():
    d = SpecDoc()

    # Cover
    d.p("โรงพยาบาลวชิระภูเก็ต", size=20, bold=True, align="center", space_after=4, color=NAVY)
    d.p("ระบบบัญชี ERP (Odoo 18)", size=16, align="center", space_after=18, color=TEAL)
    d.p("ข้อกำหนดการรับข้อมูลรายได้จากระบบ HIS", size=28, bold=True, align="center", space_after=4, color=NAVY)
    d.p("HIS Revenue Ingest Specification", size=18, align="center", space_after=20, color=TEAL)
    d.p("เอกสารร่างสำหรับผู้พัฒนาระบบ HIS, ฝ่ายบัญชี และฝ่ายเทคโนโลยีสารสนเทศ", size=16, align="center", space_after=24)

    d.table(
        ["รายการ", "รายละเอียด"],
        [
            ["รหัสเอกสาร", "VPK-HIS-SPEC-REV-001"],
            ["เวอร์ชัน", "1.1  (ร่าง)"],
            ["วันที่", "25 สิงหาคม 2569"],
            ["สถานะ", "Draft — สำหรับทบทวนร่วมกับฝ่าย HIS"],
            ["ระบบต้นทาง", "Hospital Information System (Front HIS)"],
            ["ระบบปลายทาง", "Odoo 18 ERP  โรงพยาบาลวชิระภูเก็ต"],
            ["โมดูล", "vpk_his_api"],
            ["ขอบเขตฉบับนี้", "การรับสรุปรายได้รายวัน (Revenue ingest) และการผ่านบัญชี"],
            ["ผู้ใช้เอกสาร", "ทีม HIS / บัญชีลูกหนี้ / IT"],
        ],
        col_widths=[5.5, 11.5],
    )

    d.h1("วัตถุประสงค์และขอบเขต")
    d.h2("วัตถุประสงค์")
    d.p(
        "เอกสารนี้กำหนดวิธีที่ระบบ HIS ส่งสรุปรายได้รายวันเข้าสู่ระบบบัญชีของโรงพยาบาล "
        "เพื่อให้ลงรายได้ ตั้งลูกหนี้สิทธิ์ และรับเงินสด/โอน/บัตรที่เคาน์เตอร์ได้อย่างถูกต้อง "
        "โดยไม่ส่งข้อมูลส่วนบุคคลของผู้ป่วย และไม่ต้องลงบัญชีรายตัวคนไข้ในระบบ ERP",
        align="justify",
    )
    d.h2("อยู่ในขอบเขต")
    d.bullets(
        [
            "รับสรุปยอดขายและยอดรับเงินรายวัน (หรือรายกะ) จาก HIS ผ่าน REST API",
            "รองรับผู้ป่วยนอก (OPD) และผู้ป่วยใน (IPD) รวมถึงการจ่ายหลายสิทธิ์ในใบเดียว และการจ่ายบางส่วน",
            "จัดคิวตรวจสอบในระบบบัญชี แล้วให้เจ้าหน้าที่กดผ่านบัญชี (Post)",
            "ตั้งลูกหนี้สิทธิ์ (UC / SSO / CSMBS) และลูกหนี้คนไข้ IPD ที่จ่ายไม่ครบ ให้แสดงในเมนูลูกหนี้คงค้าง",
            "ยกเลิกชุดที่ผ่านบัญชีแล้วด้วย reversal batch",
        ]
    )
    d.h2("อยู่นอกขอบเขตฉบับนี้")
    d.bullets(
        [
            "การส่งข้อมูลตัวตนผู้ป่วย (HN, ชื่อ, เลขประจำตัวประชาชน) เข้า ERP — ห้ามส่ง",
            "การเคลียร์เงินจากกองทุนเข้าบัญชีธนาคารโรงพยาบาล (ทำในหน้าใบแจ้งหนี้ด้วยสมุด Bank หลังกองทุนโอน)",
            "การตัดสต็อกยา/เวชภัณฑ์รายวัน (ใช้ endpoint คนละตัว: POST /vpk/api/v1/his/stock-issues)",
            "การออกใบเสร็จรับเงินให้ผู้ป่วยที่หน้าเคาน์เตอร์ HIS — ยังทำที่ HIS ตามเดิม",
        ]
    )
    d.h2("หลักการออกแบบ")
    d.bullets(
        [
            "สรุปรายวัน ไม่ใช่รายการคนไข้: HIS ส่งยอดรวมตามสิทธิ์ / ประเภทบริการ / วิธีรับเงิน",
            "Staging ก่อนลงบัญชี: ข้อมูลเข้าคิวสถานะ draft → ready หรือ error จากนั้นเจ้าหน้าที่ Post",
            "Idempotency: คีย์ซ้ำคือ source_system + external_id หากชุดยังไม่ผ่านบัญชี การส่งซ้ำจะแทนที่บรรทัดเดิม",
            "สิทธิ์ไม่ใช่เงินสด: ประเภทรับเงิน sso / uc / csmbs / ar_claim ไม่สร้างใบรับเงิน แต่ตั้งลูกหนี้ค้างรับ",
            "เงินล่วงหน้า: วันรับมัดจำใช้ advance_in (เงินสดเข้าหนี้สินรับล่วงหน้า) วันมารักษาใช้ advance ตัดบิล และส่งเงินสดส่วนต่างถ้ามัดจำไม่พอ",
            "เงินจริงที่เคาน์เตอร์เท่านั้นที่ลงใบรับเงินสดวันนี้: cash / transfer / credit_card",
        ]
    )

    d.h1("สถาปัตยกรรมและขั้นตอนงาน")
    d.h2("ภาพรวมการไหลของข้อมูล")
    d.bullets(
        [
            "HIS สรุปยอดปิดกะ/ปิดวัน แล้ว POST JSON ไปที่ /vpk/api/v1/his/revenue",
            "ERP ตรวจรหัสสิทธิ์ วิธีรับเงิน และยอดคุม แล้วสร้างชุดข้อมูล (HIS Batch) ในคิว HIS Revenue Queue",
            "ถ้าแมปครบ ชุดอยู่ในสถานะ Ready — เจ้าหน้าที่บัญชีเปิดตรวจแล้วกด Post",
            "ระบบสร้างใบแจ้งหนี้ตามคู่ค้าสิทธิ์ (และใบรับเงินถ้าเป็นเงินจริง) จับคู่ชำระเฉพาะส่วนที่จ่ายเงินสด/โอน/บัตร",
            "ยอดสิทธิ์และยอดคนไข้ IPD ที่ยังไม่ครบ แสดงใน บัญชีลูกหนี้ → ลูกหนี้คงค้าง",
            "เมื่อกองทุนโอนเงินเข้าธนาคาร บัญชีเปิดใบแจ้งหนี้สิทธิ์นั้น แล้วลงทะเบียนชำระที่สมุด Bank",
        ],
        numbered=True,
    )
    d.h2("บทบาทระบบ")
    d.table(
        ["ระบบ", "หน้าที่"],
        [
            ["HIS", "สรุปยอดขายและรับเงิน ไม่ส่ง HN/ชื่อคนไข้ ส่ง ticket_external_id เป็นรหัสใบเสร็จภายใน HIS"],
            ["REST API (vpk_his_api)", "รับ JSON, ตัด PII, ตรวจ schema, จัดคิว, ตอบสถานะชุดข้อมูล"],
            ["เจ้าหน้าที่บัญชี", "ตรวจคิว HIS Revenue Queue, กด Post, พิมพ์ใบนำส่งเงิน, ติดตามลูกหนี้คงค้าง"],
            ["Odoo Accounting", "ใบแจ้งหนี้, ใบรับเงิน, ผังบัญชีลูกหนี้ตามสิทธิ์ OP/IP, aging"],
        ],
        col_widths=[5, 12],
    )

    d.h1("การเชื่อมต่อและการยืนยันตัวตน")
    d.h2("โปรโตคอล")
    d.table(
        ["รายการ", "ค่า"],
        [
            ["โปรโตคอล", "HTTPS, JSON (application/json; charset=utf-8)"],
            ["เวอร์ชัน API", "v1"],
            ["Base path", "/vpk/api/v1/his"],
            ["Charset", "UTF-8"],
            ["วันที่", "YYYY-MM-DD (คริสต์ศักราช)"],
            ["จำนวนเงิน", "ตัวเลขทศนิยม 2 ตำแหน่ง, สกุลเงิน THB"],
        ],
        col_widths=[5, 12],
    )
    d.h2("การยืนยันตัวตน")
    d.p("ทุก endpoint ยกเว้น health ต้องส่ง API key อย่างใดอย่างหนึ่ง:")
    d.code("X-Api-Key: <HIS_API_KEY>\n\nหรือ\n\nAuthorization: Bearer <HIS_API_KEY>")
    d.table(
        ["HTTP", "ความหมาย"],
        [
            ["401", "ไม่มีคีย์ หรือคีย์ไม่ถูกต้อง"],
            ["503", "ปิด HIS API จากตั้งค่า หรือยังไม่กำหนดคีย์บนเซิร์ฟเวอร์"],
        ],
        col_widths=[3, 14],
    )
    d.note("คีย์ตั้งที่ การตั้งค่า → การออกใบแจ้งหนี้ → HIS API Key ห้ามฝังคีย์ในเอกสารนี้")

    d.h1("รายการ API ที่เกี่ยวข้องกับรายได้")
    d.table(
        ["Method", "Path", "Auth", "หน้าที่"],
        [
            ["GET", "/vpk/api/v1/his/health", "ไม่ต้อง", "ตรวจว่าบริการทำงานและมีการตั้งคีย์"],
            ["POST", "/vpk/api/v1/his/revenue", "ต้อง", "ส่งสรุปรายได้ / reversal รายได้"],
            ["GET", "/vpk/api/v1/his/batches/{external_id}", "ต้อง", "สอบถามสถานะชุดข้อมูล (query source_system)"],
            ["GET", "/vpk/api/v1/his/lookups/entitlements", "ต้อง", "รายการรหัสสิทธิ์ที่ ERP รับได้"],
            ["GET", "/vpk/api/v1/his/lookups/payment-methods", "ต้อง", "รายการรหัสวิธีรับเงินที่ ERP รับได้"],
        ],
        col_widths=[2.2, 7.3, 2.0, 5.5],
    )
    d.p(
        "Query ของ GET batch: source_system=front_his (จำเป็นเมื่อมีหลายระบบใช้ external_id ซ้ำรูปแบบเดียวกัน)",
        align="justify",
    )

    d.h1("สัญญาข้อมูล POST /revenue")
    d.h2("ฟิลด์ระดับชุดข้อมูล (batch)")
    d.table(
        ["ฟิลด์", "ชนิด", "จำเป็น", "คำอธิบาย"],
        [
            ["external_id", "string", "ใช่", "รหัสชุดข้อมูลฝั่ง HIS ต้องไม่ซ้ำภายใน source_system เดียวกัน"],
            ["source_system", "string", "ใช่", "รหัสระบบต้นทาง เช่น front_his"],
            ["business_date", "date", "ใช่", "วันที่ทางธุรกิจ YYYY-MM-DD"],
            ["shift", "string", "ไม่", "รหัสกะ เช่น 1, 2, night"],
            ["currency", "string", "ไม่", "ค่าเริ่มต้น THB"],
            ["batch_type", "string", "ไม่", "revenue (ค่าเริ่มต้นของ endpoint นี้) หรือ reversal"],
            ["original_external_id", "string", "เมื่อยกเลิก", "external_id ของชุดที่ Post แล้ว เมื่อต้องการกลับรายการ"],
            ["sales", "array", "ใช่*", "บรรทัดขาย (*ชุดรายได้ต้องมี sales หรือ payments)"],
            ["payments", "array", "ใช่*", "บรรทัดรับเงิน"],
            ["control_totals.sales_total", "number", "แนะนำ", "ยอดคุมขาย ถ้าส่งมาต้องเท่ากับผลรวม sales"],
            ["control_totals.payments_total", "number", "แนะนำ", "ยอดคุมรับเงิน ถ้าส่งมาต้องเท่ากับผลรวม payments"],
        ],
        col_widths=[5.2, 2.2, 2.2, 7.4],
    )
    d.h2("บรรทัดขาย (sales[])")
    d.table(
        ["ฟิลด์", "ชนิด", "จำเป็น", "คำอธิบาย"],
        [
            ["line_external_id", "string", "แนะนำ", "รหัสบรรทัดฝั่ง HIS ใช้ตอนไล่ error"],
            ["ticket_external_id", "string", "แนะนำ", "รหัสใบเสร็จ/ใบรับบริการของคนไข้ใน HIS ใช้เมื่อหนึ่งใบมีหลายสิทธิ์"],
            ["entitlement_code", "string", "ใช่", "SELF_PAY | UC | SSO | CSMBS"],
            ["service_type", "string", "ใช่", "op หรือ ip (ถ้าไม่ส่ง ค่าเริ่มต้น op)"],
            ["item_type", "string", "ไม่", "service | drug | medical_supply | lab | other (ค่าเริ่มต้น service)"],
            ["department_code", "string", "ไม่", "รหัสหน่วยงาน เช่น OPD, IPD แสดงในคำอธิบายใบแจ้งหนี้"],
            ["qty", "number", "ไม่", "ค่าเริ่มต้น 1"],
            ["amount_total", "number", "ใช่", "ยอดรวมบรรทัด ห้ามเป็นศูนย์"],
            ["amount_untaxed", "number", "ไม่", "ถ้าไม่ส่ง จะคำนวณจาก amount_total − amount_tax"],
            ["amount_tax", "number", "ไม่", "ค่าเริ่มต้น 0"],
            ["income_account_code", "string", "ไม่", "ระบุเมื่อต้องการ override บัญชีรายได้ของแมปสิทธิ์"],
        ],
        col_widths=[5.2, 2.2, 2.2, 7.4],
    )
    d.h2("บรรทัดรับเงิน (payments[])")
    d.table(
        ["ฟิลด์", "ชนิด", "จำเป็น", "คำอธิบาย"],
        [
            ["line_external_id", "string", "แนะนำ", "รหัสบรรทัดฝั่ง HIS"],
            ["ticket_external_id", "string", "แนะนำ", "ต้องตรงกับใบขายเมื่อหนึ่งใบมีหลายสิทธิ์หรือจ่ายบางส่วน"],
            ["payment_method_code", "string", "ใช่", "ดูตารางวิธีรับเงิน"],
            ["entitlement_code", "string", "ตามประเภท", "จำเป็นเมื่อเป็นสิทธิ์ (sso/uc/csmbs ระบบอนุมานให้ได้) และเมื่อ ar_claim"],
            ["service_type", "string", "ไม่", "op หรือ ip ถ้าว่างระบบอนุมานจากบรรทัดขายในใบเดียวกัน"],
            ["amount", "number", "ใช่", "ยอดรับ ห้ามเป็นศูนย์"],
            ["journal_code", "string", "ไม่", "ระบุสมุดรายวันเมื่อต้องการ override แมป"],
        ],
        col_widths=[5.2, 2.2, 2.2, 7.4],
    )
    d.h2("ยอดคุม (control_totals)")
    d.p(
        "ถ้าส่งยอดคุมมา ระบบจะเทียบกับผลรวมบรรทัดจริง ทศนิยม 2 ตำแหน่ง หากไม่ตรง ชุดข้อมูลเข้าสถานะ error "
        "ไม่บังคับให้ยอดขายเท่ากับยอดรับเงินในทางบัญชี (เช่น คนไข้จ่ายไม่ครบ) แต่เพื่อปิดยอดรายวันให้สมดุล "
        "HIS ควรส่ง ar_claim สำหรับส่วนที่ค้าง และให้ payments_total เท่ากับ sales_total",
        align="justify",
    )
    d.h2("ข้อห้าม — ข้อมูลส่วนบุคคล")
    d.p("ห้ามส่งฟิลด์ต่อไปนี้ ระบบจะตัดทิ้งอัตโนมัติถ้ามีอยู่ใน JSON ทุกระดับ:")
    d.bullets(
        [
            "hn, an, vn, cid, pid, national_id",
            "patient_id, patient_hn, patient_name, firstname, lastname",
            "birthdate, dob, visit_id, encounter_id",
        ]
    )
    d.p(
        "ใช้ ticket_external_id เป็นรหัสอ้างอิงใบเสร็จภายใน HIS เท่านั้น ไม่ใช้เป็น HN",
        align="justify",
    )

    d.h1("รหัสมาตรฐานที่ HIS ต้องใช้")
    d.h2("รหัสสิทธิ์ (entitlement_code)")
    d.table(
        ["รหัส", "ความหมาย", "OP / IP", "ผลทางบัญชีเมื่อ Post"],
        [
            ["SELF_PAY", "ชำระเงินเอง / ส่วนคนไข้", "แยกแมป OP และ IP", "ใบแจ้งหนี้คู่ค้าชำระเอง จ่ายเงินจริงแล้วปิด / ค้างส่วนที่ยังไม่จ่าย (IP)"],
            ["UC", "บัตรทอง / UC", "แยก OP และ IP", "ตั้งลูกหนี้กองทุน UC ไม่ลงเงิน"],
            ["SSO", "ประกันสังคม", "แยก OP และ IP", "ตั้งลูกหนี้กองทุน SSO ไม่ลงเงิน"],
            ["CSMBS", "เบิกจ่ายตรงกรมบัญชีกลาง", "แยก OP และ IP", "ตั้งลูกหนี้ CSMBS ไม่ลงเงิน"],
        ],
        col_widths=[3.2, 4.5, 3.3, 6.0],
    )
    d.h2("ประเภทบริการ (service_type)")
    d.table(
        ["รหัส", "ความหมาย", "หมายเหตุ"],
        [
            ["op", "ผู้ป่วยนอก", "ค่าเริ่มต้นถ้าไม่ส่งในบรรทัดขาย"],
            ["ip", "ผู้ป่วยใน", "ต้องส่งเมื่อเป็นเคส IPD เพื่อไปลูกหนี้ IP และบัญชีรายได้ IP"],
        ],
        col_widths=[3, 4, 10],
    )
    d.h2("ประเภทค่าใช้จ่าย (item_type)")
    d.table(
        ["รหัส", "ความหมาย"],
        [
            ["service", "ค่าบริการ"],
            ["drug", "ยา"],
            ["medical_supply", "เวชภัณฑ์"],
            ["lab", "แล็บ"],
            ["other", "อื่น ๆ"],
        ],
        col_widths=[5, 12],
    )
    d.h2("วิธีรับเงิน (payment_method_code)")
    d.table(
        ["รหัส", "ชื่อ", "ลงใบรับเงิน", "กฎ"],
        [
            ["cash", "เงินสด", "ใช่ (สมุด HIS เงินสดเคาน์เตอร์)", "จับคู่ใบแจ้งหนี้ชำระเองของใบนั้น"],
            ["transfer", "โอนเงิน", "ใช่ (สมุดธนาคาร)", "เช่นเดียวกัน"],
            ["credit_card", "บัตรเครดิต", "ใช่ (สมุดธนาคาร)", "เช่นเดียวกัน"],
            ["advance_in", "รับเงินล่วงหน้า", "ลงเงินสดเข้าหนี้สินรับล่วงหน้า", "ใช้วันที่รับมัดจำ ไม่ส่ง sales ไม่ลงรายได้"],
            ["advance", "ตัดเงินล่วงหน้า", "ใช่ (สมุด HADV ไม่ใช่เงินสดวันนี้)", "ตัดยอดมัดจำเข้าใบแจ้งหนี้ชำระเอง; สำรอง prepaid / deposit"],
            ["sso", "สิทธิ์ประกันสังคม", "ไม่", "อนุมาน entitlement SSO ตั้งลูกหนี้กองทุน"],
            ["uc", "สิทธิ์ UC / บัตรทอง", "ไม่", "อนุมาน entitlement UC"],
            ["csmbs", "สิทธิ์เบิกจ่ายตรง", "ไม่", "อนุมาน entitlement CSMBS"],
            ["ar_claim", "ลูกหนี้สิทธิ์ / ยอดค้าง", "ไม่", "ต้องส่ง entitlement_code เช่น SELF_PAY เมื่อเป็นยอดค้างคนไข้"],
        ],
        col_widths=[3.0, 4.2, 4.3, 5.5],
    )
    d.note(
        "เงินสดที่ไม่มี entitlement_code ระบบจะจับเป็น SELF_PAY หากใบนั้นมีบรรทัดขาย SELF_PAY "
        "และอนุมาน service_type จากบรรทัดขายใน ticket เดียวกัน (ถ้าเป็น IPD จะไปคู่ค้าชำระเอง IP)"
    )

    d.h1("กฎการผ่านบัญชี")
    d.h2("การจัดกลุ่มใบแจ้งหนี้")
    d.p(
        "เมื่อกด Post ระบบสร้างใบแจ้งหนี้แยกตาม คู่ค้าสิทธิ์ + ประเภทบริการ + แมปสิทธิ์ + ticket_external_id "
        "ดังนั้นสองใบเสร็จสิทธิ์เดียวกันในวันเดียวกันจะได้สองใบแจ้งหนี้ ไม่รวมยอด",
        align="justify",
    )
    d.h2("ผลลัพธ์ตามประเภทรับเงิน")
    d.table(
        ["สถานการณ์", "เอกสารที่สร้าง", "ลูกหนี้คงค้าง"],
        [
            ["เงินสด/โอน/บัตร ครบยอดส่วนคนไข้", "ใบแจ้งหนี้ชำระเอง + ใบรับเงิน จับคู่ปิด", "ไม่มี (หรือไม่เหลือ)"],
            ["สิทธิ์ sso / uc / csmbs", "ใบแจ้งหนี้คู่ค้ากองทุน เปิดค้าง", "มียอดกองทุนเต็ม"],
            ["คนไข้ IPD จ่ายบางส่วน", "ใบแจ้งหนี้ชำระเอง IP + ใบรับเงินบางส่วน", "ส่วนต่างคนไข้คงค้าง"],
            ["รับมัดจำล่วงหน้า (advance_in)", "เดบิตเงินสด / เครดิต 2103010103.101", "ไม่มีใบแจ้งหนี้"],
            ["ตัดมัดจำ + จ่ายเพิ่มถ้าไม่พอ", "ใบแจ้งหนี้ชำระเอง + ตัดล่วงหน้า + ใบรับเงินส่วนต่าง", "ปิดครบถ้าตัด+จ่ายเพิ่มเท่ากับยอดขาย"],
            ["ar_claim ส่วนค้างคนไข้", "ไม่ลงเงิน ใช้ให้ยอดคุมสมดุล", "ยอดนี้อยู่ในใบชำระเองที่ยังเปิด"],
        ],
        col_widths=[5.2, 6.5, 5.3],
    )
    d.h2("คู่ค้าและบัญชีลูกหนี้หลัก")
    d.table(
        ["สิทธิ์", "บริการ", "คู่ค้า", "บัญชีลูกหนี้"],
        [
            ["SELF_PAY", "OP", "HIS Payer - ชำระเงินเอง OP", "1102050102.106"],
            ["SELF_PAY", "IP", "HIS Payer - ชำระเงินเอง IP", "1102050102.107"],
            ["UC", "OP", "HIS Payer - UC OP", "1102050101.201"],
            ["UC", "IP", "HIS Payer - UC IP", "1102050101.202"],
            ["SSO", "OP", "HIS Payer - ประกันสังคม OP", "1102050101.301"],
            ["SSO", "IP", "HIS Payer - ประกันสังคม IP", "1102050101.302"],
            ["CSMBS", "OP", "HIS Payer - CSMBS OP", "1102050101.401"],
            ["CSMBS", "IP", "HIS Payer - CSMBS IP", "1102050101.402"],
        ],
        col_widths=[3.2, 2.3, 6.8, 4.7],
    )
    d.h2("เมนูติดตามลูกหนี้")
    d.bullets(
        [
            "บัญชีลูกหนี้ → HIS Revenue Queue — คิวชุดข้อมูลรอตรวจ/ผ่านบัญชี",
            "บัญชีลูกหนี้ → ลูกหนี้คงค้าง — ใบแจ้งหนี้เปิดของสิทธิ์ และของคนไข้ IPD กรองได้เป็น สิทธิ์ / คนไข้ IPD",
            "บัญชีลูกหนี้ → ใบนำส่งเงิน — สรุปยอดเงินจริงและยอดสิทธิ์รายวัน",
            "เมื่อกองทุนโอน: เปิดใบแจ้งหนี้สิทธิ์ → ลงทะเบียนชำระที่สมุด Bank (ห้าม Register Payment ไปบัญชี outstanding ที่ดึงลูกหนี้ออกจาก aging ก่อนเงินเข้าจริง)",
        ]
    )

    d.h1("สถานะชุดข้อมูลและการทำงานของเจ้าหน้าที่")
    d.table(
        ["state", "ความหมาย", "HIS ทำอะไรได้", "บัญชีทำอะไร"],
        [
            ["draft", "รับเข้าแล้วยังตรวจไม่จบ (ช่วงสั้น)", "ส่งซ้ำได้", "รอ"],
            ["ready", "แมปครบ ยอดคุมตรง พร้อมผ่านบัญชี", "ส่งซ้ำได้ (แทนที่บรรทัด)", "ตรวจแล้วกด Post"],
            ["error", "แมปไม่ครบหรือยอดคุมไม่ตรง", "ส่งซ้ำชุดเดิมหลังแก้ payload", "แก้แมปแล้วกด Validate"],
            ["posted", "ลงบัญชีแล้ว", "ส่งคีย์เดิมได้แต่ไม่เปลี่ยนข้อมูล (HTTP 200)", "ดูใบแจ้งหนี้/ใบรับเงิน; ถ้าผิดใช้ reversal"],
            ["cancelled", "ยกเลิกชุดที่ยังไม่ผ่านบัญชี", "ต้องใช้ external_id ใหม่", "ไม่ Post"],
        ],
        col_widths=[2.6, 4.6, 5.2, 4.6],
    )
    d.h2("Idempotency")
    d.bullets(
        [
            "คีย์ = source_system + external_id ในบริษัทเดียวกัน",
            "ชุดยังไม่ posted: POST ซ้ำจะลบบรรทัดเดิมแล้วใส่ชุดใหม่ แล้ว validate ใหม่ (action = updated)",
            "ชุด posted แล้ว: POST คีย์เดิมคืนชุดเดิม ไม่สร้างซ้ำ (action = unchanged, HTTP 200)",
            "ชุด cancelled: ห้ามใช้คีย์เดิม ต้องขึ้น external_id ใหม่",
        ]
    )
    d.h2("การกลับรายการ (reversal)")
    d.p("เมื่อชุดรายได้ถูก Post ไปแล้วและต้องยกเลิก HIS ส่งชุดใหม่ คนละ external_id:")
    d.code(
        '{\n'
        '  "external_id": "HIS-REV-2026-08-23-IPD-1-REV",\n'
        '  "source_system": "front_his",\n'
        '  "business_date": "2026-08-23",\n'
        '  "batch_type": "reversal",\n'
        '  "original_external_id": "HIS-REV-2026-08-23-IPD-1"\n'
        "}"
    )
    d.p(
        "ถ้ามี original_external_id ระบบถือเป็น reversal โดยอัตโนมัติ ชุดต้นทางต้องเป็น posted และต้องไม่ใช่ reversal ซ้อน",
        align="justify",
    )

    d.h1("รหัสตอบกลับ HTTP และรูปแบบผลลัพธ์")
    d.h2("รหัส HTTP ของ POST /revenue")
    d.table(
        ["HTTP", "เมื่อไร", "ความหมายต่อ HIS"],
        [
            ["202", "สร้างหรืออัปเดตชุดที่ยังไม่ posted (ready หรือ error)", "รับแล้ว อยู่ระหว่างคิว/รอ Post; อ่าน state และ errors"],
            ["200", "คีย์ซ้ำของชุดที่ posted แล้ว", "ไม่ทำซ้ำ คืนสถานะชุดเดิม"],
            ["400", "JSON ผิด, ขาดฟิลด์จำเป็น, ค่าไม่อยู่ในโดเมน", "แก้ payload แล้วส่งใหม่"],
            ["401", "คีย์ผิดหรือไม่ส่ง", "ตรวจ header"],
            ["503", "API ถูกปิด หรือยังไม่ตั้งคีย์บนเซิร์ฟเวอร์", "ติดต่อ IT โรงพยาบาล"],
            ["500", "ข้อผิดพลาดภายใน", "ส่งซ้ำได้หลังระบบกลับมา; แจ้ง IT หากซ้ำ"],
        ],
        col_widths=[2.2, 7.5, 7.3],
    )
    d.h2("โครง response สำเร็จ")
    d.code(
        '{\n'
        '  "ok": true,\n'
        '  "batch_id": 101,\n'
        '  "name": "HIS/2026/00101",\n'
        '  "external_id": "HIS-REV-2026-08-23-IPD-1",\n'
        '  "source_system": "front_his",\n'
        '  "batch_type": "revenue",\n'
        '  "business_date": "2026-08-23",\n'
        '  "state": "ready",\n'
        '  "sales_total": 95000.0,\n'
        '  "payments_total": 95000.0,\n'
        '  "invoice_ids": [],\n'
        '  "payment_ids": [],\n'
        '  "picking_ids": [],\n'
        '  "errors": [],\n'
        '  "action": "created"\n'
        "}"
    )
    d.p("action มีค่า created | updated | unchanged เมื่อชุด error ฟิลด์ errors เป็นรายการ {line_external_id, message}")
    d.h2("ข้อผิดพลาดที่พบบ่อย")
    d.table(
        ["ข้อความโดยสังเขป", "สาเหตุและการแก้"],
        [
            ["external_id is required", "ไม่ได้ส่งรหัสชุดข้อมูล"],
            ["business_date is required", "วันที่ต้องเป็น YYYY-MM-DD"],
            ["sales[].entitlement_code is required", "ทุกบรรทัดขายต้องมีสิทธิ์"],
            ["sales[].service_type must be op or ip", "ค่าอื่นไม่รับ"],
            ["No entitlement mapping", "รหัสสิทธิ์ไม่อยู่ในแมป หรือยังไม่ผูกบัญชีรายได้"],
            ["No payment method mapping", "payment_method_code ไม่ตรงรหัสมาตรฐาน"],
            ["Entitlement payment requires entitlement_code", "ar_claim ต้องระบุสิทธิ์ เช่น SELF_PAY"],
            ["Sales/Payments total does not match control", "ยอดคุมไม่เท่าผลรวมบรรทัด"],
            ["Batch is cancelled; use a new external_id", "ห้ามใช้คีย์ชุดที่ยกเลิกแล้ว"],
        ],
        col_widths=[7.5, 9.5],
    )

    d.h1("ข้อกำหนดตามสถานการณ์ธุรกิจ")
    d.h2("ผู้ป่วยนอก จ่ายสิทธิ์เดียวทั้งใบ")
    d.p(
        "ส่ง sales เป็นสิทธิ์นั้นทั้งยอด และ payments เป็น sso หรือ uc หรือ csmbs ทั้งยอด "
        "ไม่ต้องมีเงินสด ระบบจะตั้งลูกหนี้กองทุนทั้งจำนวน",
        align="justify",
    )
    d.h2("ผู้ป่วยนอก หลายสิทธิ์ในใบเดียว")
    d.p(
        "ใส่ ticket_external_id ชุดเดียวกันทุกบรรทัดขายและรับเงินของใบนั้น "
        "แยกยอดขายตาม entitlement_code และแยกยอดรับตาม payment_method_code ที่คู่กัน "
        "ส่วนคนไข้ใช้ cash/transfer/credit_card คู่ SELF_PAY ส่วนกองทุนใช้ sso/uc/csmbs",
        align="justify",
    )
    d.h2("ผู้ป่วยใน จ่ายตามสิทธิ์และจ่ายบางส่วน")
    d.bullets(
        [
            "บรรทัดขายทั้งหมด service_type = ip และใช้ ticket_external_id เดียวกัน",
            "ส่วนที่กองทุนรับผิดชอบ: ขายด้วย UC/SSO/CSMBS และรับด้วย uc/sso/csmbs ตามลำดับ",
            "ส่วนคนไข้: ขายด้วย SELF_PAY ยอดเต็มที่คนไข้ต้องออก",
            "เงินที่เก็บได้จริง: cash/transfer/credit_card เท่าจำนวนที่รับมา (ไม่ต้องใส่ entitlement_code ก็ได้)",
            "ส่วนที่คนไข้ยังไม่จ่าย: ar_claim + entitlement_code SELF_PAY ให้ยอดรับรวมเท่ายอดขาย",
            "ผล: ใบแจ้งหนี้ชำระเอง IP คงเหลือส่วนต่างคนไข้ และใบแจ้งหนี้กองทุน IP ค้างเต็ม ขึ้นในลูกหนี้คงค้างทั้งคู่",
        ]
    )
    d.h2("เงินล่วงหน้า แล้วมาใช้บริการ (ตัดมัดจำ / จ่ายเพิ่มถ้าไม่พอ)")
    d.bullets(
        [
            "วันรับมัดจำ: ส่งชุดรายได้ที่มีเฉพาะ payments รหัส advance_in ห้ามใส่บรรทัดขาย (ยังไม่เกิดรายได้) เงินสดวันนี้เพิ่ม หนี้สินรับล่วงหน้า (บัญชี 2103010103.101) เพิ่ม",
            "วันมารักษา: ส่งยอดขายส่วนคนไข้ (SELF_PAY) ตามบิลจริง",
            "ตัดมัดจำ: payments รหัส advance (หรือ prepaid) เท่ากับจำนวนที่ใช้จากเงินล่วงหน้า ไม่นับเป็นเงินสดในใบนำส่งเงินวันรักษา",
            "ถ้ามัดจำไม่พอ: ส่ง cash / transfer / credit_card เพิ่มเป็นส่วนต่างที่เก็บที่เคาน์เตอร์วันนี้",
            "ถ้ามัดจำพอทั้งบิล: ส่งแค่ advance เท่ากับยอดขาย ไม่ต้องมีเงินสด",
            "ยอดคุมวันที่รักษา: sales_total = advance + เงินเพิ่ม (+ สิทธิ์กองทุนถ้ามี)",
            "HIS เป็นระบบเก็บว่าคนไข้คนใดมีมัดจำเหลือเท่าใด ERP รับยอดรวมรายวัน ไม่รับ HN",
        ]
    )
    d.h2("ข้อกำหนดยอดรวมรายวัน")
    d.bullets(
        [
            "แนะนำ 1 ชุดข้อมูลต่อกะ หรือต่อวัน ตามที่ตกลงปิดยอดกับบัญชี",
            "external_id ควรมีวันที่และกะประกอบ เช่น HIS-REV-2026-08-25-SHIFT1",
            "control_totals ควรส่งเสมอเพื่อกันยอดตกหล่นระหว่าง HIS กับ ERP",
            "ไม่รวมใบที่ยกเลิกใน HIS ก่อนปิดยอด หากยกเลิกหลัง ERP Post แล้ว ให้ใช้ reversal",
        ]
    )

    d.h1("ตัวอย่าง Payload")
    d.h2("ตัวอย่าง A — OPD สิทธิ์ประกันสังคมทั้งใบ")
    d.code(
        """
POST /vpk/api/v1/his/revenue

{
  "external_id": "HIS-REV-2026-08-25-SSO-OP",
  "source_system": "front_his",
  "business_date": "2026-08-25",
  "shift": "1",
  "currency": "THB",
  "sales": [
    {
      "line_external_id": "S1",
      "ticket_external_id": "POS-OPD-0001",
      "entitlement_code": "SSO",
      "service_type": "op",
      "item_type": "service",
      "department_code": "OPD",
      "amount_total": 215.00
    }
  ],
  "payments": [
    {
      "line_external_id": "P1",
      "ticket_external_id": "POS-OPD-0001",
      "payment_method_code": "sso",
      "amount": 215.00
    }
  ],
  "control_totals": {
    "sales_total": 215.00,
    "payments_total": 215.00
  }
}
""".strip()
    )
    d.p("ผลหลัง Post: ใบแจ้งหนี้คู่ค้าประกันสังคม OP 215 บาท ค้างรับทั้งจำนวน ไม่มีใบรับเงินสด")

    d.h2("ตัวอย่าง B — OPD หลายสิทธิ์ในใบเดียว")
    d.code(
        """
{
  "external_id": "HIS-REV-2026-08-25-MIX-OP",
  "source_system": "front_his",
  "business_date": "2026-08-25",
  "shift": "1",
  "currency": "THB",
  "sales": [
    {
      "line_external_id": "T1-CASH",
      "ticket_external_id": "POS-OPD-0001",
      "entitlement_code": "SELF_PAY",
      "service_type": "op",
      "item_type": "service",
      "department_code": "OPD",
      "amount_total": 50.00
    },
    {
      "line_external_id": "T1-SSO",
      "ticket_external_id": "POS-OPD-0001",
      "entitlement_code": "SSO",
      "service_type": "op",
      "item_type": "drug",
      "department_code": "OPD",
      "amount_total": 200.00
    },
    {
      "line_external_id": "T1-UC",
      "ticket_external_id": "POS-OPD-0001",
      "entitlement_code": "UC",
      "service_type": "op",
      "item_type": "lab",
      "department_code": "OPD",
      "amount_total": 80.00
    }
  ],
  "payments": [
    {
      "line_external_id": "T1-PAY-CASH",
      "ticket_external_id": "POS-OPD-0001",
      "payment_method_code": "cash",
      "entitlement_code": "SELF_PAY",
      "amount": 50.00
    },
    {
      "line_external_id": "T1-PAY-SSO",
      "ticket_external_id": "POS-OPD-0001",
      "payment_method_code": "sso",
      "amount": 200.00
    },
    {
      "line_external_id": "T1-PAY-UC",
      "ticket_external_id": "POS-OPD-0001",
      "payment_method_code": "uc",
      "amount": 80.00
    }
  ],
  "control_totals": {
    "sales_total": 330.00,
    "payments_total": 330.00
  }
}
""".strip()
    )
    d.table(
        ["ส่วน", "เอกสาร", "ลูกหนี้คงค้าง"],
        [
            ["เงินสด 50", "ใบแจ้งหนี้ชำระเอง OP + ใบรับเงิน", "ไม่มี"],
            ["SSO 200", "ใบแจ้งหนี้ประกันสังคม OP", "200"],
            ["UC 80", "ใบแจ้งหนี้ UC OP", "80"],
        ],
        col_widths=[4, 8, 5],
    )

    d.h2("ตัวอย่าง C — IPD จ่ายตามสิทธิ์และจ่ายบางส่วน")
    d.p(
        "คนไข้ในมียอด 95,000 บาท กองทุน SSO รับ 80,000 คนไข้ต้องออก 15,000 จ่ายสดได้ 5,000 ค้าง 10,000"
    )
    d.code(
        """
{
  "external_id": "HIS-REV-2026-08-25-IPD-1",
  "source_system": "front_his",
  "business_date": "2026-08-25",
  "shift": "1",
  "currency": "THB",
  "sales": [
    {
      "line_external_id": "IPD-SELF",
      "ticket_external_id": "POS-IPD-0001",
      "entitlement_code": "SELF_PAY",
      "service_type": "ip",
      "item_type": "service",
      "department_code": "IPD",
      "amount_total": 15000.00
    },
    {
      "line_external_id": "IPD-SSO",
      "ticket_external_id": "POS-IPD-0001",
      "entitlement_code": "SSO",
      "service_type": "ip",
      "item_type": "drug",
      "department_code": "IPD",
      "amount_total": 80000.00
    }
  ],
  "payments": [
    {
      "line_external_id": "IPD-PAY-CASH",
      "ticket_external_id": "POS-IPD-0001",
      "payment_method_code": "cash",
      "amount": 5000.00
    },
    {
      "line_external_id": "IPD-PAY-SSO",
      "ticket_external_id": "POS-IPD-0001",
      "payment_method_code": "sso",
      "amount": 80000.00
    },
    {
      "line_external_id": "IPD-PAY-AR",
      "ticket_external_id": "POS-IPD-0001",
      "payment_method_code": "ar_claim",
      "entitlement_code": "SELF_PAY",
      "amount": 10000.00
    }
  ],
  "control_totals": {
    "sales_total": 95000.00,
    "payments_total": 95000.00
  }
}
""".strip()
    )
    d.table(
        ["ส่วน", "เอกสาร", "ลูกหนี้คงค้าง"],
        [
            ["เงินสด 5,000 จากส่วนคนไข้ 15,000", "ใบแจ้งหนี้ชำระเอง IP + ใบรับเงิน", "คนไข้ 10,000 (partial)"],
            ["สิทธิ์ SSO 80,000", "ใบแจ้งหนี้ประกันสังคม IP", "กองทุน 80,000"],
            ["ar_claim คนไข้ 10,000", "ไม่ลงเงิน", "อยู่ในใบชำระเอง IP"],
        ],
        col_widths=[6.2, 6.3, 4.5],
    )

    d.h2("ตัวอย่าง D — ตัดเงินล่วงหน้าแล้วจ่ายเพิ่ม")
    d.p(
        "คนไข้วางมัดจำไว้ 10,000 บาทก่อนหน้า (ส่ง advance_in คนละชุด) "
        "วันมารักษายอดชำระเอง 15,000 ตัดมัดจำ 10,000 จ่ายสดเพิ่ม 5,000"
    )
    d.code(
        """
{
  "external_id": "HIS-REV-2026-08-25-ADV-1",
  "source_system": "front_his",
  "business_date": "2026-08-25",
  "sales": [
    {
      "line_external_id": "ADV-SELF",
      "ticket_external_id": "POS-OPD-ADV-1",
      "entitlement_code": "SELF_PAY",
      "service_type": "op",
      "item_type": "service",
      "amount_total": 15000.00
    }
  ],
  "payments": [
    {
      "line_external_id": "ADV-APPLY",
      "ticket_external_id": "POS-OPD-ADV-1",
      "payment_method_code": "advance",
      "amount": 10000.00
    },
    {
      "line_external_id": "ADV-CASH",
      "ticket_external_id": "POS-OPD-ADV-1",
      "payment_method_code": "cash",
      "amount": 5000.00
    }
  ],
  "control_totals": {
    "sales_total": 15000.00,
    "payments_total": 15000.00
  }
}
""".strip()
    )
    d.p("ผลหลัง Post: ใบแจ้งหนี้ชำระเองปิดครบ ใบรับเงินตัดล่วงหน้า 10,000 (ไม่นับเงินสดวันนี้) และใบรับเงินสด 5,000")

    d.h1("การทดสอบร่วม (Test plan)")
    d.table(
        ["ลำดับ", "กรณี", "ผลที่คาด"],
        [
            ["1", "GET /health ไม่ส่งคีย์", "200, enabled และ api_key_configured"],
            ["2", "POST /revenue ไม่ส่งคีย์", "401"],
            ["3", "ขาด external_id", "400"],
            ["4", "ส่งตัวอย่าง A", "202, state=ready จากนั้นบัญชี Post ได้ใบแจ้งหนี้ SSO OP"],
            ["5", "POST ตัวอย่าง A ซ้ำก่อน Post", "202, action=updated แทนที่บรรทัด"],
            ["6", "POST คีย์เดิมหลัง Post", "200, action=unchanged ไม่สร้างใบซ้ำ"],
            ["7", "ส่งตัวอย่าง B", "3 ใบแจ้งหนี้ เงินสดปิด ส่วน SSO/UC ค้างในลูกหนี้คงค้าง"],
            ["8", "ส่งตัวอย่าง C", "ชำระเอง IP คงเหลือ 10,000 และ SSO IP ค้าง 80,000 ทั้งคู่ขึ้นลูกหนี้คงค้าง"],
            ["9", "รับมัดจำ advance_in 10,000 ไม่มี sales", "202 ready; Post แล้วไม่มีใบแจ้งหนี้ มีเงินสดและหนี้สินรับล่วงหน้า"],
            ["10", "วันรักษา ตัด advance 10,000 + cash 5,000 บนบิล 15,000", "ใบแจ้งหนี้ปิดครบ เงินสดวันนี้ 5,000"],
            ["11", "ยอดคุมไม่ตรง", "state=error มีข้อความ control mismatch"],
            ["12", "มีฟิลด์ hn ใน payload", "รับได้ แต่ตัด hn ทิ้ง ไม่เก็บใน ERP"],
            ["13", "reversal ชุดที่ posted", "ชุดใหม่ ready เมื่อ Post จะกลับรายการต้นทาง"],
        ],
        col_widths=[1.8, 6.5, 8.7],
    )

    d.h1("การเปิดบริการและจุดติดต่อ")
    d.bullets(
        [
            "ฝ่าย HIS ขอ API key และ URL จาก IT โรงพยาบาล (ไม่ใส่ในเอกสารนี้)",
            "ก่อนเปิดจริง ทดสอบด้วยชุดตัวอย่างในสภาพแวดล้อมที่ตกลงกัน และให้บัญชีตรวจคิว HIS Revenue Queue",
            "เมื่อเปลี่ยนรหัสสิทธิ์หรือวิธีรับเงินฝั่ง HIS ต้องให้บัญชีเพิ่มแมปใน HIS Entitlements / HIS Payment Methods ก่อนส่งยอดจริง",
            "สอบถามรหัสที่ระบบรับได้ผ่าน GET lookups ก่อนปิดยอดวันแรก",
        ]
    )

    d.h1("ประวัติเอกสาร")
    d.table(
        ["เวอร์ชัน", "วันที่", "รายละเอียด"],
        [
            ["1.0 ร่าง", "25 ส.ค. 2569", "ร่างแรกครอบคลุมการรับรายได้ OPD/IPD, หลายสิทธิ์, จ่ายบางส่วน, idempotency และ reversal"],
            ["1.1 ร่าง", "25 ส.ค. 2569", "เพิ่มเงินล่วงหน้า: advance_in รับมัดจำ และ advance ตัดเมื่อมาใช้บริการ พร้อมจ่ายเพิ่มถ้าไม่พอ"],
        ],
        col_widths=[3.5, 3.5, 10],
    )
    d.p(
        "เอกสารนี้เป็นร่างสำหรับทบทวน หากมีการเปลี่ยนแปลงรหัสสิทธิ์ วิธีรับเงิน หรือกฎยอดคุม ให้ปรับปรุงฉบับนี้ก่อนเปิดใช้งานจริง",
        align="justify",
    )

    d.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
