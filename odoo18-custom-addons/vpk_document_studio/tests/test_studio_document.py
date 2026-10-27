# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

import base64
import io

from docx import Document
from odoo.exceptions import UserError, ValidationError
from odoo.tests import TransactionCase, tagged

from odoo.addons.vpk_document_studio.models.pdf_converter import find_soffice
from odoo.addons.vpk_document_studio.models.word_layout import build_print_html


def make_docx(paragraphs, table_cells=None):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_cells:
        table = document.add_table(rows=1, cols=len(table_cells))
        for index, text in enumerate(table_cells):
            table.rows[0].cells[index].text = text
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def docx_text(content):
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


@tagged("post_install", "-at_install")
class TestStudioDocument(TransactionCase):
    def _template(self, extra_fields=None, paragraphs=None):
        fields = [
            (0, 0, {
                "name": "เรื่อง",
                "code": "subject",
                "field_type": "char",
                "required": True,
                "sequence": 10,
            }),
            (0, 0, {
                "name": "วันที่",
                "code": "doc_date",
                "field_type": "date",
                "sequence": 20,
            }),
            (0, 0, {
                "name": "จำนวน",
                "code": "qty",
                "field_type": "integer",
                "default_value": "0",
                "sequence": 30,
            }),
            (0, 0, {
                "name": "ประเภท",
                "code": "kind",
                "field_type": "selection",
                "selection_options": "ปกติ\nเร่งด่วน",
                "sequence": 40,
            }),
        ]
        if extra_fields:
            fields.extend(extra_fields)
        return self.env["vpk.document.studio.template"].create(
            {
                "name": "ใบทดสอบ",
                "code": "test_form",
                "datas": base64.b64encode(
                    make_docx(
                        paragraphs
                        or [
                            "เรื่อง {{subject}}",
                            "วันที่ {{doc_date}}",
                            "จำนวน {{qty}}",
                            "ประเภท {{kind}}",
                        ]
                    )
                ),
                "filename": "test.docx",
                "field_ids": fields,
            }
        )

    def test_invalid_field_code(self):
        with self.assertRaises(ValidationError):
            self._template(
                extra_fields=[
                    (0, 0, {
                        "name": "ผิด",
                        "code": "Bad Code",
                        "field_type": "char",
                    })
                ]
            )

    def test_document_body_copied_from_template_layout(self):
        template = self._template()
        template.layout_html = "<p>ด้วย {{subject}} มีความประสงค์จะประกวดราคา</p>"
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {"subject": "ขออนุมัติลา"},
            }
        )
        self.assertIn("ด้วย ขออนุมัติลา มีความประสงค์จะประกวดราคา", document.body or "")
        self.assertNotIn("{{subject}}", document.body or "")
        document.body = False
        document.action_reload_body_from_template()
        self.assertIn("ด้วย ขออนุมัติลา มีความประสงค์จะประกวดราคา", document.body or "")

    def test_edited_body_is_saved_and_not_overwritten(self):
        template = self._template()
        template.layout_html = "<p>{{subject}}</p>"
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {"subject": "ขออนุมัติลา"},
            }
        )
        document.write({"body": "<p>แก้ไขเนื้อหาเอง</p>"})
        document.write({"note": "บันทึกฟิลด์อื่น", "template_id": template.id})
        self.assertIn("แก้ไขเนื้อหาเอง", document.body or "")
        document.action_save_body()

    def test_print_html_keeps_editor_absolute_fill_line_coords(self):
        html = build_print_html(
            '<p>หัว</p>'
            '<span class="o_vpk_we_fill_line" data-line-cm="6" data-line-left="4.5" '
            'data-line-top="3.2" data-line-style="dotted" '
            'style="position: absolute; left: 4.5cm; top: 3.2cm; min-width: 6cm; '
            'border-bottom: 1px dotted #000;">xxx</span>',
            "TH SarabunIT๙",
            16.0,
            {"top": 2.5, "bottom": 2.0, "left": 3.0, "right": 2.0},
        )
        self.assertIn("left: 4.5cm", html)
        self.assertIn("top: 3.2cm", html)
        self.assertIn("position: relative", html)
        self.assertIn("padding: 2.5cm 2.0cm 2.0cm 3.0cm", html)
        self.assertIn("margin: 0", html)

    def test_empty_field_box_defaults_to_empty_label(self):
        from odoo.addons.vpk_document_studio.models.word_layout import (
            fill_empty_field_boxes_html,
        )

        html = fill_empty_field_boxes_html(
            '<span class="o_vpk_we_field_box" data-field-code="dept"></span>'
            '<span class="o_vpk_we_field_box" data-field-code="agency">{{agency}}</span>'
        )
        self.assertIn("ว่าง", html)
        self.assertIn("o_vpk_we_field_empty", html)
        self.assertIn("{{agency}}", html)

        printed = build_print_html(
            '<span class="o_vpk_we_field_box" data-field-code="dept"></span>',
            "TH SarabunIT๙",
            16.0,
            {"top": 2.5, "bottom": 2.0, "left": 3.0, "right": 2.0},
        )
        self.assertIn("ว่าง", printed)

    def test_print_html_keeps_absolute_field_boxes(self):
        html = build_print_html(
            '<p>หัว</p>'
            '<span class="o_vpk_we_field_box" data-field-code="dept" data-line-left="4.5" '
            'data-line-top="3.2" contenteditable="false" '
            'style="position: absolute; left: 4.5cm; top: 3.2cm;">{{dept}}</span>',
            "TH SarabunIT๙",
            16.0,
            {"top": 2.5, "bottom": 2.0, "left": 3.0, "right": 2.0},
        )
        self.assertIn("o_vpk_we_field_box", html)
        self.assertIn("{{dept}}", html)
        self.assertIn("left: 4.5cm", html)
        self.assertIn("top: 3.2cm", html)
        self.assertIn("background: transparent", html)
        self.assertIn("line-height: 1.05", html)

    def test_print_html_matches_word_editor_page(self):
        html = build_print_html(
            '<p>เรื่อง <span class="o_vpk_we_fill_line" data-line-cm="6" '
            'data-line-style="dotted" style="min-width: 6cm; border-bottom: 1px dotted #000;">xxx</span></p>',
            "TH SarabunIT๙",
            16.0,
            {"top": 2.5, "bottom": 2.0, "left": 3.0, "right": 2.0},
        )
        self.assertIn("o_vpk_we_fill_line", html)
        self.assertNotIn("[[VPK_LINE", html)
        self.assertIn("21.0cm 29.7cm", html)
        self.assertIn("line-height: 1.15", html)
        self.assertIn("padding: 2.5cm 2.0cm 2.0cm 3.0cm", html)

    def test_print_html_normalizes_word_font_size_attributes(self):
        html = build_print_html(
            '<table><tr><td height="21" class="x">'
            '<p class="western"><font face="TH Sarabun New" size="4" '
            'style="font-size: 16pt"><b>ส่วนราชการ</b></font></p>'
            "</td></tr></table>",
            "TH SarabunIT๙",
            16.0,
            {"top": 2.5, "bottom": 2.0, "left": 3.0, "right": 2.0},
        )
        self.assertNotIn("size=", html)
        self.assertNotIn('class="western"', html)
        self.assertNotIn('height="21"', html)
        self.assertIn("ส่วนราชการ", html)
        self.assertIn("font-size: 16pt", html)

    def test_print_html_keeps_memo_label_on_one_line(self):
        html = build_print_html(
            '<table width="604">'
            "<colgroup><col width=\"26\"><col width=\"21\"><col width=\"45\">"
            "<col width=\"211\"><col width=\"45\"><col width=\"244\"></colgroup>"
            "<tr><td colspan=\"3\" width=\"96\" height=\"21\">"
            "<p class=\"western\"><font size=\"4\" style=\"font-size: 16pt\">"
            "<b>ส่วนราชการ</b></font></p></td>"
            "<td colspan=\"3\" width=\"504\">หน่วยงานทดสอบ</td></tr></table>",
            "TH SarabunIT๙",
            16.0,
            {"top": 2.5, "bottom": 2.0, "left": 3.0, "right": 2.0},
        )
        self.assertNotIn('width="96"', html)
        self.assertNotIn("<colgroup>", html)
        self.assertIn("white-space: nowrap", html)
        self.assertIn("min-width: 3.6cm", html)
        self.assertIn("ส่วนราชการ", html)

    def test_parse_layout_html_converts_underline_cells_to_fill_lines(self):
        from odoo.addons.vpk_document_studio.models.word_layout import parse_layout_html

        html = (
            "<html><body><table width=\"604\"><colgroup>"
            "<col width=\"96\"><col width=\"504\"></colgroup>"
            "<tr><td><b>ส่วนราชการ</b></td>"
            "<td width=\"504\" style=\"border-top: none; border-bottom: 1px dotted #000000; "
            "border-left: none; border-right: none\"><p>{{agency}}</p></td>"
            "</tr></table></body></html>"
        )
        body, _settings = parse_layout_html(html)
        self.assertIn("o_vpk_we_fill_line", body)
        self.assertIn("data-line-style=\"dotted\"", body)
        self.assertIn("{{agency}}", body)
        self.assertIn("ส่วนราชการ", body)
        self.assertNotIn("<colgroup>", body)

    def test_print_html_rewrites_memo_header_table(self):
        html = build_print_html(
            '<table width="604"><colgroup><col width="26"><col width="21">'
            '<col width="45"><col width="211"><col width="45"><col width="244">'
            "</colgroup><tr valign=\"bottom\">"
            '<td colspan="3" width="96"><p><font size="4" style="font-size: 16pt">'
            "<b>ส่วนราชการ</b></font></p></td>"
            '<td colspan="3" width="504"><p>หน่วยงานทดสอบ โรงพยาบาลวชิระภูเก็ต '
            "โทร ๐ ๗๖๓๖ ๑๒๓๔ต่อ๖๕๔๑</p></td></tr></table>",
            "TH SarabunIT๙",
            16.0,
            {"top": 2.5, "bottom": 2.0, "left": 3.0, "right": 2.0},
        )
        self.assertNotIn('width="96"', html)
        self.assertNotIn("<colgroup>", html)
        self.assertNotIn('colspan="3"', html)
        self.assertIn(" nowrap", html)
        self.assertIn("<nobr>", html)
        self.assertIn("<b>ส่วนราชการ</b>", html)
        self.assertIn("หน่วยงานทดสอบ โรงพยาบาลวชิระภูเก็ต", html)

    def test_required_field_blocks_print(self):
        template = self._template()
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {"subject": ""},
            }
        )
        with self.assertRaises(UserError):
            document.action_generate_docx()

    def test_generate_docx_from_values(self):
        template = self._template()
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {
                    "subject": "ขออนุมัติลา",
                    "doc_date": "2026-10-24",
                    "qty": 3,
                    "kind": "เร่งด่วน",
                },
            }
        )
        document.action_generate_docx()
        self.assertEqual(document.state, "generated")
        self.assertTrue(document.docx_file)
        self.assertTrue(document.docx_filename.endswith(".docx"))
        text = docx_text(base64.b64decode(document.docx_file))
        self.assertIn("เรื่อง ขออนุมัติลา", text)
        self.assertIn("24 ตุลาคม 2569", text)
        self.assertIn("จำนวน 3", text)
        self.assertIn("ประเภท เร่งด่วน", text)
        self.assertNotIn("{{", text)

    def test_sync_fields_from_word(self):
        template = self._template(
            paragraphs=["เรื่อง {{subject}}", "ของขาด {{missing}}"]
        )
        missing = template.field_ids.filtered(lambda field: field.code == "missing")
        self.assertTrue(missing)
        self.assertEqual(missing.field_type, "char")
        report = template._placeholder_report()
        self.assertNotIn("missing", report["unknown"])
        self.assertIn("doc_date", report["unused"])

    def test_value_can_reference_other_fields(self):
        template = self._template(
            extra_fields=[
                (0, 0, {
                    "name": "ส่วนราชการ",
                    "code": "agency",
                    "field_type": "char",
                    "sequence": 5,
                }),
                (0, 0, {
                    "name": "ย่อหน้าแรก",
                    "code": "intro",
                    "field_type": "text",
                    "sequence": 50,
                }),
            ],
            paragraphs=["{{intro}}"],
        )
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {
                    "subject": "ขออนุมัติ",
                    "agency": "โรงพยาบาลวชิระภูเก็ต",
                    "intro": (
                        "ด้วย {{agency}} มีความประสงค์จะประกวดราคาซื้อเวชภัณฑ์มิใช่ยา"
                    ),
                },
            }
        )
        mapping = document._merge_mapping()
        self.assertEqual(
            mapping["intro"],
            "ด้วย โรงพยาบาลวชิระภูเก็ต มีความประสงค์จะประกวดราคาซื้อเวชภัณฑ์มิใช่ยา",
        )
        document.action_generate_docx()
        text = docx_text(base64.b64decode(document.docx_file))
        self.assertIn(
            "ด้วย โรงพยาบาลวชิระภูเก็ต มีความประสงค์จะประกวดราคาซื้อเวชภัณฑ์มิใช่ยา",
            text,
        )
        self.assertNotIn("{{agency}}", text)

    def test_display_template_references_other_fields(self):
        template = self._template(
            extra_fields=[
                (0, 0, {
                    "name": "ส่วนราชการ",
                    "code": "agency",
                    "field_type": "char",
                    "sequence": 5,
                }),
                (0, 0, {
                    "name": "รายการสินค้า",
                    "code": "items",
                    "field_type": "text",
                    "sequence": 45,
                }),
                (0, 0, {
                    "name": "ย่อหน้าแรก",
                    "code": "intro",
                    "field_type": "text",
                    "sequence": 50,
                    "display_template": (
                        "ด้วย {{agency}} มีความประสงค์จะประกวดราคา "
                        "จำนวน {{items_count_th}} รายการ {{self}}"
                    ),
                }),
            ],
            paragraphs=["{{intro}}"],
        )
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {
                    "subject": "ขออนุมัติ",
                    "agency": "โรงพยาบาลวชิระภูเก็ต",
                    "items": "1. ยาเอ\n2. ยาบี",
                    "intro": "ด้วยวิธีประกวดราคาอิเล็กทรอนิกส์ (e-bidding)",
                },
            }
        )
        mapping = document._merge_mapping()
        self.assertEqual(
            mapping["intro"],
            "ด้วย โรงพยาบาลวชิระภูเก็ต มีความประสงค์จะประกวดราคา "
            "จำนวน ๒ รายการ ด้วยวิธีประกวดราคาอิเล็กทรอนิกส์ (e-bidding)",
        )

    def test_print_allows_extra_word_placeholders(self):
        template = self._template(
            paragraphs=["เรื่อง {{subject}}", "{{item_1}}", "{{approver_name}}"]
        )
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {"subject": "ขออนุมัติลา"},
            }
        )
        document.action_generate_docx()
        self.assertTrue(document.docx_file)
        text = docx_text(base64.b64decode(document.docx_file))
        self.assertIn("เรื่อง ขออนุมัติลา", text)
        self.assertNotIn("{{", text)

    def test_generate_pdf_when_soffice_available(self):
        if not find_soffice():
            self.skipTest("LibreOffice is required to generate PDF")
        template = self._template()
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {"subject": "พิมพ์ PDF"},
            }
        )
        action = document.action_generate_pdf()
        self.assertTrue(document.pdf_file)
        self.assertTrue(base64.b64decode(document.pdf_file).startswith(b"%PDF"))
        self.assertEqual(action.get("type"), "ir.actions.act_url")
        self.assertIn("/pdf_file/", action.get("url") or "")
        self.assertEqual(action.get("target"), "new")
