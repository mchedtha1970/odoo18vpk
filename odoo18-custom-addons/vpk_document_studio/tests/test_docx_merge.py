# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

import io

from docx import Document
from odoo.tests import TransactionCase, tagged

from odoo.addons.vpk_document_studio.models.docx_merge import (
    extract_placeholders,
    extract_repeat_tokens,
    merge_docx_bytes,
)


def make_docx(paragraphs, table_cells=None):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_cells:
        table = document.add_table(rows=1, cols=len(table_cells))
        for index, text in enumerate(table_cells):
            table.rows[0].cells[index].text = text
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def docx_text(content):
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


@tagged("post_install", "-at_install")
class TestDocxMerge(TransactionCase):
    def test_merge_fills_paragraphs_and_table(self):
        template = make_docx(
            ["เรื่อง {{subject}}", "{{empty_line}}", "ลงชื่อ {{signer}}"],
            table_cells=["วันที่", "{{doc_date}}"],
        )
        content = merge_docx_bytes(
            template,
            {
                "subject": "ขออนุมัติลา",
                "empty_line": "",
                "signer": "นายทดสอบ",
                "doc_date": "24 ตุลาคม 2569",
            },
            known_codes={"subject", "empty_line", "signer", "doc_date"},
        )
        text = docx_text(content)
        self.assertIn("เรื่อง ขออนุมัติลา", text)
        self.assertIn("ลงชื่อ นายทดสอบ", text)
        self.assertIn("24 ตุลาคม 2569", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("empty_line", text)

    def test_extract_placeholders(self):
        template = make_docx(
            ["{{subject}} และ {{subject}}", "วันที่ {{doc_date}}"],
            table_cells=["{{qty}}"],
        )
        self.assertEqual(
            extract_placeholders(template),
            ["subject", "doc_date", "qty"],
        )

    def test_extract_dotted_placeholder_uses_parent_code(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "รายการ"
        table.cell(1, 0).text = "{{items.product_id}}"
        table.cell(1, 1).text = "{{items.product_qty}}"
        buf = io.BytesIO()
        document.save(buf)
        content = buf.getvalue()
        self.assertEqual(extract_placeholders(content), ["items"])
        self.assertEqual(
            extract_repeat_tokens(content)["items"],
            ["product_id", "product_qty"],
        )

    def test_merge_repeats_item_table_rows(self):
        document = Document()
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "ลำดับ"
        table.cell(0, 1).text = "รายการ"
        table.cell(0, 2).text = "จำนวน"
        table.cell(1, 0).text = "{{no}}"
        table.cell(1, 1).text = "{{items.product_id}}"
        table.cell(1, 2).text = "{{items.product_qty}}"
        buf = io.BytesIO()
        document.save(buf)
        content = merge_docx_bytes(
            buf.getvalue(),
            {"subject": "ทดสอบ"},
            rows={
                "items": [
                    {
                        "no": "1",
                        "items.product_id": "ยา A",
                        "items.product_qty": "10",
                    },
                    {
                        "no": "2",
                        "items.product_id": "ยา B",
                        "items.product_qty": "5",
                    },
                ]
            },
        )
        text = docx_text(content)
        self.assertIn("ยา A", text)
        self.assertIn("ยา B", text)
        self.assertIn("10", text)
        self.assertNotIn("{{", text)
        filled = Document(io.BytesIO(content))
        self.assertGreaterEqual(len(filled.tables[0].rows), 3)

    def test_unknown_placeholder_is_cleared(self):
        template = make_docx(["หัวข้อ {{missing}}", "{{empty_only}}"])
        content = merge_docx_bytes(
            template,
            {"subject": "ทดสอบ"},
            known_codes={"subject"},
        )
        text = docx_text(content)
        self.assertIn("หัวข้อ", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("empty_only", text)
