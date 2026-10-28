# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

from odoo.tests import TransactionCase, tagged

from odoo.addons.vpk_document_studio.models.word_layout import expand_repeat_rows_html


@tagged("post_install", "-at_install")
class TestStudioBlocks(TransactionCase):
    def _template(self):
        return self.env["vpk.document.studio.template"].create(
            {"name": "บันทึกข้อความทดสอบ", "code": "memo_blocks_test"}
        )

    def test_number_display_formats(self):
        template = self._template()
        field = self.env["vpk.document.studio.field"].create(
            {"template_id": template.id, "code": "budget", "name": "งบ", "field_type": "float"}
        )
        expected = {
            "plain": "12500",
            "comma": "12,500.00",
            "comma_thai": "๑๒,๕๐๐.๐๐",
            "thai_text": "หนึ่งหมื่นสองพันห้าร้อย",
            "baht_text": "หนึ่งหมื่นสองพันห้าร้อยบาทถ้วน",
        }
        for number_format, text in expected.items():
            field.number_format = number_format
            self.assertEqual(field.format_value(12500), text, number_format)
        field.write({"field_type": "integer", "number_format": "comma"})
        self.assertEqual(field.format_value("12,500"), "12,500")

    def test_number_field_counts_mapped_lines(self):
        template = self._template()
        partner = self.env["res.partner"].create(
            {"name": "แม่", "child_ids": [(0, 0, {"name": "ลูก 1"}), (0, 0, {"name": "ลูก 2"})]}
        )
        field = self.env["vpk.document.studio.field"].create(
            {
                "template_id": template.id,
                "code": "items_count",
                "name": "จำนวนรายการ",
                "field_type": "integer",
                "number_format": "thai_text",
                "map_model_id": self.env["ir.model"]._get_id("res.partner"),
                "map_field_id": self.env["ir.model.fields"]._get("res.partner", "child_ids").id,
            }
        )
        value = field.extract_mapped_record_value(partner)
        self.assertEqual(value, 2)
        self.assertEqual(field.format_value(value), "สอง")

    def test_table_rows_from_many2one_sub_lines(self):
        template = self._template()
        parent = self.env["res.partner"].create(
            {"name": "แม่", "child_ids": [(0, 0, {"name": "ลูก 1"}), (0, 0, {"name": "ลูก 2"})]}
        )
        child = parent.child_ids[:1]
        fields_model = self.env["ir.model.fields"]
        field = self.env["vpk.document.studio.field"].create(
            {
                "template_id": template.id,
                "code": "items",
                "name": "รายการ",
                "field_type": "many2one",
                "map_model_id": self.env["ir.model"]._get_id("res.partner"),
                "map_field_id": fields_model._get("res.partner", "parent_id").id,
                "map_sub_field_id": fields_model._get("res.partner", "child_ids").id,
            }
        )
        block = self.env["vpk.document.studio.block"].create(
            {"template_id": template.id, "block_type": "table", "table_repeat_code": "items"}
        )
        self.assertEqual(block.line_model, "res.partner")
        rows = field._relation_rows({"res_id": child.id}, ("name",))
        self.assertEqual([row["line.name"] for row in rows], ["ลูก 1", "ลูก 2"])
        self.assertEqual(field.extract_mapped_record_value(child), "1. ลูก 1\n2. ลูก 2")

    def test_table_cell_formula(self):
        html = expand_repeat_rows_html(
            '<table><tr data-vpk-repeat="items"><td>={{qty}}*{{unit_price}}</td>'
            "<td>=({{line.a}}+{{line.missing}})/2</td><td>=__import__('os')</td></tr></table>",
            {},
            {"items": [{"no": "1", "qty": "1,000", "unit_price": "699.50", "line.a": "5"}]},
        )
        self.assertIn("<td>699,500.00</td>", html)
        self.assertIn("<td>2.50</td>", html)
        self.assertIn("<td>=__import__('os')</td>", html)

    def test_custom_layout_survives_block_edits(self):
        template = self._template()
        template.action_load_memo_preset()
        template.write({"layout_custom": True, "layout_html": "<p>แก้เอง {{extra_note}}</p>"})
        self.assertIn("extra_note", template.field_ids.mapped("code"))
        template.block_ids[:1].write({"title": "บันทึกข้อความใหม่"})
        template.write({"layout_margin_top": 3.0})
        self.assertIn("แก้เอง", template.layout_html)
        template.action_regenerate_layout()
        self.assertFalse(template.layout_custom)
        self.assertNotIn("แก้เอง", template.layout_html)

    def test_import_word_as_editable_layout(self):
        import base64
        import io

        import docx

        doc = docx.Document()
        doc.add_paragraph("บันทึกข้อความ")
        doc.add_paragraph("เรื่อง {{subject}}")
        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "ลำดับ"
        table.cell(0, 1).text = "รายการ"
        buf = io.BytesIO()
        doc.save(buf)
        template = self._template()
        template.write({"datas": base64.b64encode(buf.getvalue()), "filename": "memo.docx"})
        template.action_import_word_layout()
        self.assertTrue(template.layout_custom)
        self.assertIn("บันทึกข้อความ", template.layout_html)
        self.assertIn("{{subject}}", template.layout_html)
        self.assertIn("<table", template.layout_html)
        self.assertIn("รายการ", template.layout_html)
        self.assertIn("subject", template.field_ids.mapped("code"))

    def test_insert_field_into_heading(self):
        template = self._template()
        field = self.env["vpk.document.studio.field"].create(
            {"template_id": template.id, "code": "subject", "name": "เรื่อง"}
        )
        block = self.env["vpk.document.studio.block"].new(
            {"template_id": template.id, "block_type": "heading", "heading_text": "เรื่อง"}
        )
        block.insert_field_id = field
        block._onchange_insert_field_id()
        self.assertEqual(block.heading_text, "เรื่อง {{subject}}")
        self.assertFalse(block.insert_field_id)

    def test_memo_preset_builds_layout_and_fields(self):
        template = self._template()
        template.action_load_memo_preset()
        self.assertEqual(
            template.block_ids.mapped("block_type"),
            ["memo_header", "paragraph", "table", "paragraph", "signature"],
        )
        html = template.layout_html or ""
        number_code = template.block_ids[:1].code_number
        for code in ("agency", number_code, "memo_date", "subject", "recipient"):
            self.assertIn("{{%s}}" % code, html)
        self.assertIn("บันทึกข้อความ", html)
        self.assertIn("o_vpk_we_fill_line", html)
        self.assertNotIn("position: absolute", html)
        codes = set(template.field_ids.mapped("code"))
        self.assertTrue(
            {"agency", number_code, "memo_date", "subject", "items"} <= codes
        )
        self.assertNotIn("items_count_th", codes)
        self.assertIn("{{items_count_th}}", html)

    def test_default_text_in_paragraph_uses_count_placeholder(self):
        template = self._template()
        template.action_load_memo_preset()
        items = template.field_ids.filtered(lambda f: f.code == "items")
        items.field_type = "text"
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {
                    "agency": "กลุ่มงานเภสัชกรรม",
                    "items": "ยา ก\nยา ข",
                },
            }
        )
        self.assertIn("จำนวน ๒ รายการ", document.body or "")
        self.assertIn("กลุ่มงานเภสัชกรรม", document.body or "")

    def test_header_lines_fit_printable_width(self):
        template = self._template()
        template.action_load_memo_preset()
        header = template.block_ids[0]
        html = header._render_html(template._printable_width_cm(), 16.0)
        import re

        widths = [float(w) for w in re.findall(r'data-line-cm="([0-9.]+)"', html)]
        self.assertEqual(len(widths), 5)
        self.assertLessEqual(max(widths), template._printable_width_cm())

    def test_editing_block_updates_layout(self):
        template = self._template()
        template.action_load_memo_preset()
        paragraph = template.block_ids.filtered(
            lambda b: b.block_type == "paragraph"
        )[0]
        paragraph.content = "<p>ข้อความใหม่ {{extra_note}}</p>"
        self.assertIn("ข้อความใหม่", template.layout_html)
        self.assertIn("text-indent: 2.5cm", template.layout_html)
        self.assertIn("extra_note", template.field_ids.mapped("code"))

    def test_margin_change_resizes_lines(self):
        template = self._template()
        template.action_load_memo_preset()
        before = template.layout_html
        template.layout_margin_left = 4.0
        self.assertNotEqual(before, template.layout_html)

    def test_document_uses_block_layout_and_fills_values(self):
        template = self._template()
        template.action_load_memo_preset()
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {"agency": "โรงพยาบาลวชิระภูเก็ต", "subject": "ขอซื้อ"},
            }
        )
        self.assertIn("โรงพยาบาลวชิระภูเก็ต", document.body or "")
        self.assertIn("ขอซื้อ", document.body or "")
        self.assertNotIn("{{agency}}", document.body or "")

    def test_legacy_template_without_blocks_keeps_layout(self):
        template = self._template()
        template.layout_html = "<p>อิสระ</p>"
        template.layout_margin_left = 4.0
        self.assertEqual(template.layout_html, "<p>อิสระ</p>")

    def test_table_block_repeats_rows_per_item(self):
        template = self._template()
        template.action_load_memo_preset()
        items = template.field_ids.filtered(lambda f: f.code == "items")
        items.field_type = "text"
        self.assertIn("<table", template.layout_html)
        self.assertIn('data-vpk-repeat="items"', template.layout_html)
        self.assertNotIn("line", template.field_ids.mapped("code"))
        document = self.env["vpk.document.studio.document"].create(
            {
                "template_id": template.id,
                "value_json": {"items": "1. ยา ก\n2. ยา ข\n3. ยา ค"},
            }
        )
        body = document.body or ""
        self.assertNotIn("data-vpk-repeat", body)
        self.assertEqual(body.count("<tr"), 5)  # header + 3 items + total row
        for text in ("๑", "๒", "๓", "ยา ก", "ยา ข", "ยา ค"):
            self.assertIn(text, body)
        self.assertNotIn("{{line}}", body)
        self.assertNotIn("1. ยา", body)

    def test_table_block_fixed_rows(self):
        template = self._template()
        template.write(
            {
                "block_ids": [
                    (
                        0,
                        0,
                        {
                            "block_type": "table",
                            "table_rows": 3,
                            "table_header": False,
                            "column_ids": [
                                (0, 0, {"header": "A", "content": "{{subject}}"}),
                                (0, 0, {"header": "B", "width_cm": 3.0}),
                            ],
                        },
                    )
                ],
            }
        )
        self.assertEqual(template.layout_html.count("<tr"), 3)
        self.assertIn("subject", template.field_ids.mapped("code"))
        template.block_ids.column_ids[0].content = "{{agency}}"
        self.assertIn("{{agency}}", template.layout_html)

    def test_structured_rows_fill_name_qty_amount_columns(self):
        template = self._template()
        template.action_load_memo_preset()
        html = expand_repeat_rows_html(
            template.layout_html,
            {"items_total": "300.00"},
            {
                "items": [
                    {"no": "1", "name": "ยา ก", "qty": "10", "uom": "กล่อง",
                     "amount": "100.00", "amount_value": 100.0},
                    {"no": "2", "name": "ยา ข", "qty": "5", "uom": "ขวด",
                     "amount": "200.00", "amount_value": 200.0},
                ]
            },
        )
        for text in ("ยา ก", "10 กล่อง", "100.00", "ยา ข", "5 ขวด", "200.00", "๑", "๒"):
            self.assertIn(text, html)
        self.assertNotIn("{{qty}}", html)
        self.assertNotIn("{{amount}}", html)
        self.assertNotIn("data-vpk-repeat", html)

    def test_row_tokens_not_created_as_fields(self):
        template = self._template()
        template.action_load_memo_preset()
        codes = set(template.field_ids.mapped("code"))
        for token in ("name", "qty", "uom", "amount", "no_th", "items_total"):
            self.assertNotIn(token, codes)
        self.assertIn("items", codes)

    def test_heading_block(self):
        template = self._template()
        self.env["vpk.document.studio.block"].create(
            {
                "template_id": template.id,
                "block_type": "heading",
                "heading_text": "ความเห็นของ {{committee}}",
                "heading_align": "center",
                "heading_underline": True,
            }
        )
        html = template.layout_html or ""
        self.assertIn("ความเห็นของ {{committee}}", html)
        self.assertIn("text-align: center", html)
        self.assertIn("<u>", html)
        self.assertIn("committee", set(template.field_ids.mapped("code")))
