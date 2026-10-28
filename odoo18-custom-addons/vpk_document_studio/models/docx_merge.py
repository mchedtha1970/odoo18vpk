# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

"""Fill a Word template by replacing {{placeholder}} tokens."""

import copy
import io
import re

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover
    Document = None
    qn = None
    Paragraph = None

THAI_PDF_FONT = "TH Sarabun New"
PLACEHOLDER_RE = re.compile(
    r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*(?:\.[a-zA-Z_][a-zA-Z0-9_]*)*)\s*\}\}"
)
PLACEHOLDER_ONLY_RE = re.compile(
    r"^\s*\{\{\s*[a-zA-Z_][a-zA-Z0-9_]*(?:\.[a-zA-Z_][a-zA-Z0-9_]*)*\s*\}\}\s*$"
)
DOTTED_PLACEHOLDER_RE = re.compile(
    r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)\.([a-zA-Z_][a-zA-Z0-9_]*)\s*\}\}"
)
_ROW_HINT_RE = re.compile(
    r"\{\{\s*(?:no|no_th|line|name|qty|uom|unit_price|amount|spec)\s*\}\}"
)


class DocxMergeError(ValueError):
    """Raised when a Word template cannot be filled."""


def _require_docx():
    if Document is None:
        raise DocxMergeError(
            "ไม่พบไลบรารี python-docx กรุณาติดตั้งในสภาพแวดล้อม Odoo ก่อนพิมพ์แบบฟอร์ม"
        )


def _paragraph_text(paragraph):
    return "".join(run.text or "" for run in paragraph.runs)


def _replace_in_paragraph_runs(paragraph, mapping):
    """Replace {{code}} inside individual runs when possible so run-level
    formatting (e.g. dotted underline on a fill-line) is preserved."""
    runs = paragraph.runs
    if not runs:
        return
    # First coalesce adjacent runs if a {{placeholder}} was split across runs
    i = 0
    while i < len(runs):
        txt = runs[i].text or ""
        if "{{" in txt and "}}" not in txt.split("{{")[-1]:
            j = i + 1
            while j < len(runs):
                txt += runs[j].text or ""
                runs[j].text = ""
                if "}}" in txt.split("{{")[-1]:
                    break
                j += 1
            runs[i].text = txt
        i += 1
    for run in runs:
        r_pr = run._element.find(qn("w:rPr"))
        if r_pr is not None:
            bdr = r_pr.find(qn("w:bdr"))
            if bdr is not None:
                from docx.oxml import OxmlElement

                val = (bdr.get(qn("w:val")) or "dotted").lower()
                u_val = "dotted" if val in ("dotted", "dashed", "dotdash") else "single"
                r_pr.remove(bdr)
                u_el = r_pr.find(qn("w:u"))
                if u_el is None:
                    u_el = OxmlElement("w:u")
                    r_pr.append(u_el)
                u_el.set(qn("w:val"), u_val)
        rtxt = run.text or ""
        if "{{" in rtxt:
            trailing_nbsp = len(rtxt) - len(rtxt.rstrip("\u00a0"))
            core = rtxt[: len(rtxt) - trailing_nbsp] if trailing_nbsp else rtxt
            replaced_core = _replace_in_text(core, mapping)
            if trailing_nbsp:
                placeholder_budget = 10 * len(PLACEHOLDER_RE.findall(core))
                total_budget = trailing_nbsp + placeholder_budget
                new_pad = max(2, total_budget - len(replaced_core.strip()))
                run.text = replaced_core + ("\u00a0" * new_pad)
            else:
                run.text = replaced_core


def _remove_paragraph(paragraph):
    elm = paragraph._element
    parent = elm.getparent()
    if parent is not None:
        parent.remove(elm)


def _set_run_font(run, font_name=THAI_PDF_FONT):
    from docx.oxml import OxmlElement

    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), font_name)


def _iter_container_paragraphs(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _iter_story_containers(document):
    yield document
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
        ):
            if part is not None:
                yield part


def _iter_paragraphs(document):
    seen = set()
    for container in _iter_story_containers(document):
        for paragraph in _iter_container_paragraphs(container):
            element = paragraph._element
            if element in seen:
                continue
            seen.add(element)
            yield paragraph


def _apply_thai_font(document, font_name=THAI_PDF_FONT):
    for paragraph in _iter_paragraphs(document):
        for run in paragraph.runs:
            _set_run_font(run, font_name)


def extract_placeholders(template_bytes):
    """Return unique placeholder codes found in the Word file, in order.

    `{{items.product_id}}` contributes the parent field `items`.
    """
    _require_docx()
    document = Document(io.BytesIO(template_bytes))
    found = []
    seen = set()
    for paragraph in _iter_paragraphs(document):
        for token in PLACEHOLDER_RE.findall(_paragraph_text(paragraph)):
            code = token.split(".")[0]
            if code not in seen:
                seen.add(code)
                found.append(code)
    return found


def extract_repeat_tokens(template_bytes):
    """Return {parent_code: [field, ...]} from {{items.product_id}} / {{line.qty}}."""
    _require_docx()
    document = Document(io.BytesIO(template_bytes))
    result = {}
    for paragraph in _iter_paragraphs(document):
        for parent, field in DOTTED_PLACEHOLDER_RE.findall(_paragraph_text(paragraph)):
            fields = result.setdefault(parent, [])
            if field not in fields:
                fields.append(field)
    return result


def _repeat_code_for_text(text, rows):
    if not text or "{{" not in text or not rows:
        return None
    for parent, _field in DOTTED_PLACEHOLDER_RE.findall(text):
        if parent in rows:
            return parent
        if parent == "line" and "items" in rows:
            return "items"
        if parent == "line" and len(rows) == 1:
            return next(iter(rows))
    if _ROW_HINT_RE.search(text):
        if "items" in rows:
            return "items"
        if len(rows) == 1:
            return next(iter(rows))
    return None


def _row_token_mapping(code, item):
    mapping = {}
    prefix = "%s." % code
    for key, value in (item or {}).items():
        if key == "amount_value":
            continue
        text = "" if value is None else str(value)
        mapping[str(key)] = text
        if key.startswith(prefix):
            field = key[len(prefix) :]
            mapping.setdefault(field, text)
            mapping.setdefault("line.%s" % field, text)
        elif "." not in str(key):
            mapping.setdefault("%s.%s" % (code, key), text)
            mapping.setdefault("line.%s" % key, text)
    return mapping


def _fill_element_placeholders(element, mapping):
    for p_el in element.iter(qn("w:p")):
        paragraph = Paragraph(p_el, None)
        if "{{" not in _paragraph_text(paragraph):
            continue
        _replace_in_paragraph_runs(paragraph, mapping)


def _expand_tbl(tbl, rows):
    for tr in list(tbl.findall(qn("w:tr"))):
        text = "".join(tr.itertext())
        code = _repeat_code_for_text(text, rows)
        if not code:
            continue
        items = list(rows.get(code) or []) or [{}]
        parent = tr.getparent()
        index = list(parent).index(tr)
        parent.remove(tr)
        for offset, item in enumerate(items):
            new_tr = copy.deepcopy(tr)
            _fill_element_placeholders(new_tr, _row_token_mapping(code, item))
            parent.insert(index + offset, new_tr)


def _expand_repeat_table_rows(document, rows):
    if not rows:
        return
    roots = [document.element.body]
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
        ):
            if part is not None:
                roots.append(part._element)
    seen = set()
    for root in roots:
        for tbl in root.iter(qn("w:tbl")):
            if tbl in seen:
                continue
            seen.add(tbl)
            _expand_tbl(tbl, rows)


def _replace_in_text(text, mapping):
    def replacer(match):
        code = match.group(1)
        if code not in mapping:
            return match.group(0)
        value = mapping[code]
        return "" if value is None else str(value)

    return PLACEHOLDER_RE.sub(replacer, text)


def merge_docx_bytes(template_bytes, mapping, known_codes=None, rows=None):
    """Replace {{code}} tokens and return filled .docx bytes.

    Placeholders without a form field are filled with a blank. Paragraphs
    that contain only an empty placeholder are removed.

    `rows` is {field_code: [row dict, ...]} used to clone Word table rows
    that contain {{items.product_id}} / {{line.qty}} / {{no}}.
    """
    _require_docx()
    if not template_bytes:
        raise DocxMergeError("ไม่มีไฟล์ Word สำหรับพิมพ์")
    mapping = {
        str(key): ("" if value is None else value)
        for key, value in (mapping or {}).items()
    }
    if known_codes:
        for code in known_codes:
            mapping.setdefault(code, "")
    document = Document(io.BytesIO(template_bytes))
    _expand_repeat_table_rows(document, rows or {})
    for paragraph in list(_iter_paragraphs(document)):
        for run in paragraph.runs:
            r_pr = run._element.find(qn("w:rPr"))
            if r_pr is not None:
                bdr = r_pr.find(qn("w:bdr"))
                if bdr is not None:
                    from docx.oxml import OxmlElement

                    val = (bdr.get(qn("w:val")) or "dotted").lower()
                    u_val = "dotted" if val in ("dotted", "dashed", "dotdash") else "single"
                    r_pr.remove(bdr)
                    u_el = r_pr.find(qn("w:u"))
                    if u_el is None:
                        u_el = OxmlElement("w:u")
                        r_pr.append(u_el)
                    u_el.set(qn("w:val"), u_val)
        text = _paragraph_text(paragraph)
        if "{{" not in text:
            continue
        for code in PLACEHOLDER_RE.findall(text):
            mapping.setdefault(code, "")
        placeholder_only = bool(PLACEHOLDER_ONLY_RE.match(text))
        new = _replace_in_text(text, mapping)
        in_table_cell = paragraph._element.getparent() is not None and paragraph._element.getparent().tag.endswith("}tc")
        if placeholder_only and not (new or "").strip():
            if in_table_cell:
                for idx, run in enumerate(paragraph.runs):
                    run.text = "\u00a0" if idx == 0 else ""
            else:
                _remove_paragraph(paragraph)
            continue
        if new != text:
            _replace_in_paragraph_runs(paragraph, mapping)
    _apply_thai_font(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_placeholder_guide_docx(title, fields):
    """Build a starter Word file listing each {{code}} for the user to rearrange."""
    _require_docx()
    document = Document()
    heading = document.add_paragraph(title or "แบบฟอร์ม")
    heading.runs[0].bold = True
    document.add_paragraph(
        "คัดลอก {{รหัสฟิลด์}} ไปวางในแบบฟอร์มที่ต้องการ แล้วอัปโหลดไฟล์กลับมาที่ระบบ"
    )
    if not fields:
        document.add_paragraph("ยังไม่มีฟิลด์ในแบบฟอร์มนี้")
    for field in fields:
        code = field.get("code") or ""
        label = field.get("name") or code
        field_type = field.get("field_type") or "char"
        document.add_paragraph("%s (%s)" % (label, field_type))
        document.add_paragraph("{{%s}}" % code)
    _apply_thai_font(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
