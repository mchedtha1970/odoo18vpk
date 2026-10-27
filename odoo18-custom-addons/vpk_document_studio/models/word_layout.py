# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

"""Edit a template's Word file as HTML and rebuild the Word file from it."""

import io
import copy
import re

from lxml import etree
from lxml import html as lxml_html

from .docx_merge import DocxMergeError, _require_docx

A4_WIDTH_CM = 21.0
A4_HEIGHT_CM = 29.7

_LENGTH_RE = re.compile(r"(-?\d+(?:\.\d+)?)\s*(in|cm|mm|pt)\b")
_CM_PER_UNIT = {"in": 2.54, "cm": 1.0, "mm": 0.1, "pt": 2.54 / 72}


def _to_cm(text):
    match = _LENGTH_RE.search(text or "")
    if not match:
        return None
    return float(match.group(1)) * _CM_PER_UNIT[match.group(2)]


def parse_layout_html(html_text):
    """Return (body_html, settings) from LibreOffice's HTML export."""
    settings = {}
    page = re.search(r"@page\s*\{([^}]*)\}", html_text)
    if page:
        for side in ("top", "bottom", "left", "right"):
            match = re.search(r"margin-%s\s*:\s*([^;}]+)" % side, page.group(1))
            value = _to_cm(match.group(1)) if match else None
            if value is not None:
                settings["margin_%s" % side] = round(value, 2)
    font_rule = re.search(r"p\.western\s*\{([^}]*)\}", html_text) or re.search(
        r"(?<![\w.])p\s*\{([^}]*)\}", html_text
    )
    if font_rule:
        family = re.search(r"font-family\s*:\s*\"?([^\",;]+)", font_rule.group(1))
        if family:
            settings["font"] = family.group(1).strip()
        size = re.search(r"font-size\s*:\s*(\d+(?:\.\d+)?)pt", font_rule.group(1))
        if size:
            settings["font_size"] = float(size.group(1))

    document = lxml_html.document_fromstring(html_text)
    body = document.find("body")
    if body is None:
        return html_text, settings
    for img in list(body.xpath(".//img")):
        alt = (img.get("alt") or img.get("name") or "").strip()
        if re.match(r"^Shape\d+$", alt, re.IGNORECASE):
            tail = img.tail
            parent = img.getparent()
            if tail and parent is not None:
                prev = img.getprevious()
                if prev is not None:
                    prev.tail = (prev.tail or "") + tail
                else:
                    parent.text = (parent.text or "") + tail
            if parent is not None:
                parent.remove(img)
            continue
        # Remove float/align on inline images (e.g. Garuda) so they don't
        # export as floating anchors overlapping the heading below.
        if img.get("align") in ("left", "right"):
            parent_p = img.getparent()
            while parent_p is not None and parent_p.tag not in ("p", "div", "body"):
                parent_p = parent_p.getparent()
            if parent_p is not None and parent_p.tag in ("p", "div"):
                parent_p.set("align", img.get("align"))
                p_style = parent_p.get("style") or ""
                p_style = re.sub(r"line-height\s*:\s*[^;]+;?", "", p_style)
                parent_p.set("style", p_style.strip())
            img.attrib.pop("align", None)
    for frame in list(body.xpath(".//span[starts-with(@id, 'Frame')]")):
        if not (frame.text_content() or "").strip() and len(frame) == 0:
            parent = frame.getparent()
            if parent is not None:
                parent.remove(frame)
    _convert_underline_cells_to_fill_lines(body)
    parts = [body.text or ""]
    parts += [etree.tostring(child, encoding="unicode", method="html") for child in body]
    return "".join(parts).strip(), settings


_FILL_TOKEN_RE = re.compile(r"\[\[VPK_LINE:([0-9.]+):(dotted|solid):(.*?)\]\]", re.DOTALL)
_HTML_FONT_SIZE_PT = {
    "1": 8,
    "2": 10,
    "3": 12,
    "4": 14,
    "5": 18,
    "6": 24,
    "7": 36,
}


def _normalize_html_fonts(root):
    """Make LibreOffice PDF use the same font size as the Word editor.

    Word HTML often has both `<font size="4">` and `style="font-size: 16pt"`.
    Browsers keep 16pt; LibreOffice often uses size="4" (~14pt) so labels
    like ส่วน / ที่ look smaller than ราชการ / วันที่. Drop the size
    attribute and Writer-only paragraph classes, and don't squeeze cells
    with a fixed height.
    """
    for node in root.xpath(".//*[@class]"):
        classes = [
            cls
            for cls in (node.get("class") or "").split()
            if cls not in ("western", "cjk", "ctl", "asian")
        ]
        if classes:
            node.set("class", " ".join(classes))
        else:
            node.attrib.pop("class", None)

    for font in list(root.xpath(".//font")):
        style = font.get("style") or ""
        if not re.search(r"font-size\s*:", style, re.I):
            mapped = _HTML_FONT_SIZE_PT.get((font.get("size") or "").strip())
            # size 4/5 from 16pt Word templates should follow the editor page
            if (font.get("size") or "").strip() in ("4", "5"):
                mapped = None
            if mapped:
                style = ("%s; font-size: %spt" % (style, mapped)).strip("; ")
        face = (font.get("face") or "").split(",")[0].strip()
        if face and not re.search(r"font-family\s*:", style, re.I):
            style = ('%s; font-family: "%s"' % (style, face)).strip("; ")
        font.attrib.pop("size", None)
        font.attrib.pop("face", None)
        font.attrib.pop("color", None)
        if style.strip():
            font.set("style", style.strip().strip(";"))
        font.tag = "span"

    for cell in root.xpath(".//td | .//th | .//tr"):
        cell.attrib.pop("height", None)
        style = cell.get("style") or ""
        if re.search(r"(^|;)\s*height\s*:", style, re.I):
            cell.set("style", re.sub(r"height\s*:\s*[^;]+;?", "", style).strip().strip(";"))


_MEMO_LABELS = {"ส่วนราชการ", "ที่", "วันที่", "เรื่อง", "เรียน"}
_MEMO_LABEL_MIN_WIDTH = {
    "ส่วนราชการ": "3.6cm",
    "ที่": "1.2cm",
    "วันที่": "1.8cm",
    "เรื่อง": "1.8cm",
    "เรียน": "1.8cm",
}


def _cell_plain_text(cell):
    return " ".join((cell.text_content() or "").split())


def _inner_html(element):
    parts = [element.text or ""]
    parts += [
        etree.tostring(child, encoding="unicode", method="html")
        for child in element
    ]
    return "".join(parts).strip()


def _cell_value_html(cell):
    paragraphs = cell.xpath("./p")
    if paragraphs:
        return " ".join(_inner_html(paragraph) for paragraph in paragraphs).strip()
    return _inner_html(cell)


def _append_html(parent, html_text):
    payload = html_text or "\u00a0"
    wrapper = lxml_html.fromstring("<div>%s</div>" % payload)
    parent.text = wrapper.text
    for child in list(wrapper):
        parent.append(child)


def _add_memo_label_cell(row, label):
    cell = etree.SubElement(row, "td")
    min_width = _MEMO_LABEL_MIN_WIDTH.get(label, "1.2cm")
    cell.set("width", "1%")
    cell.set("nowrap", "nowrap")
    cell.set("valign", "bottom")
    cell.set(
        "style",
        "white-space: nowrap; width: 1%; min-width: "
        + min_width
        + "; padding: 0 0.25cm 0 0; border: none; vertical-align: bottom;",
    )
    nobr = etree.SubElement(cell, "nobr")
    bold = etree.SubElement(nobr, "b")
    bold.text = label
    return cell


def _add_memo_value_cell(row, html_text):
    cell = etree.SubElement(row, "td")
    cell.set("valign", "bottom")
    has_fill_line = "o_vpk_we_fill_line" in (html_text or "")
    if has_fill_line:
        cell.set(
            "style",
            "border: none; padding: 0 0.15cm 1px 0.15cm; vertical-align: bottom;",
        )
    else:
        cell.set(
            "style",
            "border: none; border-bottom: 1px dotted #000000; "
            "padding: 0 0.15cm 1px 0.15cm; vertical-align: bottom;",
        )
    _append_html(cell, html_text)
    return cell


_UNDERLINE_BORDER_RE = re.compile(
    r"border-bottom\s*:\s*[^;]*(dotted|dashed|solid)",
    re.I,
)


def _underline_style_from_cell(cell):
    style = cell.get("style") or ""
    match = _UNDERLINE_BORDER_RE.search(style)
    if not match:
        return None
    kind = match.group(1).lower()
    if kind == "solid":
        return "solid"
    return "dotted"


def _cell_width_cm(cell):
    width = cell.get("width") or ""
    digits = re.sub(r"[^\d.]", "", width)
    if digits:
        try:
            px = float(digits)
            if px > 0:
                return max(1.0, min(18.0, round(px / 37.795 * 10) / 10))
        except ValueError:
            pass
    style = cell.get("style") or ""
    match = re.search(r"(?:min-)?width\s*:\s*([\d.]+)\s*(px|cm|in)", style, re.I)
    if match:
        value = float(match.group(1))
        unit = match.group(2).lower()
        if unit == "cm":
            cm = value
        elif unit == "in":
            cm = value * 2.54
        else:
            cm = value / 37.795
        return max(1.0, min(18.0, round(cm * 10) / 10))
    return 8.0


def _convert_underline_cells_to_fill_lines(root):
    """Turn Word table bottom-borders into editable fill-line spans.

    LibreOffice exports บันทึกข้อความ underlines as `td { border-bottom: dotted }`.
    Those cannot be selected or resized in the Word editor. Wrap the cell
    content in `.o_vpk_we_fill_line` so the user can click, drag width, and
    change dotted/solid like lines inserted from the toolbar.
    """
    converted = False
    for cell in root.xpath(".//td | .//th"):
        if cell.xpath(".//*[contains(concat(' ', normalize-space(@class), ' '), ' o_vpk_we_fill_line ')]"):
            continue
        line_style = _underline_style_from_cell(cell)
        if not line_style:
            continue
        cm = _cell_width_cm(cell)
        inner = _inner_html(cell)
        for child in list(cell):
            cell.remove(child)
        cell.text = None
        span = etree.SubElement(cell, "span")
        span.set("class", "o_vpk_we_fill_line")
        span.set("data-line-cm", "%.1f" % cm)
        span.set("data-line-style", line_style)
        span.set(
            "style",
            "display: inline-block; min-width: %.1fcm; width: %.1fcm; "
            "border-bottom: 1px %s #000000; padding: 0 4px; "
            "line-height: 1.05; text-align: left;"
            % (cm, cm, line_style),
        )
        _append_html(span, inner or "\u00a0")
        style = re.sub(
            r"border-bottom\s*:\s*[^;]+;?",
            "border-bottom: none;",
            cell.get("style") or "",
        )
        cell.set("style", style.strip().strip(";"))
        cell.attrib.pop("width", None)
        converted = True
    if not converted:
        return False
    for table in root.xpath(".//table"):
        if not table.xpath(".//*[contains(@class, 'o_vpk_we_fill_line')]"):
            continue
        for colgroup in list(table.xpath("./colgroup")):
            parent = colgroup.getparent()
            if parent is not None:
                parent.remove(colgroup)
        for col in list(table.xpath("./col")):
            parent = col.getparent()
            if parent is not None:
                parent.remove(col)
        table.attrib.pop("width", None)
    return True


def _table_column_count(table):
    count = 0
    for row in table.xpath("./tbody/tr | ./tr"):
        span_sum = 0
        for cell in row.xpath("./td | ./th"):
            try:
                span_sum += max(1, int(cell.get("colspan") or 1))
            except ValueError:
                span_sum += 1
        count = max(count, span_sum)
    return count


def _rewrite_memo_header_tables(root):
    """Rebuild บันทึกข้อความ header tables so LibreOffice cannot wrap labels.

    Word HTML uses a 6-column grid (`ส่วนราชการ` in a 96px colspan=3 cell).
    Browsers grow that cell; LibreOffice keeps the narrow grid and prints
    ส่วน / ราชการ on two lines.
    """
    for table in list(root.xpath(".//table")):
        if table.get("data-vpk-memo-header"):
            continue
        rows = table.xpath("./tbody/tr | ./tr")
        if not rows:
            continue
        parsed_rows = []
        for row in rows:
            cells = row.xpath("./td | ./th")
            pairs = []
            index = 0
            while index < len(cells):
                text = _cell_plain_text(cells[index])
                if text in _MEMO_LABELS:
                    index += 1
                    values = []
                    while index < len(cells) and _cell_plain_text(cells[index]) not in _MEMO_LABELS:
                        values.append(_cell_value_html(cells[index]))
                        index += 1
                    pairs.append((text, " ".join(part for part in values if part).strip()))
                else:
                    index += 1
            if pairs:
                parsed_rows.append(pairs)
        labels = {label for row_pairs in parsed_rows for label, _value in row_pairs}
        cramped = bool(table.xpath("./colgroup | .//col")) or _table_column_count(table) >= 3
        if "ส่วนราชการ" not in labels or not cramped:
            continue
        new_table = etree.Element("table")
        new_table.set("border", "0")
        new_table.set("cellpadding", "0")
        new_table.set("cellspacing", "0")
        new_table.set("width", "100%")
        new_table.set("data-vpk-memo-header", "1")
        new_table.set(
            "style",
            "width: 100%; table-layout: auto; border-collapse: collapse; border: none;",
        )
        body = etree.SubElement(new_table, "tbody")
        for row_pairs in parsed_rows:
            new_row = etree.SubElement(body, "tr")
            new_row.set("valign", "bottom")
            for label, value in row_pairs:
                _add_memo_label_cell(new_row, label)
                _add_memo_value_cell(new_row, value)
        table.addprevious(new_table)
        parent = table.getparent()
        if parent is not None:
            parent.remove(table)


def _relax_print_tables(root):
    """Let memo header tables size like the Word editor.

    LibreOffice honors tiny Word `width="96"` columns so 'ส่วนราชการ'
    wraps to ส่วน / ราชการ. Browsers use auto layout and keep one line.
    """
    for table in root.xpath(".//table"):
        if table.get("data-vpk-memo-header") or table.get("data-vpk-fixed-table"):
            continue
        table.attrib.pop("width", None)
        table.set("border", "0")
        style = re.sub(r"width\s*:\s*[^;]+;?", "", table.get("style") or "")
        table.set(
            "style",
            (style + "; width: 100%; border-collapse: collapse; border: none")
            .strip()
            .strip(";"),
        )
        for colgroup in list(table.xpath("./colgroup")):
            parent = colgroup.getparent()
            if parent is not None:
                parent.remove(colgroup)
        for col in list(table.xpath(".//col")):
            parent = col.getparent()
            if parent is not None:
                parent.remove(col)
        for cell in table.xpath(".//td | .//th"):
            cell.attrib.pop("width", None)
            style = re.sub(r"(^|;)\s*width\s*:\s*[^;]+", "", cell.get("style") or "")
            label = _cell_plain_text(cell)
            if label in _MEMO_LABELS:
                if "white-space" not in style.lower():
                    style = (style + "; white-space: nowrap").strip("; ")
                min_width = _MEMO_LABEL_MIN_WIDTH.get(label)
                if min_width and "min-width" not in style.lower():
                    style = (style + "; min-width: %s" % min_width).strip("; ")
            cell.set("style", style.strip().strip(";"))


def _read_style_cm(span, attr, css_prop, style):
    raw = span.get(attr)
    if raw:
        try:
            return float(raw)
        except ValueError:
            pass
    match = re.search(r"(?:^|;)\s*%s\s*:\s*([\d.]+)\s*cm" % css_prop, style or "", re.I)
    return float(match.group(1)) if match else None


def _set_style_prop(style, prop, value):
    style = re.sub(r"(?:^|;)\s*%s\s*:\s*[^;]*" % prop, "", style or "", flags=re.I)
    style = style.strip().strip(";")
    extra = "%s: %s" % (prop, value)
    return (style + "; " + extra).strip("; ") if style else extra


_PINNED_SPAN_XPATH = (
    ".//*[contains(concat(' ', normalize-space(@class), ' '), ' o_vpk_we_fill_line ') or "
    "contains(concat(' ', normalize-space(@class), ' '), ' o_vpk_we_field_box ')]"
)
_FIELD_BOX_XPATH = (
    ".//*[contains(concat(' ', normalize-space(@class), ' '), ' o_vpk_we_field_box ')]"
)
EMPTY_FIELD_LABEL = "ว่าง"


def _fill_empty_field_boxes(root):
    """Put the default label in field boxes that have no value."""
    for span in list(root.xpath(_FIELD_BOX_XPATH)):
        text = re.sub(r"[\xa0\u200b\s]", "", "".join(span.itertext()) or "")
        if text:
            continue
        span.text = EMPTY_FIELD_LABEL
        for child in list(span):
            span.remove(child)
        cls = span.get("class") or ""
        if "o_vpk_we_field_empty" not in cls.split():
            span.set("class", ("%s o_vpk_we_field_empty" % cls).strip())


def fill_empty_field_boxes_html(body_html):
    if not body_html:
        return body_html
    try:
        root = lxml_html.fromstring("<div>%s</div>" % body_html)
    except Exception:
        return body_html
    _fill_empty_field_boxes(root)
    parts = [root.text or ""]
    parts += [etree.tostring(child, encoding="unicode", method="html") for child in root]
    return "".join(parts)


_LINE_NUMBER_PREFIX_RE = re.compile(r"^\s*[0-9๐-๙]+\s*[.)]\s*")
_THAI_DIGITS = str.maketrans("0123456789", "๐๑๒๓๔๕๖๗๘๙")


_ROW_TOKEN_RE = re.compile(r"\{\{\s*([a-z_][a-z0-9_.]*)\s*\}\}")
ROW_TOKENS = (
    "no", "no_th", "line", "name", "qty", "uom", "unit_price", "amount", "spec"
)


def _replace_row_tokens(row, tokens):
    def sub(text):
        if not text or "{{" not in text:
            return text
        def pick(match):
            key = match.group(1)
            if key in tokens:
                return tokens[key]
            return "" if key.startswith("line.") else match.group(0)

        return _ROW_TOKEN_RE.sub(pick, text)

    row.text = sub(row.text)
    for node in row.iterdescendants():
        node.text = sub(node.text)
        node.tail = sub(node.tail)


def expand_repeat_rows_html(body_html, mapping, rows=None):
    """Expand `<tr data-vpk-repeat="code">` into one row per item of that field.

    `rows` holds structured items per field code (e.g. PR lines with name, qty,
    uom, unit_price, amount). Without it the field text is split by line and
    only `{{no}}`, `{{no_th}}`, `{{line}}`, `{{name}}` are filled.
    """
    if not body_html or "data-vpk-repeat" not in body_html:
        return body_html
    try:
        root = lxml_html.fromstring("<div>%s</div>" % body_html)
    except Exception:
        return body_html
    rows = rows or {}
    for row in list(root.xpath(".//tr[@data-vpk-repeat]")):
        code = row.get("data-vpk-repeat")
        items = list(rows.get(code) or [])
        if not items:
            raw = mapping.get(code)
            raw = "" if raw is None or raw is False else str(raw)
            for index, line in enumerate(
                [ln for ln in raw.splitlines() if ln.strip()], start=1
            ):
                text = _LINE_NUMBER_PREFIX_RE.sub("", line).strip()
                items.append({"no": str(index), "line": text, "name": text})
        parent = row.getparent()
        if parent is None:
            continue
        position = parent.index(row)
        row.attrib.pop("data-vpk-repeat", None)
        template_row = copy.deepcopy(row)
        parent.remove(row)
        if not items:
            items = [{}]
        for offset, item in enumerate(items):
            new_row = copy.deepcopy(template_row)
            tokens = {key: "" for key in ROW_TOKENS}
            tokens.update({k: str(v) for k, v in item.items() if isinstance(v, (str, int))})
            tokens["no_th"] = tokens.get("no", "").translate(_THAI_DIGITS)
            tail = new_row.tail
            _replace_row_tokens(new_row, tokens)
            new_row.tail = tail
            parent.insert(position + offset, new_row)
    parts = [root.text or ""]
    parts += [etree.tostring(child, encoding="unicode", method="html") for child in root]
    return "".join(parts)


def _lift_absolute_fill_lines(root):
    """Move pinned fill-lines and field boxes out of tables so print coords use the page origin."""
    for span in list(root.xpath(_PINNED_SPAN_XPATH)):
        style = span.get("style") or ""
        if "absolute" not in style.lower() and not span.get("data-line-left"):
            continue
        parent = span.getparent()
        if parent is None or parent is root:
            continue
        root.append(span)
        if parent.tag in ("td", "th"):
            leftover = re.sub(r"[\xa0\u200b\s]", "", "".join(parent.itertext()) or "")
            if not leftover and parent.find(".//*") is None:
                parent.text = "\xa0"


def _shift_absolute_fill_lines_for_print(root, margins):
    """Editor left/top are paper-relative; PDF origin is inside @page margins."""
    if not margins:
        return
    left_m = float(margins.get("left") or 0)
    top_m = float(margins.get("top") or 0)
    for span in root.xpath(_PINNED_SPAN_XPATH):
        style = span.get("style") or ""
        if "absolute" not in style.lower() and not span.get("data-line-left"):
            continue
        left = _read_style_cm(span, "data-line-left", "left", style)
        top = _read_style_cm(span, "data-line-top", "top", style)
        if left is None and top is None:
            continue
        if left is not None:
            left = max(0, round((left - left_m) * 100) / 100)
            span.set("data-line-left", str(left))
            style = _set_style_prop(style, "left", "%scm" % left)
        if top is not None:
            top = max(0, round((top - top_m) * 100) / 100)
            span.set("data-line-top", str(top))
            style = _set_style_prop(style, "top", "%scm" % top)
        style = _set_style_prop(style, "position", "absolute")
        span.set("style", style)


def _prepare_fill_lines_for_export(body_html, convert_fill_lines=True, margins=None):
    """Normalize editor HTML before LibreOffice export.

    When `convert_fill_lines` is True (Word export), fill-line spans become
    `[[VPK_LINE:...]]` tokens. For PDF we keep the CSS spans so the printout
    matches the Word editor.
    """
    if not body_html:
        return body_html
    try:
        root = lxml_html.fromstring("<div>%s</div>" % body_html)
    except Exception:
        return body_html

    _normalize_html_fonts(root)
    _convert_underline_cells_to_fill_lines(root)
    _rewrite_memo_header_tables(root)
    _relax_print_tables(root)
    _fill_empty_field_boxes(root)
    _lift_absolute_fill_lines(root)
    if not convert_fill_lines:
        _shift_absolute_fill_lines_for_print(root, margins)

    # 1. Remove empty floating Frame spans left over from Word imports
    for frame in list(root.xpath(".//span[starts-with(@id, 'Frame')]")):
        if not (frame.text_content() or "").strip() and len(frame) == 0:
            parent = frame.getparent()
            if parent is not None:
                parent.remove(frame)

    # 2. Prevent <img align="left|right"> from becoming a floating anchor that
    # overlaps the next paragraph, and clear fixed line-height on image paragraphs
    for img in root.xpath(".//img"):
        img_align = (img.get("align") or "").lower()
        parent_p = img.getparent()
        while parent_p is not None and parent_p.tag not in ("p", "div", "td", "th"):
            parent_p = parent_p.getparent()
        if img_align in ("left", "right"):
            if parent_p is not None and parent_p.tag in ("p", "div") and not (parent_p.get("style") or "").find("text-align") >= 0:
                parent_p.set("align", img_align)
            img.attrib.pop("align", None)
        if parent_p is not None and parent_p.tag in ("p", "div"):
            p_style = parent_p.get("style") or ""
            if "line-height" in p_style:
                p_style = re.sub(r"line-height\s*:\s*[^;]+;?", "", p_style)
                parent_p.set("style", p_style.strip())

    # 3. Sync inline style="text-align: ..." with HTML align="..." attribute
    # because LibreOffice Writer prefers the HTML align attribute over CSS
    for blk in root.xpath(".//p | .//div | .//h1 | .//h2 | .//h3 | .//h4"):
        st = blk.get("style") or ""
        m = re.search(r"text-align\s*:\s*(left|center|right|justify)", st, re.IGNORECASE)
        if m:
            blk.set("align", m.group(1).lower())

    # 3b. Convert `<span class="o_vpk_we_tab">` into non-breaking spaces for Word
    for tab_span in list(root.xpath(".//*[contains(@class, 'o_vpk_we_tab')]")):
        if not (tab_span.text_content() or ""):
            tab_span.text = "\u00a0" * 8

    for node in root.xpath(".//*[@style]"):
        style = node.get("style") or ""
        if "min-height" in style:
            node.set("style", re.sub(r"min-height\s*:\s*[^;]+;?", "", style).strip())

    # 4. Replace `<span class="o_vpk_we_fill_line">` with a marker token that
    # `apply_page_setup` converts into a fixed-width bottom-bordered table row
    # so lines sit cleanly below Thai descenders and align to the right margin.
    for span in list(root.xpath(".//*[contains(@class, 'o_vpk_we_fill_line')]")):
        if not convert_fill_lines:
            continue
        style_attr = span.get("style") or ""
        try:
            cm = float(span.get("data-line-cm") or 0)
        except ValueError:
            cm = 0.0
        if not cm:
            m = re.search(r"(?:min-width|width)\s*:\s*(\d+(?:\.\d+)?)cm", style_attr)
            cm = float(m.group(1)) if m else 6.0
        line_style = (span.get("data-line-style") or "").strip()
        if line_style not in ("dotted", "solid"):
            line_style = "solid" if "solid" in style_attr else "dotted"
        text = (span.text_content() or "").replace("\u00a0", " ").replace("\u200b", "").strip()
        token = "[[VPK_LINE:%.2f:%s:%s]]" % (cm, line_style, text)
        tail = (span.tail or "").replace("\u200b", "")
        parent = span.getparent()
        if parent is not None:
            prev = span.getprevious()
            if prev is not None:
                prev.tail = (prev.tail or "") + token + tail
            else:
                parent.text = (parent.text or "") + token + tail
            parent.remove(span)
    parts = [root.text or ""]
    parts += [etree.tostring(child, encoding="unicode", method="html") for child in root]
    return "".join(parts)


def build_print_html(body_html, font, font_size, margins, font_face_css=""):
    """Wrap editor HTML for Save-as-PDF using the same A4 page as the Word editor.

    Absolute fill-lines/field boxes stay paper-relative, matching the editor:
    the page padding is the margin, and @page margin is 0.
    """
    export_body = _prepare_fill_lines_for_export(
        body_html or "", convert_fill_lines=False
    )
    return """<!DOCTYPE html>
<html>
<head>
<meta http-equiv="content-type" content="text/html; charset=utf-8"/>
<style type="text/css">
{faces}
@page {{
  size: {width}cm {height}cm;
  margin: 0;
}}
html {{
  margin: 0;
  padding: 0;
  background: #fff;
}}
body {{
  width: {width}cm;
  min-height: {height}cm;
  box-sizing: border-box;
  margin: 0;
  padding: {top}cm {right}cm {bottom}cm {left}cm;
  position: relative;
  background: #fff;
  color: #000;
}}
body, p, td, th, li, span {{
  font-family: "{font}", "TH Sarabun New", sans-serif;
  font-size: {size}pt;
  line-height: 1.15;
}}
p, p.western, p.cjk, p.ctl {{
  margin-top: 0;
  margin-bottom: 0;
  font-size: {size}pt;
  font-family: "{font}", "TH Sarabun New", sans-serif;
}}
td, th {{
  height: auto;
  font-size: {size}pt;
}}
b, strong {{
  font-size: inherit;
}}
img {{
  max-width: 100%;
  border: 0;
}}
table {{
  width: 100%;
  table-layout: auto;
  border-collapse: collapse;
  border: none;
}}
.o_vpk_we_fill_line {{
  display: inline-block;
  vertical-align: baseline;
  padding: 0 4px;
  line-height: 1.05;
  min-height: 1.05em;
  border-bottom: 1px dotted #000;
}}
.o_vpk_we_fill_line[data-line-style="solid"] {{
  border-bottom-style: solid;
}}
.o_vpk_we_fill_line[data-line-left],
.o_vpk_we_fill_line[style*="absolute"],
.o_vpk_we_field_box[data-line-left],
.o_vpk_we_field_box[style*="absolute"] {{
  position: absolute;
}}
.o_vpk_we_field_box {{
  display: inline-block;
  box-sizing: border-box;
  white-space: nowrap;
  line-height: 1.05;
  padding: 0 4px;
  margin: 0;
  background: transparent !important;
  border: 1px solid transparent;
  color: inherit !important;
}}
.o_vpk_we_line_frame,
.o_vpk_we_img_frame,
.o_vpk_we_guide_line {{
  display: none !important;
}}
.o_vpk_we_caret_sink {{
  display: inline;
  width: 0;
  min-width: 0;
  padding: 0;
  margin: 0;
}}
.o_vpk_page_break {{
  page-break-before: always;
}}
</style>
</head>
<body>
{body}
</body>
</html>
""".format(
        faces=font_face_css or "",
        width=A4_WIDTH_CM,
        height=A4_HEIGHT_CM,
        top=margins["top"],
        bottom=margins["bottom"],
        left=margins["left"],
        right=margins["right"],
        font=font,
        size=font_size,
        body=export_body,
    )


def build_layout_document(body_html, font, font_size, margins):
    """Wrap editor HTML in a full document LibreOffice can import as A4."""
    export_body = _prepare_fill_lines_for_export(body_html or "", convert_fill_lines=True)
    return """<!DOCTYPE html>
<html>
<head>
<meta http-equiv="content-type" content="text/html; charset=utf-8"/>
<style type="text/css">
@page {{ size: {width}cm {height}cm; margin-top: {top}cm; margin-bottom: {bottom}cm; margin-left: {left}cm; margin-right: {right}cm }}
body, p, td, th, li {{ font-family: "{font}"; font-size: {size}pt }}
p {{ margin-top: 0; margin-bottom: 0; line-height: 100% }}
</style>
</head>
<body>
{body}
</body>
</html>
""".format(
        width=A4_WIDTH_CM,
        height=A4_HEIGHT_CM,
        top=margins["top"],
        bottom=margins["bottom"],
        left=margins["left"],
        right=margins["right"],
        font=font,
        size=font_size,
        body=export_body,
    )


def _mirror_complex_script_formatting(body, qn):
    """LibreOffice stores Thai bold/italic/size only as complex-script
    properties; copy them so Latin text in the same run matches.
    Also convert run-level `<w:bdr>` (which renders as a 4-sided box around
    text) into `<w:u w:val="dotted|single"/>` so fill-lines render as clean
    bottom underlines instead of boxes."""
    pairs = (("w:bCs", "w:b"), ("w:iCs", "w:i"), ("w:szCs", "w:sz"))
    for rpr in body.iter(qn("w:rPr")):
        for cs_tag, latin_tag in pairs:
            cs_el = rpr.find(qn(cs_tag))
            if cs_el is None or rpr.find(qn(latin_tag)) is not None:
                continue
            latin_el = etree.Element(qn(latin_tag))
            for key, value in cs_el.attrib.items():
                latin_el.set(key, value)
            cs_el.addprevious(latin_el)
        bdr = rpr.find(qn("w:bdr"))
        if bdr is not None:
            val = (bdr.get(qn("w:val")) or "dotted").lower()
            u_val = "dotted" if val in ("dotted", "dashed", "dotdash") else "single"
            rpr.remove(bdr)
            u_el = rpr.find(qn("w:u"))
            if u_el is None:
                u_el = etree.SubElement(rpr, qn("w:u"))
            u_el.set(qn("w:val"), u_val)


_THAI_COMBINE_RE = re.compile(r"[\u0e31\u0e34-\u0e3a\u0e47-\u0e4e]")


def _visible_label_width_cm(text, font_size=16.0, is_middle=False):
    clean = _THAI_COMBINE_RE.sub("", (text or "").replace("\u200b", "").replace("\u00a0", " "))
    base = len(clean.strip()) * 0.25 * (font_size / 16.0)
    extra = 0.50 if is_middle else 0.35
    return round(max(0.75, base + extra), 2)


def _convert_fill_line_paragraphs_to_tables(document, font, font_size, printable_cm):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Pt

    for paragraph in list(document.paragraphs):
        full_text = "".join(r.text or "" for r in paragraph.runs)
        if "[[VPK_LINE:" not in full_text:
            continue
        p_bold = False
        para_pt = float(font_size or 16.0)
        for r in paragraph.runs:
            if not (r.text or "").strip():
                continue
            rpr = r._element.find(qn("w:rPr"))
            if rpr is not None:
                if rpr.find(qn("w:b")) is not None or rpr.find(qn("w:bCs")) is not None:
                    p_bold = True
                sz_el = rpr.find(qn("w:sz"))
                if sz_el is None:
                    sz_el = rpr.find(qn("w:szCs"))
                if sz_el is not None and sz_el.get(qn("w:val")):
                    try:
                        para_pt = float(sz_el.get(qn("w:val"))) / 2.0
                    except ValueError:
                        pass
            if r.bold:
                p_bold = True
        segments = []
        pos = 0
        for m in _FILL_TOKEN_RE.finditer(full_text):
            before = full_text[pos : m.start()].replace("\u200b", "").replace("\u00a0", " ").strip()
            if before:
                is_mid = len(segments) > 0
                lbl_txt = ("  " + before) if is_mid else before
                segments.append(("label", lbl_txt, _visible_label_width_cm(before, para_pt, is_middle=is_mid), None))
            line_cm = max(1.0, float(m.group(1)))
            line_style = m.group(2)
            line_text = m.group(3).strip()
            segments.append(("line", line_text, line_cm, line_style))
            pos = m.end()
        after = full_text[pos:].replace("\u200b", "").replace("\u00a0", " ").strip()
        if after:
            segments.append(("label", "  " + after, _visible_label_width_cm(after, para_pt, is_middle=True), None))
        if not segments:
            continue

        # Fit total width to printable_cm so the last fill-line aligns cleanly
        # with the right margin when the user stretched it near the end of the line.
        label_sum = sum(s[2] for s in segments if s[0] == "label")
        line_indices = [idx for idx, s in enumerate(segments) if s[0] == "line"]
        line_sum = sum(segments[idx][2] for idx in line_indices)
        available_for_lines = max(2.0, printable_cm - label_sum)
        if line_indices:
            if line_sum > available_for_lines or (label_sum + line_sum) >= printable_cm * 0.78:
                if len(line_indices) == 1:
                    idx = line_indices[0]
                    kind, txt, _w, st = segments[idx]
                    segments[idx] = (kind, txt, round(available_for_lines, 2), st)
                else:
                    scale = available_for_lines / max(0.1, line_sum)
                    used = 0.0
                    for i_pos, idx in enumerate(line_indices):
                        kind, txt, w, st = segments[idx]
                        if i_pos == len(line_indices) - 1:
                            new_w = round(max(1.0, available_for_lines - used), 2)
                        else:
                            new_w = round(max(1.0, w * scale), 2)
                            used += new_w
                        segments[idx] = (kind, txt, new_w, st)

        total_cm = sum(s[2] for s in segments)
        tbl = document.add_table(rows=1, cols=len(segments))
        tbl.autofit = False
        tblPr = tbl._element.xpath("./w:tblPr")[0]
        tblPr.append(
            parse_xml(f'<w:tblW {nsdecls("w")} w:w="{int(round(total_cm * 567))}" w:type="dxa"/>')
        )
        tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        tblPr.append(
            parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                '<w:top w:val="none"/><w:left w:val="none"/>'
                '<w:bottom w:val="none"/><w:right w:val="none"/>'
                '<w:insideH w:val="none"/><w:insideV w:val="none"/>'
                "</w:tblBorders>"
            )
        )
        tblPr.append(
            parse_xml(
                f'<w:tblCellMar {nsdecls("w")}>'
                '<w:top w:w="0" w:type="dxa"/><w:left w:w="15" w:type="dxa"/>'
                '<w:bottom w:w="0" w:type="dxa"/><w:right w:w="15" w:type="dxa"/>'
                "</w:tblCellMar>"
            )
        )
        tblGrid = tbl._element.xpath("./w:tblGrid")[0]
        for c in list(tblGrid):
            tblGrid.remove(c)
        row = tbl.rows[0]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:trHeight {nsdecls("w")} w:val="340" w:hRule="atLeast"/>'))

        for idx, (kind, txt, w_cm, lstyle) in enumerate(segments):
            dxa = int(round(w_cm * 567))
            tblGrid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="{dxa}"/>'))
            cell = row.cells[idx]
            tcPr = cell._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:tcW")):
                tcPr.remove(old)
            tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{dxa}" w:type="dxa"/>'))
            tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="bottom"/>'))
            if kind == "line":
                bval = "single" if lstyle == "solid" else "dotted"
                tcPr.append(
                    parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        '<w:top w:val="none"/><w:left w:val="none"/><w:right w:val="none"/>'
                        f'<w:bottom w:val="{bval}" w:sz="6" w:space="0" w:color="000000"/>'
                        "</w:tcBorders>"
                    )
                )
            else:
                tcPr.append(
                    parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        '<w:top w:val="none"/><w:left w:val="none"/>'
                        '<w:right w:val="none"/><w:bottom w:val="none"/>'
                        "</w:tcBorders>"
                    )
                )
            cp = cell.paragraphs[0]
            cpPr = cp._element.get_or_add_pPr()
            cpPr.append(
                parse_xml(
                    f'<w:spacing {nsdecls("w")} w:before="40" w:after="0" w:line="320" w:lineRule="exact"/>'
                )
            )
            run = cp.add_run((" " + txt) if (kind == "line" and txt) else txt)
            run.bold = bool(kind == "label" and p_bold)
            run.font.name = font
            run.font.size = Pt(para_pt)
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(attr), font)
            sz_cs = rpr.find(qn("w:szCs"))
            if sz_cs is None:
                sz_cs = etree.SubElement(rpr, qn("w:szCs"))
            sz_cs.set(qn("w:val"), str(int(round(para_pt * 2))))
            if run.bold:
                if rpr.find(qn("w:bCs")) is None:
                    etree.SubElement(rpr, qn("w:bCs"))

        paragraph._element.addprevious(tbl._element)
        paragraph._element.getparent().remove(paragraph._element)


def apply_page_setup(docx_bytes, font, font_size, margins):
    """Force A4, margins and the default Thai font on a generated .docx."""
    _require_docx()
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    try:
        document = Document(io.BytesIO(docx_bytes))
    except Exception as error:
        raise DocxMergeError("ไฟล์ Word ที่สร้างจาก layout ไม่ถูกต้อง") from error
    for section in document.sections:
        section.page_width = Cm(A4_WIDTH_CM)
        section.page_height = Cm(A4_HEIGHT_CM)
        section.top_margin = Cm(margins["top"])
        section.bottom_margin = Cm(margins["bottom"])
        section.left_margin = Cm(margins["left"])
        section.right_margin = Cm(margins["right"])

    normal = document.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(font_size)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)
    size_cs = rpr.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = etree.SubElement(rpr, qn("w:szCs"))
    size_cs.set(qn("w:val"), str(int(round(font_size * 2))))
    printable_cm = max(8.0, A4_WIDTH_CM - margins["left"] - margins["right"])
    _convert_fill_line_paragraphs_to_tables(document, font, font_size, printable_cm)
    _mirror_complex_script_formatting(document.element.body, qn)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
