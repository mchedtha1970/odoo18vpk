# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

import base64
import io

from odoo.exceptions import UserError
from odoo.tests.common import TransactionCase, tagged


@tagged("post_install", "-at_install")
class TestOnlyOfficeTemplate(TransactionCase):
    def test_edit_requires_word_file(self):
        template = self.env["vpk.document.studio.template"].create(
            {"name": "OnlyOffice empty", "code": "oo_empty"}
        )
        with self.assertRaises(UserError):
            template.action_edit_with_onlyoffice()

    def test_edit_opens_docx_attachment(self):
        import docx

        document = docx.Document()
        document.add_paragraph("เรื่อง {{subject}}")
        buf = io.BytesIO()
        document.save(buf)
        template = self.env["vpk.document.studio.template"].create(
            {
                "name": "OnlyOffice memo",
                "code": "oo_memo",
                "datas": base64.b64encode(buf.getvalue()),
                "filename": "memo.docx",
            }
        )
        action = template.action_edit_with_onlyoffice()
        self.assertEqual(action["type"], "ir.actions.client")
        self.assertEqual(action["tag"], "onlyoffice_editor")
        self.assertEqual(action["target"], "current")
        self.assertTrue(action["params"]["attachment_id"])
        attachment = template._onlyoffice_word_attachment()
        self.assertTrue(attachment.name.lower().endswith(".docx"))
        self.assertFalse(attachment.res_field)
        self.assertTrue(attachment.raw)
        self.assertGreater(len(attachment.raw), 100)

    def test_onlyoffice_save_updates_word_not_html_layout(self):
        import docx

        original = docx.Document()
        original.add_paragraph("ต้นฉบับ {{subject}}")
        buf = io.BytesIO()
        original.save(buf)
        template = self.env["vpk.document.studio.template"].create(
            {
                "name": "OnlyOffice save",
                "code": "oo_save",
                "datas": base64.b64encode(buf.getvalue()),
                "filename": "memo.docx",
                "layout_html": "<p>ตัวอย่างเก่า</p>",
            }
        )
        attachment = template._onlyoffice_word_attachment()
        updated = docx.Document()
        updated.add_paragraph("แก้ไขแล้ว {{subject}}")
        out = io.BytesIO()
        updated.save(out)
        self.env["ir.attachment"].browse(attachment.id).write({"raw": out.getvalue()})
        self.assertIn("ตัวอย่างเก่า", template.layout_html or "")
        saved = docx.Document(io.BytesIO(base64.b64decode(template.datas)))
        self.assertIn("แก้ไขแล้ว", "\n".join(p.text for p in saved.paragraphs))
