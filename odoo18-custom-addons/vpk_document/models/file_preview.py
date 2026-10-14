# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

import base64
import io
import json
import os
from datetime import date, datetime

from markupsafe import Markup, escape

MAX_SHEETS = 30
MAX_ROWS = 400
MAX_COLS = 40
MAX_IMAGE_BYTES = 4 * 1024 * 1024

EXT_TO_TYPE = {
    ".pdf": "pdf",
    ".xls": "excel",
    ".xlsx": "excel",
    ".xlsm": "excel",
    ".doc": "word",
    ".docx": "word",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".webp": "image",
    ".bmp": "image",
    ".tif": "image",
    ".tiff": "image",
}

EXT_TO_MIMETYPE = {
    ".pdf": "application/pdf",
    ".xls": "application/vnd.ms-excel",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}

TYPE_LABELS = {
    "pdf": "PDF",
    "excel": "Excel",
    "word": "Word",
    "image": "รูปภาพ",
}

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def extension_of(filename):
    return os.path.splitext(filename or "")[1].lower()


def file_type_from_name(filename):
    return EXT_TO_TYPE.get(extension_of(filename)) or False


def mimetype_from_name(filename):
    return EXT_TO_MIMETYPE.get(extension_of(filename)) or False


def validate_file_content(filename, raw):
    """Raise ValueError when the upload is not a supported document."""
    if not filename:
        raise ValueError("กรุณาระบุชื่อไฟล์")
    file_type = file_type_from_name(filename)
    if not file_type:
        raise ValueError(
            "รองรับเฉพาะไฟล์ PDF, Excel (.xls, .xlsx), Word (.doc, .docx) "
            "และรูปภาพ (PNG, JPG, GIF, WEBP, BMP, TIFF)"
        )
    if not raw:
        raise ValueError("ไฟล์ว่างเปล่า")
    if not _magic_ok(extension_of(filename), raw):
        raise ValueError("เนื้อหาไฟล์ไม่ตรงกับชนิด %s" % TYPE_LABELS[file_type])
    return file_type


def excel_workbook(raw, ext):
    """Parse a spreadsheet into sheet/row/cell data for the web grid."""
    try:
        if ext == ".xls":
            sheets, truncated, sheet_count = _xls_sheets(raw)
        else:
            sheets, truncated, sheet_count = _xlsx_sheets(raw)
    except Exception as error:  # noqa: BLE001 - show the failure inside the viewer
        return {
            "title": "",
            "truncated": False,
            "sheetCount": 0,
            "maxRows": MAX_ROWS,
            "maxCols": MAX_COLS,
            "error": "ไม่สามารถเปิดไฟล์ Excel ได้: %s" % error,
            "sheets": [],
        }
    return {
        "title": "",
        "truncated": truncated,
        "sheetCount": sheet_count,
        "maxRows": MAX_ROWS,
        "maxCols": MAX_COLS,
        "error": "",
        "sheets": sheets,
    }


def sheet_viewer_page(title, workbook):
    """Google Sheets-style page. Cell text is JSON, not HTML."""
    workbook = dict(workbook)
    workbook["title"] = title or workbook.get("title") or ""
    payload = json.dumps(workbook, ensure_ascii=False)
    payload = (
        payload.replace("&", "\\u0026")
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
    )
    html = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>__TITLE__</title>
<link rel="stylesheet" href="/vpk_document/static/src/css/sheet_viewer.css"/>
</head>
<body>
<div class="gs-app">
  <header class="gs-top">
    <div class="gs-logo" aria-hidden="true"></div>
    <div class="gs-title" id="doc-title"></div>
  </header>
  <div class="gs-formula">
    <div class="gs-namebox" id="namebox">A1</div>
    <div class="gs-fx">fx</div>
    <div class="gs-formula-value" id="formula"></div>
  </div>
  <div class="gs-grid-wrap" id="grid"></div>
  <div class="gs-tabs" id="tabs"></div>
</div>
<script type="application/json" id="workbook-data">__DATA__</script>
<script src="/vpk_document/static/src/js/sheet_viewer.js"></script>
</body>
</html>
"""
    safe_title = str(escape(title or "Excel"))
    return Markup(html.replace("__TITLE__", safe_title).replace("__DATA__", payload))


def viewer_page(title, body):
    """Full HTML page so the browser displays the file instead of downloading it."""
    return Markup(
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\"/>"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\"/>"
        "<title>{}</title>"
        "<style>"
        "body {{ margin: 0; background: #f4f5f7; color: #1f2328; "
        "font-family: Sarabun, \"Segoe UI\", sans-serif; }}"
        "main {{ padding: 16px 20px 40px; }}"
        "</style></head><body><main>{}</main></body></html>"
    ).format(escape(title or ""), body)


def render_preview(filename, raw, content_url=None, download_url=None):
    file_type = file_type_from_name(filename)
    ext = extension_of(filename)
    if file_type == "pdf":
        body = _pdf_html(filename, content_url)
    elif file_type == "image":
        body = _image_html(filename, raw, content_url)
    elif file_type == "excel":
        body = _excel_html(raw, ext)
    elif file_type == "word":
        body = _word_html(raw, ext)
    else:
        body = Markup(
            '<div class="alert alert-warning">ไม่รองรับการแสดงไฟล์นี้</div>'
        )
    link = Markup("")
    if download_url:
        link = Markup(
            ' <a href="{}" target="_blank">ดาวน์โหลด</a>'
        ).format(escape(download_url))
    return Markup(
        '<div class="mb-2"><strong>{}</strong> '
        '<span class="badge text-bg-secondary">{}</span>{}</div>{}'
    ).format(
        escape(filename or ""),
        escape(TYPE_LABELS.get(file_type, "")),
        link,
        body,
    )


def _magic_ok(ext, raw):
    if ext == ".pdf":
        return raw.startswith(b"%PDF")
    if ext == ".png":
        return raw.startswith(b"\x89PNG\r\n\x1a\n")
    if ext in (".jpg", ".jpeg"):
        return raw.startswith(b"\xff\xd8\xff")
    if ext == ".gif":
        return raw.startswith((b"GIF87a", b"GIF89a"))
    if ext == ".webp":
        return len(raw) >= 12 and raw.startswith(b"RIFF") and raw[8:12] == b"WEBP"
    if ext == ".bmp":
        return raw.startswith(b"BM")
    if ext in (".tif", ".tiff"):
        return raw.startswith((b"II*\x00", b"MM\x00*"))
    if ext in (".xlsx", ".xlsm", ".docx"):
        return raw.startswith(b"PK")
    if ext in (".xls", ".doc"):
        return raw.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    return False


def _pdf_html(filename, content_url):
    if not content_url:
        return Markup(
            '<div class="alert alert-info">บันทึกเอกสารก่อนจึงจะแสดงตัวอย่าง PDF ได้</div>'
        )
    return Markup(
        '<iframe src="{}" title="{}" '
        'style="width:100%;height:75vh;border:1px solid #dee2e6;background:#fff;">'
        "</iframe>"
    ).format(escape(content_url), escape(filename or "PDF"))


def _image_html(filename, raw, content_url):
    if content_url:
        src = content_url
    else:
        mime = mimetype_from_name(filename) or "image/png"
        src = "data:%s;base64,%s" % (mime, base64.b64encode(raw).decode())
    return Markup(
        '<div style="text-align:center;background:#f8f9fa;padding:12px;">'
        '<img src="{}" alt="{}" '
        'style="max-width:100%;max-height:75vh;object-fit:contain;"/>'
        "</div>"
    ).format(escape(src), escape(filename or "image"))


def _excel_html(raw, ext):
    try:
        if ext == ".xls":
            return _xls_html(raw)
        return _xlsx_html(raw)
    except Exception as error:  # noqa: BLE001 - show a readable preview failure
        return Markup(
            '<div class="alert alert-warning">ไม่สามารถแสดงตัวอย่าง Excel ได้: {}</div>'
        ).format(escape(str(error)))


def _word_html(raw, ext):
    if ext == ".doc":
        return Markup(
            '<div class="alert alert-info">'
            "ไฟล์ Word รุ่นเก่า (.doc) ไม่สามารถแสดงตัวอย่างในหน้าจอได้ "
            "กรุณาดาวน์โหลดเพื่อเปิดด้วยโปรแกรม Word"
            "</div>"
        )
    try:
        return _docx_html(raw)
    except Exception as error:  # noqa: BLE001 - show a readable preview failure
        return Markup(
            '<div class="alert alert-warning">ไม่สามารถแสดงตัวอย่าง Word ได้: {}</div>'
        ).format(escape(str(error)))


def _xlsx_sheets(raw):
    import openpyxl
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=False)
    try:
        worksheets = workbook.worksheets
        truncated = len(worksheets) > MAX_SHEETS
        sheets = []
        for sheet in worksheets[:MAX_SHEETS]:
            rich, row_truncated = _xlsx_rich_sheet(sheet, get_column_letter)
            truncated = truncated or row_truncated
            sheets.append(rich)
        return sheets, truncated, len(worksheets)
    finally:
        workbook.close()


def _xlsx_rich_sheet(sheet, get_column_letter):
    last_row, last_col = _used_bounds(sheet)
    truncated = (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLS
    merges, covered = _merge_map(sheet, last_row, last_col)
    rows = []
    for row_index in range(1, last_row + 1):
        line = []
        for col_index in range(1, last_col + 1):
            if (row_index, col_index) in covered:
                line.append({"s": 1})
                continue
            cell = sheet.cell(row_index, col_index)
            span = merges.get((row_index, col_index), (1, 1))
            line.append(_style_cell(cell, span[0], span[1]))
        rows.append(line)
    widths = []
    for col_index in range(1, last_col + 1):
        dimension = sheet.column_dimensions.get(get_column_letter(col_index))
        widths.append(_column_px(dimension.width if dimension else None))
    heights = []
    for row_index in range(1, last_row + 1):
        dimension = sheet.row_dimensions.get(row_index)
        heights.append(_row_px(dimension.height if dimension and dimension.height else None))
    return {
        "name": sheet.title,
        "widths": widths,
        "heights": heights,
        "rows": rows,
    }, truncated


def _used_bounds(sheet):
    last_row, last_col = 1, 1
    for (row_index, col_index), cell in getattr(sheet, "_cells", {}).items():
        if row_index > MAX_ROWS or col_index > MAX_COLS:
            continue
        if cell.value not in (None, "") or _fill_hex(cell):
            last_row = max(last_row, row_index)
            last_col = max(last_col, col_index)
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= MAX_ROWS and merged.min_col <= MAX_COLS:
            last_row = max(last_row, min(merged.max_row, MAX_ROWS))
            last_col = max(last_col, min(merged.max_col, MAX_COLS))
    return last_row, last_col


def _merge_map(sheet, last_row, last_col):
    origins = {}
    covered = set()
    for merged in sheet.merged_cells.ranges:
        if merged.min_row > last_row or merged.min_col > last_col:
            continue
        row_end = min(merged.max_row, last_row)
        col_end = min(merged.max_col, last_col)
        rowspan = row_end - merged.min_row + 1
        colspan = col_end - merged.min_col + 1
        if rowspan < 1 or colspan < 1:
            continue
        origins[(merged.min_row, merged.min_col)] = (rowspan, colspan)
        for row_index in range(merged.min_row, row_end + 1):
            for col_index in range(merged.min_col, col_end + 1):
                if (row_index, col_index) != (merged.min_row, merged.min_col):
                    covered.add((row_index, col_index))
    return origins, covered


def _style_cell(cell, rowspan, colspan):
    data = {}
    text = _cell_text(cell.value)
    if text:
        data["v"] = text
    background = _fill_hex(cell)
    if background:
        data["bg"] = background
    font_color = _rgb(cell.font.color) if cell.font and cell.font.color else None
    if font_color and font_color.upper() != "#000000":
        data["fg"] = font_color
    if cell.font and cell.font.bold:
        data["b"] = 1
    if cell.font and cell.font.size and cell.font.size > 11:
        data["sz"] = cell.font.size
    horizontal = cell.alignment.horizontal
    if horizontal in ("center", "right", "left", "justify"):
        data["a"] = "justify" if horizontal == "justify" else horizontal
    vertical = cell.alignment.vertical
    if vertical == "center":
        data["va"] = "middle"
    elif vertical in ("top", "bottom"):
        data["va"] = vertical
    if cell.alignment.wrap_text:
        data["w"] = 1
    if rowspan > 1:
        data["rs"] = rowspan
    if colspan > 1:
        data["cs"] = colspan
    return data or None


def _fill_hex(cell):
    fill = cell.fill
    if not fill or fill.patternType != "solid":
        return None
    color = _rgb(fill.fgColor)
    if color and color.upper() == "#000000":
        return None
    return color


def _rgb(color):
    if not color or getattr(color, "type", None) != "rgb":
        return None
    value = getattr(color, "rgb", None)
    if not isinstance(value, str):
        return None
    if value in ("00000000", "None"):
        return None
    if len(value) == 8:
        if value[:2] == "00":
            return None
        value = value[2:]
    if len(value) != 6:
        return None
    return "#" + value


def _column_px(width):
    if not width:
        return 72
    return max(24, int(round(float(width) * 8 + 5)))


def _row_px(height):
    if not height:
        return 0
    return max(18, int(round(float(height) * 1.33)))


def _xlsx_html(raw):
    sheets, truncated, sheet_count = _xlsx_sheets(raw)
    parts = []
    for sheet in sheets:
        parts.append(
            Markup("<h4 style='margin:12px 0 8px;'>{}</h4>").format(
                escape(sheet["name"])
            )
        )
        parts.append(_rows_to_table(_plain_rows(sheet["rows"]), truncated))
    if sheet_count > MAX_SHEETS:
        parts.append(
            Markup(
                '<div class="text-muted">แสดง {} ชีตแรก จาก {} ชีต</div>'
            ).format(MAX_SHEETS, sheet_count)
        )
    return Markup("").join(parts) or Markup('<div class="text-muted">ไม่มีข้อมูล</div>')


def _plain_rows(rows):
    plain = []
    for line in rows:
        plain.append([
            cell.get("v", "") if isinstance(cell, dict) else ""
            for cell in line
        ])
    return plain


def _xlsx_rows(sheet):
    rows = []
    truncated = False
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index >= MAX_ROWS:
            truncated = True
            break
        if len(row) > MAX_COLS:
            truncated = True
        rows.append([_cell_text(value) for value in row[:MAX_COLS]])
    return rows, truncated


def _xls_sheets(raw):
    try:
        import xlrd
    except ImportError as error:
        raise ValueError(
            "ไฟล์ Excel รุ่นเก่า (.xls) เปิดบนหน้าเว็บไม่ได้ กรุณาบันทึกเป็น .xlsx"
        ) from error
    book = xlrd.open_workbook(file_contents=raw)
    truncated = book.nsheets > MAX_SHEETS
    sheets = []
    for sheet in book.sheets()[:MAX_SHEETS]:
        if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLS:
            truncated = True
        rows = []
        for row_index in range(min(sheet.nrows, MAX_ROWS)):
            rows.append(
                [
                    _cell_text(sheet.cell_value(row_index, col_index))
                    for col_index in range(min(sheet.ncols, MAX_COLS))
                ]
            )
        sheets.append({"name": sheet.name, "rows": rows})
    return sheets, truncated, book.nsheets


def _xls_html(raw):
    sheets, truncated, sheet_count = _xls_sheets(raw)
    parts = []
    for sheet in sheets:
        parts.append(
            Markup("<h4 style='margin:12px 0 8px;'>{}</h4>").format(
                escape(sheet["name"])
            )
        )
        parts.append(_rows_to_table(sheet["rows"], truncated))
    if sheet_count > MAX_SHEETS:
        parts.append(
            Markup(
                '<div class="text-muted">แสดง {} ชีตแรก จาก {} ชีต</div>'
            ).format(MAX_SHEETS, sheet_count)
        )
    return Markup("").join(parts) or Markup('<div class="text-muted">ไม่มีข้อมูล</div>')


def _rows_to_table(rows, truncated):
    if not rows or all(not any(cell for cell in row) for row in rows):
        return Markup('<div class="text-muted">ไม่มีข้อมูลในชีตนี้</div>')
    body = []
    for row in rows:
        cells = "".join(
            "<td style='border:1px solid #dee2e6;padding:4px 8px;white-space:nowrap;'>%s</td>"
            % escape(cell)
            for cell in row
        )
        body.append("<tr>%s</tr>" % cells)
    note = ""
    if truncated:
        note = (
            '<div class="text-muted">แสดงไม่เกิน %s แถว และ %s คอลัมน์</div>'
            % (MAX_ROWS, MAX_COLS)
        )
    return Markup(
        '<div style="overflow:auto;max-height:70vh;border:1px solid #dee2e6;">'
        '<table style="border-collapse:collapse;background:#fff;">%s</table></div>%s'
        % ("".join(body), note)
    )


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.strftime("%d/%m/%Y")
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)


def _docx_html(raw):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(raw))
    parts = []
    for child in document.element.body:
        if child.tag == "{%s}p" % _W_NS:
            parts.append(_paragraph_html(Paragraph(child, document)))
        elif child.tag == "{%s}tbl" % _W_NS:
            parts.append(_table_html(Table(child, document)))
    body = Markup("").join(part for part in parts if part)
    if not body:
        body = Markup('<div class="text-muted">ไม่มีข้อความในเอกสาร</div>')
    return Markup(
        '<div style="max-height:75vh;overflow:auto;padding:16px 20px;'
        'background:#fff;border:1px solid #dee2e6;">{}</div>'
    ).format(body)


def _paragraph_html(paragraph):
    inner = _runs_html(paragraph)
    if not inner:
        return Markup("")
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if style_name.startswith("Heading"):
        level = style_name.replace("Heading", "").strip()
        tag = {"1": "h3", "2": "h4", "3": "h5"}.get(level, "h5")
        return Markup("<{}>{}</{}>").format(tag, inner, tag)
    return Markup("<p>{}</p>").format(inner)


def _runs_html(paragraph):
    chunks = []
    for run in paragraph.runs:
        text = escape(run.text or "")
        if text:
            if run.bold:
                text = Markup("<strong>{}</strong>").format(text)
            if run.italic:
                text = Markup("<em>{}</em>").format(text)
            if run.underline:
                text = Markup("<u>{}</u>").format(text)
            chunks.append(text)
        chunks.extend(_run_images(run))
    return Markup("").join(chunks)


def _run_images(run):
    images = []
    rel_attr = "{%s}embed" % _R_NS
    for blip in run._element.findall(".//{%s}blip" % _A_NS):
        embed = blip.get(rel_attr)
        part = run.part.related_parts.get(embed) if embed else None
        if part is None:
            continue
        blob = part.blob or b""
        if len(blob) > MAX_IMAGE_BYTES:
            images.append(Markup('<div class="text-muted">[ข้ามรูปภาพขนาดใหญ่]</div>'))
            continue
        mime = part.content_type or "image/png"
        encoded = base64.b64encode(blob).decode()
        images.append(
            Markup(
                '<img src="data:{};base64,{}" alt="" '
                'style="max-width:100%;height:auto;"/>'
            ).format(escape(mime), encoded)
        )
    return images


def _table_html(table):
    rows = []
    for row in table.rows:
        cells = "".join(
            "<td style='border:1px solid #dee2e6;padding:4px 8px;vertical-align:top;'>%s</td>"
            % escape(cell.text or "")
            for cell in row.cells
        )
        rows.append("<tr>%s</tr>" % cells)
    return Markup(
        '<table style="border-collapse:collapse;width:100%;margin:8px 0;">%s</table>'
        % "".join(rows)
    )
