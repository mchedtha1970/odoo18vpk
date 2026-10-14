# Copyright 2026 VPK
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).

import base64
import io
from urllib.parse import unquote

import openpyxl
from docx import Document

from odoo.exceptions import UserError, ValidationError
from odoo.tests import TransactionCase, tagged

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _b64(raw):
    return base64.b64encode(raw)


def _xlsx(value):
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx(text):
    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@tagged("post_install", "-at_install")
class TestVpkDocument(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.category = cls.env["vpk.document.category"].create({
            "name": "ทดสอบหมวดเอกสาร",
            "code": "TEST-DOC",
        })

    def _create_document(self, filename, raw, **extra):
        values = {
            "name": filename,
            "category_id": self.category.id,
            "file_name": filename,
            "file": _b64(raw),
        }
        values.update(extra)
        return self.env["vpk.document"].create(values)

    def test_multi_upload_creates_one_document_per_file(self):
        first = self.env["ir.attachment"].create({
            "name": "แผน-ก.png",
            "datas": base64.b64encode(PNG_1X1),
            "type": "binary",
        })
        second = self.env["ir.attachment"].create({
            "name": "แผน-ข.png",
            "datas": base64.b64encode(PNG_1X1),
            "type": "binary",
        })
        documents = self.env["vpk.document"].create({
            "category_id": self.category.id,
            "description": "อัปโหลดพร้อมกัน",
            "upload_ids": [(6, 0, [first.id, second.id])],
        })
        self.assertEqual(len(documents), 2)
        self.assertEqual(set(documents.mapped("name")), {"แผน-ก", "แผน-ข"})
        self.assertEqual(documents.category_id, self.category)
        self.assertEqual(set(documents.mapped("description")), {"อัปโหลดพร้อมกัน"})
        self.assertTrue(all(document.file_type == "image" for document in documents))

    def test_multi_upload_rejects_unsupported_file(self):
        bad = self.env["ir.attachment"].create({
            "name": "note.txt",
            "datas": base64.b64encode(b"hello"),
            "type": "binary",
        })
        with self.assertRaises(ValidationError):
            self.env["vpk.document"].create({
                "category_id": self.category.id,
                "upload_ids": [(6, 0, [bad.id])],
            })

    def test_create_category(self):
        category = self.env["vpk.document.category"].create({
            "name": "หมวดใหม่",
            "code": "NEW-CAT",
        })
        self.assertEqual(category.name, "หมวดใหม่")

    def test_subfolder_under_category(self):
        child = self.env["vpk.document.category"].create({
            "name": "ปี 2569",
            "parent_id": self.category.id,
        })
        self.assertEqual(child.display_name, "%s / ปี 2569" % self.category.name)
        sibling = self.env["vpk.document.category"].create({"name": "ปี 2569"})
        self.assertFalse(sibling.parent_id)
        with self.assertRaises(ValidationError):
            self.env["vpk.document.category"].create({
                "name": "ปี 2569",
                "parent_id": self.category.id,
            })
        with self.assertRaises(UserError):
            self.category.parent_id = child

    def test_duplicate_category_code(self):
        with self.assertRaises(ValidationError):
            self.env["vpk.document.category"].create({
                "name": "หมวดรหัสซ้ำ",
                "code": "TEST-DOC",
            })

    def test_download_uses_form_name(self):
        document = self._create_document("source.xlsx", _xlsx("a"), name="รายงานงบ")
        self.assertEqual(document._download_filename(), "รายงานงบ.xlsx")
        action = document.action_download()
        self.assertIn("รายงานงบ.xlsx", unquote(action["url"]))
        named = self._create_document("source.xlsx", _xlsx("b"), name="รายงานงบ.xlsx")
        self.assertEqual(named._download_filename(), "รายงานงบ.xlsx")
        unsafe = self._create_document("source.pdf", b"%PDF-1.4\n%%EOF", name="โฟลเดอร์/รายงาน")
        self.assertEqual(unsafe._download_filename(), "โฟลเดอร์ รายงาน.pdf")

    def test_pdf_preview_action(self):
        document = self._create_document("report.pdf", b"%PDF-1.4\n%%EOF")
        self.assertEqual(document.file_type, "pdf")
        action = document.action_view_file()
        self.assertEqual(action["type"], "ir.actions.act_url")
        self.assertEqual(action["url"], "/vpk_document/open/%s" % document.id)
        self.assertEqual(action["target"], "new")
        body, mimetype = document._open_payload()
        self.assertEqual(mimetype, "application/pdf")
        self.assertTrue(body.startswith(b"%PDF"))
        self.assertIn("/vpk_document/open/", document.preview_html)

    def test_excel_preview_escapes_html(self):
        document = self._create_document(
            "sheet.xlsx",
            _xlsx("<script>alert(1)</script>"),
        )
        self.assertEqual(document.file_type, "excel")
        body, mimetype = document._open_payload()
        self.assertIn("text/html", mimetype)
        text = body.decode()
        self.assertIn("sheet_viewer.js", text)
        self.assertIn("\\u003cscript\\u003e", text)
        self.assertNotIn("<script>", text)

    def test_word_preview(self):
        document = self._create_document("memo.docx", _docx("บันทึกข้อความทดสอบ"))
        self.assertEqual(document.file_type, "word")
        body, mimetype = document._open_payload()
        self.assertIn("text/html", mimetype)
        self.assertIn("บันทึกข้อความทดสอบ", body.decode())

    def test_image_preview(self):
        document = self._create_document("photo.png", PNG_1X1)
        self.assertEqual(document.file_type, "image")
        body, mimetype = document._open_payload()
        self.assertEqual(mimetype, "image/png")
        self.assertTrue(body.startswith(b"\x89PNG"))

    def test_reject_unsupported_extension(self):
        with self.assertRaises(ValidationError):
            self._create_document("notes.txt", b"hello")

    def test_create_with_bin_size_context(self):
        document = self.env["vpk.document"].with_context(bin_size=True).create({
            "name": "รายงาน",
            "category_id": self.category.id,
            "file_name": "report.pdf",
            "file": _b64(b"%PDF-1.4\n%%EOF"),
        })
        self.assertEqual(document.file_type, "pdf")
        self.assertIn("iframe", document.with_context(bin_size=True).preview_html)
        self.assertTrue(document.with_context(bin_size=True)._file_bytes().startswith(b"%PDF"))

    def test_reject_mismatched_content(self):
        with self.assertRaises(ValidationError):
            self._create_document("fake.pdf", b"this is not a pdf")
