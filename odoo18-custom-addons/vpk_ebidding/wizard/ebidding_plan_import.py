# -*- coding: utf-8 -*-
import base64
import io
import re
from datetime import date, datetime

from odoo import _, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook import Workbook
except ImportError:  # pragma: no cover
    openpyxl = None
    Workbook = None


EXCEL_COLUMNS = [
    ("office_code", "รหัสสำนักงานแผนฯ", True),
    ("fiscal_year", "ปีงบประมาณ", True),
    ("project_name", "ชื่อโครงการที่จะจัดซื้อจัดจ้าง", True),
    ("expected_announce_month", "เดือน/ปีที่คาดว่าจะประกาศจัดซื้อจัดจ้าง", True),
    ("budget_act_amount", "เงินงบประมาณตามพ.ร.บ.รายจ่ายประจำปี", False),
    ("budget_act_type_code", "ประเภทเงินงบประมาณตามพ.ร.บ.รายจ่ายประจำปี", False),
    ("extra_budget_amount", "เงินนอกงบประมาณ", False),
    ("extra_budget_type_code", "ประเภทเงินนอกงบประมาณ", False),
    ("agency_budget_amount", "เงินงบประมาณหน่วยงาน", False),
    ("plan_code", "รหัสแผนการจัดซื้อจัดจ้าง", False),
    ("line_state", "สถานะแผนการจัดซื้อจัดจ้าง", False),
]

HEADER_ALIASES = {
    "office_code": ["รหัสสำนักงานแผนฯ", "รหัสสำนักงานแผน", "office_code"],
    "fiscal_year": ["ปีงบประมาณ", "fiscal_year"],
    "project_name": [
        "ชื่อโครงการที่จะจัดซื้อจัดจ้าง",
        "ชื่อโครงการ",
        "project_name",
    ],
    "expected_announce_month": [
        "เดือน/ปีที่คาดว่าจะประกาศจัดซื้อจัดจ้าง",
        "เดือน/ปีที่คาดว่าจะประกาศ",
        "expected_announce_month",
    ],
    "budget_act_amount": [
        "เงินงบประมาณตามพ.ร.บ.รายจ่ายประจำปี",
        "เงินงบประมาณตามพรบ",
        "budget_act_amount",
    ],
    "budget_act_type_code": [
        "ประเภทเงินงบประมาณตามพ.ร.บ.รายจ่ายประจำปี",
        "ประเภทเงินงบประมาณตามพรบ",
        "budget_act_type_code",
    ],
    "extra_budget_amount": ["เงินนอกงบประมาณ", "extra_budget_amount"],
    "extra_budget_type_code": [
        "ประเภทเงินนอกงบประมาณ",
        "extra_budget_type_code",
    ],
    "agency_budget_amount": ["เงินงบประมาณหน่วยงาน", "agency_budget_amount"],
    "plan_code": ["รหัสแผนการจัดซื้อจัดจ้าง", "plan_code"],
    "line_state": ["สถานะแผนการจัดซื้อจัดจ้าง", "สถานะแผน", "line_state"],
}

LINE_STATE_MAP = {
    "ยังไม่เผยแพร่ขึ้นเว็บไซต์": "not_published",
    "ยังไม่เผยแพร่": "not_published",
    "not_published": "not_published",
    "เสนอหัวหน้าอนุมัติ": "submitted",
    "submitted": "submitted",
    "เผยแพร่ขึ้นเว็บไซต์": "published",
    "published": "published",
}

NAVY = "1F4E79"
YELLOW = "FFF2CC"
HINT = "E8F1FA"
GREEN = "E2EFDA"


def _norm_header(value):
    text = str(value or "").strip().lower()
    text = re.sub(r"\s+", "", text)
    text = text.replace("*", "")
    return text


def _to_float(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        raise ValueError(_("จำนวนเงินไม่ถูกต้อง: %s") % value)


def _ad_year(year):
    return year - 543 if year >= 2400 else year


def _month_start(year, month):
    try:
        return date(_ad_year(int(year)), int(month), 1)
    except (TypeError, ValueError, OverflowError):
        return False


def _to_announce_date(value):
    """Parse month/year from e-GP Excel into the first day of that month."""
    if value in (None, ""):
        return False
    if isinstance(value, datetime):
        return _month_start(value.year, value.month)
    if isinstance(value, date):
        return _month_start(value.year, value.month)
    text = str(value).strip()
    slash = re.fullmatch(r"(0?[1-9]|1[0-2])[/\-](\d{4})", text)
    if slash:
        return _month_start(slash.group(2), slash.group(1))
    iso = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-(\d{1,2}))?", text)
    if iso:
        return _month_start(iso.group(1), iso.group(2))
    if isinstance(value, (int, float)):
        digits = str(int(value))
    else:
        digits = re.sub(r"\D", "", text)
    padded = digits.zfill(6) if digits.isdigit() else ""
    if re.fullmatch(r"(0[1-9]|1[0-2])\d{4}", padded):
        return _month_start(padded[2:], padded[:2])
    return False


def _to_year(value):
    if value in (None, ""):
        return False
    if isinstance(value, (int, float)):
        return str(int(value))
    text = re.sub(r"\D", "", str(value))
    return text if text else False


def build_import_template():
    if Workbook is None:
        raise UserError(_("ต้องติดตั้ง python package openpyxl ก่อนนำเข้า Excel"))
    wb = Workbook()
    ws = wb.active
    ws.title = "แผนจัดซื้อจัดจ้าง"
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    example_font = Font(name="Arial", size=10, color="7A4E00")
    fill_navy = PatternFill("solid", fgColor=NAVY)
    fill_yellow = PatternFill("solid", fgColor=YELLOW)
    fill_hint = PatternFill("solid", fgColor=HINT)
    fill_green = PatternFill("solid", fgColor=GREEN)
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    ws.merge_cells("A1:K1")
    ws["A1"] = (
        "เทมเพลตนำเข้าแผนจัดซื้อจัดจ้างจาก e-GP — คอลัมน์ตามคู่มือกรมบัญชีกลาง "
        "วันที่ 8 ก.พ. 2562 ข้อ 5.1 (เครื่องหมาย * บังคับ)"
    )
    ws["A1"].font = Font(name="Arial", size=12, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32

    for col, (_key, title, required) in enumerate(EXCEL_COLUMNS, start=1):
        cell = ws.cell(2, col, ("* " if required else "") + title)
        cell.font = header_font
        cell.fill = fill_navy
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.row_dimensions[2].height = 36

    example = [
        "67000",
        2570,
        "จ้างบริการพนักงานรักษาความปลอดภัย",
        "10/2569",
        1000000,
        "งบกรม",
        0,
        "",
        0,
        "",
        "ยังไม่เผยแพร่ขึ้นเว็บไซต์",
    ]
    for col, value in enumerate(example, start=1):
        cell = ws.cell(3, col, value)
        cell.font = example_font
        cell.fill = fill_yellow
        cell.border = thin
        if col in (5, 7, 9):
            cell.number_format = "#,##0.00"
    ws.row_dimensions[3].height = 22

    guide = wb.create_sheet("คู่มือ")
    guide["A1"] = "วิธีกรอกไฟล์นำเข้าแผน e-GP"
    guide["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    rows = [
        ("แถว 1", "คำอธิบายเทมเพลต ห้ามลบ"),
        ("แถว 2", "หัวคอลัมน์ ห้ามเปลี่ยนชื่อ"),
        ("แถว 3", "แถวตัวอย่าง สีเหลือง ลบหรือเขียนทับก่อนนำเข้าจริง"),
        ("แถว 4 เป็นต้นไป", "ข้อมูลจริง"),
        ("รหัสสำนักงานแผนฯ *", "รหัสสำนักงานแผนตาม e-GP"),
        ("ปีงบประมาณ *", "พ.ศ. เช่น 2570"),
        ("ชื่อโครงการ *", "ชื่อโครงการที่จะจัดซื้อจัดจ้าง"),
        ("เดือน/ปีที่คาดว่าจะประกาศ *", "เดือน/ปี พ.ศ. เช่น 10/2569 หรือ 102569 ตามไฟล์ e-GP"),
        ("เงินงบประมาณตาม พ.ร.บ.", "ตัวเลข ไม่ใส่เครื่องหมาย เช่น 850000"),
        ("ประเภทเงินงบประมาณตาม พ.ร.บ.", "รหัสประเภทเงินตามแท็บ e-GP"),
        ("เงินนอกงบประมาณ", "ตัวเลข หรือเว้นว่าง"),
        ("ประเภทเงินนอกงบประมาณ", "รหัสประเภทเงินนอกงบประมาณ"),
        ("เงินงบประมาณหน่วยงาน", "ตัวเลข หรือเว้นว่าง"),
        ("รหัสแผนการจัดซื้อจัดจ้าง", "กรอกถ้า e-GP ออกรหัสให้แล้ว"),
        ("สถานะแผน", "ยังไม่เผยแพร่ขึ้นเว็บไซต์ / เสนอหัวหน้าอนุมัติ / เผยแพร่ขึ้นเว็บไซต์"),
        ("แหล่งที่มา", "คู่มือจัดทำแผนการจัดซื้อจัดจ้าง กรมบัญชีกลาง 8 ก.พ. 2562 ข้อ 5.1"),
    ]
    guide["A3"] = "หัวข้อ"
    guide["B3"] = "รายละเอียด"
    for cell in (guide["A3"], guide["B3"]):
        cell.fill = fill_navy
        cell.font = header_font
        cell.border = thin
    for idx, (title, detail) in enumerate(rows, start=4):
        guide.cell(idx, 1, title).fill = fill_hint
        guide.cell(idx, 1).font = Font(name="Arial", size=10, bold=True)
        guide.cell(idx, 1).border = thin
        guide.cell(idx, 2, detail).font = body_font
        guide.cell(idx, 2).border = thin
        if idx == 4:
            guide.cell(idx, 2).fill = fill_green
    guide.column_dimensions["A"].width = 36
    guide.column_dimensions["B"].width = 78
    guide.freeze_panes = "A4"
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A2:K3"
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


class EbiddingPlanImportWizard(models.TransientModel):
    _name = "ebidding.plan.import.wizard"
    _description = "นำเข้าแผนจัดซื้อจัดจ้างจาก Excel e-GP"

    plan_id = fields.Many2one(
        comodel_name="ebidding.plan",
        string="ชุดแผนที่จะนำเข้า",
        domain="[('state', '=', 'draft')]",
        help="เว้นว่างถ้าต้องการให้ระบบสร้างชุดแผนใหม่จากปีงบประมาณในไฟล์",
    )
    data_file = fields.Binary(string="ไฟล์ Excel")
    filename = fields.Char(string="ชื่อไฟล์")
    template_file = fields.Binary(string="เทมเพลต")
    template_filename = fields.Char(string="ชื่อไฟล์เทมเพลต")
    state = fields.Selection(
        selection=[
            ("draft", "รอไฟล์"),
            ("done", "เสร็จสิ้น"),
        ],
        default="draft",
    )
    success_count = fields.Integer(string="จำนวนรายการที่นำเข้าสำเร็จ", readonly=True)
    fail_count = fields.Integer(string="จำนวนรายการที่นำเข้าไม่สำเร็จ", readonly=True)
    created_plan_id = fields.Many2one(
        comodel_name="ebidding.plan",
        string="ชุดแผนที่บันทึก",
        readonly=True,
    )
    row_ids = fields.One2many(
        comodel_name="ebidding.plan.import.row",
        inverse_name="wizard_id",
        string="ผลการตรวจสอบ",
        readonly=True,
    )

    def action_download_template(self):
        self.ensure_one()
        content = build_import_template()
        self.write({
            "template_file": base64.b64encode(content),
            "template_filename": "template-แผนจัดซื้อจัดจ้าง-e-GP.xlsx",
        })
        return {
            "type": "ir.actions.act_url",
            "url": (
                "/web/content/?model=%s&id=%s&field=template_file"
                "&filename_field=template_filename&download=true"
                % (self._name, self.id)
            ),
            "target": "self",
        }

    def action_import(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_("ต้องติดตั้ง python package openpyxl ก่อนนำเข้า Excel"))
        if not self.data_file:
            raise UserError(_("กรุณาเลือกไฟล์ Excel"))
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.data_file)),
                data_only=True,
                read_only=True,
            )
        except Exception as exc:
            raise UserError(_("เปิดไฟล์ Excel ไม่ได้: %s") % exc) from exc

        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        header_row_index, mapping = self._find_header(rows)
        if not mapping:
            raise UserError(
                _("ไม่พบหัวคอลัมน์ตามเทมเพลต e-GP ในชีตแรก")
            )

        parsed = []
        for offset, raw in enumerate(rows[header_row_index + 1 :], start=header_row_index + 2):
            values = self._row_values(raw, mapping)
            if self._is_empty_row(values):
                continue
            parsed.append((offset, values))

        if not parsed:
            raise UserError(_("ไม่พบรายการแผนในไฟล์"))

        plan = self.plan_id
        target_year = plan.fiscal_year if plan else False
        result_rows = []
        to_create = []
        for excel_row, values in parsed:
            office = str(values.get("office_code") or "").strip()
            project = str(values.get("project_name") or "").strip()
            if office == "67000" and project == "จ้างบริการพนักงานรักษาความปลอดภัย":
                continue
            error = self._validate_row(values, target_year)
            if error:
                result_rows.append((0, 0, {
                    "excel_row": excel_row,
                    "project_name": values.get("project_name") or "",
                    "status": "fail",
                    "message": error,
                }))
                continue
            to_create.append((excel_row, values))
            if not target_year:
                target_year = values["fiscal_year"]
            result_rows.append((0, 0, {
                "excel_row": excel_row,
                "project_name": values.get("project_name") or "",
                "status": "ok",
                "message": _("ผ่านการตรวจสอบ"),
            }))

        if not to_create and not result_rows:
            raise UserError(_("ไม่พบรายการแผนในไฟล์ (แถวตัวอย่างไม่ถูกนำเข้า)"))

        if to_create and not plan:
            first_year = to_create[0][1]["fiscal_year"]
            plan = self.env["ebidding.plan"].search([
                ("fiscal_year", "=", first_year),
                ("company_id", "=", self.env.company.id),
            ], limit=1)
            if not plan:
                plan = self.env["ebidding.plan"].create({
                    "fiscal_year": first_year,
                })

        created = 0
        if plan:
            for _excel_row, values in to_create:
                self.env["ebidding.plan.line"].create(
                    self._prepare_line_vals(plan, values)
                )
                created += 1

        fail = sum(1 for _cmd, _id, rec in result_rows if rec["status"] == "fail")
        self.write({
            "state": "done",
            "success_count": created,
            "fail_count": fail,
            "created_plan_id": plan.id if plan else False,
            "row_ids": [(5, 0, 0)] + result_rows,
        })
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    def action_open_plan(self):
        self.ensure_one()
        if not self.created_plan_id:
            raise UserError(_("ยังไม่มีชุดแผนที่บันทึกจากไฟล์นี้"))
        return {
            "type": "ir.actions.act_window",
            "res_model": "ebidding.plan",
            "res_id": self.created_plan_id.id,
            "view_mode": "form",
            "target": "current",
        }

    def _find_header(self, rows):
        alias_map = {}
        for key, aliases in HEADER_ALIASES.items():
            for alias in aliases:
                alias_map[_norm_header(alias)] = key
        for index, row in enumerate(rows[:10]):
            mapping = {}
            for col, cell in enumerate(row or []):
                key = alias_map.get(_norm_header(cell))
                if key and key not in mapping:
                    mapping[key] = col
            required = {"office_code", "fiscal_year", "project_name", "expected_announce_month"}
            if required.issubset(mapping):
                return index, mapping
        return 0, {}

    def _row_values(self, raw, mapping):
        raw = raw or []
        values = {}
        for key, col in mapping.items():
            values[key] = raw[col] if col < len(raw) else None
        return values

    def _is_empty_row(self, values):
        return not any(
            str(values.get(key) or "").strip()
            for key in (
                "office_code",
                "project_name",
                "expected_announce_month",
                "plan_code",
            )
        )

    def _validate_row(self, values, fiscal_year):
        office = str(values.get("office_code") or "").strip()
        if not office:
            return _("ต้องระบุรหัสสำนักงานแผนฯ")
        year = _to_year(values.get("fiscal_year"))
        if not year:
            return _("ต้องระบุปีงบประมาณ")
        if fiscal_year and year != fiscal_year:
            return _(
                "ปีงบประมาณในไฟล์ (%s) ไม่ตรงกับชุดแผน (%s)"
            ) % (year, fiscal_year)
        project = str(values.get("project_name") or "").strip()
        if not project:
            return _("ต้องระบุชื่อโครงการที่จะจัดซื้อจัดจ้าง")
        month = _to_announce_date(values.get("expected_announce_month"))
        if not month:
            return _("เดือน/ปีที่คาดว่าจะประกาศต้องเป็น ดด/ปปปป เช่น 10/2569 หรือ 102569")
        try:
            budget_act = _to_float(values.get("budget_act_amount"))
            extra = _to_float(values.get("extra_budget_amount"))
            agency = _to_float(values.get("agency_budget_amount"))
        except ValueError as exc:
            return str(exc)
        if budget_act + extra + agency <= 0:
            return _("ต้องระบุจำนวนเงินอย่างน้อยหนึ่งแหล่ง")
        state_label = str(values.get("line_state") or "").strip()
        if state_label and state_label not in LINE_STATE_MAP:
            return _("สถานะแผนไม่รู้จัก: %s") % state_label
        plan_code = str(values.get("plan_code") or "").strip()
        if plan_code:
            existing = self.env["ebidding.plan.line"].search_count([
                ("company_id", "=", self.env.company.id),
                ("fiscal_year", "=", year),
                ("plan_code", "=", plan_code),
            ])
            if existing:
                return _("รหัสแผนการจัดซื้อจัดจ้าง %s มีในระบบแล้ว") % plan_code
        values["office_code"] = office
        values["fiscal_year"] = year
        values["project_name"] = project
        values["expected_announce_month"] = month
        values["budget_act_amount"] = budget_act
        values["extra_budget_amount"] = extra
        values["agency_budget_amount"] = agency
        values["budget_act_type_code"] = str(values.get("budget_act_type_code") or "").strip()
        values["extra_budget_type_code"] = str(values.get("extra_budget_type_code") or "").strip()
        values["plan_code"] = plan_code
        values["line_state"] = LINE_STATE_MAP.get(state_label, "not_published")
        return False

    def _prepare_line_vals(self, plan, values):
        return {
            "plan_id": plan.id,
            "office_code": values["office_code"],
            "plan_code": values.get("plan_code") or False,
            "project_name": values["project_name"],
            "expected_announce_month": values["expected_announce_month"],
            "budget_act_amount": values["budget_act_amount"],
            "budget_act_type_code": values.get("budget_act_type_code") or False,
            "extra_budget_amount": values["extra_budget_amount"],
            "extra_budget_type_code": values.get("extra_budget_type_code") or False,
            "agency_budget_amount": values["agency_budget_amount"],
            "import_type": "excel",
            "line_state": values.get("line_state") or "not_published",
        }


class EbiddingPlanImportRow(models.TransientModel):
    _name = "ebidding.plan.import.row"
    _description = "ผลตรวจสอบแถวนำเข้าแผน e-GP"
    _order = "excel_row"

    wizard_id = fields.Many2one(
        comodel_name="ebidding.plan.import.wizard",
        ondelete="cascade",
    )
    excel_row = fields.Integer(string="แถวในไฟล์")
    project_name = fields.Char(string="ชื่อโครงการ")
    status = fields.Selection(
        selection=[
            ("ok", "นำเข้าสำเร็จ"),
            ("fail", "นำเข้าไม่สำเร็จ"),
        ],
        string="ผลการตรวจสอบ",
    )
    message = fields.Char(string="สาเหตุ / หมายเหตุ")
