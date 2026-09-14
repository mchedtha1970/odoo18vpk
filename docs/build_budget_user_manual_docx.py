#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Thai Word training manual for the VPK budget system."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUT = Path("/opt/odoo18vpk/docs/คู่มือการใช้งาน-ระบบงบประมาณ.docx")
FIG = Path("/opt/odoo18vpk/docs/manual_assets/figures")

NAVY = RGBColor(0x0B, 0x68, 0x48)
TEAL = RGBColor(0x0B, 0x68, 0x48)
SLATE = RGBColor(0x2A, 0x33, 0x3A)
MUTED = RGBColor(0x7D, 0x8A, 0x9E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC6, 0x28, 0x28)
AMBER = RGBColor(0xC4, 0x7E, 0x14)
HEADER_FILL = "0B6848"
META_FILL = "EAF3F1"
LOGO = FIG / "moph_logo.png"


def set_run(run, size=12, bold=False, color=SLATE, font_name="Prompt"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font_name)


def add_text(p, text, size=12, bold=False, color=SLATE):
    run = p.add_run(text)
    set_run(run, size, bold, color)
    return run


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, color="D7DEE6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, size=11, bold=False, color=SLATE, align="left", fill=None):
    if fill:
        shade_cell(cell, fill)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.clear()
    add_text(p, text, size, bold, color)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)


def set_narrow_cell_margin(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), "80")
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def make_table(doc, headers, rows, col_widths=None, header_fill=HEADER_FILL):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, 11, True, WHITE, "center", header_fill)
        set_narrow_cell_margin(table.rows[0].cells[i])
    for r, row in enumerate(rows, 1):
        fill = "F6F8FA" if r % 2 == 0 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, 11, False, SLATE, "left", fill)
            set_narrow_cell_margin(table.rows[r].cells[c])
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_border_line(paragraph, edge="bottom", color="0B6848", size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), "6")
    el.set(qn("w:color"), color)
    pBdr.append(el)
    pPr.append(pBdr)


def add_bottom_line(paragraph, color="0B6848", size="12"):
    add_border_line(paragraph, "bottom", color, size)


def page_title(doc, text, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    add_text(p, text, 18, True, TEAL)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, 18 if level == 1 else 14 if level == 2 else 13, True, TEAL)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    if level == 1:
        add_bottom_line(p)
    return p


def body(doc, text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text, size)
    return p


def bullets(doc, items, numbered=False):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        p.clear()
        add_text(p, item, 12)
        p.paragraph_format.space_after = Pt(3)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    label = text if text.startswith("▲") else f"▲ {text}"
    add_text(p, label, 10, False, MUTED)


def picture(doc, name, width_cm=16.2, cap=""):
    path = FIG / name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    if cap:
        caption(doc, cap)


def callout(doc, title, text, tone="info"):
    fill = {"info": "E5F5F3", "warn": "FDF3E0", "danger": "FBECEA", "ok": "E5F5EA"}[tone]
    title_color = {"info": TEAL, "warn": AMBER, "danger": RED, "ok": RGBColor(0x2C, 0x7A, 0x45)}[tone]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell, "D7DEE6")
    p1 = cell.paragraphs[0]
    add_text(p1, title, 12, True, title_color)
    p2 = cell.add_paragraph()
    add_text(p2, text, 12, False, SLATE)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(4)
    doc.add_paragraph()


def steps(doc, items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_text(p, f"ขั้นที่ {i}  ", 12, True, TEAL)
        add_text(p, item, 12, False, SLATE)


def header_footer(doc):
    section = doc.sections[0]
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(1.0)
    header = section.header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = ""
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_border_line(fp, "top", "C5CDD6", "6")
    add_text(fp, "ระบบงบประมาณ · ระบบ ERP รพ.วชิระภูเก็ต | หน้า ", 9, False, MUTED)
    run = fp.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run(run, 9, False, MUTED)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Prompt"
    normal.font.size = Pt(12)
    normal.font.color.rgb = SLATE
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), "Prompt")
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15


def build():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.2)
    header_footer(doc)

    # Cover — document control, same pattern as UM-ERP-MEET-001
    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.paragraph_format.space_before = Pt(12)
    logo_p.paragraph_format.space_after = Pt(8)
    logo_p.add_run().add_picture(str(LOGO), width=Cm(3.4))

    hosp = doc.add_paragraph()
    hosp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hosp.paragraph_format.space_after = Pt(0)
    add_text(hosp, "โรงพยาบาลวชิระภูเก็ต", 16, True, SLATE)
    hosp_en = doc.add_paragraph()
    hosp_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hosp_en.paragraph_format.space_after = Pt(18)
    add_text(hosp_en, "VACHIRAPHUKET HOSPITAL", 11, True, MUTED)

    dtype = doc.add_paragraph()
    dtype.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dtype.paragraph_format.space_after = Pt(6)
    add_text(dtype, "เอกสารคู่มือการใช้งาน (User Manual)", 13, False, MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    add_text(title, "ระบบงบประมาณ", 28, True, TEAL)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(22)
    add_text(sub, "ระบบบริหารทรัพยากรองค์กร (ERP)", 14, True, SLATE)

    info = [
        ("รหัสเอกสาร", "UM-ERP-BUD-001"),
        ("ชื่อระบบ", "งบประมาณ · คำของบ · ใช้จ่ายงบ (Budget)"),
        ("เวอร์ชันเอกสาร", "1.0"),
        ("วันที่จัดทำ", "24 สิงหาคม พ.ศ. 2569"),
        ("สถานะเอกสาร", "ฉบับส่งมอบ (Released)"),
        ("กลุ่มผู้อ่าน", "เจ้าหน้าที่ผู้ใช้งานระบบ"),
        ("จัดทำโดย", "ทีมพัฒนาระบบ ERP"),
        ("จัดทำเพื่อ", "โรงพยาบาลวชิระภูเก็ต"),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(info):
        cell_text(table.rows[i].cells[0], k, 11, True, TEAL, "left", META_FILL)
        cell_text(table.rows[i].cells[1], v, 11, False, SLATE, "left", "FFFFFF")
        table.rows[i].cells[0].width = Cm(4.8)
        table.rows[i].cells[1].width = Cm(11.4)
        set_narrow_cell_margin(table.rows[i].cells[0])
        set_narrow_cell_margin(table.rows[i].cells[1])

    doc.add_page_break()

    page_title(doc, "ประวัติการแก้ไขเอกสาร (Document Revision History)", center=True)
    make_table(
        doc,
        ["เวอร์ชัน", "วันที่", "รายละเอียดการแก้ไข", "ผู้จัดทำ"],
        [
            ["1.0", "24 ส.ค. 2569", "จัดทำเอกสารคู่มือฉบับแรก สำหรับฝึกอบรมผู้ใช้ระบบงบประมาณ", "ทีมพัฒนาระบบ ERP"],
        ],
        [2.4, 3.0, 7.4, 3.4],
    )
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    add_text(
        note,
        "เอกสารฉบับนี้จัดทำขึ้นเพื่อเป็นคู่มือการใช้งานระบบงบประมาณ สำหรับเจ้าหน้าที่ของโรงพยาบาลวชิระภูเก็ต "
        "โปรดอ่านและปฏิบัติตามขั้นตอนที่ระบุไว้ หากพบปัญหาการใช้งานให้ติดต่อผู้ดูแลระบบ",
        12,
        False,
        MUTED,
    )

    doc.add_page_break()

    page_title(doc, "สารบัญ")
    toc_items = [
        "บทที่ 1 เข้าสู่ระบบและบัญชีฝึกอบรม",
        "บทที่ 2 ภาพรวมระบบงบประมาณ",
        "บทที่ 3 บทบาทผู้ใช้และแผนที่เมนู",
        "บทที่ 4 หน่วยงาน — สร้างและส่งคำของบ",
        "บทที่ 5 งานงบประมาณ — รับเรื่อง จัดสรร อนุมัติ",
        "บทที่ 6 ใช้งบผ่านใบขอซื้อ ใบสั่งซื้อ และใบตั้งเจ้าหนี้",
        "บทที่ 7 สูตรงบคงเหลือ",
        "บทที่ 8 Dashboard และการติดตาม",
        "บทที่ 9 ข้อผิดพลาดที่พบบ่อย",
        "บทที่ 10 แบบฝึกหัดในห้องอบรม",
        "ภาคผนวก ก คำศัพท์ในระบบ",
        "ภาคผนวก ข ช่องทางสอบถาม",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.space_before = Pt(2)
        add_text(p, item, 13, False, SLATE)

    doc.add_page_break()

    # Ch1
    heading(doc, "บทที่ 1 เข้าสู่ระบบและบัญชีฝึกอบรม", 1)
    body(
        doc,
        "ก่อนเริ่มอบรม ให้เปิดเบราว์เซอร์แล้วเข้าสู่ระบบ Odoo ตามลิงก์ที่งานเทคโนโลยีสารสนเทศแจ้ง "
        "ใช้บัญชีตามบทบาทของตนเอง ห้ามใช้บัญชีผู้ดูแลระบบคนเดียวกดครบทุกขั้นตอน "
        "เพราะจะไม่เห็นว่าเมนูและปุ่มถูกจำกัดตามสิทธิ์อย่างไร",
    )
    picture(doc, "fig_login.png", 16.2, "ภาพที่ 1  ขั้นตอนเข้าสู่ระบบและบัญชีฝึกอบรม")
    make_table(
        doc,
        ["บทบาท", "บัญชี", "รหัสผ่าน", "ใช้ทำอะไร"],
        [
            ["หน่วยงาน", "uat.dept", "VpkUat@2569", "สร้างคำของบ และสร้างใบขอซื้อ"],
            ["งานงบประมาณ", "uat.budget", "VpkUat@2569", "รับเรื่อง จัดสรร อนุมัติคำของบ"],
            ["พัสดุ / จัดซื้อ", "uat.procurement", "VpkUat@2569", "สร้างและยืนยันใบสั่งซื้อ"],
            ["การเงิน", "ผู้ใช้บัญชีของโรงพยาบาล", "ตามที่ได้รับ", "ตั้งเจ้าหนี้และโพสต์บิล"],
        ],
        [3.5, 3.4, 3.4, 6.0],
    )
    steps(
        doc,
        [
            "เปิดเบราว์เซอร์ แล้วไปที่ระบบ Odoo ของโรงพยาบาล",
            "ใส่บัญชีตามบทบาท จากนั้นเข้าแอป งบประมาณ",
            "ตรวจว่าเมนูที่เห็นตรงกับสิทธิ์ของตน หน่วยงานจะไม่เห็นเมนูการกำหนดค่า",
            "ถ้าเข้าไม่ได้ ให้แจ้งผู้ประสานอบรมทันที อย่าใช้บัญชีของผู้อื่น",
        ],
    )
    callout(
        doc,
        "ข้อควรจำตอนเข้าสู่ระบบ",
        "รหัสผ่านชุดฝึกอบรมเป็นชุดเดียวกันทั้งสามบัญชีหน่วยงาน งานงบ และพัสดุ "
        "หลังอบรมจริงควรเปลี่ยนรหัสผ่านตามนโยบายความปลอดภัยของโรงพยาบาล",
        "warn",
    )

    # Ch2
    heading(doc, "บทที่ 2 ภาพรวมระบบงบประมาณ", 1)
    body(
        doc,
        "ระบบงบประมาณของโรงพยาบาลวชิระภูเก็ตทำงานบน Odoo 18 แบ่งเป็น 2 ช่วงใหญ่ "
        "ช่วงตั้งงบประจำปี และช่วงใช้งบผ่านเอกสารจัดซื้อ ห้ามสลับลำดับ "
        "หน่วยงานยังซื้อของไม่ได้ จนกว่าคำของบจะได้รับอนุมัติและมียอดจัดสรรเข้าแผนประจำปี",
    )
    picture(doc, "fig_two_phases.png", 16.2, "ภาพที่ 2  สองช่วงของระบบงบประมาณ")
    heading(doc, "2.1 วงจรงานทั้งระบบ", 2)
    picture(doc, "fig_cycle.png", 16.2, "ภาพที่ 3  วงจรงานจากข้อมูลหลักถึงการดูคงเหลือ")
    body(doc, "จำประโยคนี้ไว้ก่อนลงมือกดในระบบ")
    bullets(
        doc,
        [
            "แผนปีเป็นกรอบ หน่วยงานของบเข้าไปในกรอบนี้",
            "งานงบประมาณเป็นคนบอกว่าใช้ได้เท่าไร ยอดใช้ได้คือยอดจัดสรร ไม่ใช่ยอดที่ขอ",
            "การซื้อของต้องมีใบขอซื้อ ระบบจะกันเงินให้เองเมื่อส่งอนุมัติ",
            "งบคงเหลือดูจากสูตรเดียวทั้งระบบ ไม่ต้องคำนวณคนละแบบตามหน้าจอ",
        ],
    )

    heading(doc, "2.2 แบบฟอร์มคำของบ 4 ใบ", 2)
    body(
        doc,
        "เมนูคำของบหน่วยงานเหลือแบบฟอร์ม 4 ใบเท่านั้น เลือกใบให้ถูกก่อนกรอก "
        "ถ้าเลือกผิด หน้าจอรายการจะไม่ตรงกับของที่ต้องการ และงานงบพิจารณาต่อไม่ได้ง่าย",
    )
    picture(doc, "fig_forms.png", 16.2, "ภาพที่ 4  แบบฟอร์มคำของบทั้ง 4 ใบ")
    make_table(
        doc,
        ["แบบฟอร์ม", "ใช้เมื่อ", "สิ่งที่ต้องกรอกเพิ่ม"],
        [
            ["แบบฟอร์มตั้งงบวัสดุ", "ยา เวชภัณฑ์ วัสดุทั่วไป", "กลุ่มงบประมาณ-วัสดุ ประเภทวัสดุย่อย และรายละเอียดสินค้า"],
            ["แบบฟอร์มครุภัณฑ์", "เครื่องมือ เครื่องใช้ คอมพิวเตอร์", "กลุ่มและประเภทครุภัณฑ์ จำนวน ราคาต่อหน่วย"],
            ["แบบฟอร์ม ก่อสร้าง", "งานซ่อม ต่อเติม ปรับปรุงอาคาร", "รายการงานก่อสร้าง สถานที่ และเหตุผล"],
            ["แบบฟอร์มโครงการ", "ค่าใช้จ่ายโครงการ / กิจกรรม", "รายการโครงการและประมาณการค่าใช้จ่าย"],
        ],
        [4.2, 5.0, 7.0],
    )
    callout(
        doc,
        "เลือกแบบฟอร์มผิดแล้วต้องทำอย่างไร",
        "ถ้ายังเป็นร่าง ให้สร้างใบใหม่ด้วยแบบฟอร์มที่ถูกต้อง อย่าฝืนกรอกรายการคนละประเภทในใบเดิม "
        "คอมพิวเตอร์และครุภัณฑ์ใช้แบบฟอร์มครุภัณฑ์ ไม่ใช้ใบวัสดุ",
        "warn",
    )

    # Ch3
    heading(doc, "บทที่ 3 บทบาทผู้ใช้และแผนที่เมนู", 1)
    picture(doc, "fig_roles.png", 16.2, "ภาพที่ 5  บทบาทและบัญชีฝึกอบรม")
    picture(doc, "fig_menus.png", 16.2, "ภาพที่ 6  แผนที่เมนูแอปงบประมาณ")
    heading(doc, "3.1 เมนูที่แต่ละบทบาทใช้บ่อย", 2)
    make_table(
        doc,
        ["บทบาท", "เมนูหลัก", "หมายเหตุ"],
        [
            ["หน่วยงาน", "งบประมาณ → บันทึกคำของบ → ทั้งหมด", "สร้างคำของบจากแบบฟอร์ม 4 ใบ"],
            ["หน่วยงาน", "จัดซื้อ → ใบขอซื้อ", "สร้าง PR เช็คงบ แล้วส่งอนุมัติ"],
            ["งานงบประมาณ", "งบประมาณ → รวบรวมคำของบ (หน่วยงบประมาณ)", "รับเรื่อง พิจารณา จัดสรร อนุมัติ"],
            ["งานงบประมาณ", "งบประมาณ → การกำหนดค่า", "แหล่งเงิน หมวดงบ กลุ่มวัสดุ แผนปี"],
            ["ทุกคนที่มีสิทธิ์งบ", "Dashboard งบประมาณ", "ดูสถานะคำขอ ภาพรวมรายปี และรายแผนก"],
            ["พัสดุ", "จัดซื้อ → ใบสั่งซื้อ", "สร้าง PO จาก PR ที่อนุมัติแล้ว"],
        ],
        [3.6, 7.2, 5.4],
    )

    # Ch4
    heading(doc, "บทที่ 4 หน่วยงาน — สร้างและส่งคำของบ", 1)
    body(
        doc,
        "บทนี้ใช้บัญชี uat.dept ตัวอย่างด้านล่างเป็นงบวัสดุ เพราะมีกติกาเพิ่มเรื่องรายละเอียดสินค้า "
        "ครุภัณฑ์ ก่อสร้าง และโครงการใช้ลำดับเดียวกัน ต่างกันที่แท็บรายการ",
    )
    heading(doc, "4.1 สร้างใบคำของบ", 2)
    steps(
        doc,
        [
            "เข้าแอป งบประมาณ แล้วเปิด บันทึกคำของบ → ทั้งหมด จากนั้นกดสร้าง",
            "เลือกหน่วยงาน / ศูนย์ต้นทุน ให้ตรงหน่วยงานของตน",
            "เลือกงบประมาณประจำปี และแหล่งเงิน เช่น เงินบำรุง",
            "เลือกแบบฟอร์มคำของบ ระบบจะเปิดแท็บรายการตามแบบฟอร์มนั้น",
            "ใส่เหตุผลประกอบคำขอตั้ง ห้ามว่าง",
            "เพิ่มบรรทัดรายการ หมวดงบ กลุ่มวัสดุ ประเภทวัสดุย่อย และยอดที่ขอ",
            "ถ้าเป็นงบวัสดุ ให้เปิดแท็บรายละเอียดสินค้า แล้วใส่รายการสินค้าให้ครบ",
            "ตรวจยอดงบประมาณที่ขอที่หัวใบ ว่ารวมจากบรรทัดถูกต้อง แล้วบันทึก",
        ],
    )
    picture(doc, "fig_request_form.png", 16.2, "ภาพที่ 7  หน้าจอคำของบวัสดุขณะยังเป็นร่าง")
    heading(doc, "4.2 ส่งใบคำขอ", 2)
    body(doc, "เมื่อกรอกครบแล้ว กดปุ่ม ส่งใบคำขอ ที่แถบหัวเอกสาร สถานะจะเปลี่ยนจาก ร่าง เป็น ส่งงานงบ")
    callout(
        doc,
        "ส่งไม่ได้ เพราะอะไร",
        "ระบบจะกันการส่งถ้ายังไม่มีเหตุผลประกอบคำขอ ไม่มีบรรทัดรายการ "
        "หรือเป็นงบวัสดุแต่ยังไม่มีรายละเอียดสินค้าในแท็บถัดไป "
        "ให้อ่านข้อความสีแดงบนหน้าจอ แล้วแก้ในใบเดิม อย่าสร้างใบใหม่ซ้ำโดยไม่จำเป็น",
        "danger",
    )
    heading(doc, "4.3 สิ่งที่หน่วยงานทำได้และทำไม่ได้หลังส่ง", 2)
    make_table(
        doc,
        ["สถานะ", "หน่วยงานทำได้", "หน่วยงานทำไม่ได้"],
        [
            ["ร่าง", "แก้รายการ ลบใบ บันทึก และส่ง", "อนุมัติเองไม่ได้"],
            ["ส่งงานงบ / รับเรื่อง / พิจารณา", "เปิดดูสถานะ", "แก้รายการหลักไม่ได้ ต้องรองานงบตีกลับ"],
            ["อนุมัติ", "เปิดดู และนำไปใช้ตอนสร้างใบขอซื้อ", "แก้รายการหลักไม่ได้"],
            ["ตีกลับ", "ตั้งเป็นร่าง แก้ แล้วส่งใหม่", "ใช้วงเงินจากใบนี้ไม่ได้จนกว่าจะอนุมัติใหม่"],
        ],
        [4.2, 6.0, 6.0],
    )

    # Ch5
    heading(doc, "บทที่ 5 งานงบประมาณ — รับเรื่อง จัดสรร อนุมัติ", 1)
    body(
        doc,
        "บทนี้ใช้บัญชี uat.budget เปิดเมนู รวบรวมคำของบ (หน่วยงบประมาณ) "
        "งานงบเป็นคนปิดงานตั้งงบ หน่วยงานใช้จ่ายได้ก็ต่อเมื่อใบนี้ได้รับอนุมัติ",
    )
    picture(doc, "fig_request_flow.png", 16.2, "ภาพที่ 8  สถานะคำของบและปุ่มที่งานงบต้องกด")
    heading(doc, "5.1 ลำดับปุ่มบนใบคำของบ", 2)
    make_table(
        doc,
        ["ลำดับ", "ปุ่ม", "เมื่อไหร่", "ผลลัพธ์"],
        [
            ["1", "รับเรื่อง", "สถานะ ส่งงานงบ", "ยืนยันว่าเอกสารเข้าคิวงานงบแล้ว"],
            ["2", "เริ่มพิจารณา", "สถานะ รับเรื่อง", "เริ่มตรวจรายการและใส่ยอดจัดสรรได้"],
            ["3", "อนุมัติ", "สถานะ พิจารณา และใส่ยอดจัดสรรครบ", "ยอดเข้าแผนประจำปี ใช้จ่ายได้"],
            ["—", "ตีกลับ", "ยังไม่อนุมัติ", "คืนหน่วยงานให้แก้ ใบนี้ยังไม่เข้าแผน"],
            ["—", "ตั้งเป็นร่าง", "เมื่อต้องการดึงกลับ", "ใช้เมื่อต้องการให้หน่วยงานแก้ทั้งใบ"],
        ],
        [1.8, 3.0, 5.4, 6.0],
    )
    heading(doc, "5.2 ใส่ยอดจัดสรร", 2)
    body(
        doc,
        "ยอดงบประมาณที่ได้รับในแต่ละบรรทัดคือวงเงินใช้จ่ายจริง "
        "ใส่น้อยกว่าหรือเท่ากับยอดที่ขอได้ ถ้าบางบรรทัดว่าง ระบบจะไม่อนุมัติ "
        "ตัวอย่างชุดอบรม: หน่วยงานขอวัสดุ 10,000 บาท งานงบจัดสรร 8,000 บาท "
        "หน่วยงานจะซื้อได้ไม่เกิน 8,000 บาท",
    )
    callout(
        doc,
        "กติกาอนุมัติ",
        "อนุมัติได้เมื่อระบุงบประมาณที่ได้รับครบทุกบรรทัด "
        "หลังอนุมัติ ระบบซิงก์บรรทัดงบเข้าแผนประจำปีตามหน่วยงาน แหล่งเงิน และประเภทงบ "
        "ถ้าตั้งกลับเป็นร่าง บรรทัดงบที่เคยซิงก์ของใบนั้นจะถูกลบออกจากแผน",
        "info",
    )
    heading(doc, "5.3 ข้อมูลหลักที่งานงบต้องเตรียมก่อนเปิดใช้จริง", 2)
    bullets(
        doc,
        [
            "แหล่งเงิน เช่น เงินบำรุง งบประมาณ เงินบริจาค",
            "หน่วยงาน / บัญชีวิเคราะห์ ของทุกแผนกที่ใช้ระบบ",
            "แผนงบประมาณประจำปี และช่วงวันที่",
            "หมวดงบประมาณ และกลุ่มงบประมาณ-วัสดุ",
            "แบบฟอร์มคำของบทั้ง 4 ใบ พร้อมใช้งาน",
        ],
    )

    # Ch6
    heading(doc, "บทที่ 6 ใช้งบผ่านใบขอซื้อ ใบสั่งซื้อ และใบตั้งเจ้าหนี้", 1)
    body(
        doc,
        "เมื่อมีงบที่อนุมัติในแผนแล้ว จึงสร้างใบขอซื้อได้ "
        "บทนี้ต่อจากตัวอย่างวัสดุที่จัดสรร 8,000 บาท แล้วซื้อด้วยใบขอซื้อ 3,000 บาท",
    )
    picture(doc, "fig_spend_flow.png", 16.2, "ภาพที่ 9  ลำดับใช้งบจาก PR ถึงใบตั้งเจ้าหนี้")
    heading(doc, "6.1 หน่วยงานสร้างใบขอซื้อ", 2)
    steps(
        doc,
        [
            "เข้าเมนูจัดซื้อ → ใบขอซื้อ แล้วกดสร้าง ใช้บัญชี uat.dept",
            "เลือกหน่วยงานงบประมาณ แหล่งเงิน หมวดงบประมาณ ให้ตรงกับงบที่อนุมัติ",
            "ถ้าซื้อวัสดุ ต้องเลือกกลุ่มงบประมาณด้วย",
            "เลือกคลังรับของเฉพาะประเภท Receipts ของคลังนั้น",
            "เพิ่มสินค้า จำนวน และราคา รวมยอดให้ไม่เกินงบคงเหลือ",
            "บันทึก ระบบจะออกเลขที่อัตโนมัติ เช่น PR6908001 แก้เลขเองไม่ได้",
        ],
    )
    heading(doc, "6.2 กดเช็คงบประมาณ", 2)
    body(
        doc,
        "ที่หัวใบขอซื้อมีปุ่ม เช็คงบประมาณ กดตอนยังเป็นร่าง "
        "ระบบจะบอกว่างบเพียงพอหรือไม่ แต่ยังไม่กันเงิน "
        "ส่งอนุมัติได้เฉพาะเมื่อผลการเช็คงบเป็น งบเพียงพอ",
    )
    picture(doc, "fig_pr_check.png", 16.2, "ภาพที่ 10  แท็บตรวจสอบงบประมาณบนใบขอซื้อ")
    callout(
        doc,
        "เช็คงบผ่าน ยังไม่เท่ากับกันเงิน",
        "เช็คงบเป็นการถามว่ารอบนี้ซื้อได้ไหม เมื่อกดส่งอนุมัติ ระบบจึงจองงบ "
        "ใบอื่นจะใช้วงเงินซ้ำไม่ได้จนกว่าใบนี้จะถูกยกเลิก หรือถูกครอบคลุมด้วยใบสั่งซื้อ",
        "warn",
    )
    heading(doc, "6.3 พัสดุยืนยันใบสั่งซื้อ", 2)
    steps(
        doc,
        [
            "ใช้บัญชี uat.procurement เปิดใบขอซื้อที่อนุมัติแล้ว",
            "สร้างใบสั่งซื้อจาก PR ใส่ผู้ขาย ราคา ภาษี แล้วตรวจจำนวนให้ตรง",
            "กดยืนยันใบสั่งซื้อ สถานะเป็น Purchase Order",
            "ยอดงบจะย้ายจากจอง PR เป็นผูกพัน PO คงเหลือไม่เปลี่ยน ณ จุดนี้",
        ],
    )
    heading(doc, "6.4 การเงินตั้งเจ้าหนี้", 2)
    body(
        doc,
        "เมื่อได้รับใบแจ้งหนี้ผู้ขาย การเงินสร้างบิลจากใบสั่งซื้อแล้วโพสต์เอกสาร "
        "ยอดผูกพัน PO ลดลง ยอดใช้จริงเพิ่มขึ้น คงเหลือลดตามยอดบิลที่โพสต์ "
        "ถ้ายังไม่โพสต์ ตัวเลขใช้จริงจะยังไม่ขยับ",
    )

    # Ch7
    heading(doc, "บทที่ 7 สูตรงบคงเหลือ", 1)
    body(
        doc,
        "ทุกหน้าจอที่แสดงงบคงเหลือในระบบนี้ใช้สูตรเดียวกัน "
        "ถ้าตัวเลขไม่ตรง ให้เปิดรายการงบประมาณของหน่วยงานนั้น แล้วไล่ทีละช่องตามสูตร อย่าเพิ่งสรุปว่าระบบผิด",
    )
    picture(doc, "fig_formula.png", 16.2, "ภาพที่ 11  สูตรคงเหลือและตัวอย่างชุดอบรม")
    make_table(
        doc,
        ["รายการ", "ความหมาย", "เกิดเมื่อ"],
        [
            ["งบจัดสรร", "วงเงินใช้จ่ายที่งานงบอนุมัติเข้าแผน", "อนุมัติคำของบ"],
            ["ใช้จริง", "เงินที่ตัดจากบิลผู้ขายที่โพสต์แล้ว", "การเงินโพสต์ใบตั้งเจ้าหนี้"],
            ["จอง PR", "เงินที่กันไว้จากใบขอซื้อที่ส่งแล้ว", "หน่วยงานส่งอนุมัติ PR"],
            ["ผูกพัน PO", "เงินที่ผูกกับใบสั่งซื้อที่ยืนยันแล้ว", "พัสดุยืนยัน PO"],
            ["คงเหลือ", "ซื้อเพิ่มได้อีกเท่านี้", "คำนวณอัตโนมัติทั้งระบบ"],
        ],
        [3.2, 6.4, 6.6],
    )
    heading(doc, "7.1 ตัวอย่างตัวเลขชุดอบรม", 2)
    body(doc, "สมมติจัดสรรวัสดุ 8,000 บาท แล้วซื้อด้วย PR 3,000 บาท ทั้งจำนวน")
    make_table(
        doc,
        ["จังหวะ", "จัดสรร", "จอง PR", "ผูกพัน PO", "ใช้จริง", "คงเหลือ"],
        [
            ["เพิ่งอนุมัติคำของบ", "8,000", "0", "0", "0", "8,000"],
            ["เช็คงบบน PR ร่าง", "8,000", "0", "0", "0", "8,000"],
            ["ส่งอนุมัติ PR 3,000", "8,000", "3,000", "0", "0", "5,000"],
            ["ยืนยัน PO ทั้งจำนวน", "8,000", "0", "3,000", "0", "5,000"],
            ["โพสต์บิล 3,000", "8,000", "0", "0", "3,000", "5,000"],
        ],
        [4.4, 2.2, 2.2, 2.4, 2.2, 2.4],
    )
    heading(doc, "7.2 ถ้าเช็คงบไม่ผ่าน", 2)
    picture(doc, "fig_match_keys.png", 16.2, "ภาพที่ 12  สี่มิติที่ต้องตรงกับแผน")
    bullets(
        doc,
        [
            "หน่วยงาน / ศูนย์ต้นทุน ตรงกับงบในแผนหรือไม่",
            "แหล่งเงิน ตรงหรือไม่ ห้ามใช้เงินคนละแหล่งแม้หน่วยงานเดียวกัน",
            "หมวดงบประมาณ ตรงหรือไม่ วัสดุใช้หมวดวัสดุ ครุภัณฑ์ใช้หมวดครุภัณฑ์",
            "ถ้าเป็นวัสดุ เลือกกลุ่มงบประมาณ-วัสดุแล้วหรือยัง",
            "ยอดในใบขอซื้อเกินคงเหลือหรือไม่",
        ],
        numbered=True,
    )

    # Ch8
    heading(doc, "บทที่ 8 Dashboard และการติดตาม", 1)
    body(
        doc,
        "เข้าแอปงบประมาณ แล้วเปิดกลุ่มเมนู Dashboard งบประมาณ "
        "ใช้ดูภาพรวมหลังตั้งงบและหลังใช้งบ ไม่ใช่หน้าสำหรับสร้างเอกสาร",
    )
    picture(doc, "fig_dashboard.png", 16.2, "ภาพที่ 13  ภาพรวมงบประมาณรายจ่ายประจำปี")
    make_table(
        doc,
        ["เมนู", "ดูอะไร", "เหมาะกับใคร"],
        [
            ["ภาพรวมงบประมาณรายจ่ายประจำปี", "ยอดจัดสรร ใช้จริง คงเหลือ และสัดส่วนตามแบบฟอร์ม", "ผู้บริหาร งานงบ"],
            ["บริการงบประมาณรายแผนก", "เจาะแผนก โครงการ และการเบิกจ่ายสะสม", "หัวหน้าหน่วยงาน งานงบ"],
            ["สถานะคำของบประมาณ", "คิวคำขอตามสถานะ", "งานงบ หน่วยงาน"],
            ["รายการงบประมาณ", "บรรทัดงบในแผนปี", "งานงบ พัสดุ เมื่อต้องไล่ยอด"],
        ],
        [5.2, 6.2, 4.8],
    )
    callout(
        doc,
        "อ่านตัวเลขบน Dashboard อย่างไร",
        "เลือกปีงบประมาณให้ตรงปีที่กำลังใช้ แล้วเทียบกับสูตรคงเหลือ "
        "ถ้าเพิ่งส่ง PR แต่ Dashboard ยังไม่ขยับ ให้รีเฟรชหน้าจอ และตรวจว่าใบนั้นส่งอนุมัติแล้วจริง ไม่ใช่ค้างร่าง",
        "info",
    )

    # Ch9
    heading(doc, "บทที่ 9 ข้อผิดพลาดที่พบบ่อย", 1)
    picture(doc, "fig_do_dont.png", 16.2, "ภาพที่ 14  สิ่งที่ควรทำและไม่ควรทำ")
    make_table(
        doc,
        ["อาการ", "สาเหตุที่พบบ่อย", "วิธีแก้"],
        [
            ["กดส่งคำของบไม่ได้", "ไม่มีเหตุผล หรืองบวัสดุไม่มีรายละเอียดสินค้า", "กลับไปกรอกแท็บที่ขาด แล้วส่งอีกครั้ง"],
            ["งานงบกดอนุมัติไม่ได้", "ยอดจัดสรรบางบรรทัดว่าง", "ใส่ยอดงบประมาณที่ได้รับให้ครบ"],
            ["สร้าง PR แล้วเช็คงบไม่เจอวงเงิน", "หน่วยงาน แหล่งเงิน หมวด หรือกลุ่มวัสดุไม่ตรงแผน", "เปิดงบที่อนุมัติแล้วเทียบทีละช่อง"],
            ["เช็คงบผ่าน แต่ส่ง PR ไม่ได้", "ยังไม่ได้กดเช็คงบรอบล่าสุดหลังแก้ยอด", "กดเช็คงบอีกครั้ง แล้วส่ง"],
            ["ส่ง PR แล้วงบยังไม่ลด", "สับสนระหว่างเช็คงบกับจองงบ", "จองงบเกิดตอนส่งอนุมัติ ไม่ใช่ตอนเช็คงบ"],
            ["ยืนยัน PO แล้วยอดจองยังอยู่", "PO ยังไม่ครอบคลุมทั้งใบ หรือยังไม่ยืนยัน", "ตรวจว่า Confirm แล้ว และยอด PO ครอบคลุม PR"],
            ["รับของแล้วแต่ใช้จริงไม่ขึ้น", "ยังไม่ได้โพสต์ใบตั้งเจ้าหนี้", "การใช้จริงมาจากบิล ไม่ใช่จากใบรับของ"],
            ["หน่วยงานแก้ใบที่อนุมัติไม่ได้", "ระบบล็อกตามสิทธิ์", "ถูกต้องแล้ว ถ้าต้องแก้ ให้งานงบตั้งเป็นร่าง"],
        ],
        [4.4, 5.6, 6.2],
    )

    # Ch10
    heading(doc, "บทที่ 10 แบบฝึกหัดในห้องอบรม", 1)
    body(
        doc,
        "ทำตามลำดับนี้ในรอบอบรม 3 ชั่วโมง ใช้หน่วยงานทดสอบ แหล่งเงินเงินบำรุง และสินค้าชุดเดียวกันทั้งรอบ "
        "วิทยากรเป็นคนเดียวที่เปลี่ยนชุดข้อมูลกลาง",
    )
    make_table(
        doc,
        ["ข้อ", "ผู้ทำ", "งาน", "ผลที่ต้องเห็น"],
        [
            ["1", "ทุกคน", "เข้าสู่ระบบด้วยบัญชีตามบทบาท", "เห็นเมนูตามสิทธิ์"],
            ["2", "หน่วยงาน", "สร้างคำของบวัสดุ ขอ 10,000 บาท แล้วส่ง", "สถานะ ส่งงานงบ"],
            ["3", "งานงบ", "รับเรื่อง เริ่มพิจารณา จัดสรร 8,000 แล้วอนุมัติ", "มียอดในแผนปี 8,000"],
            ["4", "หน่วยงาน", "สร้าง PR 3,000 บาท กดเช็คงบ", "งบเพียงพอ ยังไม่กันเงิน"],
            ["5", "หน่วยงาน", "ส่งอนุมัติ PR", "จอง 3,000 คงเหลือ 5,000"],
            ["6", "หน่วยงาน", "ลองสร้าง PR เกินคงเหลือ", "งบไม่พอ ส่งไม่ได้"],
            ["7", "พัสดุ", "สร้างและยืนยัน PO จาก PR ข้อ 5", "ยอดย้ายเป็นผูกพัน PO"],
            ["8", "ทุกคน", "เปิด Dashboard ปี 2569 เทียบสูตร", "ตัวเลขตรงกับข้อ 5–7"],
        ],
        [1.4, 2.6, 6.2, 6.0],
    )
    callout(
        doc,
        "จบรอบอบรมเมื่อไร",
        "ผู้เรียนอธิบายสูตรคงเหลือได้ด้วยคำพูดของตนเอง สร้างคำของบวัสดุจนอนุมัติได้ "
        "และสร้างใบขอซื้อจนเช็คงบและจองงบได้ โดยไม่ต้องให้วิทยากรกดแทน",
        "ok",
    )

    heading(doc, "ภาคผนวก ก  คำศัพท์ในระบบ", 1)
    make_table(
        doc,
        ["คำในระบบ", "พูดกับผู้ใช้", "เอกสารที่เกี่ยวข้อง"],
        [
            ["คำของบ / Departmental Budget Request", "ใบตั้งงบของหน่วยงาน", "แบบฟอร์ม 4 ใบ"],
            ["งบประมาณที่ขอ", "ยอดที่หน่วยงานอยากได้", "หัวใบคำของบ"],
            ["งบประมาณที่ได้รับ / จัดสรร", "ยอดที่ใช้จ่ายได้จริง", "บรรทัดแผนปี"],
            ["ใบขอซื้อ / PR", "ใบขอซื้อของหน่วยงาน", "Purchase Request"],
            ["ใบสั่งซื้อ / PO", "ใบสั่งซื้อที่พัสดุยืนยันกับผู้ขาย", "Purchase Order"],
            ["จองงบ", "กันเงินไว้ระหว่างรอซื้อ", "ส่ง PR แล้ว"],
            ["ผูกพันงบ", "ผูกเงินกับคำสั่งซื้อ", "ยืนยัน PO แล้ว"],
            ["ใช้จริง", "ตัดงบจากบิลที่โพสต์", "ใบตั้งเจ้าหนี้"],
            ["Receipts", "ประเภทรับของเข้าคลังจากผู้ขาย", "คลังสินค้าบน PR"],
        ],
        [6.0, 5.2, 5.0],
    )

    heading(doc, "ภาคผนวก ข  ช่องทางสอบถาม", 1)
    bullets(
        doc,
        [
            "ติดขัดเรื่องสิทธิ์หรือเข้าสู่ระบบไม่ได้: งานเทคโนโลยีสารสนเทศ",
            "ติดขัดเรื่องวงเงิน แหล่งเงิน หมวดงบ: งานงบประมาณ",
            "ติดขัดเรื่องใบขอซื้อ ใบสั่งซื้อ ผู้ขาย: งานพัสดุ",
            "ติดขัดเรื่องใบตั้งเจ้าหนี้และการลงบัญชี: งานการเงิน",
        ],
    )
    body(
        doc,
        "เอกสารคู่กัน: แผนการทดสอบระบบงบประมาณ (UAT) และใบบันทึกผล UAT-งบประมาณ-ใบบันทึกผล.xlsx "
        "ใช้เมื่อต้องการตรวจว่าระบบพร้อม Go-Live ไม่ใช่เอกสารสอนในห้องอบรมผู้ใช้ทั่วไป",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    add_text(p, "— จบบทคู่มือ —", 12, True, TEAL)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p2, "โรงพยาบาลวชิระภูเก็ต  ·  ระบบงบประมาณบน Odoo 18  ·  เวอร์ชันเอกสาร 1.0", 11, False, MUTED)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
