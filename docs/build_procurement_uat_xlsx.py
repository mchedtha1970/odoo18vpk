#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Create the UAT result workbook for purchase request, purchasing, and warehouse."""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "/opt/odoo18vpk/docs/UAT-จัดซื้อคลัง-ใบบันทึกผล.xlsx"
FONT = "Prompt"

NAVY = "123A56"
TEAL = "0F7A72"
RED = "B4433A"
GREEN = "2C7A45"
LINE = "D7DEE6"
WHITE = "FFFFFF"
SOFT = "E5F5F3"
P1 = "FDECDC"
P2 = "E8F1F8"
PASS = "E5F5EA"
FAIL = "FBECEA"
SKIP = "F6F8FA"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


def font(size=11, bold=False, color="2F3A48"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def fill(color):
    return PatternFill("solid", fgColor=color)


def style_cell(cell, size=11, bold=False, color="2F3A48", bg=None, align="left", wrap=True):
    cell.font = font(size, bold, color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border = thin
    if bg:
        cell.fill = fill(bg)


def header_row(ws, row, values, bg=NAVY):
    for i, value in enumerate(values, 1):
        cell = ws.cell(row, i, value)
        style_cell(cell, 11, True, WHITE, bg, "center")
    ws.row_dimensions[row].height = 28


def landscape(ws, fit_height=1):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.page_setup.leftMargin = 0.4
    ws.page_setup.rightMargin = 0.4
    ws.page_setup.topMargin = 0.5
    ws.page_setup.bottomMargin = 0.5
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "โรงพยาบาลวชิระภูเก็ต · UAT ขอซื้อ จัดซื้อ คลังสินค้า"
    ws.oddFooter.left.text = "เอกสารภายใน สำหรับทดสอบระบบ"
    ws.oddFooter.right.text = "หน้า &P / &N"


def build_cover(wb):
    ws = wb.active
    ws.title = "ปกแผน"
    ws.sheet_view.showGridLines = False
    for col, width in enumerate([4, 22, 28, 22, 18, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.merge_cells("B2:F2")
    ws["B2"].value = "โรงพยาบาลวชิระภูเก็ต"
    ws["B2"].font = font(14, True, TEAL)
    ws.merge_cells("B3:F3")
    ws["B3"].value = "ใบบันทึกผล UAT ขอซื้อ จัดซื้อ งานคลังสินค้า"
    ws["B3"].font = font(20, True, NAVY)
    ws.merge_cells("B4:F4")
    ws["B4"].value = "ใช้คู่กับเอกสารแผนการทดสอบ · วงจร PR → PO → รับของเข้าคลัง"
    ws["B4"].font = font(12, False, "5B6878")

    info = [
        ("ระบบ", "Odoo 18 · Purchase Request / Purchase / Inventory"),
        ("ฐานข้อมูล", "VPK-S1"),
        ("รอบทดสอบ", "UAT ครั้งที่ 1"),
        ("วันที่เริ่ม", ""),
        ("วันที่จบ", ""),
        ("ผู้ประสาน", ""),
        ("สินค้าตัวอย่าง", ""),
        ("คลังรับของ", ""),
    ]
    for i, (label, value) in enumerate(info):
        row = 6 + i
        label_cell = ws.cell(row, 2, label)
        style_cell(label_cell, 11, True, WHITE, NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        value_cell = ws.cell(row, 3, value)
        style_cell(value_cell, 11, False, "2F3A48", WHITE)

    ws.merge_cells("B16:F16")
    ws["B16"].value = "เกณฑ์ผ่านรอบนี้: เคส P1 ชุด A–E ผ่านทั้งหมด · ไม่มี Blocker ค้าง · ผู้แทน 3 ฝ่ายลงนาม"
    ws["B16"].font = font(12, True, TEAL)

    header_row(ws, 18, ["ลำดับ", "ฝ่าย", "ผู้ทดสอบ", "ลงชื่อ", "วันที่", "ผ่านภาพรวม"])
    for i, role in enumerate(["หน่วยงาน", "พัสดุ / จัดซื้อ", "งานคลังสินค้า"], 1):
        row = 18 + i
        style_cell(ws.cell(row, 1, i), align="center", bg=SOFT)
        style_cell(ws.cell(row, 2, role), bold=True)
        for col in range(3, 7):
            style_cell(ws.cell(row, col, ""))

    ws.merge_cells("B23:F26")
    ws["B23"].value = (
        "วิธีใช้\n"
        "1) กรอกปกนี้ก่อนเริ่ม   2) เดินเคสตามลำดับในชีต “บันทึกผลเคส” ห้ามข้ามชุด A\n"
        "3) ถ้าไม่ผ่าน ให้เปิดรายการในชีต “บั๊ก”   4) ดูจำนวนผ่านที่ชีต “สรุป”\n"
        "บัญชี UAT: uat.dept / uat.procurement / uat.committee1–3   รหัสผ่าน: VpkUat@2569"
    )
    ws["B23"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B23"].font = font(12)
    ws["B23"].fill = fill(SOFT)
    landscape(ws, 1)


def build_cases(wb):
    ws = wb.create_sheet("บันทึกผลเคส")
    headers = [
        "รหัส", "ชุด", "P", "หัวข้อ", "ขั้นตอนโดยย่อ", "ผลที่ต้องเห็น",
        "ผล", "ผู้ทดสอบ", "วันที่", "หลักฐาน/หมายเหตุ",
    ]
    widths = [10, 18, 6, 32, 52, 50, 12, 16, 14, 28]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    header_row(ws, 1, headers)
    cases = [
        ("A01", "A สิทธิ์/ข้อมูลหลัก", "P1", "หน่วยงานสร้าง PR ได้",
         "login uat.dept เข้าเมนูใบขอซื้อ แล้วสร้างใหม่",
         "เข้าได้ บันทึกใบร่างได้"),
        ("A02", "A สิทธิ์/ข้อมูลหลัก", "P1", "หน่วยงานสร้าง PO เองไม่ได้",
         "login หน่วยงาน พยายามเข้าเมนูใบสั่งซื้อ",
         "เข้าไม่ได้ หรือสร้าง PO เองไม่ได้"),
        ("A03", "A สิทธิ์/ข้อมูลหลัก", "P1", "พัสดุเห็น PR และสร้าง PO ได้",
         "login uat.procurement เข้าใบขอซื้อและใบสั่งซื้อ",
         "เห็น PR ของหน่วยงาน และสร้าง PO ได้"),
        ("A04", "A สิทธิ์/ข้อมูลหลัก", "P1", "งานคลังเข้าใบรับสินค้าได้",
         "เข้าเมนูคลัง / ใบรับสินค้า",
         "เห็นคลังที่รับผิดชอบ และสร้างใบรับได้"),
        ("A05", "A สิทธิ์/ข้อมูลหลัก", "P1", "สินค้าทดสอบพร้อม",
         "เปิดสินค้าที่ใช้ทดสอบ",
         "มีหน่วยนับ และเส้นทางคลังพร้อมรับของ"),
        ("A06", "A สิทธิ์/ข้อมูลหลัก", "P1", "ผู้ขายทดสอบพร้อม",
         "เปิดผู้ขายที่ใช้ใน PO",
         "เลือกในใบสั่งซื้อได้"),
        ("A07", "A สิทธิ์/ข้อมูลหลัก", "P2", "หน่วยงานเปิด PR ได้โดยไม่ติดกรรมการ",
         "เปิดใบขอซื้อที่มีช่องคณะกรรมการ/สเปก",
         "เปิด PR ได้ ไม่บังคับสิทธิ์จัดซื้อ"),
        ("B01", "B ขอซื้อผ่าน", "P1", "สร้าง PR วัสดุ 3,000 บาท",
         "เลือกหน่วยงาน แหล่งเงิน หมวดงบ กลุ่มวัสดุ คลัง Receipt ใส่สินค้า แล้วบันทึก",
         "ได้เลขอัตโนมัติ เช่น PR6908xxx สถานะร่าง แก้เลขที่ไม่ได้"),
        ("B02", "B ขอซื้อผ่าน", "P1", "dropdown คลังเป็น Receipt เท่านั้น",
         "เปิดช่องประเภทการรับสินค้า",
         "เห็นเฉพาะ Receipts ไม่มี Internal Transfers / Delivery / Manufacturing"),
        ("B03", "B ขอซื้อผ่าน", "P1", "เช็คงบบน PR ร่าง",
         "กดเช็คงบ ยอด 3,000",
         "งบเพียงพอ ยังไม่กันเงิน"),
        ("B04", "B ขอซื้อผ่าน", "P1", "ส่งอนุมัติ PR แล้วจองงบ",
         "กดส่ง / ขออนุมัติ",
         "สถานะรออนุมัติ และจองงบ 3,000"),
        ("B05", "B ขอซื้อผ่าน", "P1", "อนุมัติ PR ครบขั้น",
         "ผู้อนุมัติตาม tier อนุมัติครบ",
         "สถานะ Approved สร้าง PO จาก PR ได้"),
        ("B06", "B ขอซื้อผ่าน", "P2", "หน่วยงานแก้ PR ที่อนุมัติแล้ว",
         "login หน่วยงาน เปิดใบ Approved แล้วพยายามแก้รายการ",
         "แก้รายการหลักไม่ได้"),
        ("C01", "C กติกาขอซื้อ", "P1", "ห้ามพิมพ์เลข PR เอง",
         "พยายามพิมพ์เลขที่ PR เอง",
         "ช่องเลขที่ล็อก ระบบออกเลขให้"),
        ("C02", "C กติกาขอซื้อ", "P1", "ห้ามเลือกคลังที่ไม่ใช่ Receipt",
         "พยายามเลือก Internal Transfers / Delivery",
         "เลือกไม่ได้ เพราะถูกกรองออก"),
        ("C03", "C กติกาขอซื้อ", "P1", "PR เกินงบคงเหลือ",
         "สร้าง PR ยอดเกินคงเหลือ แล้วเช็คงบ / ส่ง",
         "งบไม่พอ ส่งอนุมัติไม่ได้"),
        ("C04", "C กติกาขอซื้อ", "P1", "PR วัสดุไม่ครบมิติงบ",
         "ไม่ใส่หน่วยงาน / แหล่งเงิน / กลุ่มวัสดุ แล้วเช็คงบ",
         "ไม่เจอวงเงิน หรือบังคับกรอก"),
        ("C05", "C กติกาขอซื้อ", "P2", "ส่ง PR ไม่มีบรรทัดสินค้า",
         "สร้าง PR ไม่ใส่สินค้า แล้วส่ง",
         "ส่งไม่ได้"),
        ("C06", "C กติกาขอซื้อ", "P2", "PR ครุภัณฑ์ไม่บังคับกลุ่มวัสดุ",
         "สร้าง PR ครุภัณฑ์ ไม่ใส่กลุ่มวัสดุ แล้วเช็คงบ",
         "สร้างและเช็คงบได้"),
        ("D01", "D จัดซื้อ", "P1", "สร้าง PO จาก PR ที่อนุมัติ",
         "พัสดุสร้างใบสั่งซื้อจาก PR ใน B05",
         "ได้ RFQ/PO อ้างอิง PR สินค้าและจำนวนตรง PR"),
        ("D02", "D จัดซื้อ", "P1", "ยืนยัน PO",
         "ใส่ผู้ขาย ราคา แล้ว Confirm",
         "สถานะ Purchase Order ยอดย้ายจากจอง PR เป็นผูกพัน PO"),
        ("D03", "D จัดซื้อ", "P1", "PR ต้นทางถูกครอบคลุมด้วย PO",
         "เปิด PR หลังยืนยัน PO",
         "เห็นว่าถูกครอบคลุม / คงเหลือที่ต้องสั่งลดลง"),
        ("D04", "D จัดซื้อ", "P1", "ดูใบสั่งซื้อ",
         "พิมพ์หรือเปิดฟอร์ม PO",
         "เปิดได้ มีผู้ขาย สินค้า จำนวน ราคา"),
        ("D05", "D จัดซื้อ", "P2", "หน่วยงานยืนยัน PO เองไม่ได้",
         "login หน่วยงาน พยายาม Confirm PO",
         "ทำไม่ได้ตามสิทธิ์"),
        ("D06", "D จัดซื้อ", "P2", "สร้าง PO โดยไม่ผ่าน PR",
         "พยายามสร้าง PO ตรงโดยไม่ผ่าน PR",
         "ระบบกัน หรือมีคำเตือนตามนโยบายโรงพยาบาล"),
        ("E01", "E คลังสินค้า", "P1", "เปิดใบรับจาก PO",
         "งานคลังเปิด Receipt จาก PO ใน D02",
         "มีใบ Receipt ของคลังที่เลือก จำนวนตาม PO"),
        ("E02", "E คลังสินค้า", "P1", "รับของครบแล้ว Validate",
         "รับครบจำนวน แล้ว Validate",
         "สถานะ Done สต็อกในคลังนั้นเพิ่มตามจำนวนที่รับ"),
        ("E03", "E คลังสินค้า", "P1", "ตรวจคงเหลือสินค้า",
         "เปิดสินค้า ดู On Hand ของคลังที่รับ",
         "จำนวนตรงกับที่เพิ่งรับ"),
        ("E04", "E คลังสินค้า", "P1", "ของเข้าคลังตาม Receipt ที่เลือก",
         "เทียบคลังในใบรับกับคลังที่เลือกบน PR",
         "ไม่ไปเข้าคลังอื่น"),
        ("E05", "E คลังสินค้า", "P2", "โอนภายในคลัง",
         "โอนจากคลังรับไปคลังหน่วยงาน (ถ้าใช้จริง)",
         "ต้นทางลด ปลายทางเพิ่ม"),
        ("E06", "E คลังสินค้า", "P2", "จ่ายของออกจากคลัง",
         "ทำใบจ่าย / Delivery ตามกระบวนการโรงพยาบาล",
         "สต็อกลดตามจำนวนที่จ่าย"),
        ("E07", "E คลังสินค้า", "P2", "ผู้ไม่มีสิทธิ์ Validate ใบรับไม่ได้",
         "login user ที่ไม่มีสิทธิ์คลัง แล้วพยายาม Validate",
         "ทำไม่ได้"),
        ("F01", "F รับบางส่วน/ยกเลิก", "P1", "รับของน้อยกว่าที่สั่ง",
         "สร้าง PR+PO ใหม่ แล้วรับน้อยกว่าจำนวน PO",
         "สต็อกเพิ่มเท่าที่รับ PO ยังค้างรับส่วนที่เหลือ"),
        ("F02", "F รับบางส่วน/ยกเลิก", "P1", "รับส่วนที่เหลือ",
         "รับส่วนที่ค้างในใบถัดไป",
         "คงเหลือค้างรับเป็น 0"),
        ("F03", "F รับบางส่วน/ยกเลิก", "P1", "ห้ามรับเกิน PO",
         "พยายามรับเกินจำนวนที่สั่ง",
         "รับไม่ได้ หรือระบบเตือน"),
        ("F04", "F รับบางส่วน/ยกเลิก", "P1", "ยกเลิก PR ร่างที่เช็คงบแล้ว",
         "เช็คงบแล้วยกเลิกโดยยังไม่ส่ง",
         "ไม่กันงบ"),
        ("F05", "F รับบางส่วน/ยกเลิก", "P1", "ยกเลิก PR ที่จองงบแล้ว",
         "ยกเลิก PR ที่ส่งแล้วและยังไม่มี PO",
         "คืนวงเงินจอง"),
        ("F06", "F รับบางส่วน/ยกเลิก", "P2", "ยกเลิก PO ที่ยังไม่รับของ",
         "ยกเลิก PO ที่ยืนยันแล้วแต่ยังไม่รับของ",
         "คืนผูกพันงบ ตามกติกาที่ตั้งไว้"),
        ("F07", "F รับบางส่วน/ยกเลิก", "P2", "คืนของเข้าคลังหลังรับแล้ว",
         "ทำ Return จากใบรับ",
         "สต็อกลด และเอกสารคืนของอ้างอิงต้นทาง"),
        ("G01", "G เชื่อมงบ", "P1", "เทียบจองหลังส่ง PR",
         "ดูยอดงบหลัง B04",
         "จอง PR เพิ่ม คงเหลือลด"),
        ("G02", "G เชื่อมงบ", "P1", "เทียบยอดหลังยืนยัน PO",
         "ดูยอดงบหลัง D02",
         "จองลด ผูกพันเพิ่ม คงเหลือเท่าเดิม"),
        ("G03", "G เชื่อมงบ", "P1", "โพสต์ใบตั้งเจ้าหนี้",
         "ตั้งเจ้าหนี้ตาม PO แล้วโพสต์",
         "ผูกพันลด ใช้จริงเพิ่ม"),
        ("G04", "G เชื่อมงบ", "P1", "เทียบสูตรคงเหลือ",
         "คำนวณมือ: จัดสรร − ใช้จริง − จอง PR − ผูกพัน PO",
         "ตัวเลขตรงหน้าจอ"),
    ]
    for r, row in enumerate(cases, 2):
        for col, value in enumerate(row, 1):
            bg = P1 if (col == 3 and value == "P1") else P2 if (col == 3 and value == "P2") else None
            align = "center" if col in (1, 3, 7, 8, 9) else "left"
            style_cell(ws.cell(r, col, value), bg=bg, align=align)
        ws.row_dimensions[r].height = 46
    dv = DataValidation(type="list", formula1='"ผ่าน,ไม่ผ่าน,ข้าม"', allow_blank=True)
    dv.error = "เลือก ผ่าน / ไม่ผ่าน / ข้าม"
    dv.errorTitle = "ผลทดสอบ"
    ws.add_data_validation(dv)
    dv.add("G2:G80")
    ws.conditional_formatting.add(
        "G2:G80",
        CellIsRule(operator="equal", formula=['"ผ่าน"'], fill=fill(PASS), font=font(11, True, GREEN)),
    )
    ws.conditional_formatting.add(
        "G2:G80",
        CellIsRule(operator="equal", formula=['"ไม่ผ่าน"'], fill=fill(FAIL), font=font(11, True, RED)),
    )
    ws.conditional_formatting.add(
        "G2:G80",
        CellIsRule(operator="equal", formula=['"ข้าม"'], fill=fill(SKIP), font=font(11, False, "5B6878")),
    )
    ws.freeze_panes = "A2"
    last = 1 + len(cases)
    ws.auto_filter.ref = f"A1:J{last}"
    landscape(ws, 0)
    return len(cases)


def build_bugs(wb):
    ws = wb.create_sheet("บั๊ก")
    headers = [
        "รหัสบั๊ก", "เคสที่พบ", "ระดับ", "สรุปปัญหา", "ขั้นตอนที่เกิด",
        "ผลที่คาด", "ผลที่เกิดจริง", "ผู้พบ", "วันที่", "สถานะแก้",
    ]
    widths = [12, 12, 14, 28, 32, 26, 26, 14, 14, 18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    header_row(ws, 1, headers, RED)
    for row in range(2, 21):
        for col in range(1, 11):
            style_cell(
                ws.cell(row, col, f"BUG-{row-1:02d}" if col == 1 else ""),
                align="center" if col in (1, 3, 9, 10) else "left",
            )
    sev = DataValidation(type="list", formula1='"Blocker,Major,Minor,Cosmetic"', allow_blank=True)
    st = DataValidation(type="list", formula1='"เปิด,กำลังแก้,แก้แล้วรอทดสอบซ้ำ,ปิด"', allow_blank=True)
    ws.add_data_validation(sev)
    ws.add_data_validation(st)
    sev.add("C2:C20")
    st.add("J2:J20")
    ws.freeze_panes = "A2"
    landscape(ws, 1)


def build_summary(wb):
    ws = wb.create_sheet("สรุป")
    for col, width in enumerate([6, 24, 12, 12, 12, 12, 12, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.merge_cells("B2:H2")
    ws["B2"].value = "สรุปผลการทดสอบขอซื้อ จัดซื้อ งานคลังสินค้า"
    ws["B2"].font = font(18, True, NAVY)
    header_row(ws, 4, ["", "ชุด", "ทั้งหมด", "ผ่าน", "ไม่ผ่าน", "ข้าม", "ยังไม่ทำ", "% ผ่าน P1"])
    groups = [
        ("A สิทธิ์/ข้อมูลหลัก", "A"),
        ("B ขอซื้อผ่าน", "B"),
        ("C กติกาขอซื้อ", "C"),
        ("D จัดซื้อ", "D"),
        ("E คลังสินค้า", "E"),
        ("F รับบางส่วน/ยกเลิก", "F"),
        ("G เชื่อมงบ", "G"),
    ]
    for i, (name, prefix) in enumerate(groups):
        row = 5 + i
        style_cell(ws.cell(row, 1, i + 1), align="center", bg=SOFT)
        style_cell(ws.cell(row, 2, name), bold=True)
        formulas = {
            3: f'=COUNTIF(\'บันทึกผลเคส\'!A:A,"{prefix}*")',
            4: f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",\'บันทึกผลเคส\'!G:G,"ผ่าน")',
            5: f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",\'บันทึกผลเคส\'!G:G,"ไม่ผ่าน")',
            6: f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",\'บันทึกผลเคส\'!G:G,"ข้าม")',
            7: f"=C{row}-D{row}-E{row}-F{row}",
            8: (
                f'=IFERROR(COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",'
                f'\'บันทึกผลเคส\'!C:C,"P1",\'บันทึกผลเคส\'!G:G,"ผ่าน")/'
                f'COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",\'บันทึกผลเคส\'!C:C,"P1"),0)'
            ),
        }
        for col, formula in formulas.items():
            style_cell(ws.cell(row, col, formula), align="center")
        ws.cell(row, 8).number_format = "0%"

    row = 13
    style_cell(ws.cell(row, 2, "รวมทุกชุด"), 12, True, WHITE, NAVY)
    for col in range(3, 8):
        letter = get_column_letter(col)
        style_cell(ws.cell(row, col, f"=SUM({letter}5:{letter}11)"), 12, True, WHITE, NAVY, "center")
    style_cell(
        ws.cell(
            row,
            8,
            '=IFERROR(COUNTIFS(\'บันทึกผลเคส\'!C:C,"P1",\'บันทึกผลเคส\'!G:G,"ผ่าน")/'
            'COUNTIF(\'บันทึกผลเคส\'!C:C,"P1"),0)',
        ),
        12,
        True,
        WHITE,
        NAVY,
        "center",
    )
    ws.cell(row, 8).number_format = "0%"

    ws.merge_cells("B16:H19")
    ws["B16"].value = (
        "ตัดสินใจ:  ผ่าน UAT และไป Go-Live ได้   /   ผ่านแบบมีเงื่อนไข   /   ต้องทดสอบซ้ำ\n"
        "เงื่อนไขที่ค้าง: .................................................................\n"
        "ลงนามผู้ประสาน: ....................................     วันที่: ................"
    )
    ws["B16"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B16"].font = font(12)
    landscape(ws, 1)


def main():
    wb = Workbook()
    build_cover(wb)
    count = build_cases(wb)
    build_bugs(wb)
    build_summary(wb)
    wb.save(OUT)
    print(count, OUT)


if __name__ == "__main__":
    main()
