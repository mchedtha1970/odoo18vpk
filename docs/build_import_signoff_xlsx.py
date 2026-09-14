#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export data-migration inventory + sign-off workbooks from live VPK-S1 counts."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from psycopg2.extras import RealDictCursor

OUT_DIR = Path("/opt/odoo18vpk/docs/UAT-Data-Migrate")
DBNAME = "VPK-S1"
FONT = "Prompt"
NAVY, TEAL, RED, GREEN = "123A56", "0F7A72", "B4433A", "2C7A45"
LINE, WHITE, SOFT = "D7DEE6", "FFFFFF", "E5F5F3"
P1, PASS, FAIL, SKIP = "FDECDC", "E5F5EA", "FBECEA", "F6F8FA"
AMBER = "F4E4C8"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
TH_MONTHS = "ม.ค. ก.พ. มี.ค. เม.ย. พ.ค. มิ.ย. ก.ค. ส.ค. ก.ย. ต.ค. พ.ย. ธ.ค.".split()
CONFIRM = '"ถูกต้อง,ต้องแก้ไข,ยังไม่ครบ,ข้าม"'


def fnt(size=11, bold=False, color="2F3A48"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def fill(c):
    return PatternFill("solid", fgColor=c)


def sc(cell, size=11, bold=False, color="2F3A48", bg=None, align="left"):
    cell.font = fnt(size, bold, color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = thin
    if bg:
        cell.fill = fill(bg)


def hdr(ws, row, values, bg=NAVY):
    for i, v in enumerate(values, 1):
        sc(ws.cell(row, i, v), 11, True, WHITE, bg, "center")
    ws.row_dimensions[row].height = 28


def page(ws, title, fit=1):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit
    ws.page_setup.leftMargin = ws.page_setup.rightMargin = 0.35
    ws.page_setup.topMargin = ws.page_setup.bottomMargin = 0.45
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = f"โรงพยาบาลวชิระภูเก็ต · {title}"
    ws.oddFooter.left.text = "เอกสารยืนยันข้อมูลนำเข้า (Data Migration Sign-off) · ใช้ภายใน"
    ws.oddFooter.right.text = "หน้า &P / &N"


def thai_dt(dt: datetime) -> str:
    return f"{dt.day} {TH_MONTHS[dt.month - 1]} {dt.year + 543}  {dt.strftime('%H:%M')}"


def file_stamp(dt: datetime | None = None) -> str:
    dt = dt or datetime.now()
    return f"{dt.year + 543}-{dt.month:02d}-{dt.day:02d}"


def dated_xlsx(name: str, stamp: str | None = None) -> str:
    p = Path(name)
    return f"{p.stem}-{stamp or file_stamp()}{p.suffix}"


def jname_sql(alias="name"):
    return f"COALESCE({alias}->>'th_TH', {alias}->>'en_US', {alias}::text)"


def account_code_sql(alias):
    return (
        f"(SELECT v FROM jsonb_each_text(COALESCE({alias}.code_store::jsonb, '{{}}'::jsonb))"
        f" t(k, v) LIMIT 1)"
    )


def connect():
    return psycopg2.connect(
        dbname=DBNAME,
        user="odoo18vpk",
        password="odoo18vpk",
        host="/var/run/postgresql",
        connect_timeout=15,
    )


def fetch_all(cur, sql):
    cur.execute(sql)
    rows = cur.fetchall()
    cols = [d[0] for d in cur.description]
    return cols, rows


def count_sql(cur, sql):
    cur.execute(sql)
    row = cur.fetchone()
    if row is None:
        return 0
    return int(next(iter(row.values() if isinstance(row, dict) else row)))


def status_of(kind, n):
    if n <= 0:
        if kind == "คงค้าง":
            return "ยังไม่นำเข้า (คงค้าง)"
        return "ยังไม่นำเข้า"
    if kind == "คงค้าง":
        return "มีรายการคงค้าง"
    if kind == "ธุรกรรม":
        return "มีข้อมูลในระบบ (ธุรกรรม)"
    return "นำเข้าแล้ว"


def status_bg(st):
    if st == "นำเข้าแล้ว":
        return PASS
    if st.startswith("มีข้อมูล") or st == "มีรายการคงค้าง":
        return AMBER
    return FAIL


# Live imported vendor master — exclude DEMO / ตัวอย่าง / API sample codes.
REAL_VENDOR_SQL = """
    p.supplier_rank > 0
    AND p.parent_id IS NULL
    AND p.name NOT ILIKE '%ตัวอย่าง%'
    AND p.name NOT ILIKE '%DEMO%'
    AND p.name NOT ILIKE '%เอกสารทางการค้า%'
    AND COALESCE(p.ref, '') NOT IN (
        'VENDOR-DEMO-001', 'EXT-001', 'EXT-002', 'HTTP-EXT-100',
        'DOC-EXT-200', 'SUP-PRICE-001', 'SUP-PRICE-002', 'VENDOR-RED-CROSS-SAMPLE'
    )
    AND COALESCE(p.vpk_vendor_external_id, '') NOT IN (
        'EXT-001', 'EXT-002', 'HTTP-EXT-100', 'DOC-EXT-200', 'TRCS'
    )
"""


DATASETS = [
    # ---- งบประมาณ ----
    {
        "id": "fund",
        "mod": "budget",
        "kind": "Master",
        "name": "แหล่งเงิน",
        "related": "งบประมาณ",
        "model": "vpk.budget.fund.source",
        "key": "code",
        "sheet": "ยืนยัน-แหล่งเงิน",
        "count": "SELECT COUNT(*) FROM vpk_budget_fund_source",
        "list": f"""
            SELECT sequence AS ลำดับ, code AS รหัส, name AS ชื่อ,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM vpk_budget_fund_source ORDER BY sequence, id
        """,
    },
    {
        "id": "bgroup",
        "mod": "budget",
        "kind": "Master",
        "name": "กลุ่มงบประมาณ-วัสดุ",
        "related": "งบประมาณ",
        "model": "vpk.budget.group",
        "key": "code",
        "sheet": "ยืนยัน-กลุ่มงบ",
        "count": "SELECT COUNT(*) FROM vpk_budget_group",
        "list": """
            SELECT sequence AS ลำดับ, code AS รหัส, name AS ชื่อ,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM vpk_budget_group ORDER BY sequence, id
        """,
    },
    {
        "id": "msub",
        "mod": "budget",
        "kind": "Master",
        "name": "ประเภทวัสดุย่อย",
        "related": "งบประมาณ",
        "model": "vpk.budget.material.sub.type",
        "key": "code + กลุ่มงบ",
        "sheet": "ยืนยัน-วัสดุย่อย",
        "count": "SELECT COUNT(*) FROM vpk_budget_material_sub_type",
        "list": """
            SELECT s.sequence AS ลำดับ, s.code AS รหัส, s.name AS ชื่อ,
                   g.code AS รหัสกลุ่มงบ, g.name AS กลุ่มงบ,
                   CASE WHEN s.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM vpk_budget_material_sub_type s
            LEFT JOIN vpk_budget_group g ON g.id = s.budget_group_id
            ORDER BY g.sequence, s.sequence, s.id
        """,
    },
    {
        "id": "btype",
        "mod": "budget",
        "kind": "Master",
        "name": "ประเภทงบประมาณ",
        "related": "งบประมาณ",
        "model": "vpk.budget.type",
        "key": "code",
        "sheet": "ยืนยัน-ประเภทงบ",
        "count": "SELECT COUNT(*) FROM vpk_budget_type",
        "list": f"""
            SELECT sequence AS ลำดับ, code AS รหัส, {jname_sql()} AS ชื่อ,
                   section_key AS กลุ่มแบบฟอร์ม,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM vpk_budget_type ORDER BY sequence, id
        """,
    },
    {
        "id": "acat",
        "mod": "budget",
        "kind": "Master",
        "name": "กลุ่มครุภัณฑ์",
        "related": "งบประมาณ",
        "model": "vpk.budget.asset.category",
        "key": "code",
        "sheet": "ยืนยัน-กลุ่มครุภัณฑ์",
        "count": "SELECT COUNT(*) FROM vpk_budget_asset_category",
        "list": """
            SELECT a.sequence AS ลำดับ, a.code AS รหัส, a.name AS ชื่อ,
                   pc.complete_name AS หมวดสินค้า,
                   CASE WHEN a.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM vpk_budget_asset_category a
            LEFT JOIN product_category pc ON pc.id = a.product_categ_id
            ORDER BY a.sequence, a.id
        """,
    },
    {
        "id": "form",
        "mod": "budget",
        "kind": "Master",
        "name": "แบบฟอร์มคำของบ",
        "related": "งบประมาณ",
        "model": "vpk.budget.request.form.type",
        "key": "code",
        "sheet": "ยืนยัน-แบบฟอร์ม",
        "count": "SELECT COUNT(*) FROM vpk_budget_request_form_type",
        "list": f"""
            SELECT sequence AS ลำดับ, code AS รหัส, {jname_sql()} AS ชื่อ,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM vpk_budget_request_form_type ORDER BY sequence, id
        """,
    },
    {
        "id": "formline",
        "mod": "budget",
        "kind": "Master",
        "name": "ประเภทงบในแบบฟอร์ม",
        "related": "งบประมาณ",
        "model": "vpk.budget.request.form.type.line",
        "key": "form + ประเภทงบ",
        "sheet": "ยืนยัน-แท็บแบบฟอร์ม",
        "count": "SELECT COUNT(*) FROM vpk_budget_request_form_type_line",
        "list": f"""
            SELECT f.code AS รหัสแบบฟอร์ม, {jname_sql('f.name')} AS แบบฟอร์ม,
                   t.code AS รหัสประเภทงบ, {jname_sql('t.name')} AS ประเภทงบ,
                   l.sequence AS ลำดับ
            FROM vpk_budget_request_form_type_line l
            JOIN vpk_budget_request_form_type f ON f.id = l.form_type_id
            LEFT JOIN vpk_budget_type t ON t.id = l.budget_type_id
            ORDER BY f.sequence, l.sequence
        """,
    },
    {
        "id": "aplan",
        "mod": "budget",
        "kind": "Master",
        "name": "แผนวิเคราะห์ / มิติหน่วยงาน",
        "related": "งบประมาณ, บัญชี",
        "model": "account.analytic.plan",
        "key": "name",
        "sheet": "ยืนยัน-แผนวิเคราะห์",
        "count": "SELECT COUNT(*) FROM account_analytic_plan",
        "list": f"""
            SELECT id AS รหัสภายใน, {jname_sql()} AS ชื่อ, complete_name AS ชื่อเต็ม
            FROM account_analytic_plan ORDER BY id
        """,
    },
    {
        "id": "aacc",
        "mod": "budget",
        "kind": "Master",
        "name": "บัญชีวิเคราะห์ / หน่วยงาน",
        "related": "งบประมาณ, จัดซื้อ, บัญชี",
        "model": "account.analytic.account",
        "key": "code",
        "sheet": "ยืนยัน-หน่วยงานวิเคราะห์",
        "count": "SELECT COUNT(*) FROM account_analytic_account",
        "list": f"""
            SELECT code AS รหัส, {jname_sql('a.name')} AS ชื่อ,
                   {jname_sql('p.name')} AS แผน,
                   CASE WHEN a.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM account_analytic_account a
            LEFT JOIN account_analytic_plan p ON p.id = a.plan_id
            ORDER BY a.code NULLS LAST, a.id
        """,
    },
    {
        "id": "bplan",
        "mod": "budget",
        "kind": "Master",
        "name": "แผนงบประมาณประจำปี",
        "related": "งบประมาณ",
        "model": "budget.budget",
        "key": "name + ช่วงวันที่",
        "sheet": "ยืนยัน-แผนปี",
        "count": "SELECT COUNT(*) FROM budget_budget",
        "list": """
            SELECT name AS ชื่อแผน, date_from AS เริ่ม, date_to AS สิ้นสุด, state AS สถานะ
            FROM budget_budget ORDER BY date_from, id
        """,
    },
    {
        "id": "bpost",
        "mod": "budget",
        "kind": "Master",
        "name": "หมวดงบประมาณ",
        "related": "งบประมาณ, จัดซื้อ",
        "model": "account.budget.post",
        "key": "name",
        "sheet": "ยืนยัน-หมวดงบ",
        "count": "SELECT COUNT(*) FROM account_budget_post",
        "list": """
            SELECT name AS ชื่อหมวด FROM account_budget_post ORDER BY name
        """,
    },
    {
        "id": "bline",
        "mod": "budget",
        "kind": "ตั้งค่า",
        "name": "บรรทัดวงเงินงบประมาณ",
        "related": "งบประมาณ",
        "model": "budget.lines",
        "key": "แผน + หมวด + หน่วยงาน",
        "sheet": "ยืนยัน-วงเงินงบ",
        "count": "SELECT COUNT(*) FROM budget_lines",
        "list": None,
    },
    {
        "id": "deptreq",
        "mod": "budget",
        "kind": "ธุรกรรม",
        "name": "คำของบหน่วยงาน",
        "related": "งบประมาณ",
        "model": "departmental.budget.request",
        "key": "เลขที่เอกสาร",
        "sheet": None,
        "count": "SELECT COUNT(*) FROM departmental_budget_request",
        "list": None,
    },
    # ---- จัดซื้อ ----
    {
        "id": "vendor",
        "mod": "purchase",
        "kind": "Master",
        "name": "ผู้จำหน่าย / เจ้าหนี้",
        "related": "จัดซื้อ, เจ้าหนี้",
        "model": "res.partner (ผู้ขายที่นำเข้า, ไม่รวมข้อมูลตัวอย่าง)",
        "key": "ref / vat / ชื่อ",
        "sheet": "ยืนยัน-ผู้จำหน่าย",
        "count": f"SELECT COUNT(*) FROM res_partner p WHERE {REAL_VENDOR_SQL}",
        "list": f"""
            SELECT p.id AS รหัสภายใน,
                   p.ref AS รหัสอ้างอิง,
                   p.vpk_vendor_external_id AS รหัสจากไฟล์นำเข้า,
                   p.name AS ชื่อ,
                   p.vat AS เลขผู้เสียภาษี,
                   CASE WHEN p.is_company THEN 'นิติบุคคล' ELSE 'บุคคล' END AS ประเภท,
                   p.street AS ที่อยู่,
                   p.street2 AS ตำบล,
                   p.city AS อำเภอ,
                   COALESCE(s.name, '') AS จังหวัด,
                   p.zip AS รหัสไปรษณีย์,
                   COALESCE(p.phone, p.mobile) AS โทร,
                   p.email AS อีเมล,
                   COALESCE(g.name, '') AS กลุ่มผู้จำหน่าย,
                   CASE WHEN p.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM res_partner p
            LEFT JOIN res_country_state s ON s.id = p.state_id
            LEFT JOIN procurement_vendor_group g ON g.id = p.procurement_vendor_group_id
            WHERE {REAL_VENDOR_SQL}
            ORDER BY g.sequence NULLS LAST, p.name
        """,
    },
    {
        "id": "vbank",
        "mod": "purchase",
        "kind": "Master",
        "name": "บัญชีธนาคารผู้ขาย",
        "related": "จัดซื้อ, การเงิน",
        "model": "res.partner.bank",
        "key": "partner + เลขบัญชี",
        "sheet": "ยืนยัน-บัญชีผู้ขาย",
        "count": f"""
            SELECT COUNT(*) FROM res_partner_bank b
            JOIN res_partner p ON p.id = b.partner_id
            WHERE {REAL_VENDOR_SQL}
        """,
        "list": f"""
            SELECT p.ref AS รหัสผู้ขาย, p.vpk_vendor_external_id AS รหัสจากไฟล์นำเข้า,
                   p.name AS ผู้ขาย, b.acc_number AS เลขบัญชี,
                   CASE WHEN b.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM res_partner_bank b
            JOIN res_partner p ON p.id = b.partner_id
            WHERE {REAL_VENDOR_SQL}
            ORDER BY p.name
        """,
    },
    {
        "id": "sinfo",
        "mod": "purchase",
        "kind": "Master",
        "name": "ราคา / MOQ ผู้ขาย",
        "related": "จัดซื้อ, สินค้า",
        "model": "product.supplierinfo",
        "key": "ผู้ขาย + สินค้า",
        "sheet": "ยืนยัน-ราคาผู้ขาย",
        "count": f"""
            SELECT COUNT(*) FROM product_supplierinfo s
            LEFT JOIN res_partner p ON p.id = s.partner_id
            WHERE p.id IS NULL OR ({REAL_VENDOR_SQL})
        """,
        "list": f"""
            SELECT p.name AS ผู้ขาย, t.default_code AS รหัสสินค้า,
                   {jname_sql('t.name')} AS สินค้า,
                   s.min_qty AS จำนวนขั้นต่ำ, s.price AS ราคา, s.delay AS ระยะเวลานำ
            FROM product_supplierinfo s
            LEFT JOIN res_partner p ON p.id = s.partner_id
            LEFT JOIN product_template t ON t.id = s.product_tmpl_id
            WHERE p.id IS NULL OR ({REAL_VENDOR_SQL})
            ORDER BY p.name, t.default_code
        """,
    },
    {
        "id": "vgroup",
        "mod": "purchase",
        "kind": "Master",
        "name": "กลุ่มผู้จำหน่ายจัดซื้อ",
        "related": "จัดซื้อ, เจ้าหนี้",
        "model": "procurement.vendor.group",
        "key": "code",
        "sheet": "ยืนยัน-กลุ่มผู้จำหน่าย",
        "count": "SELECT COUNT(*) FROM procurement_vendor_group",
        "list": f"""
            SELECT sequence AS ลำดับ, code AS รหัส, name AS ชื่อ,
                   description AS รายละเอียด,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ,
                   (SELECT COUNT(*) FROM res_partner p
                    WHERE p.procurement_vendor_group_id = g.id
                      AND ({REAL_VENDOR_SQL})) AS จำนวนผู้ขาย
            FROM procurement_vendor_group g
            ORDER BY sequence, id
        """,
    },
    {
        "id": "pr",
        "mod": "purchase",
        "kind": "ธุรกรรม",
        "name": "ใบขอซื้อ (PR)",
        "related": "จัดซื้อ",
        "model": "purchase.request",
        "key": "name",
        "sheet": "ยืนยัน-PR",
        "count": "SELECT COUNT(*) FROM purchase_request",
        "list": """
            SELECT name AS เลขที่, date_start AS วันที่, state AS สถานะ
            FROM purchase_request ORDER BY id
        """,
    },
    {
        "id": "po",
        "mod": "purchase",
        "kind": "ธุรกรรม",
        "name": "ใบสั่งซื้อ (PO) ทั้งหมด",
        "related": "จัดซื้อ, คลัง, เจ้าหนี้",
        "model": "purchase.order",
        "key": "name",
        "sheet": "ยืนยัน-POทั้งหมด",
        "count": "SELECT COUNT(*) FROM purchase_order",
        "list": """
            SELECT po.name AS เลขที่PO, po.date_order AS วันที่สั่ง,
                   COALESCE(rp.name, '') AS ผู้ขาย, po.state AS สถานะ,
                   po.invoice_status AS สถานะใบแจ้งหนี้,
                   po.amount_total AS มูลค่ารวม
            FROM purchase_order po
            LEFT JOIN res_partner rp ON rp.id = po.partner_id
            ORDER BY po.date_order NULLS LAST, po.id
        """,
    },
    {
        "id": "po_open",
        "mod": "purchase",
        "kind": "คงค้าง",
        "name": "PO คงค้าง (ค้างรับ / ค้างตั้งหนี้)",
        "related": "จัดซื้อ, คลัง, เจ้าหนี้",
        "model": "purchase.order (คงค้าง)",
        "key": "name",
        "sheet": "ยืนยัน-POคงค้าง",
        "count": """
            SELECT COUNT(*) FROM (
                SELECT po.id
                FROM purchase_order po
                LEFT JOIN purchase_order_line pol ON pol.order_id = po.id
                WHERE po.state IN ('purchase', 'done')
                GROUP BY po.id, po.invoice_status
                HAVING SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                THEN pol.product_qty - COALESCE(pol.qty_received, 0) ELSE 0 END) > 0
                    OR SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                THEN COALESCE(pol.qty_received, 0) - COALESCE(pol.qty_invoiced, 0) ELSE 0 END) > 0
                    OR po.invoice_status IS DISTINCT FROM 'invoiced'
            ) t
        """,
        "list": """
            SELECT po.name AS เลขที่PO, po.date_order AS วันที่สั่ง,
                   COALESCE(rp.name, '') AS ผู้ขาย, po.state AS สถานะ,
                   po.invoice_status AS สถานะใบแจ้งหนี้,
                   po.amount_total AS มูลค่ารวม,
                   COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                     THEN pol.product_qty ELSE 0 END), 0) AS จำนวนสั่ง,
                   COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                     THEN COALESCE(pol.qty_received, 0) ELSE 0 END), 0) AS จำนวนรับแล้ว,
                   COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                     THEN COALESCE(pol.qty_invoiced, 0) ELSE 0 END), 0) AS จำนวนตั้งหนี้แล้ว,
                   COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                     THEN pol.product_qty - COALESCE(pol.qty_received, 0) ELSE 0 END), 0) AS ค้างรับ,
                   COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                     THEN COALESCE(pol.qty_received, 0) - COALESCE(pol.qty_invoiced, 0) ELSE 0 END), 0) AS ค้างตั้งหนี้
            FROM purchase_order po
            LEFT JOIN res_partner rp ON rp.id = po.partner_id
            LEFT JOIN purchase_order_line pol ON pol.order_id = po.id
            WHERE po.state IN ('purchase', 'done')
            GROUP BY po.id, po.name, po.date_order, rp.name, po.state, po.invoice_status, po.amount_total
            HAVING SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                            THEN pol.product_qty - COALESCE(pol.qty_received, 0) ELSE 0 END) > 0
                OR SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                            THEN COALESCE(pol.qty_received, 0) - COALESCE(pol.qty_invoiced, 0) ELSE 0 END) > 0
                OR po.invoice_status IS DISTINCT FROM 'invoiced'
            ORDER BY po.date_order NULLS LAST, po.name
        """,
        "footer_note": "PO คงค้าง = สถานะ purchase/done ที่ยังรับของไม่ครบ หรือยังตั้งหนี้ไม่ครบ หรือ invoice_status ไม่ใช่ invoiced",
    },
    # ---- คลัง ----
    {
        "id": "uomcat",
        "mod": "stock",
        "kind": "Master",
        "name": "หมวดหน่วยนับ",
        "related": "คลัง, สินค้า",
        "model": "uom.category",
        "key": "name",
        "sheet": "ยืนยัน-หมวดหน่วยนับ",
        "count": "SELECT COUNT(*) FROM uom_category",
        "list": f"""
            SELECT id AS รหัสภายใน, {jname_sql()} AS ชื่อ
            FROM uom_category ORDER BY id
        """,
    },
    {
        "id": "uom",
        "mod": "stock",
        "kind": "Master",
        "name": "หน่วยนับ",
        "related": "คลัง, สินค้า, จัดซื้อ",
        "model": "uom.uom",
        "key": "name + หมวด",
        "sheet": "ยืนยัน-หน่วยนับ",
        "count": "SELECT COUNT(*) FROM uom_uom",
        "list": f"""
            SELECT {jname_sql('u.name')} AS หน่วยนับ, {jname_sql('c.name')} AS หมวด,
                   CASE WHEN u.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM uom_uom u
            LEFT JOIN uom_category c ON c.id = u.category_id
            ORDER BY c.id, u.id
        """,
    },
    {
        "id": "pcateg",
        "mod": "stock",
        "kind": "Master",
        "name": "หมวดหมู่สินค้า",
        "related": "คลัง, สินค้า",
        "model": "product.category",
        "key": "complete_name",
        "sheet": "ยืนยัน-หมวดสินค้า",
        "count": "SELECT COUNT(*) FROM product_category",
        "list": """
            SELECT id AS รหัสภายใน, complete_name AS ชื่อเต็ม, name AS ชื่อ
            FROM product_category ORDER BY complete_name
        """,
    },
    {
        "id": "product",
        "mod": "stock",
        "kind": "Master",
        "name": "สินค้า (Product Master)",
        "related": "คลัง, จัดซื้อ, งบประมาณ",
        "model": "product.template",
        "key": "default_code",
        "sheet": "ยืนยัน-สินค้า",
        "count": "SELECT COUNT(*) FROM product_template",
        "list": f"""
            SELECT t.default_code AS รหัสสินค้า,
                   {jname_sql('t.name')} AS ชื่อ,
                   pc.complete_name AS หมวดหมู่,
                   {jname_sql('u.name')} AS หน่วยนับ,
                   COALESCE(cost.ต้นทุน, 0) AS ต้นทุน,
                   t.type AS ประเภท,
                   CASE WHEN t.is_storable THEN 'คุมสต็อก' ELSE 'ไม่คุมสต็อก' END AS คุมสต็อก,
                   CASE t.tracking
                     WHEN 'lot' THEN 'เช็ค lot'
                     WHEN 'serial' THEN 'เช็ค serial'
                     ELSE 'ไม่เช็ค'
                   END AS ติดตาม,
                   CASE WHEN COALESCE(t.use_expiration_date, FALSE) THEN 'เช็ค'
                        ELSE 'ไม่เช็ค' END AS หมดอายุ,
                   CASE WHEN COALESCE(t.use_expiration_date, FALSE)
                        THEN t.expiration_time END AS "อายุ (วัน)",
                   CASE WHEN t.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM product_template t
            LEFT JOIN product_category pc ON pc.id = t.categ_id
            LEFT JOIN uom_uom u ON u.id = t.uom_id
            LEFT JOIN (
                SELECT DISTINCT ON (pp.product_tmpl_id)
                       pp.product_tmpl_id,
                       COALESCE(NULLIF(pp.standard_price ->> rc.id::text, '')::numeric, 0) AS ต้นทุน
                FROM product_product pp
                CROSS JOIN res_company rc
                WHERE rc.name = 'โรงพยาบาลวชิระภูเก็ต'
                ORDER BY pp.product_tmpl_id, pp.id
            ) cost ON cost.product_tmpl_id = t.id
            ORDER BY t.default_code NULLS LAST, t.id
        """,
    },
    {
        "id": "wh",
        "mod": "stock",
        "kind": "Master",
        "name": "คลังสินค้า (เปิดใช้ · รับเข้า 1 ขั้น)",
        "related": "คลัง",
        "model": "stock.warehouse",
        "key": "code",
        "sheet": "ยืนยัน-คลัง",
        "count": """
            SELECT COUNT(*) FROM stock_warehouse
            WHERE active AND reception_steps = 'one_step'
        """,
        "list": """
            SELECT code AS รหัส, name AS ชื่อ,
                   CASE reception_steps
                     WHEN 'one_step' THEN 'รับเข้า 1 ขั้น'
                     WHEN 'two_steps' THEN 'รับเข้า 2 ขั้น'
                     WHEN 'three_steps' THEN 'รับเข้า 3 ขั้น'
                     ELSE reception_steps
                   END AS รับเข้า,
                   CASE delivery_steps
                     WHEN 'ship_only' THEN 'จ่ายออก 1 ขั้น'
                     WHEN 'pick_ship' THEN 'จ่ายออก 2 ขั้น'
                     WHEN 'pick_pack_ship' THEN 'จ่ายออก 3 ขั้น'
                     ELSE delivery_steps
                   END AS จ่ายออก,
                   CASE WHEN active THEN 'เปิดใช้' ELSE 'ปิด' END AS สถานะ
            FROM stock_warehouse
            WHERE active AND reception_steps = 'one_step'
            ORDER BY id
        """,
    },
    {
        "id": "loc",
        "mod": "stock",
        "kind": "Master",
        "name": "ตำแหน่งจัดเก็บ (internal · เปิดใช้)",
        "related": "คลัง",
        "model": "stock.location",
        "key": "complete_name",
        "sheet": "ยืนยัน-ตำแหน่ง",
        "count": """
            SELECT COUNT(*) FROM stock_location loc
            LEFT JOIN stock_warehouse wh ON wh.id = loc.warehouse_id
            WHERE loc.usage = 'internal'
              AND loc.active
              AND (wh.id IS NULL OR (wh.active AND wh.reception_steps = 'one_step'))
        """,
        "list": """
            SELECT loc.complete_name AS ตำแหน่ง,
                   COALESCE(wh.code, '') AS คลัง,
                   loc.usage AS ประเภท,
                   CASE WHEN loc.active THEN 'เปิดใช้' ELSE 'ปิด' END AS สถานะ
            FROM stock_location loc
            LEFT JOIN stock_warehouse wh ON wh.id = loc.warehouse_id
            WHERE loc.usage = 'internal'
              AND loc.active
              AND (wh.id IS NULL OR (wh.active AND wh.reception_steps = 'one_step'))
            ORDER BY loc.complete_name
        """,
    },
    {
        "id": "quant",
        "mod": "stock",
        "kind": "ธุรกรรม",
        "name": "ยกยอดสต็อก (On Hand ≠ 0)",
        "related": "คลัง",
        "model": "stock.quant",
        "key": "สินค้า + ตำแหน่ง + lot",
        "sheet": None,
        "count": "SELECT COUNT(*) FROM stock_quant WHERE quantity <> 0",
        "list": None,
    },
    {
        "id": "lot",
        "mod": "stock",
        "kind": "ธุรกรรม",
        "name": "Lot / Serial",
        "related": "คลัง",
        "model": "stock.lot",
        "key": "name + สินค้า",
        "sheet": None,
        "count": "SELECT COUNT(*) FROM stock_lot",
        "list": None,
    },
    # ---- การเงิน / บัญชี / เจ้าหนี้ ----
    {
        "id": "coa",
        "mod": "finance",
        "kind": "Master",
        "name": "ผังบัญชี",
        "related": "บัญชี",
        "model": "account.account",
        "key": "code",
        "sheet": "ยืนยัน-ผังบัญชี",
        "count": "SELECT COUNT(*) FROM account_account",
        "list": f"""
            SELECT code_store->>'1' AS รหัสบัญชี,
                   {jname_sql()} AS ชื่อบัญชี,
                   account_type AS ประเภทบัญชี,
                   CASE WHEN COALESCE(deprecated, FALSE) THEN 'เลิกใช้' ELSE 'ใช้งาน' END AS สถานะ
            FROM account_account
            ORDER BY code_store->>'1', id
        """,
    },
    {
        "id": "journal",
        "mod": "finance",
        "kind": "Master",
        "name": "สมุดรายวัน",
        "related": "บัญชี, การเงิน",
        "model": "account.journal",
        "key": "code",
        "sheet": "ยืนยัน-สมุดรายวัน",
        "count": "SELECT COUNT(*) FROM account_journal",
        "list": f"""
            SELECT code AS รหัส, {jname_sql()} AS ชื่อ, type AS ประเภท,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM account_journal ORDER BY id
        """,
    },
    {
        "id": "company_bank",
        "mod": "finance",
        "kind": "Master",
        "name": "บัญชีธนาคารของโรงพยาบาล",
        "related": "การเงิน, บัญชี",
        "model": "res.partner.bank (บริษัท) + account.journal",
        "key": "เลขบัญชี / รหัสสมุดรายวัน",
        "sheet": "ยืนยัน-บัญชีธนาคารรพ",
        "count": """
            SELECT COUNT(*) FROM res_partner_bank pb
            JOIN res_company c ON c.partner_id = pb.partner_id
        """,
        "list": """
            SELECT COALESCE(j.code, '') AS รหัสสมุดรายวัน,
                   pb.acc_number AS เลขบัญชี,
                   COALESCE(pb.acc_holder_name, '') AS ชื่อบัญชี,
                   COALESCE(b.name, '') AS ธนาคาร,
                   COALESCE(b.bic, '') AS SWIFT,
                   CASE WHEN pb.allow_out_payment THEN 'จ่ายออกได้' ELSE 'ยังไม่เปิดจ่าย' END AS จ่ายออก,
                   CASE WHEN pb.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM res_partner_bank pb
            JOIN res_company c ON c.partner_id = pb.partner_id
            LEFT JOIN res_bank b ON b.id = pb.bank_id
            LEFT JOIN account_journal j ON j.bank_account_id = pb.id
            ORDER BY j.code NULLS LAST, pb.id
        """,
    },
    {
        "id": "tax",
        "mod": "finance",
        "kind": "Master",
        "name": "ภาษี (VAT)",
        "related": "บัญชี, เจ้าหนี้",
        "model": "account.tax",
        "key": "name",
        "sheet": "ยืนยัน-ภาษี",
        "count": "SELECT COUNT(*) FROM account_tax",
        "list": f"""
            SELECT {jname_sql()} AS ชื่อ, type_tax_use AS ใช้กับ, amount AS อัตรา,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM account_tax ORDER BY id
        """,
    },
    {
        "id": "wht",
        "mod": "finance",
        "kind": "Master",
        "name": "ภาษีหัก ณ ที่จ่าย (WHT)",
        "related": "บัญชี, การเงิน",
        "model": "account.withholding.tax",
        "key": "name",
        "sheet": None,
        "count": "SELECT COUNT(*) FROM account_withholding_tax",
        "list": None,
    },
    {
        "id": "pterm",
        "mod": "finance",
        "kind": "Master",
        "name": "เงื่อนไขการชำระ",
        "related": "เจ้าหนี้, การเงิน",
        "model": "account.payment.term",
        "key": "name",
        "sheet": "ยืนยัน-เงื่อนไขชำระ",
        "count": "SELECT COUNT(*) FROM account_payment_term",
        "list": f"""
            SELECT {jname_sql()} AS ชื่อ,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM account_payment_term ORDER BY id
        """,
    },
    {
        "id": "aprofile",
        "mod": "finance",
        "kind": "Master",
        "name": "โปรไฟล์สินทรัพย์ถาวร",
        "related": "บัญชี, ครุภัณฑ์",
        "model": "account.asset.profile",
        "key": "name",
        "sheet": "ยืนยัน-โปรไฟล์สินทรัพย์",
        "count": "SELECT COUNT(*) FROM account_asset_profile",
        "list": """
            SELECT name AS ชื่อ, method_number AS จำนวนงวด,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM account_asset_profile ORDER BY id
        """,
    },
    {
        "id": "asset",
        "mod": "finance",
        "kind": "Master",
        "name": "บัตรสินทรัพย์ / ครุภัณฑ์",
        "related": "บัญชี, งบประมาณ",
        "model": "account.asset",
        "key": "code",
        "sheet": "ยืนยัน-บัตรสินทรัพย์",
        "count": "SELECT COUNT(*) FROM account_asset",
        "list": """
            SELECT a.code AS รหัส, a.name AS ชื่อ, p.name AS โปรไฟล์,
                   a.state AS สถานะเอกสาร, a.date_start AS วันเริ่ม,
                   a.purchase_value AS มูลค่าซื้อ,
                   CASE WHEN a.active THEN 'ใช้งาน' ELSE 'ปิด' END AS ใช้งาน
            FROM account_asset a
            LEFT JOIN account_asset_profile p ON p.id = a.profile_id
            ORDER BY a.code NULLS LAST, a.id
        """,
    },
    {
        "id": "equip",
        "mod": "finance",
        "kind": "Master",
        "name": "ครุภัณฑ์ซ่อมบำรุง (Equipment)",
        "related": "ซ่อมบำรุง, สินทรัพย์",
        "model": "maintenance.equipment",
        "key": "name",
        "sheet": "ยืนยัน-ครุภัณฑ์ซ่อม",
        "count": "SELECT COUNT(*) FROM maintenance_equipment",
        "list": f"""
            SELECT {jname_sql()} AS ชื่อ,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM maintenance_equipment ORDER BY id
        """,
    },
    {
        "id": "wa",
        "mod": "finance",
        "kind": "ธุรกรรม",
        "name": "ใบตรวจรับงาน (WA)",
        "related": "จัดซื้อ, เจ้าหนี้",
        "model": "work.acceptance",
        "key": "name",
        "sheet": None,
        "count": "SELECT COUNT(*) FROM work_acceptance",
        "list": None,
    },
    {
        "id": "bill",
        "mod": "finance",
        "kind": "ธุรกรรม",
        "name": "ใบแจ้งหนี้ผู้ขาย / ใบลดหนี้ ทั้งหมด",
        "related": "เจ้าหนี้",
        "model": "account.move (in_invoice / in_refund)",
        "key": "name / ref",
        "sheet": "ยืนยัน-บิลผู้ขาย",
        "count": "SELECT COUNT(*) FROM account_move WHERE move_type IN ('in_invoice','in_refund')",
        "list": """
            SELECT am.name AS เลขที่เอกสาร, am.ref AS อ้างอิง,
                   am.invoice_date AS วันที่บิล, am.invoice_date_due AS วันครบกำหนด,
                   COALESCE(rp.name, '') AS เจ้าหนี้, am.state AS สถานะ,
                   am.payment_state AS สถานะชำระ,
                   am.amount_total AS มูลค่ารวม, am.amount_residual AS ยอดคงค้าง
            FROM account_move am
            LEFT JOIN res_partner rp ON rp.id = am.partner_id
            WHERE am.move_type IN ('in_invoice', 'in_refund')
            ORDER BY am.invoice_date NULLS LAST, am.id
        """,
    },
    {
        "id": "bill_open",
        "mod": "finance",
        "kind": "คงค้าง",
        "name": "เจ้าหนี้คงค้าง (บิลที่ยังไม่จ่ายครบ)",
        "related": "เจ้าหนี้, การเงิน",
        "model": "account.move (posted, amount_residual ≠ 0)",
        "key": "name / เจ้าหนี้",
        "sheet": "ยืนยัน-เจ้าหนี้คงค้าง",
        "count": """
            SELECT COUNT(*) FROM account_move
            WHERE move_type IN ('in_invoice', 'in_refund')
              AND state = 'posted'
              AND COALESCE(amount_residual, 0) <> 0
        """,
        "list": """
            SELECT am.name AS เลขที่เอกสาร, am.ref AS อ้างอิง,
                   am.invoice_date AS วันที่บิล, am.invoice_date_due AS วันครบกำหนด,
                   COALESCE(rp.name, '') AS เจ้าหนี้, am.state AS สถานะ,
                   am.payment_state AS สถานะชำระ,
                   am.amount_total AS มูลค่ารวม, am.amount_residual AS ยอดคงค้าง
            FROM account_move am
            LEFT JOIN res_partner rp ON rp.id = am.partner_id
            WHERE am.move_type IN ('in_invoice', 'in_refund')
              AND am.state = 'posted'
              AND COALESCE(am.amount_residual, 0) <> 0
            ORDER BY am.invoice_date NULLS LAST, am.name
        """,
        "footer_note": "เจ้าหนี้คงค้าง = ใบแจ้งหนี้/ใบลดหนี้ผู้ขายที่ผ่านรายการแล้ว และมียอดค้างชำระ (amount_residual ≠ 0)",
    },
    {
        "id": "pay",
        "mod": "finance",
        "kind": "ธุรกรรม",
        "name": "การจ่ายชำระ",
        "related": "การเงิน",
        "model": "account.payment",
        "key": "name",
        "sheet": None,
        "count": "SELECT COUNT(*) FROM account_payment",
        "list": None,
    },
    {
        "id": "stmt",
        "mod": "finance",
        "kind": "ธุรกรรม",
        "name": "Bank Statement",
        "related": "การเงิน",
        "model": "account.bank.statement",
        "key": "name",
        "sheet": None,
        "count": "SELECT COUNT(*) FROM account_bank_statement",
        "list": None,
    },
    {
        "id": "je",
        "mod": "finance",
        "kind": "ธุรกรรม",
        "name": "รายการบัญชีทั่วไป / ยกยอด GL",
        "related": "บัญชี",
        "model": "account.move (entry)",
        "key": "name + วันที่",
        "sheet": "ยืนยัน-รายการบัญชีทั่วไป",
        "count": "SELECT COUNT(*) FROM account_move WHERE move_type = 'entry'",
        "list": """
            SELECT am.name AS เลขที่, am.date AS วันที่, am.ref AS อ้างอิง,
                   am.state AS สถานะ, am.amount_total AS มูลค่า
            FROM account_move am
            WHERE am.move_type = 'entry'
            ORDER BY am.date NULLS LAST, am.id
        """,
    },
    {
        "id": "tb2569",
        "mod": "finance",
        "kind": "คงค้าง",
        "name": "งบทดลอง (TB) ปีงบประมาณ 2569",
        "related": "บัญชี",
        "model": "account.move.line (posted) 1 ต.ค. 2568 – 30 ก.ย. 2569",
        "key": "รหัสบัญชี",
        "sheet": "ยืนยัน-TB2569",
        "count": """
            SELECT COUNT(*) FROM (
                SELECT aml.account_id
                FROM account_move_line aml
                JOIN account_move am ON am.id = aml.move_id
                WHERE am.state = 'posted'
                  AND aml.account_id IS NOT NULL
                  AND COALESCE(aml.display_type, '') NOT IN ('line_section', 'line_note', 'cogs')
                  AND am.date BETWEEN DATE '2025-10-01' AND DATE '2026-09-30'
                GROUP BY aml.account_id
                HAVING SUM(aml.debit) <> 0 OR SUM(aml.credit) <> 0
            ) t
        """,
        "list": f"""
            SELECT COALESCE(aa.code_store->>'1', {account_code_sql('aa')}) AS รหัสบัญชี,
                   {jname_sql('aa.name')} AS ชื่อบัญชี,
                   aa.account_type AS ประเภทบัญชี,
                   SUM(aml.debit) AS เดบิต,
                   SUM(aml.credit) AS เครดิต,
                   SUM(aml.debit - aml.credit) AS ยอดคงเหลือเดบิตสุทธิ
            FROM account_move_line aml
            JOIN account_move am ON am.id = aml.move_id
            JOIN account_account aa ON aa.id = aml.account_id
            WHERE am.state = 'posted'
              AND aml.account_id IS NOT NULL
              AND COALESCE(aml.display_type, '') NOT IN ('line_section', 'line_note', 'cogs')
              AND am.date BETWEEN DATE '2025-10-01' AND DATE '2026-09-30'
            GROUP BY aa.id, aa.code_store, aa.name, aa.account_type
            HAVING SUM(aml.debit) <> 0 OR SUM(aml.credit) <> 0
            ORDER BY 1
        """,
        "footer_note": (
            "งบทดลองปีงบประมาณ 2569 = รวมยอดบัญชีที่ผ่านรายการ ระหว่าง 1 ต.ค. 2568 ถึง 30 ก.ย. 2569 "
            "(2025-10-01 ถึง 2026-09-30)  ยังไม่มียอดยกเข้าจากระบบเดิม — รายการที่เห็นตอนนี้มาจากเอกสารที่บันทึกใน Odoo "
            "ไม่ใช่ TB ยกยอดที่นำเข้า  ให้ผู้ใช้ยืนยันว่าจะนำเข้า TB ภายหลัง"
        ),
    },
    {
        "id": "whtc",
        "mod": "finance",
        "kind": "ธุรกรรม",
        "name": "หนังสือรับรองหัก ณ ที่จ่าย",
        "related": "การเงิน, บัญชี",
        "model": "withholding.tax.cert",
        "key": "number",
        "sheet": None,
        "count": "SELECT COUNT(*) FROM withholding_tax_cert",
        "list": None,
    },
    # ---- สินทรัพย์ถาวร ----
    {
        "id": "fa_group",
        "mod": "asset",
        "kind": "Master",
        "name": "กลุ่มสินทรัพย์",
        "related": "บัญชี, ครุภัณฑ์",
        "model": "account.asset.group",
        "key": "code",
        "sheet": "ยืนยัน-กลุ่มสินทรัพย์",
        "count": "SELECT COUNT(*) FROM account_asset_group",
        "list": """
            SELECT g.code AS รหัส, g.name AS ชื่อ,
                   p.code AS รหัสกลุ่มแม่, p.name AS กลุ่มแม่,
                   CASE WHEN g.parent_id IS NULL THEN 'กลุ่มหลัก' ELSE 'กลุ่มย่อย' END AS ระดับ
            FROM account_asset_group g
            LEFT JOIN account_asset_group p ON p.id = g.parent_id
            ORDER BY COALESCE(p.code, g.code), g.code, g.id
        """,
    },
    {
        "id": "fa_profile",
        "mod": "asset",
        "kind": "Master",
        "name": "โปรไฟล์สินทรัพย์ถาวร",
        "related": "บัญชี, ครุภัณฑ์",
        "model": "account.asset.profile",
        "key": "name",
        "sheet": "ยืนยัน-โปรไฟล์สินทรัพย์",
        "count": "SELECT COUNT(*) FROM account_asset_profile",
        "list": f"""
            SELECT p.name AS ชื่อ, p.method AS วิธีคิดค่าเสื่อม,
                   p.method_time AS วิธีนับเวลา, p.method_number AS จำนวนปี,
                   p.method_period AS งวด, p.salvage_value AS มูลค่าซาก,
                   {account_code_sql('aa')} AS บัญชีสินทรัพย์,
                   {account_code_sql('ad')} AS บัญชีค่าเสื่อมสะสม,
                   {account_code_sql('ae')} AS บัญชีค่าใช้จ่ายค่าเสื่อม, j.code AS สมุดรายวัน,
                   CASE WHEN p.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM account_asset_profile p
            LEFT JOIN account_account aa ON aa.id = p.account_asset_id
            LEFT JOIN account_account ad ON ad.id = p.account_depreciation_id
            LEFT JOIN account_account ae ON ae.id = p.account_expense_depreciation_id
            LEFT JOIN account_journal j ON j.id = p.journal_id
            ORDER BY p.name, p.id
        """,
        "list_fallback": """
            SELECT name AS ชื่อ, method AS วิธีคิดค่าเสื่อม,
                   method_time AS วิธีนับเวลา, method_number AS จำนวนปี,
                   method_period AS งวด, salvage_value AS มูลค่าซาก,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM account_asset_profile ORDER BY name, id
        """,
    },
    {
        "id": "fa_substate",
        "mod": "asset",
        "kind": "Master",
        "name": "สถานะย่อยสินทรัพย์",
        "related": "บัญชี, ครุภัณฑ์",
        "model": "account.asset.sub.state",
        "key": "name",
        "sheet": "ยืนยัน-สถานะย่อย",
        "count": "SELECT COUNT(*) FROM account_asset_sub_state",
        "list": """
            SELECT sequence AS ลำดับ, name AS ชื่อ,
                   CASE WHEN draft THEN 'ใช่' ELSE '' END AS ร่าง,
                   CASE WHEN open THEN 'ใช่' ELSE '' END AS ใช้งาน,
                   CASE WHEN close THEN 'ใช่' ELSE '' END AS ปิด,
                   CASE WHEN removed THEN 'ใช่' ELSE '' END AS จำหน่าย
            FROM account_asset_sub_state ORDER BY sequence, id
        """,
    },
    {
        "id": "fa_parent",
        "mod": "asset",
        "kind": "Master",
        "name": "ชุดสินทรัพย์ (Asset Set)",
        "related": "บัญชี, ครุภัณฑ์",
        "model": "account.asset.parent",
        "key": "code",
        "sheet": "ยืนยัน-ชุดสินทรัพย์",
        "count": "SELECT COUNT(*) FROM account_asset_parent",
        "list": """
            SELECT code AS รหัส, name AS ชื่อ
            FROM account_asset_parent ORDER BY code NULLS LAST, id
        """,
    },
    {
        "id": "fa_card",
        "mod": "asset",
        "kind": "Master",
        "name": "บัตรสินทรัพย์ / ครุภัณฑ์",
        "related": "บัญชี, งบประมาณ, ซ่อมบำรุง",
        "model": "account.asset",
        "key": "number / code",
        "sheet": "ยืนยัน-บัตรสินทรัพย์",
        "count": "SELECT COUNT(*) FROM account_asset",
        "list": f"""
            SELECT a.number AS หมายเลขทรัพย์สิน, a.code AS รหัสอ้างอิง, a.name AS ชื่อ,
                   p.name AS โปรไฟล์,
                   CASE a.state
                     WHEN 'draft' THEN 'ร่าง'
                     WHEN 'open' THEN 'ใช้งาน'
                     WHEN 'close' THEN 'ปิด'
                     WHEN 'removed' THEN 'จำหน่าย'
                     ELSE a.state
                   END AS สถานะเอกสาร,
                   ss.name AS สถานะย่อย,
                   a.date_start AS วันเริ่มคิดค่าเสื่อม, a.acquisition_date AS วันที่ได้มา,
                   a.fiscal_year AS ปีงบประมาณ, a.useful_life_years AS อายุใช้งานปี,
                   a.purchase_value AS มูลค่าซื้อ, a.salvage_value AS มูลค่าซาก,
                   a.depreciation_base AS ฐานค่าเสื่อม, a.value_depreciated AS ค่าเสื่อมสะสม,
                   a.value_residual AS มูลค่าคงเหลือ,
                   rp.ref AS รหัสผู้ขาย, rp.name AS ผู้ขาย,
                   aa.code AS รหัสหน่วยงาน, {jname_sql('aa.name')} AS หน่วยงานเจ้าของ,
                   a.project_code AS รหัสโครงการ, a.project_name AS ชื่อโครงการ,
                   fs.name AS แหล่งเงิน, a.vcr_number AS รหัสVCR,
                   pt.default_code AS รหัสสินค้า, {jname_sql('pt.name')} AS สินค้าต้นทาง,
                   po.name AS ใบสั่งซื้อ, am.name AS ใบแจ้งหนี้,
                   a.equipment_model AS โมเดล, a.equipment_serial_no AS ซีเรียล,
                   a.equipment_location AS พิกัด,
                   e.name AS ครุภัณฑ์ซ่อมบำรุง,
                   CASE WHEN a.active THEN 'ใช้งาน' ELSE 'ปิด' END AS เปิดใช้
            FROM account_asset a
            LEFT JOIN account_asset_profile p ON p.id = a.profile_id
            LEFT JOIN account_asset_sub_state ss ON ss.id = a.asset_sub_state_id
            LEFT JOIN res_partner rp ON rp.id = a.partner_id
            LEFT JOIN account_analytic_account aa ON aa.id = a.owning_analytic_account_id
            LEFT JOIN vpk_budget_fund_source fs ON fs.id = a.fund_source_id
            LEFT JOIN product_product pp ON pp.id = a.product_id
            LEFT JOIN product_template pt ON pt.id = pp.product_tmpl_id
            LEFT JOIN purchase_order po ON po.id = a.purchase_order_id
            LEFT JOIN account_move am ON am.id = a.vendor_bill_id
            LEFT JOIN maintenance_equipment e ON e.id = a.equipment_id
            ORDER BY a.number NULLS LAST, a.code NULLS LAST, a.id
        """,
        "list_fallback": """
            SELECT a.code AS รหัส, a.name AS ชื่อ, p.name AS โปรไฟล์,
                   a.state AS สถานะเอกสาร, a.date_start AS วันเริ่ม,
                   a.purchase_value AS มูลค่าซื้อ,
                   CASE WHEN a.active THEN 'ใช้งาน' ELSE 'ปิด' END AS เปิดใช้
            FROM account_asset a
            LEFT JOIN account_asset_profile p ON p.id = a.profile_id
            ORDER BY a.code NULLS LAST, a.id
        """,
    },
    {
        "id": "fa_line",
        "mod": "asset",
        "kind": "ตั้งค่า",
        "name": "ตารางค่าเสื่อม",
        "related": "บัญชี",
        "model": "account.asset.line",
        "key": "สินทรัพย์ + วันที่ + ประเภท",
        "sheet": "ยืนยัน-ตารางค่าเสื่อม",
        "count": "SELECT COUNT(*) FROM account_asset_line",
        "list": """
            SELECT COALESCE(a.number, a.code) AS รหัสสินทรัพย์, a.name AS ชื่อสินทรัพย์,
                   CASE l.type
                     WHEN 'create' THEN 'ฐานค่าเสื่อม'
                     WHEN 'depreciate' THEN 'ค่าเสื่อม'
                     WHEN 'remove' THEN 'จำหน่าย'
                     ELSE l.type
                   END AS ประเภท,
                   l.line_date AS วันที่, l.amount AS จำนวนเงิน,
                   l.depreciated_value AS ค่าเสื่อมสะสม, l.remaining_value AS คงเหลืองวดถัดไป,
                   CASE WHEN l.init_entry THEN 'ยกยอด' ELSE '' END AS รายการยกยอด,
                   CASE WHEN l.move_check THEN 'ลงบัญชีแล้ว' ELSE '' END AS ลงบัญชี
            FROM account_asset_line l
            JOIN account_asset a ON a.id = l.asset_id
            ORDER BY COALESCE(a.number, a.code), l.line_date, l.id
        """,
    },
    {
        "id": "fa_eqcat",
        "mod": "asset",
        "kind": "Master",
        "name": "หมวดครุภัณฑ์ซ่อมบำรุง",
        "related": "ซ่อมบำรุง, สินทรัพย์",
        "model": "maintenance.equipment.category",
        "key": "name",
        "sheet": "ยืนยัน-หมวดครุภัณฑ์ซ่อม",
        "count": "SELECT COUNT(*) FROM maintenance_equipment_category",
        "list": f"""
            SELECT {jname_sql()} AS ชื่อ
            FROM maintenance_equipment_category ORDER BY id
        """,
    },
    {
        "id": "fa_equip",
        "mod": "asset",
        "kind": "Master",
        "name": "ครุภัณฑ์ซ่อมบำรุง (Equipment)",
        "related": "ซ่อมบำรุง, สินทรัพย์",
        "model": "maintenance.equipment",
        "key": "asset_number / serial_no / ชื่อ",
        "sheet": "ยืนยัน-ครุภัณฑ์ซ่อม",
        "count": "SELECT COUNT(*) FROM maintenance_equipment",
        "list": f"""
            SELECT e.asset_number AS หมายเลขทรัพย์สิน, a.code AS รหัสสินทรัพย์,
                   {jname_sql('e.name')} AS ชื่อ, e.serial_no AS ซีเรียล,
                   {jname_sql('c.name')} AS หมวด,
                   CASE WHEN e.asset_id IS NOT NULL THEN 'ลิงก์แล้ว' ELSE 'ยังไม่ลิงก์' END AS ลิงก์สินทรัพย์,
                   CASE WHEN e.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM maintenance_equipment e
            LEFT JOIN account_asset a ON a.id = e.asset_id
            LEFT JOIN maintenance_equipment_category c ON c.id = e.category_id
            ORDER BY e.asset_number NULLS LAST, e.id
        """,
        "list_fallback": f"""
            SELECT {jname_sql()} AS ชื่อ,
                   CASE WHEN active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
            FROM maintenance_equipment ORDER BY id
        """,
    },
]

MODULES = {
    "budget": {
        "file": "UAT-นำเข้าข้อมูล-งบประมาณ.xlsx",
        "doc": "DM-SIGNOFF-BUD-001",
        "title": "ยืนยันข้อมูลนำเข้า ระบบงบประมาณ",
        "system": "Odoo 18 · vpk_budget และบัญชีวิเคราะห์",
        "owner": "งานงบประมาณ",
        "roles": ["งานงบประมาณ (เจ้าของข้อมูล)", "หน่วยงานผู้ใช้", "IT / ผู้ดูแลระบบ", "ผู้บริหารโครงการ"],
    },
    "purchase": {
        "file": "UAT-นำเข้าข้อมูล-จัดซื้อ.xlsx",
        "doc": "DM-SIGNOFF-PUR-001",
        "title": "ยืนยันข้อมูลนำเข้า ระบบจัดซื้อ",
        "system": "Odoo 18 · Purchase Request / Purchase Order / Vendor",
        "owner": "พัสดุ / จัดซื้อ",
        "roles": ["พัสดุ / จัดซื้อ (เจ้าของข้อมูล)", "งานงบประมาณ", "IT / ผู้ดูแลระบบ", "ผู้บริหารโครงการ"],
    },
    "stock": {
        "file": "UAT-นำเข้าข้อมูล-คลังสินค้า.xlsx",
        "doc": "DM-SIGNOFF-STK-001",
        "title": "ยืนยันข้อมูลนำเข้า ระบบคลังสินค้า",
        "system": "Odoo 18 · Inventory / Product Master",
        "owner": "งานคลัง / พัสดุ",
        "roles": ["งานคลัง (เจ้าของข้อมูล)", "พัสดุ", "IT / ผู้ดูแลระบบ", "ผู้บริหารโครงการ"],
    },
    "finance": {
        "file": "UAT-นำเข้าข้อมูล-เจ้าหนี้บัญชีการเงิน.xlsx",
        "doc": "DM-SIGNOFF-FIN-001",
        "title": "ยืนยันข้อมูลนำเข้า เจ้าหนี้ บัญชี การเงิน",
        "system": "Odoo 18 · Accounting / AP / Asset",
        "owner": "บัญชี / การเงิน",
        "roles": ["บัญชี (เจ้าของผังบัญชี/สินทรัพย์)", "การเงิน (เจ้าของจ่ายชำระ)", "พัสดุ (ตรวจรับ/เจ้าหนี้)", "IT / ผู้ดูแลระบบ", "ผู้บริหารโครงการ"],
    },
    "asset": {
        "file": "UAT-นำเข้าข้อมูล-สินทรัพย์ถาวร.xlsx",
        "doc": "DM-SIGNOFF-AST-001",
        "title": "ยืนยันข้อมูลนำเข้า ระบบสินทรัพย์ถาวร",
        "system": "Odoo 18 · account_asset_management / ทะเบียนครุภัณฑ์",
        "owner": "บัญชีสินทรัพย์ / พัสดุครุภัณฑ์",
        "roles": [
            "บัญชีสินทรัพย์ (เจ้าของบัตรสินทรัพย์)",
            "พัสดุครุภัณฑ์",
            "งานซ่อมบำรุง",
            "IT / ผู้ดูแลระบบ",
            "ผู้บริหารโครงการ",
        ],
    },
}


def load_counts(cur):
    out = {}
    for ds in DATASETS:
        try:
            n = count_sql(cur, ds["count"])
        except Exception as exc:
            print(f"COUNT FAIL {ds['id']}: {exc}")
            cur.connection.rollback()
            n = 0
        out[ds["id"]] = n
        ds["n"] = n
        ds["status"] = status_of(ds["kind"], n)
    return out


def write_cover(wb, spec, snapshot, items):
    ws = wb.active
    ws.title = "ปกและลงนาม"
    ws.sheet_view.showGridLines = False
    widths = [4, 26, 36, 22, 18, 18, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("B2:F2")
    ws["B2"].value, ws["B2"].font = "โรงพยาบาลวชิระภูเก็ต", fnt(14, True, TEAL)
    ws.merge_cells("B3:F3")
    ws["B3"].value, ws["B3"].font = spec["title"], fnt(18, True, NAVY)
    ws.merge_cells("B4:F4")
    ws["B4"].value, ws["B4"].font = (
        "เอกสารสรุปประเภทข้อมูลนำเข้า · จำนวน record ในระบบ · รายการให้ผู้ใช้ยืนยันและลงนาม (Sign-off)"
    ), fnt(12, False, "5B6878")

    masters = [d for d in items if d["kind"] == "Master"]
    imported = [d for d in masters if d["n"] > 0]
    empty = [d for d in masters if d["n"] == 0]
    total_master = sum(d["n"] for d in masters)
    outstanding = [d for d in items if d["kind"] == "คงค้าง"]
    outstanding_n = sum(d["n"] for d in outstanding)

    info = [
        ("ระบบ", spec["system"]),
        ("รหัสเอกสาร", spec["doc"]),
        ("ฐานข้อมูล", f"{DBNAME} · บริษัท โรงพยาบาลวชิระภูเก็ต"),
        ("วัน-เวลาที่ดึงข้อมูล", snapshot),
        ("เจ้าของข้อมูล", spec["owner"]),
        ("ประเภท Master ที่เกี่ยวข้อง", f"{len(masters)} ประเภท"),
        ("Master ที่นำเข้าแล้ว", f"{len(imported)} ประเภท  รวม {total_master:,} รายการ"),
        ("Master ที่ยังไม่มีข้อมูล", f"{len(empty)} ประเภท" if empty else "ไม่มี"),
        (
            "ข้อมูลคงค้าง",
            (
                f"{len(outstanding)} ประเภท · รวม {outstanding_n:,} รายการ"
                if outstanding
                else "ไม่มีชีตคงค้างในโมดูลนี้"
            ),
        ),
    ]
    for i, (k, v) in enumerate(info):
        r = 6 + i
        sc(ws.cell(r, 2, k), 11, True, WHITE, NAVY)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        sc(ws.cell(r, 3, v))

    r = 16
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"].value = (
        "วิธีใช้  1) ดูชีต “สรุปประเภทข้อมูล” ว่าแต่ละประเภทมีกี่ record  "
        "2) เปิดชีต “ยืนยัน-…” ตรวจรหัส/ชื่อที่นำเข้าแล้ว  "
        "3) ตรวจชีตคงค้าง (PO คงค้าง / เจ้าหนี้คงค้าง / TB 2569) แล้วกรอกผลยืนยัน  "
        "เกณฑ์รับรอง: จำนวนและรายการ Master ถูกต้อง · คงค้างที่ยอดเป็น 0 ให้ลงนามว่ายังไม่นำเข้า"
    )
    ws[f"B{r}"].font = fnt(11)
    ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 48

    hdr(ws, 18, ["ลำดับ", "บทบาท", "ชื่อ-สกุล", "ผลยืนยัน", "ลงชื่อ", "วันที่", "หมายเหตุ"])
    dv = DataValidation(type="list", formula1=CONFIRM, allow_blank=True)
    ws.add_data_validation(dv)
    for i, role in enumerate(spec["roles"], 1):
        rr = 18 + i
        sc(ws.cell(rr, 1, i), align="center", bg=SOFT)
        sc(ws.cell(rr, 2, role), bold=True)
        for c in range(3, 8):
            sc(ws.cell(rr, c, ""))
        dv.add(ws.cell(rr, 4))
        ws.row_dimensions[rr].height = 26
    last = 19 + len(spec["roles"])
    ws.merge_cells(start_row=last, start_column=2, end_row=last + 2, end_column=6)
    cell = ws.cell(
        last, 2,
        "คำรับรอง: ข้าพเจ้าได้ตรวจประเภทข้อมูล จำนวน record และรายการที่ export แล้ว "
        "และยืนยันตามผลที่ระบุ  หากเลือก “ต้องแก้ไข/ยังไม่ครบ” ให้ระบุรายการในหมายเหตุ\n"
        "ผลภาพรวมโมดูลนี้:  รับรองให้นำไปใช้ต่อ   /   รับรองแบบมีเงื่อนไข   /   ต้องนำเข้าใหม่",
    )
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font = fnt(12)
    cell.fill = fill(SOFT)
    page(ws, spec["title"])


def write_summary(wb, spec, snapshot, items):
    ws = wb.create_sheet("สรุปประเภทข้อมูล")
    headers = [
        "ลำดับ", "ประเภทข้อมูล", "กลุ่ม", "โมดูลที่เกี่ยวข้อง", "โมเดลในระบบ", "คีย์จับคู่",
        "จำนวนในระบบ", "สถานะนำเข้า", "ชีตรายการยืนยัน",
        "จำนวนต้นทาง (กรอก)", "ผลต่าง", "ผลยืนยัน", "ผู้ยืนยัน", "ลงชื่อ", "วันที่", "หมายเหตุ",
    ]
    hdr(ws, 1, headers, TEAL)
    widths = [8, 32, 12, 28, 36, 22, 14, 22, 24, 16, 12, 16, 16, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    dv = DataValidation(type="list", formula1=CONFIRM, allow_blank=True)
    ws.add_data_validation(dv)
    for i, d in enumerate(items, 1):
        r = i + 1
        sheet_name = d["sheet"] if d["list"] else "—"
        vals = [
            i, d["name"], d["kind"], d["related"], d["model"], d["key"],
            d["n"], d["status"], sheet_name, "", f'=IF(J{r}="","",G{r}-J{r})',
            "", "", "", "", "",
        ]
        bg_status = status_bg(d["status"])
        for c, v in enumerate(vals, 1):
            bg = bg_status if c == 8 else (
                SOFT if d["kind"] == "Master" and c == 3 else (
                    AMBER if d["kind"] == "คงค้าง" and c == 3 else None
                )
            )
            align = "center" if c in (1, 3, 7, 8, 10, 11, 12, 15) else "left"
            sc(ws.cell(r, c, v), bg=bg, align=align)
        ws.cell(r, 7).number_format = "#,##0"
        ws.cell(r, 10).number_format = "#,##0"
        ws.cell(r, 11).number_format = "#,##0"
        dv.add(ws.cell(r, 12))
        ws.row_dimensions[r].height = 32
    last = 1 + len(items)
    tot = last + 1
    sc(ws.cell(tot, 2, "รวมจำนวนทุกประเภทในชีตนี้"), 12, True, WHITE, NAVY)
    sc(ws.cell(tot, 7, f"=SUM(G2:G{last})"), 12, True, WHITE, NAVY, "center")
    ws.cell(tot, 7).number_format = "#,##0"
    for c in range(3, 7):
        sc(ws.cell(tot, c, ""), bg=NAVY)
    for c in range(8, 17):
        sc(ws.cell(tot, c, ""), bg=NAVY)
    note = last + 3
    ws.merge_cells(start_row=note, start_column=1, end_row=note + 2, end_column=16)
    cell = ws.cell(
        note, 1,
        f"ข้อมูล ณ {snapshot} จากฐาน {DBNAME}  คอลัมน์ “จำนวนต้นทาง” ให้เจ้าของข้อมูลกรอกจากไฟล์ที่ส่งมา "
        "ถ้าผลต่าง ≠ 0 ให้เลือกผลยืนยันเป็น ต้องแก้ไข หรือ ยังไม่ครบ  "
        "กลุ่ม Master = ข้อมูลหลักที่ใช้ทั้งระบบ · ตั้งค่า = วงเงิน/ค่าคงที่ · "
        "ธุรกรรม = เอกสารปฏิบัติงาน · คงค้าง = PO ค้างรับ/ค้างตั้งหนี้, เจ้าหนี้ค้างจ่าย, งบทดลอง TB ปี 2569 "
        "(ปีงบประมาณ 1 ต.ค. 2568 – 30 ก.ย. 2569)",
    )
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font = fnt(11)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:P{last}"
    rng = f"L2:L{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ถูกต้อง"'], fill=fill(PASS), font=fnt(11, True, GREEN)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ต้องแก้ไข"'], fill=fill(FAIL), font=fnt(11, True, RED)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ยังไม่ครบ"'], fill=fill(AMBER), font=fnt(11, True, "8A5A00")))
    page(ws, spec["title"], 0)


def write_listing(wb, spec, ds, cols, rows):
    title = ds["sheet"] or ds["name"][:28]
    ws = wb.create_sheet(title[:31])
    add_confirm = ds.get("kind") == "คงค้าง" or len(rows) <= 2000
    headers = ["ลำดับ"] + [c.replace("_", " ") for c in cols]
    if add_confirm:
        headers += ["ผลยืนยัน", "หมายเหตุ"]
    hdr(ws, 1, headers)
    n_col = len(headers)
    for i in range(1, n_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18 if i > 1 else 8
    if add_confirm:
        ws.column_dimensions[get_column_letter(n_col - 1)].width = 16
        ws.column_dimensions[get_column_letter(n_col)].width = 28
    if cols:
        long = {
            "ชื่อ", "ชื่อบัญชี", "ชื่อเต็ม", "ตำแหน่ง", "หมวดหมู่", "สินค้า", "ผู้ขาย",
            "หมวดสินค้า", "ชื่อสินทรัพย์", "สินค้าต้นทาง", "ครุภัณฑ์ซ่อมบำรุง", "หน่วยงานเจ้าของ",
            "เจ้าหนี้", "กลุ่มผู้จำหน่าย", "รายละเอียด",
        }
        short = {
            "หมดอายุ", "อายุ (วัน)", "ติดตาม", "คุมสต็อก", "สถานะ", "ประเภท",
            "ต้นทุน", "รับเข้า", "จ่ายออก", "เดบิต", "เครดิต",
        }
        for i, c in enumerate(cols, 2):
            if c in long or "ชื่อ" in c:
                ws.column_dimensions[get_column_letter(i)].width = 42
            elif c in short:
                ws.column_dimensions[get_column_letter(i)].width = 14
    if not rows:
        msg = "ไม่มีข้อมูลในระบบ ณ วันที่ export — ให้ลงนามในสรุปว่ายังไม่นำเข้า"
        if ds.get("kind") == "คงค้าง":
            msg = "ยังไม่มีรายการคงค้างในระบบ ณ วันที่ export — ให้ผู้ใช้ยืนยันว่าถูกต้อง หรือระบุว่าจะนำเข้าภายหลัง"
        sc(ws.cell(2, 1, msg), bg=AMBER)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(n_col, 4))
        if ds.get("footer_note"):
            ws.cell(4, 1, ds["footer_note"])
            ws.cell(4, 1).font = fnt(10, False, "5B6878")
            ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max(n_col, 4))
        page(ws, spec["title"], 0)
        return
    large = len(rows) > 800
    dv = None
    if add_confirm:
        dv = DataValidation(type="list", formula1=CONFIRM, allow_blank=True)
        ws.add_data_validation(dv)
    for i, row in enumerate(rows, 1):
        r = i + 1
        values = list(row.values()) if isinstance(row, dict) else list(row)
        cell0 = ws.cell(r, 1, i)
        if not large:
            sc(cell0, align="center")
        else:
            cell0.font = fnt(10)
        for c, v in enumerate(values, 2):
            if hasattr(v, "isoformat") and not isinstance(v, datetime):
                v = v.isoformat() if v else ""
            elif isinstance(v, datetime):
                v = v.strftime("%Y-%m-%d %H:%M")
            cell = ws.cell(r, c, v)
            if not large:
                sc(cell)
            else:
                cell.font = fnt(10)
        if add_confirm:
            c_ok = 2 + len(values)
            sc(ws.cell(r, c_ok, ""), align="center")
            sc(ws.cell(r, c_ok + 1, ""))
            if dv is not None:
                dv.add(ws.cell(r, c_ok))
    last = 1 + len(rows)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_col)}{last}"
    ws.cell(last + 2, 1, f"รวม {len(rows):,} รายการ · {ds['name']} · โมเดล {ds['model']}")
    ws.cell(last + 2, 1).font = fnt(11, True, TEAL)
    if ds.get("footer_note"):
        ws.cell(last + 3, 1, ds["footer_note"])
        ws.cell(last + 3, 1).font = fnt(10, False, "5B6878")
    page(ws, spec["title"], 0)


def write_howto(wb, spec):
    ws = wb.create_sheet("วิธีตรวจและเกณฑ์")
    for i, w in enumerate([6, 28, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("B2:C2")
    ws["B2"].value, ws["B2"].font = "วิธีตรวจก่อนลงนาม", fnt(16, True, NAVY)
    hdr(ws, 4, ["ขั้น", "ตรวจอะไร", "ผ่านเมื่อ"], TEAL)
    steps = [
        ("1", "ประเภทข้อมูลครบตามโมดูล", "ในชีตสรุปมีครบ Master ที่เกี่ยวข้อง และระบุโมเดล/คีย์จับคู่"),
        ("2", "จำนวน record", "จำนวนในระบบตรงกับไฟล์ต้นทางที่หน่วยงานส่ง (กรอกคอลัมน์จำนวนต้นทาง แล้วดูผลต่าง = 0)"),
        ("3", "รายการตัวอย่าง", "สุ่มตรวจรหัสและชื่อในชีตยืนยัน ไม่สลับ ไม่ซ้ำ ไม่สะกดผิดจนใช้ไม่ได้"),
        ("4", "ความสัมพันธ์ข้ามโมดูล", "เช่น สินค้ามีหมวดและหน่วยนับ · ผู้ขายถูกใช้ตั้งเจ้าหนี้ได้ · แหล่งเงินถูกใช้ในแผนงบ"),
        ("5", "ข้อมูลคงค้าง", "ตรวจชีต PO คงค้าง, เจ้าหนี้คงค้าง, TB ปีงบ 2569 (1 ต.ค. 2568–30 ก.ย. 2569) ถ้าเป็น 0 ให้ยืนยันว่ายังไม่นำเข้า"),
        ("6", "ลงนาม", "เจ้าของข้อมูล + IT ลงนามในปก และกรอกผลยืนยันรายประเภทในสรุป"),
    ]
    if spec.get("file", "").endswith("สินทรัพย์ถาวร.xlsx"):
        steps[2:2] = [
            (
                "2ก",
                "หมายเลขทรัพย์สินและรหัสอ้างอิง",
                "ชีตยืนยัน-บัตรสินทรัพย์ คีย์หลักคือ number (หมายเลขทรัพย์สิน) ไม่ซ้ำ และ code ใช้เมื่อยังไม่มี number",
            ),
            (
                "2ข",
                "โปรไฟล์และบัญชี",
                "ทุกบัตรมีโปรไฟล์ และโปรไฟล์มีบัญชีสินทรัพย์ / ค่าเสื่อมสะสม / ค่าใช้จ่ายค่าเสื่อม + สมุดรายวันครบ",
            ),
            (
                "2ค",
                "ลิงก์ครุภัณฑ์ซ่อมบำรุง",
                "รายการที่ต้องซ่อมบำรุงมี Equipment ลิงก์กับบัตรสินทรัพย์ (หมายเลขทรัพย์สินตรงกันทั้งสองฝั่ง)",
            ),
        ]
    if spec.get("file", "").endswith("จัดซื้อ.xlsx"):
        steps[2:2] = [
            (
                "2ก",
                "กลุ่มผู้จำหน่าย",
                "ชีตยืนยัน-กลุ่มผู้จำหน่าย มีกลุ่ม เช่น เจ้าหนี้ตรง / ส่วนสาธารณูปโภค และชีตผู้จำหน่ายคอลัมน์กลุ่มตรงกับที่นำเข้า",
            ),
            (
                "2ข",
                "PO คงค้าง",
                "ชีตยืนยัน-POคงค้าง แสดงเฉพาะ PO ที่ยังรับของไม่ครบหรือยังตั้งหนี้ไม่ครบ ถ้าว่างให้ยืนยันว่ายังไม่นำเข้า",
            ),
        ]
    if spec.get("file", "").endswith("เจ้าหนี้บัญชีการเงิน.xlsx"):
        steps[2:2] = [
            (
                "2ก",
                "บัญชีธนาคารของโรงพยาบาล",
                "ชีตยืนยัน-บัญชีธนาคารรพ มีชื่อบัญชี เลขที่บัญชี และรหัสสมุดรายวัน BK* ครบตามที่นำเข้า",
            ),
            (
                "2ข",
                "เจ้าหนี้คงค้าง",
                "ชีตยืนยัน-เจ้าหนี้คงค้าง มียอด amount_residual ≠ 0 หลังผ่านรายการ ถ้าว่างให้ยืนยันว่ายังไม่นำเข้าบิลค้างจ่าย",
            ),
            (
                "2ค",
                "งบทดลอง TB ปี 2569",
                "ชีตยืนยัน-TB2569 รวมยอดบัญชีที่ผ่านรายการ ระหว่าง 1 ต.ค. 2568 ถึง 30 ก.ย. 2569 ถ้าว่างให้ยืนยันว่ายังไม่นำเข้า TB / ยอดยกเข้า",
            ),
        ]
    if spec.get("file", "").endswith("คลังสินค้า.xlsx"):
        steps[2:2] = [
            (
                "2ก",
                "คลังที่เปิดใช้และรับเข้า 1 ขั้น",
                "ชีตยืนยัน-คลังมีเฉพาะคลังที่เปิดใช้และรับเข้า 1 ขั้น (one step) ให้หน่วยงานยืนยันก่อน คลังที่ปิดหรือหลายขั้นไม่แสดงในรอบนี้",
            ),
            (
                "2ข",
                "ตำแหน่ง internal ที่เปิดใช้",
                "ชีตยืนยัน-ตำแหน่งมีเฉพาะ usage=internal และเปิดใช้ ของคลังข้างต้น ไม่รวม Input/Output/ตำแหน่งที่ปิด",
            ),
        ]
    for i, row in enumerate(steps, 5):
        for c, v in enumerate(row, 1):
            sc(ws.cell(i, c, v), align="center" if c == 1 else "left")
        ws.row_dimensions[i].height = 36
    page(ws, spec["title"])


def build_module(mod_key, snapshot, cur):
    spec = MODULES[mod_key]
    items = [d for d in DATASETS if d["mod"] == mod_key]
    wb = Workbook()
    write_cover(wb, spec, snapshot, items)
    write_summary(wb, spec, snapshot, items)
    for d in items:
        if not d.get("list"):
            continue
        try:
            cols, rows = fetch_all(cur, d["list"])
        except Exception as exc:
            print(f"LIST FAIL {d['id']}: {exc}")
            cur.connection.rollback()
            if d.get("list_fallback"):
                try:
                    cols, rows = fetch_all(cur, d["list_fallback"])
                except Exception as exc2:
                    print(f"LIST FALLBACK FAIL {d['id']}: {exc2}")
                    cur.connection.rollback()
                    cols, rows = ["ข้อความ"], [{"ข้อความ": f"ดึงรายการไม่สำเร็จ: {exc}"}]
            else:
                cols, rows = ["ข้อความ"], [{"ข้อความ": f"ดึงรายการไม่สำเร็จ: {exc}"}]
        write_listing(wb, spec, d, cols, rows)
    write_howto(wb, spec)
    path = OUT_DIR / dated_xlsx(spec["file"])
    wb.save(path)
    n_list = sum(1 for d in items if d.get("list"))
    print(f"{len(items):2d} ประเภท  {n_list:2d} ชีตรายการ  {path.name}")
    return path


def build_overview(snapshot):
    spec = {
        "title": "สรุปภาพรวมการนำเข้าข้อมูล Master และรายการคงค้าง · ลงนามรับรอง",
        "system": "Odoo 18 · โรงพยาบาลวชิระภูเก็ต",
        "doc": "DM-SIGNOFF-ALL-001",
        "owner": "ผู้ประสานงานโครงการ ERP",
        "roles": [
            "งานงบประมาณ",
            "พัสดุ / จัดซื้อ",
            "งานคลัง",
            "บัญชี",
            "การเงิน",
            "IT / ผู้ดูแลระบบ",
            "ผู้บริหารโครงการ",
        ],
    }
    wb = Workbook()
    write_cover(wb, spec, snapshot, DATASETS)
    write_summary(wb, spec, snapshot, DATASETS)
    # module rollup
    ws = wb.create_sheet("สรุปตามโมดูล")
    hdr(ws, 1, ["โมดูล", "ไฟล์รายละเอียด", "Master (ประเภท)", "Master ที่นำเข้าแล้ว", "รวม record Master", "คงค้างที่ยังเป็น 0", "ผู้ลงนาม", "วันที่"], TEAL)
    for i, w in enumerate([22, 42, 16, 18, 20, 22, 18, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    labels = {
        "budget": "งบประมาณ",
        "purchase": "จัดซื้อ",
        "stock": "คลังสินค้า",
        "finance": "เจ้าหนี้ / บัญชี / การเงิน",
        "asset": "สินทรัพย์ถาวร",
    }
    r = 2
    for key, meta in MODULES.items():
        items = [d for d in DATASETS if d["mod"] == key]
        masters = [d for d in items if d["kind"] == "Master"]
        empty_tx = [
            d["name"]
            for d in items
            if d["kind"] in ("ธุรกรรม", "คงค้าง") and d["n"] == 0
        ]
        vals = [
            labels[key],
            dated_xlsx(meta["file"]),
            len(masters),
            sum(1 for d in masters if d["n"] > 0),
            sum(d["n"] for d in masters),
            ", ".join(empty_tx) or "—",
            "",
            "",
        ]
        for c, v in enumerate(vals, 1):
            sc(ws.cell(r, c, v), align="center" if c in (3, 4, 5, 8) else "left")
        ws.cell(r, 5).number_format = "#,##0"
        ws.row_dimensions[r].height = 36
        r += 1
    page(ws, spec["title"])
    write_howto(wb, spec)
    path = OUT_DIR / dated_xlsx("UAT-นำเข้าข้อมูล-สรุปและลงนาม.xlsx")
    wb.save(path)
    print(f"ภาพรวม  {path.name}")
    return path


def main():
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--only",
        choices=["budget", "purchase", "stock", "finance", "asset", "overview", "all"],
        default="all",
        help="สร้างเฉพาะไฟล์โมดูลที่ระบุ",
    )
    args = parser.parse_args()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    snapshot = thai_dt(datetime.now())
    conn = connect()
    try:
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            load_counts(cur)
            keys = list(MODULES) if args.only in ("all", "overview") else [args.only]
            if args.only != "overview":
                for key in keys:
                    build_module(key, snapshot, cur)
        if args.only in ("all", "overview"):
            build_overview(snapshot)
    finally:
        conn.close()


if __name__ == "__main__":
    main()
