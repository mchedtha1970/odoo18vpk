#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Build the NEW Thai Word User Manual for Group 3: พัสดุ (Procurement)
Strictly aligned with:
  ภาพรวมการทำรายการในระบบ ERP ตามกลุ่มงาน.pdf

Covers:
  (2.2) สร้าง PR ตามแผน / นอกแผน  — งานของหน่วยงานผู้ขอซื้อ
  (3.1) รับ PR ออกเอกสารส่งอนุมัติ
  (3.2) ดำเนินการ e-GP
  (3.3) ออกใบสั่งซื้อ บันทึกทะเบียนสัญญา
  (3.4) สร้างใบตรวจรับ (WA) และใบรับสินค้า (Auto) รอให้คลัง
  (3.5) ประเมินผู้ขาย
  (3.6) การจัดการข้อมูลผู้ขาย

Reference codes appear in parentheses at the end of every heading.
Every numbered step is followed immediately by a screenshot.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DOC_VERSION = "1.0"
OUT_DIR = ROOT / "คู่มือการใช้งาน-ตามผังภาพรวม"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "02-คู่มือการใช้งาน-กลุ่มงานพัสดุ.docx"
FIG = ROOT / "manual_assets" / "figures"

PRIMARY = RGBColor(0x0B, 0x68, 0x48)
DARK_NAVY = RGBColor(0x0F, 0x29, 0x42)
SLATE = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC6, 0x28, 0x28)
AMBER = RGBColor(0xB4, 0x53, 0x09)
GREEN = RGBColor(0x15, 0x80, 0x3D)
TEAL = RGBColor(0x0F, 0x76, 0x6E)

HEADER_FILL = "0B6848"
SUBHEADER_FILL = "0F2942"
META_FILL = "EAF3F1"
LOGO = FIG / "moph_logo.png"


def set_run(run, size=12, bold=False, color=SLATE, font_name="Prompt"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font_name)


def add_text(p, text, size=12, bold=False, color=SLATE):
    run = p.add_run(text)
    set_run(run, size, bold, color)
    return run


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, color="CBD5E1"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_narrow_cell_margin(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), "90")
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def cell_text(cell, text, size=11, bold=False, color=SLATE, align="left", fill=None):
    if fill:
        shade_cell(cell, fill)
    set_cell_border(cell)
    set_narrow_cell_margin(cell)
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.clear()
    add_text(p, text, size, bold, color)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)


def make_table(doc, headers, rows, col_widths=None, header_fill=HEADER_FILL):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, 11, True, WHITE, "center", header_fill)
    for r, row in enumerate(rows, 1):
        fill = "F8FAFC" if r % 2 == 0 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, 11, False, SLATE, "left", fill)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    return table


def add_border_line(paragraph, edge="bottom", color="0B6848", size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), "6")
    el.set(qn("w:color"), color)
    pBdr.append(el)
    pPr.append(pBdr)


def page_title(doc, text, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    add_text(p, text, 18, True, PRIMARY)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(
            run,
            17 if level == 1 else 14 if level == 2 else 12.5,
            True,
            PRIMARY if level == 1 else DARK_NAVY,
        )
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    if level == 1:
        add_border_line(p, "bottom", "0B6848", "12")
    return p


def body(doc, text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    add_text(p, text, size)
    return p


def bullets(doc, items, numbered=False):
    for item in items:
        p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        p.clear()
        add_text(p, item, 12)
        p.paragraph_format.space_after = Pt(3)


_FIG_COUNTER = 0


def caption(doc, text):
    global _FIG_COUNTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = False
    if re.search(r"ภาพที่\s+\d+", text):
        _FIG_COUNTER += 1
        text = re.sub(r"ภาพที่\s+\d+", f"ภาพที่ {_FIG_COUNTER}", text, count=1)
    label = text if text.startswith("▲") else f"▲ {text}"
    add_text(p, label, 10, False, MUTED)


def picture(doc, name_or_path, width_cm=15.8, cap=""):
    path = Path(name_or_path)
    if not path.is_absolute():
        path = FIG / name_or_path
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    if cap:
        p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    if cap:
        caption(doc, cap)


def pictures(doc, images, cap=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    if cap:
        p.paragraph_format.keep_with_next = True
    for index, (name, width_cm) in enumerate(images):
        path = FIG / name
        if not path.exists():
            continue
        if index:
            p.add_run(" ")
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width_cm))
    if cap:
        caption(doc, cap)


def step_with_pic(doc, num, text, img_name=None, width_cm=15.8, cap="", sub_bullets=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    if img_name or sub_bullets:
        p.paragraph_format.keep_with_next = True
    add_text(p, f"ขั้นตอนที่ {num}  ", 12, True, PRIMARY)
    add_text(p, text, 12, False, SLATE)
    if sub_bullets:
        bullets(doc, sub_bullets, numbered=False)
    if img_name:
        picture(doc, img_name, width_cm, cap)


def callout(doc, title, text, tone="info"):
    fill = {"info": "EAF3F1", "warn": "FEF3C7", "danger": "FEE2E2", "ok": "DCFCE7"}[tone]
    border = {"info": "6EE7B7", "warn": "FCD34D", "danger": "FCA5A5", "ok": "86EFAC"}[tone]
    title_color = {"info": PRIMARY, "warn": AMBER, "danger": RED, "ok": GREEN}[tone]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell, border)
    set_narrow_cell_margin(cell)
    p1 = cell.paragraphs[0]
    add_text(p1, title, 11.5, True, title_color)
    p2 = cell.add_paragraph()
    add_text(p2, text, 11.5, False, SLATE)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(3)
    doc.add_paragraph()


def task_summary_card(doc, code, title, desc, doc_name, menu_path, role, header_fill=HEADER_FILL):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows_data = [
        ("หัวข้องานตามผังภาพรวม", f"{title} ({code})", True, header_fill, WHITE),
        ("รายละเอียดงาน (Task)", desc, False, META_FILL, SLATE),
        ("เอกสารในระบบ (Document)", doc_name, True, "FFFFFF", DARK_NAVY),
        ("เมนูใช้งาน / ผู้รับผิดชอบ", f"เมนู: {menu_path}   |   ผู้ปฏิบัติงาน: {role}", False, "F8FAFC", SLATE),
    ]
    for idx, (k, v, is_bold, bg, fg) in enumerate(rows_data):
        c0 = table.rows[idx].cells[0]
        c1 = table.rows[idx].cells[1]
        c0.width = Cm(4.8)
        c1.width = Cm(11.6)
        if idx == 0:
            cell_text(c0, k, 11, True, WHITE, "left", SUBHEADER_FILL)
            cell_text(c1, v, 11.5, True, fg, "left", bg)
        else:
            cell_text(c0, k, 10.5, True, DARK_NAVY, "left", "F1F5F9")
            cell_text(c1, v, 11, is_bold, fg, "left", bg)
    doc.add_paragraph()


def header_footer(doc):
    section = doc.sections[0]
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(1.0)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(hp, "คู่มือการใช้งานระบบ ERP (VPK) — กลุ่มงานพัสดุ (2.2, 3.1 – 3.6)", 8.5, False, MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_border_line(fp, "top", "CBD5E1", "6")
    add_text(fp, "โรงพยาบาลวชิระภูเก็ต · คู่มือปฏิบัติงานระบบพัสดุ (2.2, 3.1 – 3.6) | หน้า ", 9, False, MUTED)
    run = fp.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run(run, 9, False, MUTED)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Prompt"
    normal.font.size = Pt(12)
    normal.font.color.rgb = SLATE
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), "Prompt")
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.18


def build():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    header_footer(doc)

    # ==================== COVER ====================
    if LOGO.exists():
        logo_p = doc.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_p.paragraph_format.space_before = Pt(4)
        logo_p.paragraph_format.space_after = Pt(4)
        logo_p.add_run().add_picture(str(LOGO), width=Cm(2.9))

    hosp = doc.add_paragraph()
    hosp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hosp.paragraph_format.space_after = Pt(0)
    add_text(hosp, "โรงพยาบาลวชิระภูเก็ต", 16, True, DARK_NAVY)

    hosp_en = doc.add_paragraph()
    hosp_en.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hosp_en.paragraph_format.space_after = Pt(8)
    add_text(hosp_en, "VACHIRAPHUKET HOSPITAL — ERP (VPK)", 10.5, True, MUTED)

    badge_p = doc.add_paragraph()
    badge_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    badge_p.paragraph_format.space_after = Pt(4)
    add_text(badge_p, "คู่มือการปฏิบัติงานตามผังสรุปภาพรวมการบันทึกรายการในระบบ ERP", 12.5, True, PRIMARY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    add_text(title, "คู่มือการใช้งาน : กลุ่มงานพัสดุ", 24, True, DARK_NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(10)
    add_text(
        sub,
        "การสร้าง PR ตามแผน / นอกแผน (2.2) · รับ PR ออกเอกสารส่งอนุมัติ (3.1)\n"
        "ดำเนินการ e-GP (3.2) · ออกใบสั่งซื้อและบันทึกทะเบียนสัญญา (3.3)\n"
        "สร้างใบตรวจรับ (WA) และใบรับสินค้าอัตโนมัติ (3.4) · ประเมินผู้ขาย (3.5) · จัดการข้อมูลผู้ขาย (3.6)",
        11.0,
        False,
        SLATE,
    )

    info = [
        ("รหัสเอกสาร", "UM-ERP-VPK-G02"),
        ("เอกสารอ้างอิงหลัก", "ภาพรวมการทำรายการในระบบ ERP ตามกลุ่มงาน.pdf"),
        (
            "หัวข้อที่ครอบคลุม",
            "สร้าง PR ตามแผน/นอกแผน (2.2), รับ PR ออกเอกสารส่งอนุมัติ (3.1), ดำเนินการ e-GP (3.2), "
            "ออกใบสั่งซื้อ บันทึกทะเบียนสัญญา (3.3), สร้างใบตรวจรับ (WA) และใบรับสินค้าอัตโนมัติ (3.4), "
            "ประเมินผู้ขาย (3.5), การจัดการข้อมูลผู้ขาย (3.6)",
        ),
        (
            "โมดูลในระบบ",
            "purchase_request, vpk_official_document, vpk_purchase_agreement_egp, "
            "purchase, vpk_work_acceptance, vpk_vendor_evaluation, vpk_vendor_api",
        ),
        ("กลุ่มผู้ใช้งาน", "เจ้าหน้าที่หน่วยงานผู้ขอซื้อ, เจ้าหน้าที่พัสดุ/จัดซื้อ, ผู้ขาย (แอปยืนยัน PO)"),
        ("รูปแบบคู่มือ", "Step-by-Step พร้อมภาพหน้าจอประกอบใต้ทุกขั้นตอน"),
        ("เวอร์ชันเอกสาร", f"{DOC_VERSION} (วันที่ 20 ตุลาคม พ.ศ. 2569)"),
        (
            "เอกสารคู่กัน",
            "คู่มือกลุ่มงานงบประมาณ UM-ERP-VPK-G01 (ทำแผนตั้งงบ 2.1 และการตรวจงบ 1.4) · "
            "คู่มือระบบจัดซื้อจัดจ้าง UM-ERP-PROC-001",
        ),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(info):
        cell_text(table.rows[i].cells[0], k, 10.5, True, PRIMARY, "left", META_FILL)
        cell_text(table.rows[i].cells[1], v, 10.5, False, SLATE, "left", "FFFFFF")
        table.rows[i].cells[0].width = Cm(4.6)
        table.rows[i].cells[1].width = Cm(11.8)

    doc.add_page_break()

    # ==================== TOC ====================
    page_title(doc, "สารบัญและโครงสร้างคู่มือตามผังภาพรวม")
    body(
        doc,
        "คู่มือฉบับนี้เป็นเล่มที่ 2 ของชุดคู่มือตามผังสรุปภาพรวม ครอบคลุมงานของหน่วยงานผู้ขอซื้อในข้อ 2.2 "
        "และงานของกลุ่มงานพัสดุทั้งข้อ 3.1 ถึง 3.6 โดยระบุหมายเลขหัวข้ออ้างอิงจากผังภาพรวมไว้ในวงเล็บต่อท้ายชื่อหัวข้อ "
        "และแสดงภาพหน้าจอประกอบใต้ทุกขั้นตอน",
    )
    callout(
        doc,
        "ข้อ 2.1 ทำแผนตั้งของบประมาณประจำปี อยู่ในเล่มงบประมาณ",
        "ขั้นตอนจัดทำคำขอตั้งงบประจำปีทั้ง 4 แบบฟอร์ม (วัสดุ/ครุภัณฑ์/ก่อสร้าง/โครงการ) ทั้งในแผนและนอกแผน "
        "อยู่ในคู่มือการใช้งานกลุ่มงานงบประมาณ (UM-ERP-VPK-G01) บทที่ 1 (2.1) "
        "เล่มนี้เริ่มจากหน่วยงานสร้างใบขอซื้อ (PR) หลังจากมีวงเงินในแผนแล้ว",
        "info",
    )

    make_table(
        doc,
        ["บทในคู่มือ", "หัวข้อในคู่มือ (รหัสอ้างอิงผังภาพรวม)", "งานที่ปฏิบัติในระบบ (Tasks)", "เอกสารในระบบ (Documents)"],
        [
            [
                "บทที่ 1",
                "สร้าง PR ตามแผน / นอกแผน (2.2)",
                "ขอซื้อ/จ้าง/เช่า ตามแผนหรือนอกแผน ระบุ Spec และกรรมการ",
                "ใบขอซื้อ/จ้าง (PR), ใบ Spec",
            ],
            [
                "บทที่ 2",
                "รับ PR ตรวจสอบความถูกต้อง และส่งเข้า e-GP (3.1)",
                "รับ PR ตรวจสอบความถูกต้อง ส่งกลับแก้ไขหรือยืนยันการตรวจสอบเพื่อส่งเข้า e-GP",
                "ใบขอซื้อ/จ้าง (PR), แท็บเอกสารราชการที่เกี่ยวข้อง (เรียกดูจาก e-GP)",
            ],
            [
                "บทที่ 3",
                "ดำเนินการ e-GP ส่งอนุมัติ และลงนามเอกสาร (3.2)",
                "ใส่ชื่อและเลขโครงการ เพิ่มเอกสารแนบไฟล์จาก e-GP กำหนดผู้อนุมัติส่งเข้ากล่องอนุมัติ ลงนามโดยตีกรอบตำแหน่งลายเซ็น บันทึกเปรียบเทียบราคา และเลือกผู้ชนะ",
                "คำสั่งตั้งกรรมการ, รายงานขออนุมัติ, แบบแสดงความบริสุทธิ์ใจ, รายงานผลพิจารณา, ประกาศผู้ชนะ",
            ],
            [
                "บทที่ 4",
                "ออกใบสั่งซื้อ บันทึกทะเบียนสัญญา (3.3)",
                "สร้าง PO บันทึกทะเบียนสัญญา ส่ง PO เข้าแอปให้ผู้ขายยืนยัน",
                "ใบสั่งซื้อ/จ้าง (PO), สัญญาซื้อขาย/จะซื้อจะขาย",
            ],
            [
                "บทที่ 5",
                "สร้างใบตรวจรับ (WA) และใบรับสินค้าอัตโนมัติ (3.4)",
                "สร้างใบตรวจรับพัสดุ รอคณะกรรมการตรวจรับ ใบรับสินค้าถูกสร้างอัตโนมัติให้คลัง",
                "ใบตรวจรับพัสดุ (WA), ใบรับสินค้า (Auto)",
            ],
            [
                "บทที่ 6",
                "ประเมินผู้ขาย (3.5)",
                "บันทึกการประเมินผู้ขายเมื่อรับสินค้าแล้ว",
                "แบบประเมินผู้ขาย",
            ],
            [
                "บทที่ 7",
                "การจัดการข้อมูลผู้ขาย (3.6)",
                "สร้าง ปรับปรุง ข้อมูลที่เกี่ยวข้องกับผู้ขาย",
                "แฟ้มผู้ขาย, รายการราคาจากผู้ขาย",
            ],
        ],
        [2.0, 5.2, 5.0, 4.2],
    )

    picture(
        doc,
        "fig_erp_overview_page1.png",
        16.4,
        "ภาพที่ 1  ผังสรุปภาพรวมการบันทึกรายการในระบบ ERP ของกลุ่มงานต่างๆ (อ้างอิงคอลัมน์ 2 ข้อ 2.2 และคอลัมน์ 3 พัสดุ ข้อ 3.1 – 3.6)",
    )

    doc.add_page_break()

    # ==================== CHAPTER 1 (2.2) ====================
    heading(doc, "บทที่ 1 การสร้างใบขอซื้อ PR ตามแผน / นอกแผน (2.2)", 1)
    task_summary_card(
        doc,
        code="2.2",
        title="สร้าง PR (ตามแผน / นอกแผน)",
        desc="ขอซื้อ/จ้าง/เช่า ตามแผนหรืองบนอกแผน ระบุคุณลักษณะเฉพาะ (Spec) และรายชื่อคณะกรรมการ แล้วส่งใบขอซื้อให้พัสดุ",
        doc_name="ใบขอซื้อ/จ้าง (PR), ใบ Spec",
        menu_path="ใบขอซื้อ → สร้าง",
        role="เจ้าหน้าที่หน่วยงานผู้ขอซื้อ",
        header_fill="0F766E",
    )

    body(
        doc,
        "หน่วยงานเปิดใบขอซื้อได้เมื่อมีวงเงินในแผนงบประมาณประจำปีแล้ว ทั้งจากคำขอตั้งงบ “ในแผน” ที่ได้รับอนุมัติตามรอบปี "
        "หรือจากคำขอตั้งงบ “นอกแผน” ที่หน่วยงบประมาณอนุมัติเพิ่มเติมเมื่องบเดิมไม่พอ "
        "(ดูรายละเอียดการขอตั้งงบในคู่มือกลุ่มงานงบประมาณ UM-ERP-VPK-G01 บทที่ 1 ข้อ 2.1 และบทที่ 2–4 ข้อ 1.1–1.3)",
    )
    callout(
        doc,
        "ตามแผน กับ นอกแผน ต่างกันอย่างไร (2.2)",
        "• ตามแผน: ใช้วงเงินจากคำขอตั้งงบประจำปีที่หน่วยงบประมาณอนุมัติเข้าแผนแล้ว\n"
        "• นอกแผน: ใช้วงเงินจากคำขอตั้งงบประเภท “นอกแผน” ที่อนุมัติเพิ่มเติมระหว่างปี "
        "เมื่อหน่วยงานกดเช็คงบบนใบขอซื้อ ระบบจะรวมวงเงินทั้งในแผนและนอกแผนของหมวดงบนั้นมาคำนวณคงเหลือ",
        "info",
    )

    heading(doc, "1.1 การเข้าสู่ระบบและสร้างใบขอซื้อ (2.2)", 2)
    step_with_pic(
        doc,
        1,
        "เปิดเว็บบราว์เซอร์เข้าสู่ระบบ ERP ของโรงพยาบาล กรอกชื่อผู้ใช้และรหัสผ่านของหน่วยงาน แล้วคลิกเข้าสู่ระบบ",
        "Screenshot_login.png",
        13.5,
        "ภาพที่ 2  หน้าจอเข้าสู่ระบบ ERP สำหรับผู้ใช้งานหน่วยงาน (2.2)",
    )
    step_with_pic(
        doc,
        2,
        "เข้าเมนู “ใบขอซื้อ” แล้วกดปุ่ม “สร้าง” ระบบจะออกเลขที่อัตโนมัติรูปแบบ PR + ปี พ.ศ. 2 หลัก + เดือน + เลขวิ่ง 3 หลัก เช่น PR6910005 ช่องเลขที่แก้เองไม่ได้",
        "Screenshot_PR6910005.png",
        15.8,
        "ภาพที่ 3  ตัวอย่างใบขอซื้อที่ระบบออกเลขที่ให้แล้ว (2.2)",
    )
    picture(doc, "Screenshot_New_PR.png", 15.8, "ภาพที่ 4  หน้าจอฟอร์มใบขอซื้อฉบับร่าง (2.2)")

    heading(doc, "1.2 กรอกหัวเอกสารใบขอซื้อ (2.2)", 2)
    step_with_pic(
        doc,
        1,
        "เลือกประเภทพัสดุ เช่น ซื้อ/จ้าง/เช่า (หรือจ้างที่ปรึกษา / จ้างออกแบบและควบคุมงานก่อสร้าง ตามงานจริง)",
        "Screenshot_PR_procurement_type.png",
        13.5,
        "ภาพที่ 5  เลือกประเภทพัสดุ (2.2)",
    )
    step_with_pic(
        doc,
        2,
        "เลือกประเภทการจัดซื้อเป็น “จัดซื้อจัดจ้างผ่านพัสดุ” เพื่อให้ใบนี้เข้ากระบวนการ e-GP ของกลุ่มงานพัสดุ "
        "เมื่อเลือกแล้วระบบจะแสดงช่องชื่อสำหรับซื้อใน e-GP ให้ ถ้ามีการกำหนดชื่อสำหรับซื้อใน e-GP ไว้ในฐานข้อมูลสินค้านั้นๆ",
        "Screenshot_PR_purchase_type.png",
        13.5,
        "ภาพที่ 6  เลือกประเภทการจัดซื้อ (2.2)",
    )
    step_with_pic(
        doc,
        3,
        "เลือกวิธีการจัดซื้อจัดจ้าง เช่น เฉพาะเจาะจง, E-bidding, คัดเลือก หรือประกาศเชิญชวนทั่วไป ให้ตรงวิธีที่ใช้จริง",
        "Screenshot_PR_procurement_method.png",
        13.5,
        "ภาพที่ 7  เลือกวิธีการจัดซื้อจัดจ้าง (2.2)",
    )
    step_with_pic(
        doc,
        4,
        "เลือกประเภทสิ่งที่จะซื้อ หน่วยงานงบประมาณ แหล่งเงิน และหมวดงบประมาณ ให้ตรงงบที่ได้รับอนุมัติ (ในแผนหรือนอกแผน)",
        "Screenshot_PR_budget_keys.png",
        14.0,
        "ภาพที่ 8  ประเภทสิ่งที่จะซื้อ หน่วยงานงบประมาณ แหล่งเงิน และหมวดงบประมาณ (2.2)",
    )
    step_with_pic(
        doc,
        5,
        "ถ้าซื้อวัสดุ ต้องเลือกกลุ่มงบประมาณด้วย เพื่อให้ระบบจับคู่ยอดจัดสรรรายสินค้าได้ถูกต้อง",
        "Screenshot_PR_budget_group.png",
        13.5,
        "ภาพที่ 9  เลือกกลุ่มงบประมาณกรณีซื้อวัสดุ (2.2)",
    )
    step_with_pic(
        doc,
        6,
        "เลือกรับสินค้าเข้าคลังประเภท Receipts ของคลังที่ต้องการให้ของเข้า",
        "Screenshot_PR_picking_type.png",
        15.8,
        "ภาพที่ 10  เลือกประเภทการรับสินค้าเข้าคลัง (2.2)",
    )
    step_with_pic(
        doc,
        7,
        "ใส่คำขอจัดซื้อในช่องคำอธิบาย จากนั้นใส่รายการวัสดุหรือครุภัณฑ์ในแท็บสินค้าตามหัวข้อถัดไป",
        "Screenshot_PR_description.png",
        15.8,
        "ภาพที่ 11  ช่องคำอธิบายของใบขอซื้อ (2.2)",
    )

    make_table(
        doc,
        ["ช่องในระบบ", "ตัวอย่าง", "ทำไมต้องถูก"],
        [
            ["ประเภทพัสดุ", "ซื้อ/จ้าง/เช่า", "กำหนดลักษณะงานจัดซื้อ"],
            ["ประเภทการจัดซื้อ", "จัดซื้อจัดจ้างผ่านพัสดุ", "เปิดเส้นทาง e-GP ให้พัสดุรับงานต่อ"],
            ["วิธีการจัดซื้อจัดจ้าง", "E-bidding / เฉพาะเจาะจง", "ต้องตรงวิธีที่หน่วยงานใช้จริง"],
            ["ประเภทสิ่งที่จะซื้อ", "วัสดุ / ครุภัณฑ์ / โครงการ / ก่อสร้าง", "ใช้จับคู่กับหมวดงบ"],
            ["หน่วยงานงบประมาณ", "หน่วยงานของตน", "ต้องตรงงบในแผนหรือนอกแผน"],
            ["แหล่งเงิน", "เงินบำรุง", "ห้ามใช้เงินคนละแหล่ง"],
            ["หมวดงบประมาณ", "วัสดุ", "ต้องตรงบรรทัดงบที่อนุมัติ"],
            ["ประเภทการรับสินค้า", "คลังยา: Receipts", "เลือกได้เฉพาะ Receipts"],
        ],
        [4.4, 5.2, 6.6],
    )

    heading(doc, "1.3 แท็บสินค้า — ใส่รายการและไฟล์คุณลักษณะ (2.2)", 2)
    body(
        doc,
        "แท็บสินค้าเป็นแท็บแรกและสำคัญที่สุด ถ้าไม่มีบรรทัดสินค้า ระบบจะส่งขออนุมัติไม่ได้ "
        "กดเพิ่มรายการที่ท้ายตาราง แล้วกรอกทีละบรรทัดขณะใบยังเป็นร่าง",
    )
    picture(doc, "Screenshot_TAB_dept_tabs.png", 15.8, "ภาพที่ 12  แท็บบนใบขอซื้อที่หน่วยงานต้องกรอก (2.2)")
    picture(doc, "Screenshot_TAB_overview_items.png", 15.8, "ภาพที่ 13  ภาพรวมแท็บสินค้า (2.2)")

    step_with_pic(doc, 1, "เปิดแท็บสินค้า แล้วกด “เพิ่มรายการ” ที่ท้ายตาราง", "Screenshot_TAB_ITEM_add_line.png", 15.8, "ภาพที่ 14  ปุ่มเพิ่มรายการที่ท้ายตารางสินค้า (2.2)")
    step_with_pic(doc, 2, "เลือกสินค้าจากรหัสหรือชื่อ เมื่อเลือกแล้วระบบจะเติมชื่อ หน่วยนับ และชื่อสำหรับซื้อใน e-GP ให้ ถ้ามีการกำหนดชื่อสำหรับซื้อใน e-GP ไว้ในฐานข้อมูลสินค้านั้นๆ", "Screenshot_TAB_ITEM_product.png", 15.8, "ภาพที่ 15  เลือกสินค้าจากข้อมูลหลัก (2.2)")
    step_with_pic(doc, 3, "ใส่จำนวนที่ขอ และตรวจหน่วยวัดให้ถูกต้อง เช่น กล่อง ขวด ชิ้น", "Screenshot_TAB_ITEM_qty_uom.png", 5.0, "ภาพที่ 16  จำนวนและหน่วยวัดของรายการสินค้า (2.2)")
    step_with_pic(doc, 4, "ใส่วันที่ต้องการของ ถ้าไม่ใส่ระบบจะใช้วันที่สร้างใบ", "Screenshot_TAB_ITEM_date.png", 5.0, "ภาพที่ 17  วันที่ร้องขอของรายการสินค้า (2.2)")
    step_with_pic(doc, 5, "ใส่งบประมาณของบรรทัดนั้น เป็นยอดรวมของรายการที่ใช้เช็คงบ", "Screenshot_TAB_ITEM_budget.png", 4.5, "ภาพที่ 18  งบประมาณของรายการสินค้า (2.2)")
    step_with_pic(doc, 6, "ตรวจชื่อสำหรับซื้อใน e-GP ถ้าชื่อในคลังไม่ตรงประกาศ ให้แก้ในช่องนี้", "Screenshot_TAB_ITEM_egp_name.png", 11.0, "ภาพที่ 19  ชื่อสำหรับซื้อใน e-GP (2.2)")
    step_with_pic(doc, 7, "ถ้ามีคุณลักษณะเฉพาะ กดปุ่มดู/แนบไฟล์ PDF ของรายการนั้น", "Screenshot_TAB_ITEM_spec.png", 7.0, "ภาพที่ 20  ปุ่มดู/แนบไฟล์คุณลักษณะของรายการ (2.2)")
    picture(doc, "Screenshot_TAB_ITEM_spec_popup.png", 15.8, "ภาพที่ 21  หน้าต่างแนบไฟล์คุณลักษณะเฉพาะ (2.2)")

    callout(
        doc,
        "อย่าส่งใบที่ยังไม่มีสินค้า (2.2)",
        "ถ้าตารางสินค้าว่าง หรือจำนวนเป็น 0 ระบบจะไม่ส่งขออนุมัติ "
        "กรณีมีการแก้รายการสินค้าแล้วต้องกดเช็คงบประมาณอีกครั้งก่อนส่ง",
        "warn",
    )

    heading(doc, "1.4 แท็บรายชื่อคณะกรรมการ (2.2)", 2)
    picture(doc, "Screenshot_TAB_overview_committee.png", 15.8, "ภาพที่ 22  ภาพรวมแท็บรายชื่อคณะกรรมการ (2.2)")
    step_with_pic(doc, 1, "เปิดแท็บรายชื่อคณะกรรมการ แล้วกดเพิ่มรายการในกลุ่มคณะกรรมการที่เกี่ยวข้อง จำนวนชุดกรรมการขึ้นอยู่กับวิธีการจัดซื้อจัดจ้าง", "Screenshot_TAB_committee_add.png", 15.8, "ภาพที่ 23  กดปุ่มเพิ่มรายการกรรมการ (2.2)")
    step_with_pic(doc, 2, "เลือกรายชื่อกรรมการจากพนักงานที่มีในระบบ", "Screenshot_TAB_committee_roles.png", 5.5, "ภาพที่ 24  เลือกชื่อกรรมการ (2.2)")
    step_with_pic(doc, 3, "ระบบจะเติมชื่อ แผนก และอีเมลให้ ตรวจว่ามีประธานเพียง 1 คน แล้วบันทึกใบขอซื้อ", "Screenshot_TAB_committee_filled.png", 15.8, "ภาพที่ 25  รายชื่อกรรมการที่กรอกครบ (2.2)")
    picture(doc, "Screenshot_TAB_committee_save.png", 6.5, "ภาพที่ 26  ปุ่มบันทึกใบขอซื้อ (2.2)")

    heading(doc, "1.5 แท็บตรวจสอบงบประมาณ (2.2)", 2)
    body(
        doc,
        "หน่วยงานต้องกดเช็คงบขณะใบยังเป็นฉบับร่าง ยอดที่ขอต้องตรงรายการสินค้า และสถานะต้องเป็น “งบเพียงพอ” จึงส่งพัสดุได้ "
        "รายละเอียดสูตรงบคงเหลือ ตัวอย่างตัวเลข และการแก้กรณีงบไม่พอ อธิบายไว้ในคู่มือกลุ่มงานงบประมาณ (UM-ERP-VPK-G01) บทที่ 5 หัวข้อ 5.2 "
        "และคู่มือระบบงบประมาณ (UM-ERP-BUD-001) บทที่ 6–7",
    )
    picture(doc, "Screenshot_TAB_overview_budget.png", 15.8, "ภาพที่ 27  ภาพรวมแท็บตรวจสอบงบประมาณ (2.2)")
    step_with_pic(doc, 1, "กดปุ่ม “เช็คงบประมาณ” ที่หัวเอกสาร ขณะสถานะยังเป็นฉบับร่าง", "Screenshot_TAB_44_check_btn.png", 4.2, "ภาพที่ 28  ปุ่มตรวจสอบงบประมาณที่หัวเอกสาร (2.2)")
    step_with_pic(doc, 2, "เปิดแท็บตรวจสอบงบประมาณ ดูสถานะเช็คงบ งบประมาณที่จอง และงบคงเหลือหลังหักใบนี้", "Screenshot_TAB_44_enough.png", 15.8, "ภาพที่ 29  ผลเช็คงบกรณีงบเพียงพอ (2.2)")
    picture(doc, "Screenshot_TAB_44_insufficient.png", 10.0, "ภาพที่ 30  ผลเช็คงบกรณีงบไม่เพียงพอ ให้กลับไปขอตั้งงบนอกแผนตามคู่มืองบประมาณ บทที่ 1–4 แล้วกดเช็คอีกครั้ง (2.2)")

    heading(doc, "1.6 การส่งใบขอซื้อให้กลุ่มงานพัสดุ (2.2)", 2)
    step_with_pic(
        doc,
        1,
        "เมื่อใส่รายการสินค้า รายชื่อคณะกรรมการครบถ้วน และแท็บตรวจสอบงบประมาณแสดงสถานะ “งบเพียงพอ” แล้ว ให้กดปุ่ม “ส่งพัสดุ” ที่หัวเอกสาร สถานะใบขอซื้อจะเปลี่ยนเป็น “ส่งพัสดุแล้ว” พร้อมกันงบประมาณอัตโนมัติ เพื่อให้กลุ่มงานพัสดุตรวจสอบความถูกต้องของรายการใบขอซื้อและดำเนินการต่อในระบบ e-GP ตามบทที่ 2",
        "Screenshot_TAB_44_send_btn.png",
        4.5,
        "ภาพที่ 31  ปุ่มส่งพัสดุที่หัวเอกสาร (แสดงเมื่อตรวจสอบงบประมาณผ่านแล้ว) (2.2)",
    )
    picture(doc, "Screenshot_PR_status_sent.png", 15.8, "ภาพที่ 32  สถานะใบขอซื้อเปลี่ยนเป็นส่งพัสดุแล้ว (2.2)")

    doc.add_page_break()

    # ==================== CHAPTER 2 (3.1) ====================
    heading(doc, "บทที่ 2 การรับ PR ตรวจสอบรายการ และออกเอกสารส่งอนุมัติ (3.1)", 1)
    task_summary_card(
        doc,
        code="3.1",
        title="รับ PR ตรวจสอบความถูกต้องของรายการ และออกเอกสารส่งอนุมัติ",
        desc="รับใบขอซื้อที่หน่วยงานส่งมา ตรวจสอบความถูกต้องของรายการใบขอซื้อ (หากต้องแก้ไขกดส่งกลับให้แก้ไขพร้อมระบุเหตุผล) เมื่อถูกต้องกดยืนยันการตรวจสอบ แล้วส่งเข้าระบบ e-GP รวมถึงออกชุดหนังสือขออนุมัติ",
        doc_name="ใบขอซื้อ (PR), บันทึกขอซื้อ, รายงานขอซื้อ, คำสั่งตั้งกรรมการ",
        menu_path="ระบบจัดซื้อจัดจ้าง → ข้อมูลใบขอซื้อ → Dashboard พัสดุรับงาน",
        role="เจ้าหน้าที่พัสดุ / จัดซื้อ",
    )

    heading(doc, "2.1 รับงานจาก Dashboard พัสดุ และตรวจสอบความถูกต้องของรายการใบขอซื้อ (3.1)", 2)
    step_with_pic(
        doc,
        1,
        "เข้าเมนูระบบจัดซื้อจัดจ้าง → ข้อมูลใบขอซื้อ → Dashboard พัสดุรับงาน เปิดใบขอซื้อที่สถานะ “ส่งพัสดุแล้ว” แล้วกดปุ่ม “รับงาน” หรือ “เปิดทำรายการ”",
        "Screenshot_ch5_dashboard_work.png",
        15.8,
        "ภาพที่ 41  หน้าจอ Dashboard พัสดุรับงาน (3.1)",
    )
    step_with_pic(
        doc,
        2,
        "ตรวจสอบความถูกต้องของรายการสินค้า จำนวน หน่วยนับ ราคา งบประมาณ ชื่อสำหรับซื้อใน e-GP และไฟล์คุณลักษณะเฉพาะ (Spec) บนใบขอซื้อ ทั้งนี้ในสถานะ “ส่งพัสดุแล้ว” ระบบจะล็อกข้อมูลและรายการสินค้าไม่ให้แก้ไขโดยตรง โดยที่หัวเอกสารจะมีปุ่ม “ยืนยันการตรวจสอบ” (สีเขียว) และปุ่ม “ส่งกลับให้แก้ไข” (สีเหลือง)",
        "Screenshot_ch5_verify_return_btns.png",
        15.8,
        "ภาพที่ 42  หน้าจอใบขอซื้อสถานะส่งพัสดุแล้ว แสดงปุ่ม “ยืนยันการตรวจสอบ” และ “ส่งกลับให้แก้ไข” (3.1)",
    )
    step_with_pic(
        doc,
        3,
        "กรณีพบรายการไม่ถูกต้องหรือข้อมูลไม่ครบถ้วน ให้กดปุ่ม “ส่งกลับให้แก้ไข” ที่หัวเอกสาร ระบบจะแสดงหน้าต่างป๊อปอัปให้ระบุเหตุผลหรือรายละเอียดที่ต้องการให้หน่วยงานแก้ไข แล้วกดปุ่ม “ยืนยันส่งกลับให้แก้ไข”",
        "Screenshot_ch5_return_wizard.png",
        15.0,
        "ภาพที่ 43  หน้าต่างระบุเหตุผลในการส่งกลับใบขอซื้อให้หน่วยงานแก้ไข (3.1)",
    )
    body(
        doc,
        "เมื่อกดยืนยันส่งกลับ ระบบจะเปลี่ยนสถานะใบขอซื้อกลับเป็น “ฉบับร่าง” คืนงบประมาณที่จองไว้ชั่วคราว และแสดงแถบแจ้งเตือนสีเหลืองพร้อมเหตุผลที่หัวใบขอซื้อ เพื่อให้ผู้ขอซื้อของหน่วยงานตรวจสอบ แก้ไขรายการ กดตรวจสอบงบประมาณ และกดส่งพัสดุเข้ามาใหม่",
    )
    picture(
        doc,
        "Screenshot_ch5_returned_alert.png",
        13.8,
        "ภาพที่ 44  ใบขอซื้อที่ถูกส่งกลับสถานะฉบับร่าง พร้อมแถบแจ้งเตือนสีเหลืองแสดงเหตุผลที่ต้องแก้ไข (3.1)",
    )
    step_with_pic(
        doc,
        4,
        "กรณีตรวจสอบแล้วพบว่ารายการใบขอซื้อถูกต้องครบถ้วน ให้กดปุ่ม “ยืนยันการตรวจสอบ” ที่หัวเอกสาร ระบบจะบันทึกชื่อผู้ตรวจสอบและวันที่ตรวจสอบ แสดงป้ายสถานะสีเขียว “พัสดุตรวจสอบความถูกต้องแล้ว พร้อมส่งเข้าระบบ e-GP” และเปิดปุ่ม “ส่งรายการเข้าระบบ e-GP” ให้ดำเนินการในขั้นตอนถัดไป",
        "Screenshot_ch5_verified_ready_egp.png",
        14.2,
        "ภาพที่ 45  ใบขอซื้อหลังกดยืนยันการตรวจสอบ แสดงป้ายสถานะสีเขียวและเปิดปุ่ม “ส่งรายการเข้าระบบ e-GP” (3.1)",
    )
    callout(
        doc,
        "สรุปเงื่อนไขการตรวจสอบความถูกต้องของรายการใบขอซื้อก่อนส่งเข้า e-GP (3.1)",
        "• สถานะ “ส่งพัสดุแล้ว” จะล็อกไม่ให้แก้ไขรายการสินค้าโดยตรง เพื่อให้ผ่านการตรวจสอบจากเจ้าหน้าที่พัสดุอย่างเป็นระบบ\n"
        "• ปุ่ม “ส่งรายการเข้าระบบ e-GP” จะแสดงขึ้นมาเฉพาะเมื่อเจ้าหน้าที่พัสดุกดปุ่ม “ยืนยันการตรวจสอบ” เรียบร้อยแล้วเท่านั้น\n"
        "• หากกด “ส่งกลับให้แก้ไข” ระบบจะบันทึกประวัติและเหตุผลลงใน Chatter ของใบขอซื้อโดยอัตโนมัติ",
        "info",
    )

    heading(doc, "2.2 การส่งรายการเข้าระบบ e-GP และการดูเอกสารราชการที่เกี่ยวข้อง (3.1)", 2)
    body(
        doc,
        "เอกสารราชการในขั้นตอนจัดซื้อจัดจ้าง (เช่น คำสั่งแต่งตั้งคณะกรรมการตรวจรับพัสดุ, แบบแสดงความบริสุทธิ์ใจ, "
        "ขออนุมัติแต่งตั้งคณะกรรมการกำหนดรายละเอียดคุณลักษณะ และรายงานขออนุมัติจัดซื้อจัดจ้างโดยวิธีเฉพาะเจาะจง) "
        "จะถูกจัดทำจากระบบ e-GP ของกรมบัญชีกลางโดยตรง แล้วนำไฟล์มาเพิ่มและแนบในเอกสารกระบวนการ e-GP (ตามบทที่ 3) "
        "โดยที่แท็บ “เอกสารราชการที่เกี่ยวข้อง” บนใบขอซื้อ (PR) จะแสดงรายการเอกสารที่แนบไว้ใน e-GP ให้เปิดดูได้อย่างเดียว",
    )
    step_with_pic(
        doc,
        1,
        "เมื่อพัสดุกดปุ่ม “ยืนยันการตรวจสอบ” รายการใบขอซื้อเรียบร้อยแล้ว กดปุ่ม “ส่งรายการเข้าระบบ e-GP” ที่หัวเอกสารเพื่อสร้างเอกสารกระบวนการ e-GP และเดินกระบวนการจัดซื้อตามบทที่ 3",
        "Screenshot_ch5_send_egp_btn.png",
        15.8,
        "ภาพที่ 46  ปุ่มส่งรายการเข้าระบบ e-GP หลังพัสดุยืนยันการตรวจสอบ (3.1)",
    )
    step_with_pic(
        doc,
        2,
        "เมื่อเจ้าหน้าที่พัสดุได้เพิ่มและแนบไฟล์เอกสารที่สร้างจากระบบ e-GP ในเอกสารกระบวนการ e-GP แล้ว สามารถเปิดแท็บ “เอกสารราชการที่เกี่ยวข้อง” บนใบขอซื้อ (PR) เพื่อเรียกดูรายการเอกสาร เลขที่อ้างอิง วันที่เอกสาร และกดเปิดดูไฟล์ PDF ได้ทันที (สำหรับเปิดดูเท่านั้น)",
        "Screenshot_ch5_related_docs_tab.png",
        15.8,
        "ภาพที่ 47  แท็บเอกสารราชการที่เกี่ยวข้องบนใบขอซื้อ แสดงรายการเอกสารที่เพิ่มและแนบจากกระบวนการ e-GP สำหรับเปิดดูเท่านั้น (3.1)",
    )

    doc.add_page_break()

    # ==================== CHAPTER 3 (3.2) ====================
    heading(doc, "บทที่ 3 การดำเนินการ e-GP ขั้นตอนการส่งอนุมัติ และขั้นตอนการลงนาม (3.2)", 1)
    task_summary_card(
        doc,
        code="3.2",
        title="ดำเนินการ e-GP ส่งอนุมัติ และลงนามเอกสาร",
        desc="ใส่ชื่อและหมายเลขโครงการ กดปุ่มเพิ่มเอกสารและแนบไฟล์เอกสารที่สร้างจากระบบ e-GP กำหนดผู้อนุมัติส่งเข้ากล่องอนุมัติเอกสาร ลงนามโดยกำหนดกรอบตำแหน่งลายเซ็น (Mark ตำแหน่ง) บันทึกเปรียบเทียบราคา และเลือกผู้ชนะ",
        doc_name="คำสั่งแต่งตั้งกรรมการ, รายงานขออนุมัติ, แบบแสดงความบริสุทธิ์ใจ, รายงานผลพิจารณา, ประกาศผู้ชนะ",
        menu_path="จัดซื้อ → กระบวนการ eGP / กล่องอนุมัติเอกสาร",
        role="เจ้าหน้าที่พัสดุ / หัวหน้าเจ้าหน้าที่ / เลขานุการ / ผู้อำนวยการ",
    )
    body(
        doc,
        "กระบวนการ e-GP ในระบบนี้ใช้รองรับและเก็บข้อมูลที่ได้จากการทำในระบบ e-GP จริง เช่น ไฟล์เอกสาร PDF ที่สร้างจากระบบ e-GP "
        "เลขที่เอกสาร หรือชื่อโครงการ เพื่อใช้อ้างอิงในการทำงาน เชื่อมโยงให้เปิดดูได้จากใบขอซื้อ (PR) "
        "และรองรับการส่งอนุมัติและลงนามอิเล็กทรอนิกส์ผ่าน “กล่องอนุมัติเอกสาร” โดยตรง",
    )

    heading(doc, "3.1 ส่งใบขอซื้อเข้ากระบวนการ e-GP (3.2)", 2)
    step_with_pic(doc, 1, "เปิดใบขอซื้อที่พัสดุตรวจสอบความถูกต้องของรายการเรียบร้อยแล้ว", "Screenshot_ch7_approved_pr.png", 15.8, "ภาพที่ 48  ใบขอซื้อที่พร้อมส่งเข้าระบบ e-GP (3.2)")
    step_with_pic(doc, 2, "กดปุ่ม “ส่งรายการเข้าระบบ e-GP” ที่หัวใบขอซื้อ", "Screenshot_ch7_send_egp_header.png", 5.0, "ภาพที่ 49  ปุ่มส่งรายการเข้าระบบ e-GP (3.2)")
    step_with_pic(doc, 3, "ตรวจรายการสินค้าในหน้าต่างกระบวนการ eGP ว่าครบตามใบขอซื้อ", "Screenshot_ch7_egp_wizard_lines.png", 15.8, "ภาพที่ 50  หน้าต่างกระบวนการ eGP แสดงรายการสินค้าจากใบขอซื้อ (3.2)")
    step_with_pic(doc, 4, "กดปุ่ม “กระบวนการ eGP” ระบบจะสร้างเอกสารกระบวนการ e-GP จากใบขอซื้อที่อนุมัติแล้ว", "Screenshot_ch7_egp_wizard_confirm.png", 5.5, "ภาพที่ 51  ปุ่มยืนยันสร้างเอกสารกระบวนการ e-GP (3.2)")
    step_with_pic(doc, 5, "เปิดเอกสารที่สร้างได้จากปุ่มสถิติกระบวนการ e-GP บนใบขอซื้อ หรือจากเมนูจัดซื้อ → กระบวนการ eGP", "Screenshot_ch7_egp_stat_btn.png", 15.8, "ภาพที่ 52  ปุ่มสถิติกระบวนการ e-GP บนใบขอซื้อ (3.2)")

    heading(doc, "3.2 บันทึกข้อมูลโครงการและกดปุ่มเพิ่มเอกสารแนบไฟล์จากระบบ e-GP (3.2)", 2)
    step_with_pic(doc, 1, "บันทึกเลขที่โครงการ และชื่อโครงการที่ได้จากระบบ e-GP จริง", "Screenshot_ch7_project_fields.png", 15.8, "ภาพที่ 53  ช่องเลขที่โครงการและชื่อโครงการ e-GP (3.2)")
    step_with_pic(
        doc,
        2,
        "เข้าแท็บ “เอกสารที่เปิดในระบบ e-GP” กดปุ่ม “เพิ่มเอกสาร” แล้วเลือกหัวข้อประเภทเอกสารจากรายการ เช่น คำสั่งแต่งตั้งคณะกรรมการตรวจรับพัสดุ, แบบแสดงความบริสุทธิ์ใจ (วงเงินมากกว่า ๑๐๐,๐๐๐ บาท), ขออนุมัติแต่งตั้งคณะกรรมการกำหนดรายละเอียดคุณลักษณะ หรือรายงานขออนุมัติจัดซื้อจัดจ้างโดยวิธีเฉพาะเจาะจง",
        "Screenshot_egp_doc_type_dropdown.png",
        15.8,
        "ภาพที่ 54  การเลือกหัวข้อประเภทเอกสารที่สร้างจากระบบ e-GP (3.2)",
    )
    step_with_pic(
        doc,
        3,
        "ระบุเลขอ้างอิง e-GP วันที่เอกสาร และกดอัปโหลดไฟล์แนบ (PDF) ที่ได้จากระบบ e-GP ตามหัวข้อในรายการ โดยสามารถกดปุ่ม “เพิ่มเอกสาร” เพื่อแนบเอกสารหัวข้อถัดไปได้ครบทุกฉบับ",
        "Screenshot_ch7_egp_files.png",
        15.8,
        "ภาพที่ 55  รายการเอกสารที่แนบในแท็บ “เอกสารที่เปิดในระบบ e-GP” พร้อมปุ่ม “เพิ่มเอกสาร” (3.2)",
    )
    make_table(
        doc,
        ["หัวข้อประเภทเอกสารในรายการ", "รายละเอียดการใช้งาน"],
        [
            ["คำสั่งแต่งตั้งคณะกรรมการตรวจรับพัสดุ", "แนบไฟล์คำสั่งแต่งตั้งคณะกรรมการตรวจรับพัสดุที่สร้างจากระบบ e-GP"],
            ["แบบแสดงความบริสุทธิ์ใจ (วงเงินมากกว่า ๑๐๐,๐๐๐ บาท)", "แนบไฟล์แบบแสดงความบริสุทธิ์ใจในการจัดซื้อจัดจ้างที่ได้จากระบบ e-GP"],
            ["ขออนุมัติแต่งตั้งคณะกรรมการกำหนดรายละเอียดคุณลักษณะ", "แนบไฟล์หนังสือขออนุมัติแต่งตั้งกรรมการกำหนดคุณลักษณะเฉพาะ/ราคากลางจาก e-GP"],
            ["รายงานขออนุมัติจัดซื้อจัดจ้างโดยวิธีเฉพาะเจาะจง", "แนบไฟล์รายงานขอซื้อ/ขอจ้างโดยวิธีเฉพาะเจาะจงที่สร้างจากระบบ e-GP"],
            ["ประกาศร่าง / ประกาศเชิญชวน / TOR", "แนบไฟล์ประกาศร่าง ประกาศเชิญชวน หรือขอบเขตงาน (TOR) จากระบบ e-GP"],
            ["รายงานผลการพิจารณา / รายงานผลการพิจารณาจัดซื้อจัดจ้าง", "แนบไฟล์รายงานผลการพิจารณาที่ได้จากระบบ e-GP แล้วส่งลงนาม — เมื่อลงนามครบและมีประกาศผู้ชนะ ปุ่มสร้าง PO จะพร้อมใช้งาน"],
            ["ประกาศผู้ชนะ / สัญญา", "แนบไฟล์ประกาศผู้ชนะ หรือสัญญาที่ได้จากกระบวนการจัดซื้อ"],
        ],
        [6.8, 9.2],
    )
    callout(
        doc,
        "การเชื่อมโยงไปแสดงที่ใบขอซื้อ (PR) อัตโนมัติ (3.2)",
        "เอกสารทุกฉบับที่กดปุ่ม “เพิ่มเอกสาร” และแนบไฟล์ไว้ในแท็บ “เอกสารที่เปิดในระบบ e-GP” "
        "จะไปแสดงในแท็บ “เอกสารราชการที่เกี่ยวข้อง” ของใบขอซื้อ (PR) ต้นทางโดยอัตโนมัติ เพื่อให้ผู้เกี่ยวข้องเปิดดูได้อย่างเดียว",
        "info",
    )

    heading(doc, "3.3 ขั้นตอนการส่งอนุมัติเอกสาร (กำหนดผู้อนุมัติและส่งเข้ากล่องอนุมัติเอกสาร) (3.2)", 2)
    body(
        doc,
        "เมื่อเจ้าหน้าที่พัสดุแนบไฟล์ PDF ของเอกสารราชการในแท็บ “เอกสารที่เปิดในระบบ e-GP” เรียบร้อยแล้ว "
        "สามารถกำหนดรายชื่อและลำดับผู้อนุมัติแยกตามแต่ละบรรทัดเอกสาร เพื่อส่งไฟล์ PDF เข้าสู่ “กล่องอนุมัติเอกสาร” ของผู้มีอำนาจลงนามตามลำดับ",
    )
    step_with_pic(
        doc,
        1,
        "ในตารางแท็บ “เอกสารที่เปิดในระบบ e-GP” หลังจากแนบไฟล์ PDF ในบรรทัดเอกสารแล้ว ให้กดปุ่ม “เพิ่มผู้อนุมัติ” ในบรรทัดเอกสารนั้น เพื่อเปิดหน้าต่างกำหนดรายชื่อและลำดับผู้อนุมัติ",
        "Screenshot_egp_add_approver_btn.png",
        15.8,
        "ภาพที่ 55.1  ปุ่ม “เพิ่มผู้อนุมัติ” คอลัมน์ “ลงนาม/อนุมัติ” ปุ่ม “ส่งเพื่อลงนาม” และสถานะการลงนามของแต่ละบรรทัดเอกสาร (3.2)",
    )
    step_with_pic(
        doc,
        2,
        "ในหน้าต่าง “กำหนดผู้อนุมัติเอกสาร” ระบบจะดึงรายชื่อผู้อนุมัติเริ่มต้นตามประเภทเอกสารมาให้อัตโนมัติ โดยเจ้าหน้าที่พัสดุสามารถกด “เพิ่มบรรทัด” เพื่อเพิ่มผู้อนุมัติ ลบ หรือจัดลำดับใหม่ และเลือกบทบาทในแต่ละลำดับว่าเป็น “ตรวจสอบ” หรือ “ลงนาม” เมื่อตรวจสอบครบถ้วนแล้วให้กดปุ่ม “บันทึกและส่งเข้ากล่องผู้อนุมัติ” (หรือกด “บันทึกรายชื่อผู้อนุมัติ” แล้วมากดปุ่ม “ส่งเพื่อลงนาม” ที่บรรทัดเอกสารภายหลังได้)",
        "Screenshot_egp_approver_modal.png",
        15.8,
        "ภาพที่ 55.2  หน้าต่างกำหนดผู้อนุมัติเอกสาร แสดงรายชื่อตามลำดับ ตัวเลือกตรวจสอบ/ลงนาม และปุ่มบันทึกและส่งเข้ากล่องผู้อนุมัติ (3.2)",
    )
    step_with_pic(
        doc,
        3,
        "ระบบจะแสดงหน้าต่างยืนยันการส่งเอกสาร ตรวจสอบชื่อเรื่องและเลขที่เอกสาร แล้วกดปุ่ม “ส่งเอกสาร” ระบบจะเปลี่ยนสถานะของบรรทัดเอกสารเป็น “รอลงนาม” และส่งการ์ดงานเข้าสู่ “กล่องอนุมัติเอกสาร” ของผู้อนุมัติลำดับที่ 1 ทันที",
        "Screenshot_egp_send_wizard.png",
        15.8,
        "ภาพที่ 55.3  หน้าต่างยืนยันการกดปุ่ม “ส่งเอกสาร” เพื่อส่งเข้ากล่องอนุมัติเอกสารตามลำดับ (3.2)",
    )
    make_table(
        doc,
        ["ปุ่ม/คอลัมน์ในบรรทัดเอกสาร", "ความหมายและการทำงาน"],
        [
            ["ปุ่ม “เพิ่มผู้อนุมัติ”", "เปิดหน้าต่างกำหนดรายชื่อ ลำดับ และบทบาท (ตรวจสอบ / ลงนาม) ของเอกสารบรรทัดนั้น"],
            ["คอลัมน์ “ลงนาม/อนุมัติ”", "แสดงรายชื่อผู้ลงนาม/อนุมัติทั้งหมดที่ตั้งไว้ในบรรทัดเอกสารนั้นตามลำดับ"],
            ["ปุ่ม “ส่งเพื่อลงนาม”", "ส่งเอกสารบรรทัดที่แนบไฟล์ PDF และกำหนดผู้อนุมัติแล้ว เข้ากล่องอนุมัติเอกสาร"],
            ["สถานะ “รอลงนาม” / “ลงนามครบแล้ว”", "แสดงสถานะปัจจุบันของเอกสาร เมื่อผู้อนุมัติครบทุกขั้นจะเปลี่ยนเป็น “ลงนามครบแล้ว” (สีเขียว)"],
        ],
        [5.5, 10.5],
    )

    heading(doc, "3.4 ขั้นตอนการลงนามเอกสาร และการตรวจสอบรายการที่อนุมัติแล้วในกล่องอนุมัติเอกสาร (3.2)", 2)
    body(
        doc,
        "ผู้มีอำนาจตรวจสอบและลงนาม (เช่น เจ้าหน้าที่พัสดุ, หัวหน้าเจ้าหน้าที่พัสดุ, เลขานุการ, ผู้อำนวยการ หรือคณะกรรมการ) "
        "สามารถเข้าเมนู “กล่องอนุมัติเอกสาร” เพื่อเปิดดูไฟล์ PDF กำหนดกรอบตำแหน่งลายเซ็น (Mark ตำแหน่ง) ลงนามอิเล็กทรอนิกส์ "
        "รวมถึงเรียกดูรายการเอกสารที่ตนได้อนุมัติไปแล้วและติดตามสถานะการอนุมัติของเอกสารทุกฉบับได้ในที่เดียว",
    )
    step_with_pic(
        doc,
        1,
        "เข้าเมนู “กล่องอนุมัติเอกสาร” → “แดชบอร์ด” หน้าจอจะแสดงทั้งการ์ดรายการที่ “รอดำเนินการ” (ป้ายสีเหลือง มีปุ่ม “เปิด PDF” เพื่อลงนาม) และการ์ดรายการที่ “อนุมัติแล้ว” (ป้ายสีเขียว แสดงวันเวลาที่อนุมัติ พร้อมปุ่ม “เปิดดู PDF ที่อนุมัติแล้ว”) โดยสามารถกดปุ่มตัวกรองที่มุมขวาบนเพื่อสลับดู “ทั้งหมด”, “รออนุมัติ” หรือ “อนุมัติแล้ว” ได้ทันที",
        "Screenshot_egp_inbox_all_and_approved.png",
        15.8,
        "ภาพที่ 55.4  หน้าจอแดชบอร์ดกล่องอนุมัติเอกสาร แสดงปุ่มสลับตัวกรอง (ทั้งหมด / รออนุมัติ / อนุมัติแล้ว) พร้อมการ์ดรอดำเนินการและการ์ดที่อนุมัติแล้ว (3.2)",
    )
    step_with_pic(
        doc,
        2,
        "หากต้องการดูเฉพาะประวัติเอกสารที่ตนได้ลงนามหรืออนุมัติไปแล้ว ให้คลิกปุ่มตัวกรอง “อนุมัติแล้ว” บนแดชบอร์ด หรือเข้าเมนูย่อย “รายการที่อนุมัติแล้วของฉัน” ทางแถบเมนูด้านซ้าย และกดปุ่ม “เปิดดู PDF ที่อนุมัติแล้ว” บนการ์ดเพื่อเปิดดูไฟล์ PDF ที่ประทับลายเซ็นแล้ว",
        "Screenshot_egp_inbox_approved_only.png",
        15.8,
        "ภาพที่ 55.5  การกรองดูเฉพาะรายการที่อนุมัติแล้ว และเมนู “รายการที่อนุมัติแล้วของฉัน” (3.2)",
    )
    step_with_pic(
        doc,
        3,
        "สำหรับการลงนามเอกสารที่รออนุมัติ ให้กดปุ่ม “เปิด PDF” บนการ์ดที่แสดงสถานะ “รอดำเนินการ” ระบบจะเปิดหน้าต่างแสดงไฟล์เอกสาร PDF ขึ้นมา จากนั้นให้กดปุ่ม “กำหนดกรอบลายเซ็น (Mark ตำแหน่ง)” ที่แถบเครื่องมือด้านบนหรือแถบด้านล่าง",
        "Screenshot_egp_inbox_mark_toolbar.png",
        15.8,
        "ภาพที่ 55.6  หน้าต่างเปิดอ่านไฟล์ PDF จากกล่องอนุมัติเอกสาร พร้อมปุ่ม “กำหนดกรอบลายเซ็น (Mark ตำแหน่ง)” (3.2)",
    )
    step_with_pic(
        doc,
        4,
        "เลื่อนหน้าเอกสารไปยังจุดที่ต้องการวางลายเซ็น แล้วคลิกลากตีกรอบ (หรือคลิก 1 ครั้ง) บนหน้าเอกสาร PDF ระบบจะแสดงกรอบเส้นประสีเขียว “กรอบตำแหน่งลายเซ็น” ณ ตำแหน่งที่เลือก (สามารถคลิกลากใหม่เพื่อเปลี่ยนตำแหน่งได้ตามต้องการ) จากนั้นกดปุ่มสีเขียว “ลงนาม”",
        "Screenshot_egp_mark_signature_box.png",
        15.8,
        "ภาพที่ 55.7  การคลิกลากตีกรอบตำแหน่งลายเซ็น (กรอบเส้นประสีเขียว) บนหน้าเอกสาร PDF และปุ่ม “ลงนาม” (3.2)",
    )
    step_with_pic(
        doc,
        5,
        "ในหน้าต่าง “ใช้ลายเซ็นของคุณ” ให้เลือกวาดลายเซ็น (วาด) พิมพ์ชื่อ (อัตโนมัติ) หรืออัปโหลดไฟล์รูปลายเซ็น (โหลด) แล้วกดปุ่ม “ยอมรับและลงนาม”",
        "Screenshot_egp_draw_signature_in_box.png",
        15.8,
        "ภาพที่ 55.8  หน้าต่างวาดลายเซ็นและปุ่ม “ยอมรับและลงนาม” (3.2)",
    )
    step_with_pic(
        doc,
        6,
        "ระบบจะนำรูปลายเซ็นไปประทับลงในกรอบตำแหน่งที่กำหนดไว้บนไฟล์ PDF โดยตรง บันทึกไฟล์ PDF ฉบับที่ลงนามแล้วซิงก์กลับไปยังแท็บ “เอกสารที่เปิดในระบบ e-GP” และแท็บ “เอกสารราชการที่เกี่ยวข้อง” ของใบขอซื้อ (PR) อัตโนมัติ พร้อมส่งต่อให้ผู้อนุมัติลำดับถัดไปทันที",
        "Screenshot_egp_stamped_signature_result.png",
        15.8,
        "ภาพที่ 55.9  ผลลัพธ์เอกสาร PDF ที่ประทับลายเซ็นลงตรงตามตำแหน่งกรอบที่กำหนดไว้เรียบร้อยแล้ว (3.2)",
    )
    step_with_pic(
        doc,
        7,
        "ผู้ใช้งานทุกคนสามารถติดตามความคืบหน้าการลงนามของเอกสารแต่ละฉบับได้ที่เมนู “กล่องอนุมัติเอกสาร” → “สถานะการอนุมัติ” ซึ่งจะแสดงการ์ดสถานะของแต่ละเอกสารว่าปัจจุบันรออยู่ที่ลำดับใด ใครลงนามแล้ว และเอกสารใดอนุมัติครบทุกขั้นตอนแล้ว",
        "Screenshot_egp_approval_status_board.png",
        15.8,
        "ภาพที่ 55.10  หน้าจอ “สถานะการอนุมัติ” แสดงความคืบหน้าการลงนามรายลำดับของเอกสารแต่ละฉบับ (3.2)",
    )

    heading(doc, "3.5 บันทึกเปรียบเทียบราคา เลือกผู้ชนะ และแนบรายงานผลการพิจารณา / ประกาศผู้ชนะจากระบบ e-GP (3.2)", 2)
    step_with_pic(doc, 1, "เข้าแท็บ “เปรียบเทียบราคา” กดเพิ่มรายการเพื่อบันทึกข้อมูลผู้เสนอราคา เลขอ้างอิงใบเสนอราคา e-GP วันที่เสนอราคา และราคาที่เสนอ", "Screenshot_ch7_compare_prices.png", 14.5, "ภาพที่ 57  แท็บเปรียบเทียบราคา (3.2)")
    step_with_pic(doc, 2, "เลือกผู้ชนะโดยติ๊กช่อง “ผู้ชนะ” ในรายการของผู้เสนอราคาที่ได้รับการคัดเลือก (เลือกได้ทีละ 1 ราย)", "Screenshot_ch7_pick_winner_clean.png", 14.5, "ภาพที่ 58  การติ๊กเลือกผู้ชนะในแท็บเปรียบเทียบราคา (3.2)")
    callout(
        doc,
        "ยังสร้าง PO ไม่ได้ตอนนี้ (3.2)",
        "แม้ติ๊กเลือกผู้ชนะแล้ว ปุ่ม “สร้าง PO” จะยังไม่พร้อมใช้งานจนกว่าจะแนบไฟล์รายงานผลการพิจารณา (หรือรายงานผลการพิจารณาจัดซื้อจัดจ้าง) และประกาศผู้ชนะที่ได้จากระบบ e-GP ของภาครัฐ แล้วส่งอนุมัติ/ลงนามครบตามลำดับเรียบร้อยแล้ว",
        "warn",
    )
    step_with_pic(
        doc,
        3,
        "กลับมาที่แท็บ “เอกสารที่เปิดในระบบ e-GP” กดปุ่ม “เพิ่มเอกสาร” แล้วเลือกประเภทเอกสาร “รายงานผลการพิจารณา” หรือ “รายงานผลการพิจารณาจัดซื้อจัดจ้าง” และ “ประกาศผู้ชนะ” เพื่อระบุเลขอ้างอิง วันที่เอกสาร และแนบไฟล์ PDF ที่จัดทำจากระบบ e-GP ของกรมบัญชีกลางโดยตรง (ไม่มีการกดปุ่มสร้างเอกสารใหม่ในระบบ ERP)",
        "Screenshot_egp_award_winner_attached.png",
        13.2,
        "ภาพที่ 59  การเพิ่มเอกสารและแนบไฟล์ PDF รายงานผลการพิจารณา และประกาศผู้ชนะ ที่ได้จากระบบ e-GP ของรัฐ ในแท็บ “เอกสารที่เปิดในระบบ e-GP” (3.2)",
    )
    step_with_pic(
        doc,
        4,
        "ดำเนินการส่งอนุมัติและลงนามสำหรับบรรทัดเอกสาร “รายงานผลการพิจารณา” และ “ประกาศผู้ชนะ” ด้วยขั้นตอนเดียวกับหัวข้อ 3.3 และ 3.4 คือ กดปุ่ม “เพิ่มผู้อนุมัติ” → ตรวจสอบลำดับผู้อนุมัติในหน้าต่างกำหนดผู้อนุมัติเอกสาร → กดปุ่ม “บันทึกและส่งเข้ากล่องผู้อนุมัติ” → กดยืนยัน “ส่งเอกสาร” เพื่อให้ผู้มีอำนาจเปิดเอกสารจากเมนู “กล่องอนุมัติเอกสาร” กำหนดกรอบตำแหน่งลายเซ็น (Mark ตำแหน่ง) และกดลงนามตามลำดับ",
        "Screenshot_egp_approver_modal.png",
        11.8,
        "ภาพที่ 60  การกำหนดผู้อนุมัติและส่งเอกสารเข้ากล่องอนุมัติเอกสาร (ใช้ Flow การส่งอนุมัติและลงนามเดียวกับหัวข้อ 3.3 และ 3.4) (3.2)",
    )
    make_table(
        doc,
        ["ประเภทเอกสารที่แนบจากระบบ e-GP", "ลำดับผู้อนุมัติเริ่มต้นในกล่องอนุมัติเอกสาร"],
        [
            ["รายงานผลการพิจารณา / รายงานผลการพิจารณาจัดซื้อจัดจ้าง", "เจ้าหน้าที่พัสดุ (ลงนาม) → หัวหน้าเจ้าหน้าที่พัสดุ (ลงนาม) → เลขาผู้อำนวยการ (ตรวจสอบ) → ผู้อำนวยการ (ลงนาม)"],
            ["ประกาศผู้ชนะ", "เลขาผู้อำนวยการ (ตรวจสอบ) → ผู้อำนวยการ (ลงนาม)"],
        ],
        [5.5, 10.5],
    )

    doc.add_page_break()

    # ==================== CHAPTER 4 (3.3) ====================
    heading(doc, "บทที่ 4 การออกใบสั่งซื้อและบันทึกทะเบียนสัญญา (3.3)", 1)
    task_summary_card(
        doc,
        code="3.3",
        title="ออกใบสั่งซื้อ บันทึกทะเบียนสัญญา",
        desc="สร้าง PO บันทึกทะเบียนสัญญา (แนบไฟล์สัญญา) ส่ง PO เข้าแอปให้ผู้ขายยืนยัน",
        doc_name="ใบสั่งซื้อ/จ้าง (PO), สัญญาซื้อขาย/จะซื้อจะขาย",
        menu_path="จัดซื้อ → ใบสั่งซื้อ / จัดการสั่งซื้อ → ทะเบียนสัญญา",
        role="เจ้าหน้าที่พัสดุ / จัดซื้อ",
    )

    heading(doc, "4.1 สร้างใบสั่งซื้อจากเอกสาร e-GP (3.3)", 2)
    step_with_pic(doc, 1, "เปิดเอกสารกระบวนการ e-GP ที่ประกาศผู้ชนะและลงนามครบแล้ว แล้วกดปุ่ม “สร้าง PO”", "Screenshot_ch9_create_po_btn.png", 7.0, "ภาพที่ 66  ปุ่มสร้าง PO บนเอกสารกระบวนการ e-GP (3.3)")
    step_with_pic(doc, 2, "ตรวจหน้าต่างว่าผู้ขายคือผู้ชนะที่เลือกไว้ เปลี่ยนเป็นรายอื่นไม่ได้", "Screenshot_ch8_po_wizard.png", 12.8, "ภาพที่ 67  หน้าต่างสร้าง PO แสดงผู้ขายผู้ชนะ (3.3)")
    step_with_pic(doc, 3, "ใส่เลขอ้างอิงใบเสนอราคา e-GP ถ้ามี แล้วกดสร้าง PO ตรวจส่วนหัวและส่วนท้ายของใบสั่งซื้อ", "Screenshot_ch8_po_header.png", 15.8, "ภาพที่ 68  ส่วนหัวของเอกสารใบสั่งซื้อ (3.3)")
    picture(doc, "Screenshot_ch8_po_footer.png", 15.8, "ภาพที่ 69  ส่วนท้ายของเอกสารใบสั่งซื้อ (3.3)")

    heading(doc, "4.2 ส่งใบสั่งซื้อให้ผู้ขายยืนยันผ่านแอป (3.3)", 2)
    body(
        doc,
        "หลังสร้างใบสั่งซื้อแล้ว พัสดุพิมพ์ไฟล์และส่งเข้าแอปผู้ขาย ผู้ขายเห็นเฉพาะใบที่ส่งเข้าแอป "
        "เมื่อลงลายเซ็นและกดบันทึก สถานะบนใบสั่งซื้อจะกลับมาเป็น “ผู้ขายยืนยันแล้ว” และระบบสร้างใบรอรับสินค้าให้อัตโนมัติ",
    )
    step_with_pic(doc, 1, "เปิดฟอร์มใบสั่งซื้อ กดปุ่มพิมพ์ใบสั่งซื้อ ระบบสร้างไฟล์ PDF แนบในแท็บไฟล์ใบสั่งซื้อ", "Screenshot_ch8_po_print_btn.png", 11.0, "ภาพที่ 70  ปุ่มพิมพ์ใบสั่งซื้อ (3.3)")
    picture(doc, "Screenshot_ch8_po_files_tab.png", 15.8, "ภาพที่ 71  แท็บไฟล์แนบของใบสั่งซื้อ (3.3)")
    picture(doc, "Screenshot_ch8_po_print_sample.png", 13.9, "ภาพที่ 72  ตัวอย่างใบสั่งซื้อ (3.3)")
    step_with_pic(doc, 2, "กดส่งเข้าแอปผู้ขายเพื่อยืนยัน เมื่อส่งสำเร็จ สถานะการยืนยันเปลี่ยนจาก “ยังไม่ส่ง” เป็น “รอผู้ขายยืนยัน”", "Screenshot_ch8_po_send_btn.png", 5.0, "ภาพที่ 73  ปุ่มส่งเข้าแอปผู้ขาย (3.3)")
    picture(doc, "Screenshot_ch8_po_waiting.png", 6.0, "ภาพที่ 74  สถานะใบสั่งซื้อรอผู้ขายยืนยัน (3.3)")
    step_with_pic(
        doc,
        3,
        "ผู้ขายเปิดแอประบบบริการผู้จำหน่ายด้วยบัญชี User Portal จะเห็นการ์ดรอยืนยัน กดเปิดใบสั่งซื้อ ลงนามยืนยัน แล้วกดบันทึก",
        None,
    )
    pictures(
        doc,
        [
            ("Screenshot_ch8_app_1.png", 4.3),
            ("Screenshot_ch8_app_2.png", 4.3),
            ("Screenshot_ch8_app_3.png", 4.3),
        ],
        "ภาพที่ 75  แอปผู้ขายแสดงการ์ดรอยืนยันใบสั่งซื้อ (3.3)",
    )
    pictures(
        doc,
        [
            ("Screenshot_ch8_app_4.png", 4.3),
            ("Screenshot_ch8_app_5.png", 4.3),
            ("Screenshot_ch8_app_6.png", 4.3),
        ],
        "ภาพที่ 76  ผู้ขายเปิดเอกสารและลงนามยืนยัน (3.3)",
    )
    pictures(
        doc,
        [
            ("Screenshot_ch8_app_7.png", 4.3),
            ("Screenshot_ch8_app_8.png", 4.3),
        ],
        "ภาพที่ 77  ยืนยันและลงนามใบสั่งซื้อสำเร็จ (3.3)",
    )
    picture(doc, "Screenshot_ch8_po_confirmed.png", 8.0, "ภาพที่ 78  สถานะใบสั่งซื้อเป็นผู้ขายยืนยันแล้ว และระบบสร้างใบรอรับสินค้าอัตโนมัติ (3.3)")

    heading(doc, "4.3 บันทึกทะเบียนสัญญาจากกระบวนการ e-GP (3.3)", 2)
    body(
        doc,
        "งานที่ต้องทำสัญญากับผู้ชนะ เช่น งานจ้าง งานเช่า หรือซื้อเป็นงวด พัสดุบันทึกสัญญาในทะเบียนสัญญา "
        "โดยเริ่มจากเอกสารกระบวนการ e-GP ของงานนั้น ระบบดึงคู่ค้า มูลค่า และเลขที่โครงการ e-GP มาให้ "
        "สัญญาที่ยืนยันแล้วใช้เป็นฐานตรวจวงเงินเมื่อเปิดใบขอซื้อและใบสั่งซื้อในงวดถัดไป",
    )
    callout(
        doc,
        "ใครทำได้ (3.3)",
        "ปุ่มเปิดทะเบียนสัญญา และยืนยันสัญญา แสดงเฉพาะผู้ใช้สิทธิ์ผู้จัดการจัดซื้อ "
        "และต้องเลือกผู้ชนะในกระบวนการ e-GP แล้ว จึงจะได้คู่ค้าและมูลค่าตามผู้ชนะ",
        "info",
    )
    step_with_pic(doc, 1, "เปิดเอกสารกระบวนการ e-GP ที่เลือกผู้ชนะแล้ว กดปุ่มเปิดทะเบียนสัญญาที่แถบปุ่มหัวเอกสาร", "Screenshot_ch8_ctr_egp_btn.png", 15.8, "ภาพที่ 79  ปุ่มเปิดทะเบียนสัญญาที่หัวเอกสาร e-GP (3.3)")
    step_with_pic(
        doc,
        2,
        "ระบบสร้างสัญญาสถานะแบบร่าง พร้อมรหัสสัญญา CTR/ปี/เลขวิ่ง คู่ค้าคือผู้ชนะ มูลค่าคือยอดเสนอราคาของผู้ชนะ และอ้างอิงเอกสาร e-GP ให้อัตโนมัติ",
        "Screenshot_ch8_ctr_form_draft.png",
        15.8,
        "ภาพที่ 80  ทะเบียนสัญญาแบบร่างที่สร้างจากเอกสาร e-GP (3.3)",
    )
    step_with_pic(
        doc,
        3,
        "ถ้าสัญญาแบ่งจ่ายเป็นงวด เปิดแท็บเงื่อนไขแบ่งงวด กดสร้างรายการงวดงาน เลือกจำนวนงวดและวิธีแบ่ง แล้วแก้ชื่องวด มูลค่า ให้ผลรวมเท่ากับมูลค่ารวมสัญญา",
        "Screenshot_ch8_ctr_wizard.png",
        12.0,
        "ภาพที่ 81  หน้าต่างสร้างรายการงวดงาน (3.3)",
    )
    picture(doc, "Screenshot_ch8_ctr_installments.png", 15.8, "ภาพที่ 82  รายการงวดงานของสัญญา (3.3)")
    step_with_pic(doc, 4, "แท็บเอกสารสัญญา แนบสัญญาที่ลงนามแล้ว TOR และเอกสารประกอบ แท็บหลักประกันใส่หนังสือค้ำประกันถ้ามี", "Screenshot_ch8_ctr_docs.png", 15.8, "ภาพที่ 83  แท็บเอกสารสัญญา (3.3)")
    step_with_pic(
        doc,
        5,
        "ตรวจข้อมูลครบแล้วกดยืนยันสัญญา สถานะเปลี่ยนจากแบบร่างเป็นมีผลบังคับใช้ ระบบบันทึกวันที่ยืนยันและคำนวณวันสิ้นสุดสัญญาให้ถ้าเว้นว่างไว้",
        "Screenshot_ch8_ctr_confirmed.png",
        15.8,
        "ภาพที่ 84  สัญญาสถานะมีผลบังคับใช้ (3.3)",
    )
    picture(doc, "Screenshot_ch8_ctr_stat_btn.png", 15.8, "ภาพที่ 85  ปุ่มสัญญาในเมนูเพิ่มเติมของเอกสาร e-GP (3.3)")
    callout(
        doc,
        "แก้ไขหลังยืนยัน (3.3)",
        "สัญญาที่มีผลบังคับใช้แก้ข้อมูลหัวสัญญาไม่ได้ ถ้าต้องแก้ ให้กดยกเลิก แล้วตั้งเป็นแบบร่าง "
        "สัญญาที่ยังมี PR หรือ PO ใช้งานอยู่ยกเลิกไม่ได้ เมื่อสัญญาสิ้นสุดแล้วกดปิดสัญญา",
        "warn",
    )

    doc.add_page_break()

    # ==================== CHAPTER 5 (3.4) ====================
    heading(doc, "บทที่ 5 การสร้างใบตรวจรับ (WA) และใบรับสินค้าอัตโนมัติ (3.4)", 1)
    task_summary_card(
        doc,
        code="3.4",
        title="สร้างใบตรวจรับ (WA) และใบรับสินค้า (Auto) รอให้คลัง",
        desc="สร้างใบตรวจรับพัสดุ (WA) รอคณะกรรมการตรวจรับ และใบรับสินค้าถูกสร้างอัตโนมัติให้คลังเมื่อผู้ขายยืนยันใบสั่งซื้อแล้ว",
        doc_name="ใบตรวจรับพัสดุ (WA), ใบรับสินค้า (Auto)",
        menu_path="จัดซื้อ → ใบสั่งซื้อ (ปุ่มสร้างใบตรวจรับ) / จัดซื้อ → บันทึกรับมอบงาน",
        role="เจ้าหน้าที่พัสดุ / คณะกรรมการตรวจรับ",
    )
    body(
        doc,
        "เมื่อผู้ขายยืนยันใบสั่งซื้อจากแอป ระบบสร้างใบรอรับสินค้าให้อัตโนมัติเพื่อส่งต่อให้คลังตามข้อ 5 ของผังภาพรวม "
        "ส่วนใบตรวจรับพัสดุ (Work Acceptance) พัสดุสร้างจากใบสั่งซื้อที่ยืนยันแล้ว แล้วส่งให้คณะกรรมการตรวจรับบันทึกผลและลงนาม "
        "รายละเอียดการลงนามของกรรมการอยู่ในงานของกลุ่มที่ 4 (คณะกรรมการตรวจรับ) ข้อ 4.1–4.3",
    )

    heading(doc, "5.1 สร้างใบตรวจรับจากใบสั่งซื้อ (3.4)", 2)
    step_with_pic(
        doc,
        1,
        "เปิดใบสั่งซื้อที่เป็นคำสั่งซื้อแล้ว (ผู้ขายยืนยันแล้ว) กดปุ่มสร้างใบตรวจรับ ระบบเปิดบันทึกรับมอบงานฉบับร่าง ออกเลขที่ให้ และดึงรายการสินค้าจากใบสั่งซื้อพร้อมรายชื่อกรรมการจากคำสั่งแต่งตั้ง",
        "Screenshot_ch9_create_inspection_btn.png",
        15.8,
        "ภาพที่ 86  ปุ่มสร้างใบตรวจรับบนใบสั่งซื้อ (3.4)",
    )
    callout(
        doc,
        "สร้างซ้ำไม่ได้ (3.4)",
        "ใบสั่งซื้อที่มีใบตรวจรับแล้วยังคงเปิดใบเดิมได้ แต่ปุ่มสร้างใบตรวจรับจะไม่แสดงซ้ำ",
        "warn",
    )

    heading(doc, "5.2 กรอกผลการตรวจรับ (3.4)", 2)
    body(doc, "กรอกขณะสถานะยังเป็น Draft หลังกดตรวจรับแล้วช่องเหล่านี้จะแก้ไม่ได้ ต้องกดกลับเป็นร่างก่อนจึงจะแก้ได้")
    make_table(
        doc,
        ["ช่อง", "ใส่อะไร"],
        [
            ["ผลการตรวจรับ", "ตรวจรับครบถ้วน หรือตรวจรับบางส่วนถ้าของส่งไม่ครบ"],
            ["วันที่ตรวจรับ", "วันที่คณะกรรมการตรวจรับ"],
            ["วันที่ส่งมอบงาน", "วันที่ผู้ขายส่งของจริง"],
            ["วันครบกำหนดตามสัญญา", "วันสิ้นสุดตามสัญญา ถ้าว่างระบบจะไม่คิดว่าส่งช้า"],
            ["ค่าปรับ", "ใส่เมื่อส่งมอบช้าและมีค่าปรับ"],
            ["คณะกรรมการตรวจรับ", "ตรวจชื่อ บทบาทประธาน กรรมการ และเลขานุการ ที่ดึงจากคำสั่งแต่งตั้ง"],
        ],
        [5.2, 11.0],
    )
    picture(doc, "Screenshot_ch10_form.png", 15.8, "ภาพที่ 87  บันทึกรับมอบงานหลังตรวจรับ (3.4)")
    picture(doc, "Screenshot_ch10_committee.png", 15.8, "ภาพที่ 88  แท็บคณะกรรมการตรวจรับ ดึงชื่อจากคำสั่งแต่งตั้ง (3.4)")
    callout(
        doc,
        "ส่งมอบช้า (3.4)",
        "ถ้าวันที่ส่งมอบงานหลังวันครบกำหนดตามสัญญา ระบบขึ้นแถบส่งมอบงานล่าช้าและแสดงจำนวนวัน ให้ตรวจค่าปรับก่อนกดตรวจรับ",
        "warn",
    )

    heading(doc, "5.3 กดตรวจรับและส่งต่อคลัง (3.4)", 2)
    bullets(
        doc,
        [
            "กดตรวจรับ แล้วเลือกวันตรวจรับในหน้าต่างที่ยืนยัน",
            "สถานะเปลี่ยนจาก Draft เป็น Accepted",
            "ใบรับสินค้าที่ระบบสร้างอัตโนมัติเมื่อผู้ขายยืนยัน PO จะรอให้คลังตรวจคุณภาพและรับเข้าคลังตามข้อ 5.2–5.3 ของผังภาพรวม",
            "ถ้าผู้ขอซื้อมีอีเมล ระบบเข้าคิวส่งอีเมลแจ้งผู้ขอ",
        ],
        numbered=True,
    )
    callout(
        doc,
        "ผลไม่ผ่าน (3.4)",
        "ถ้าเลือกผลการตรวจรับเป็นไม่ผ่านการตรวจรับ ปุ่มตรวจรับจะกดไม่ผ่าน ให้บันทึกเหตุผลในแท็บหมายเหตุ อย่ากดตรวจรับเพื่อปิดงาน",
        "danger",
    )

    doc.add_page_break()

    # ==================== CHAPTER 6 (3.5) ====================
    heading(doc, "บทที่ 6 การประเมินผู้ขาย (3.5)", 1)
    task_summary_card(
        doc,
        code="3.5",
        title="ประเมินผู้ขาย",
        desc="บันทึกการประเมินผู้ขายเมื่อรับสินค้าแล้ว ตามเกณฑ์ 5 ด้าน (จัดส่ง คุณภาพ ราคา บริการ เอกสาร) ระบบคำนวณคะแนนถ่วงน้ำหนักและเกรดอัตโนมัติ",
        doc_name="แบบประเมินผู้ขาย (vendor.evaluation)",
        menu_path="จัดซื้อ → ประเมินผู้จำหน่าย → ใบประเมินผู้จำหน่าย",
        role="เจ้าหน้าที่พัสดุ / จัดซื้อ",
    )

    heading(doc, "6.1 เกณฑ์ประเมินมาตรฐาน 5 ด้าน (3.5)", 2)
    body(
        doc,
        "ระบบกำหนดเกณฑ์ประเมินผู้จำหน่ายมาตรฐาน 5 ด้าน น้ำหนักรวม 100% แต่ละด้านให้คะแนนเต็ม 5 "
        "หัวหน้างานพัสดุปรับเกณฑ์ได้ที่เมนูจัดซื้อ → ประเมินผู้จำหน่าย → เกณฑ์ประเมิน",
    )
    picture(doc, "Screenshot_proc_35_criteria.png", 15.8, "ภาพที่ 89  เกณฑ์ประเมินผู้จำหน่ายมาตรฐาน 5 ด้าน (3.5)")
    make_table(
        doc,
        ["รหัส", "เกณฑ์", "น้ำหนัก", "คะแนนเต็ม"],
        [
            ["DELIVERY", "ระยะเวลาการจัดส่ง", "25%", "5"],
            ["QUALITY", "คุณภาพสินค้า/บริการ", "30%", "5"],
            ["PRICE", "ราคาและเงื่อนไขการชำระ", "20%", "5"],
            ["SERVICE", "การให้บริการ/การสนับสนุน", "15%", "5"],
            ["DOCUMENT", "ความถูกต้องของเอกสาร", "10%", "5"],
        ],
        [3.0, 6.0, 3.4, 3.8],
    )
    make_table(
        doc,
        ["คะแนนรวม", "เกรด", "ผลประเมิน"],
        [
            ["90% ขึ้นไป", "A", "ดีเยี่ยม"],
            ["80% – 89.99%", "B", "ดี"],
            ["70% – 79.99%", "C", "พอใช้"],
            ["60% – 69.99%", "D", "ต้องปรับปรุง"],
            ["ต่ำกว่า 60%", "F", "ไม่ผ่านเกณฑ์"],
        ],
        [4.4, 3.6, 8.2],
    )

    heading(doc, "6.2 ขั้นตอนบันทึกใบประเมินผู้ขาย (3.5)", 2)
    step_with_pic(
        doc,
        1,
        "เข้าเมนูจัดซื้อ → ประเมินผู้จำหน่าย → ใบประเมินผู้จำหน่าย ระบบแสดงรายการใบประเมินทั้งหมด พร้อมคะแนน เกรด และสถานะ",
        "Screenshot_proc_35_list.png",
        15.8,
        "ภาพที่ 90  หน้ารายการใบประเมินผู้จำหน่าย (3.5)",
    )
    step_with_pic(
        doc,
        2,
        "กดปุ่ม “ใหม่” เลือกผู้จำหน่าย และเลือกใบสั่งซื้อที่ยืนยันแล้วของรายนั้น (ถ้ามี) ใส่วันที่ประเมินและช่วงประเมิน แล้วกดปุ่ม “โหลดเกณฑ์ประเมิน” เพื่อดึงเกณฑ์ทั้ง 5 ด้านเข้าตาราง",
        "Screenshot_proc_35_form_draft.png",
        15.8,
        "ภาพที่ 91  ใบประเมินสถานะร่าง หลังโหลดเกณฑ์และใส่คะแนน (3.5)",
    )
    step_with_pic(
        doc,
        3,
        "ให้คะแนนแต่ละเกณฑ์ (0 ถึง 5) และใส่หมายเหตุรายบรรทัด ระบบคำนวณร้อยละและคะแนนถ่วงน้ำหนักทันที เมื่อครบแล้วกดยืนยัน จากนั้นหัวหน้างานพัสดุกดอนุมัติ เกรดและคะแนนรวมจะถูกบันทึกลงแฟ้มผู้ขาย",
        "Screenshot_proc_35_form_approved.png",
        15.8,
        "ภาพที่ 92  ใบประเมินที่อนุมัติแล้ว แสดงคะแนนรวมและเกรด (3.5)",
    )
    step_with_pic(
        doc,
        4,
        "ดูสรุปผลประเมินที่อนุมัติแล้วจากเมนูจัดซื้อ → การรายงาน → รายงานประเมินผู้จำหน่าย แยกคะแนนรายด้าน และจัดกลุ่มตามผู้ขายหรือเกรดได้",
        "Screenshot_proc_35_report.png",
        15.8,
        "ภาพที่ 93  รายงานประเมินผู้จำหน่าย (แสดงเฉพาะใบที่อนุมัติแล้ว) (3.5)",
    )
    callout(
        doc,
        "เมื่อไรควรประเมิน (3.5)",
        "ประเมินเมื่อรับสินค้าหรือตรวจรับงานแล้ว โดยผูกกับใบสั่งซื้อใบนั้น "
        "เกรดล่าสุดและคะแนนเฉลี่ยจะไปแสดงบนแฟ้มผู้ขายตามบทที่ 7 เพื่อใช้ประกอบการคัดเลือกผู้ขายครั้งถัดไป",
        "ok",
    )

    doc.add_page_break()

    # ==================== CHAPTER 7 (3.6) ====================
    heading(doc, "บทที่ 7 การจัดการข้อมูลผู้ขาย (3.6)", 1)
    task_summary_card(
        doc,
        code="3.6",
        title="การจัดการข้อมูลผู้ขาย",
        desc="สร้างและปรับปรุงข้อมูลที่เกี่ยวข้องกับผู้ขาย ได้แก่ แฟ้มผู้ขาย บัญชี Vendor Portal รายการราคาจากผู้ขาย และกลุ่มผู้จำหน่าย",
        doc_name="แฟ้มผู้ขาย (res.partner), รายการราคาผู้ขาย (product.supplierinfo)",
        menu_path="จัดซื้อ → ผู้จำหน่าย / ผลิตภัณฑ์ → รายการราคาผู้ขาย",
        role="เจ้าหน้าที่พัสดุ / จัดซื้อ",
    )

    heading(doc, "7.1 สร้างและปรับปรุงแฟ้มผู้ขาย (3.6)", 2)
    step_with_pic(
        doc,
        1,
        "เข้าเมนูผู้จำหน่าย ระบบแสดงรายชื่อผู้ขายทั้งหมด พร้อมเกรดผู้ขายจากผลการประเมินล่าสุด กดปุ่ม “ใหม่” เพื่อสร้างรายใหม่ หรือคลิกชื่อเพื่อเปิดแก้ไข",
        "Screenshot_proc_36_vendor_list.png",
        15.8,
        "ภาพที่ 94  หน้ารายการผู้จำหน่าย แสดงเกรดผู้ขาย (3.6)",
    )
    step_with_pic(
        doc,
        2,
        "กรอกชื่อผู้ขาย ที่อยู่ เลขประจำตัวผู้เสียภาษี อีเมล กลุ่มคู่ค้า และประเภท VAT ให้ครบ เลือกประเภทเป็นบริษัทหรือบุคคลให้ตรงกับเอกสารจดทะเบียน",
        "Screenshot_proc_36_vendor_form.png",
        15.8,
        "ภาพที่ 95  แฟ้มผู้ขายส่วนหัวและที่อยู่ (3.6)",
    )
    step_with_pic(
        doc,
        3,
        "เปิดแท็บการขายและการซื้อ เพื่อกำหนดเงื่อนไขการชำระเงินฝั่งจัดซื้อ และผู้ซื้อที่รับผิดชอบรายนี้ จากนั้นบันทึก",
        "Screenshot_proc_36_vendor_purchase_tab.png",
        15.8,
        "ภาพที่ 96  แท็บการขายและการซื้อของแฟ้มผู้ขาย (3.6)",
    )
    callout(
        doc,
        "สร้างบัญชี Vendor Portal (3.6)",
        "ผู้ขายต้องมีบัญชี User Portal จึงจะยืนยันใบสั่งซื้อในแอปได้ ตามขั้นตอนในบทที่ 4 "
        "กดปุ่ม User Portal (หรือ Action → สร้าง User Portal) เพื่อใส่ Email / User ID / Password โดยไม่ต้องพึ่งอีเมลระบบ",
        "info",
    )

    heading(doc, "7.2 รายการราคาจากผู้ขาย (3.6)", 2)
    step_with_pic(
        doc,
        1,
        "เข้าเมนูรายการราคาผู้ขาย เพื่อดูราคาที่ผู้ขายเสนอต่อสินค้าแต่ละรายการ สกุลเงิน และเวลานำการจัดส่ง ใช้เป็นข้อมูลอ้างอิงตอนเปรียบเทียบราคาและสร้างใบสั่งซื้อ",
        "Screenshot_proc_36_pricelist.png",
        15.8,
        "ภาพที่ 97  รายการราคาผู้ขาย (Vendor Pricelist) (3.6)",
    )
    make_table(
        doc,
        ["ข้อมูลที่ดูแลในข้อ 3.6", "ใช้ทำอะไร", "เมนู"],
        [
            ["แฟ้มผู้ขาย", "ชื่อ ที่อยู่ เลขภาษี กลุ่มคู่ค้า เกรดประเมิน บัญชี Portal", "จัดซื้อ → ผู้จำหน่าย"],
            ["รายการราคาผู้ขาย", "ราคาต่อสินค้า เวลานำส่ง สกุลเงิน", "ผลิตภัณฑ์ → รายการราคาผู้ขาย"],
            ["กลุ่มผู้จำหน่าย", "จัดกลุ่มผู้ขายเพื่อกรองและรายงาน", "แท็ก / กลุ่มคู่ค้าบนแฟ้มผู้ขาย"],
            ["เกณฑ์ประเมิน", "น้ำหนักและคะแนนเต็มของแบบประเมิน", "จัดซื้อ → ประเมินผู้จำหน่าย → เกณฑ์ประเมิน"],
        ],
        [4.0, 6.6, 5.6],
    )

    heading(doc, "ภาคผนวก ก  คำศัพท์ในระบบ", 1)
    make_table(
        doc,
        ["คำในระบบ", "พูดกับผู้ใช้", "หมายเหตุ"],
        [
            ["ใบขอซื้อ / PR", "ใบขอซื้อของหน่วยงาน", "Purchase Request ข้อ 2.2"],
            ["ส่งพัสดุ", "ปุ่มส่งใบขอซื้อจากหน่วยงานให้พัสดุรับงาน", "สถานะส่งพัสดุแล้ว"],
            ["กระบวนการ e-GP", "แฟ้มจัดซื้อหลังอนุมัติ PR", "ข้อ 3.2"],
            ["รายงานผลการพิจารณา", "ขออนุมัติสั่งซื้อ/สั่งจ้าง", "AR-EGP/ปี/เลขวิ่ง"],
            ["ประกาศผู้ชนะ", "ประกาศผลการคัดเลือก", "WIN-EGP/ปี/เลขวิ่ง"],
            ["ใบสั่งซื้อ / PO", "คำสั่งซื้อที่ยืนยันกับผู้ขาย", "ข้อ 3.3"],
            ["ทะเบียนสัญญา", "สัญญาซื้อขาย/จะซื้อจะขาย", "CTR/ปี/เลขวิ่ง"],
            ["บันทึกรับมอบงาน / WA", "ใบตรวจรับพัสดุ", "ข้อ 3.4"],
            ["ใบรับสินค้า (Auto)", "ใบรอรับที่สร้างเมื่อผู้ขายยืนยัน PO", "ส่งต่อคลังข้อ 5"],
            ["ใบประเมินผู้จำหน่าย", "แบบประเมินผู้ขาย 5 ด้าน", "ข้อ 3.5"],
            ["ขอเลขจากระบบสารบรรณ", "ปุ่มออกเลขหนังสือที่ฟอร์มเอกสารที่เกี่ยวข้องแต่ละฉบับ", "ไม่กดจากฟอร์มใบขอซื้อ"],
            ["Vendor Portal", "บัญชีผู้ขายสำหรับยืนยัน PO ในแอป", "ข้อ 3.3 และ 3.6"],
        ],
        [5.0, 5.6, 5.6],
    )

    heading(doc, "ภาคผนวก ข  ช่องทางสอบถาม", 1)
    bullets(
        doc,
        [
            "ติดขัดเรื่องสิทธิ์หรือเข้าสู่ระบบไม่ได้: งานเทคโนโลยีสารสนเทศ",
            "ติดขัดเรื่องวงเงิน แหล่งเงิน หมวดงบ และการขอตั้งงบนอกแผน: งานงบประมาณ (ดูคู่มือ UM-ERP-VPK-G01)",
            "ติดขัดเรื่องใบขอซื้อ กระบวนการ e-GP ผู้ขาย ใบสั่งซื้อ และใบตรวจรับ: งานพัสดุ",
            "ติดขัดเรื่องรับสินค้าเข้าคลัง: งานคลัง",
            "ติดขัดเรื่องใบตั้งเจ้าหนี้และการลงบัญชี: งานการเงิน",
        ],
    )

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(18)
    add_text(p_end, "— จบคู่มือการใช้งาน : กลุ่มงานพัสดุ (2.2, 3.1 – 3.6) —", 12, True, PRIMARY)

    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    build()
