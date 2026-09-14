#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""UAT workbooks for data import — one Excel file per module."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path("/opt/odoo18vpk/docs")
FONT = "Prompt"
NAVY, TEAL, RED, GREEN = "123A56", "0F7A72", "B4433A", "2C7A45"
LINE, WHITE, SOFT = "D7DEE6", "FFFFFF", "E5F5F3"
P1, P2, PASS, FAIL, SKIP = "FDECDC", "E8F1F8", "E5F5EA", "FBECEA", "F6F8FA"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


def fnt(size=11, bold=False, color="2F3A48"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def fill(color):
    return PatternFill("solid", fgColor=color)


def sc(cell, size=11, bold=False, color="2F3A48", bg=None, align="left"):
    cell.font = fnt(size, bold, color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = thin
    if bg:
        cell.fill = fill(bg)


def hdr(ws, row, values, bg=NAVY):
    for i, v in enumerate(values, 1):
        sc(ws.cell(row, i, v), 11, True, WHITE, bg, "center")
    ws.row_dimensions[row].height = 28


def page(ws, title, fit=1):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit
    ws.page_setup.leftMargin = ws.page_setup.rightMargin = 0.4
    ws.page_setup.topMargin = ws.page_setup.bottomMargin = 0.5
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = f"โรงพยาบาลวชิระภูเก็ต · {title}"
    ws.oddFooter.left.text = "เอกสารภายใน สำหรับทดสอบการนำเข้าข้อมูล"
    ws.oddFooter.right.text = "หน้า &P / &N"


def pass_dv(ws, last):
    dv = DataValidation(type="list", formula1='"ผ่าน,ไม่ผ่าน,ข้าม"', allow_blank=True)
    dv.error, dv.errorTitle = "เลือก ผ่าน / ไม่ผ่าน / ข้าม", "ผลทดสอบ"
    ws.add_data_validation(dv)
    dv.add(f"G2:G{last}")
    rng = f"G2:G{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ผ่าน"'], fill=fill(PASS), font=fnt(11, True, GREEN)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ไม่ผ่าน"'], fill=fill(FAIL), font=fnt(11, True, RED)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ข้าม"'], fill=fill(SKIP), font=fnt(11, False, "5B6878")))


def cover(wb, spec):
    ws = wb.active
    ws.title = "ปกแผน"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([4, 24, 32, 22, 18, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("B2:F2")
    ws["B2"].value, ws["B2"].font = "โรงพยาบาลวชิระภูเก็ต", fnt(14, True, TEAL)
    ws.merge_cells("B3:F3")
    ws["B3"].value, ws["B3"].font = spec["title"], fnt(20, True, NAVY)
    ws.merge_cells("B4:F4")
    ws["B4"].value, ws["B4"].font = spec["subtitle"], fnt(12, False, "5B6878")
    info = [
        ("ระบบ", spec["system"]),
        ("รหัสเอกสาร", spec["doc"]),
        ("ฐานข้อมูล", "VPK-S1 / สำเนา UAT"),
        ("รอบทดสอบ", "UAT นำเข้าข้อมูล ครั้งที่ 1"),
        ("วันที่เริ่ม", ""),
        ("วันที่จบ", ""),
        ("ผู้ประสาน", ""),
        ("ชุดข้อมูลตัวอย่าง", spec["sample"]),
    ]
    for i, (k, v) in enumerate(info):
        r = 6 + i
        sc(ws.cell(r, 2, k), 11, True, WHITE, NAVY)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        sc(ws.cell(r, 3, v))
    ws.merge_cells("B16:F16")
    ws["B16"].value = "เกณฑ์ผ่าน: เคส P1 ผ่านทั้งหมด · ไม่มี Blocker · จำนวนในระบบตรงไฟล์ต้นทาง"
    ws["B16"].font = fnt(12, True, TEAL)
    hdr(ws, 18, ["ลำดับ", "ฝ่าย", "ผู้ทดสอบ", "ลงชื่อ", "วันที่", "ผ่านภาพรวม"])
    for i, role in enumerate(spec["roles"], 1):
        r = 18 + i
        sc(ws.cell(r, 1, i), align="center", bg=SOFT)
        sc(ws.cell(r, 2, role), bold=True)
        for c in range(3, 7):
            sc(ws.cell(r, c, ""))
    end = 20 + len(spec["roles"])
    ws.merge_cells(start_row=end, start_column=2, end_row=end + 3, end_column=6)
    cell = ws.cell(end, 2, spec["howto"])
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font, cell.fill = fnt(12), fill(SOFT)
    page(ws, spec["head"])


def sequence(wb, spec):
    ws = wb.create_sheet("ลำดับนำเข้า")
    hdr(ws, 1, ["ขั้น", "ข้อมูลที่นำเข้า", "เมนู / วิซาร์ด", "ต้องมีก่อน", "จุดตรวจหลังนำเข้า"], TEAL)
    for i, w in enumerate([8, 34, 44, 38, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(spec["seq"], 2):
        for c, v in enumerate(row, 1):
            sc(ws.cell(r, c, v), align="center" if c == 1 else "left")
        ws.row_dimensions[r].height = 42
    last = 3 + len(spec["seq"])
    ws.merge_cells(start_row=last, start_column=1, end_row=last + 2, end_column=5)
    cell = ws.cell(last, 1, spec["seq_note"])
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font, cell.fill = fnt(12), fill(SOFT)
    page(ws, spec["head"])


def objects(wb, spec):
    ws = wb.create_sheet("รายการข้อมูล")
    hdr(ws, 1, ["ชีต / โมเดล", "ชื่อไทย", "คีย์จับคู่", "ฟิลด์สำคัญ", "หมายเหตุ"])
    for i, w in enumerate([24, 28, 24, 48, 46], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(spec["objs"], 2):
        for c, v in enumerate(row, 1):
            sc(ws.cell(r, c, v))
        ws.row_dimensions[r].height = 40
    page(ws, spec["head"], 0)


def cases(wb, spec):
    ws = wb.create_sheet("บันทึกผลเคส")
    hdr(ws, 1, ["รหัส", "ชุด", "P", "หัวข้อ", "ขั้นตอนโดยย่อ", "ผลที่ต้องเห็น", "ผล", "ผู้ทดสอบ", "วันที่", "หลักฐาน/หมายเหตุ"])
    for i, w in enumerate([10, 20, 6, 32, 52, 50, 12, 16, 14, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(spec["cases"], 2):
        for c, v in enumerate(row, 1):
            bg = P1 if c == 3 and v == "P1" else P2 if c == 3 and v == "P2" else None
            sc(ws.cell(r, c, v), bg=bg, align="center" if c in (1, 3, 7, 8, 9) else "left")
        ws.row_dimensions[r].height = 46
    last = 1 + len(spec["cases"])
    pass_dv(ws, max(last, 80))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{last}"
    page(ws, spec["head"], 0)


def bugs(wb, spec):
    ws = wb.create_sheet("บั๊ก")
    hdr(ws, 1, ["รหัสบั๊ก", "เคสที่พบ", "ระดับ", "สรุปปัญหา", "ขั้นตอนที่เกิด", "ผลที่คาด", "ผลที่เกิดจริง", "ผู้พบ", "วันที่", "สถานะแก้"], RED)
    for i, w in enumerate([12, 12, 14, 28, 32, 26, 26, 14, 14, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, 21):
        for c in range(1, 11):
            sc(ws.cell(r, c, f"BUG-{r-1:02d}" if c == 1 else ""), align="center" if c in (1, 3, 9, 10) else "left")
    sev = DataValidation(type="list", formula1='"Blocker,Major,Minor,Cosmetic"', allow_blank=True)
    st = DataValidation(type="list", formula1='"เปิด,กำลังแก้,แก้แล้วรอทดสอบซ้ำ,ปิด"', allow_blank=True)
    ws.add_data_validation(sev)
    ws.add_data_validation(st)
    sev.add("C2:C20")
    st.add("J2:J20")
    ws.freeze_panes = "A2"
    page(ws, spec["head"])


def summary(wb, spec):
    ws = wb.create_sheet("สรุป")
    for i, w in enumerate([6, 26, 12, 12, 12, 12, 12, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("B2:H2")
    ws["B2"].value, ws["B2"].font = spec["sum_title"], fnt(18, True, NAVY)
    hdr(ws, 4, ["", "ชุด", "ทั้งหมด", "ผ่าน", "ไม่ผ่าน", "ข้าม", "ยังไม่ทำ", "% ผ่าน P1"])
    groups = spec["groups"]
    for i, (name, pfx) in enumerate(groups):
        r = 5 + i
        sc(ws.cell(r, 1, i + 1), align="center", bg=SOFT)
        sc(ws.cell(r, 2, name), bold=True)
        sc(ws.cell(r, 3, f'=COUNTIF(\'บันทึกผลเคส\'!A:A,"{pfx}*")'), align="center")
        sc(ws.cell(r, 4, f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{pfx}*",\'บันทึกผลเคส\'!G:G,"ผ่าน")'), align="center")
        sc(ws.cell(r, 5, f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{pfx}*",\'บันทึกผลเคส\'!G:G,"ไม่ผ่าน")'), align="center")
        sc(ws.cell(r, 6, f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{pfx}*",\'บันทึกผลเคส\'!G:G,"ข้าม")'), align="center")
        sc(ws.cell(r, 7, f"=C{r}-D{r}-E{r}-F{r}"), align="center")
        sc(
            ws.cell(
                r, 8,
                f'=IFERROR(COUNTIFS(\'บันทึกผลเคส\'!A:A,"{pfx}*",\'บันทึกผลเคส\'!C:C,"P1",\'บันทึกผลเคส\'!G:G,"ผ่าน")/'
                f'COUNTIFS(\'บันทึกผลเคส\'!A:A,"{pfx}*",\'บันทึกผลเคส\'!C:C,"P1"),0)',
            ),
            align="center",
        )
        ws.cell(r, 8).number_format = "0%"
    tr, a, b = 5 + len(groups) + 1, 5, 4 + len(groups)
    sc(ws.cell(tr, 2, "รวมทุกชุด"), 12, True, WHITE, NAVY)
    for c in range(3, 8):
        L = get_column_letter(c)
        sc(ws.cell(tr, c, f"=SUM({L}{a}:{L}{b})"), 12, True, WHITE, NAVY, "center")
    sc(
        ws.cell(
            tr, 8,
            '=IFERROR(COUNTIFS(\'บันทึกผลเคส\'!C:C,"P1",\'บันทึกผลเคส\'!G:G,"ผ่าน")/COUNTIF(\'บันทึกผลเคส\'!C:C,"P1"),0)',
        ),
        12, True, WHITE, NAVY, "center",
    )
    ws.cell(tr, 8).number_format = "0%"
    box = tr + 3
    ws.merge_cells(start_row=box, start_column=2, end_row=box + 3, end_column=8)
    cell = ws.cell(
        box, 2,
        "ตัดสินใจ:  ผ่าน UAT นำเข้าและไป Go-Live ได้   /   ผ่านแบบมีเงื่อนไข   /   ต้องทดสอบซ้ำ\n"
        "จำนวนต้นทาง vs ในระบบ: .................................................................\n"
        "เงื่อนไขที่ค้าง: .................................................................\n"
        "ลงนามผู้ประสาน: ....................................     วันที่: ................",
    )
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font = fnt(12)
    page(ws, spec["head"])


def save(spec):
    wb = Workbook()
    cover(wb, spec)
    sequence(wb, spec)
    objects(wb, spec)
    cases(wb, spec)
    bugs(wb, spec)
    summary(wb, spec)
    path = OUT_DIR / spec["file"]
    wb.save(path)
    print(f"{len(spec['cases']):3d}  {path}")


BUDGET = {
    "file": "UAT-นำเข้าข้อมูล-งบประมาณ.xlsx",
    "head": "UAT นำเข้าข้อมูล ระบบงบประมาณ",
    "title": "แผน UAT นำเข้าข้อมูล ระบบงบประมาณ",
    "subtitle": "ใช้คู่กับวิซาร์ด Import Budget Master · Template มีหลายชีตในไฟล์เดียว",
    "system": "Odoo 18 · โมดูล vpk_budget",
    "doc": "UAT-IMP-BUD-001",
    "sample": "ปีงบ 2569 · หน่วยงานทดสอบ · เงินบำรุง",
    "roles": ["งานงบประมาณ", "หน่วยงาน", "IT"],
    "howto": (
        "วิธีใช้\n1) กรอกปกนี้  2) เดินตามชีต “ลำดับนำเข้า” ห้ามข้ามขั้นที่พึ่งพากัน\n"
        "3) บันทึกผลใน “บันทึกผลเคส”  4) บั๊กใส่ชีต “บั๊ก”\n"
        "เมนู: งบประมาณ → การกำหนดค่า → Import Budget Master\n"
        "บัญชีอบรมงานงบ: uat.budget   รหัสผ่าน VpkUat@2569"
    ),
    "sum_title": "สรุปผลการนำเข้าข้อมูลระบบงบประมาณ",
    "groups": [("A Template / สิทธิ์", "A"), ("B นำเข้าผ่าน", "B"), ("C กติกา", "C"), ("D ตรวจหน้าจอ", "D"), ("E อัปเดต / แผนปี", "E")],
    "seq": [
        ["1", "แหล่งเงิน (FundSource)", "งบประมาณ → Import Budget Master", "บริษัท / สกุลเงิน THB", "มีเงินบำรุง เงินงบประมาณ ตาม code"],
        ["2", "กลุ่มงบประมาณ-วัสดุ (BudgetGroup)", "ชีต BudgetGroup", "แหล่งเงินพร้อม", "มีกลุ่มยา เวชภัณฑ์ วัสดุทั่วไป"],
        ["3", "ประเภทวัสดุย่อย (MaterialSubType)", "ชีต MaterialSubType", "BudgetGroup.code ตรง", "ประเภทย่อยผูกกลุ่มถูก"],
        ["4", "ประเภทงบ (BudgetType)", "ชีต BudgetType", "section_key ครบ 4 ค่า", "วัสดุ / ครุภัณฑ์ / ก่อสร้าง / โครงการ"],
        ["5", "ประเภทและกลุ่มครุภัณฑ์", "ชีต AssetType, AssetCategory", "BudgetType ครุภัณฑ์มีแล้ว", "เลือกได้ตอนตั้งงบครุภัณฑ์"],
        ["6", "แบบฟอร์มคำของบ + บรรทัดประเภท", "ชีต FormType, FormTypeLine", "BudgetType ครบ", "เมนูเหลือ 4 แบบฟอร์ม"],
        ["7", "หน่วยงานวิเคราะห์ / ศูนย์ต้นทุน", "บัญชีวิเคราะห์ (นำเข้าแยกหรือสร้างในระบบ)", "แผน Analytic หน่วยงาน", "หน่วยงานทดสอบมี code"],
        ["8", "หมวดงบ + แผนประจำปี", "งบประมาณประจำปี", "แหล่งเงิน หน่วยงาน ประเภทงบ", "แผนปี 2569 ช่วงวันที่ถูกต้อง และ validate แล้ว"],
    ],
    "seq_note": (
        "ลำดับชีตในไฟล์: FundSource → BudgetGroup → BudgetType → AssetType → AssetCategory → "
        "MaterialSubType → FormType → FormTypeLine   จับคู่ด้วย code "
        "(MaterialSubType จับคู่ร่วมกลุ่มงบ, FormTypeLine จับคู่ร่วมแบบฟอร์ม)"
    ),
    "objs": [
        ["FundSource", "แหล่งเงิน", "code", "code, name, sequence, active", "เช่น HOSP_MAINT = เงินบำรุง"],
        ["BudgetGroup", "กลุ่มงบประมาณ-วัสดุ", "code", "code, name, active", "ใช้เฉพาะงบวัสดุ"],
        ["MaterialSubType", "ประเภทวัสดุย่อย", "code + กลุ่มงบ", "code, name, budget_group_code", "ต้องชี้ BudgetGroup ที่มี"],
        ["BudgetType", "ประเภทงบประมาณ", "code", "code, name, section_key", "material / asset / construction / project"],
        ["AssetType / AssetCategory", "ประเภทและกลุ่มครุภัณฑ์", "code", "code, name", "ใช้กับแบบฟอร์มครุภัณฑ์"],
        ["FormType", "แบบฟอร์มคำของบ", "code", "code, name, sequence", "ต้องเหลือ 4 ใบ"],
        ["FormTypeLine", "ประเภทงบในแบบฟอร์ม", "form + budget type", "form_type_code, budget_type_code", "กำหนดแท็บในใบคำของบ"],
        ["account.analytic.account", "หน่วยงาน / ศูนย์ต้นทุน", "code", "code, name, plan", "ต้องตรงตอนสร้างคำของบและ PR"],
        ["crossovered.budget", "งบประมาณประจำปี", "ชื่อแผน + ปี", "name, date_from, date_to, state", "validate แล้วจึงใช้งบได้"],
        ["account.budget.post", "หมวดงบประมาณ", "ชื่อ", "name, account_ids", "ใช้จับคู่ตอนเช็คงบบน PR"],
    ],
    "cases": [
        ["A01", "A Template/สิทธิ์", "P1", "ดาวน์โหลด Template", "login uat.budget เปิด Import Budget Master กดดาวน์โหลด Template", "ได้ไฟล์ xlsx ครบชีต FundSource ถึง FormTypeLine"],
        ["A02", "A Template/สิทธิ์", "P1", "หัวคอลัมน์จับคู่ได้", "เปิด Template ดูแถวหัว", "มีทั้งชื่อไทยและอังกฤษ เช่น code/รหัส name/ชื่อ"],
        ["A03", "A Template/สิทธิ์", "P1", "หน่วยงานเข้าวิซาร์ดไม่ได้", "login uat.dept พยายามเปิด Import Budget Master", "เข้าไม่ได้ ตามสิทธิ์งานงบ"],
        ["A04", "A Template/สิทธิ์", "P2", "ตัวอย่างใน Template อ่านรู้เรื่อง", "ดูแถวตัวอย่างเงินบำรุง / กลุ่มยา", "มีตัวอย่างครบทุกชีต"],
        ["B01", "B นำเข้าผ่าน", "P1", "นำเข้าแหล่งเงิน", "กรอก FundSource อย่างน้อย 2 แหล่ง แล้ว Import", "สร้างสำเร็จ จำนวนตรงไฟล์ error = 0"],
        ["B02", "B นำเข้าผ่าน", "P1", "นำเข้ากลุ่มวัสดุและประเภทย่อย", "BudgetGroup + MaterialSubType ที่ชี้กลุ่มถูก", "ประเภทย่อยอยู่กลุ่มที่ถูกต้อง"],
        ["B03", "B นำเข้าผ่าน", "P1", "นำเข้าประเภทงบ 4 กลุ่ม", "BudgetType พร้อม section_key ครบ", "เลือกได้ครบตอนสร้างคำของบ"],
        ["B04", "B นำเข้าผ่าน", "P1", "นำเข้าครุภัณฑ์ master", "AssetType + AssetCategory", "ปรากฏในแบบฟอร์มครุภัณฑ์"],
        ["B05", "B นำเข้าผ่าน", "P1", "นำเข้าแบบฟอร์ม 4 ใบ", "FormType 4 ใบ และ FormTypeLine", "เมนูคำของบเหลือ 4 แบบฟอร์มตามชื่อที่กำหนด"],
        ["B06", "B นำเข้าผ่าน", "P1", "สรุปจำนวนหลัง Import", "ดู created / updated / skipped / error", "error = 0 และจำนวนตรงแถวที่มีข้อมูล"],
        ["C01", "C กติกา", "P1", "รหัสแหล่งเงินซ้ำ", "ใส่ FundSource code ซ้ำ แล้ว Import โหมดอัปเดต", "ไม่สร้างสองรายการ แถวหลังอัปเดตหรือข้าม"],
        ["C02", "C กติกา", "P1", "ประเภทย่อยชี้กลุ่มที่ไม่มี", "budget_group_code ไม่มีในระบบ", "error แถวนั้น แถวอื่นยังเข้าได้"],
        ["C03", "C กติกา", "P1", "section_key ผิด", "ใส่ค่าที่ไม่ใช่ 4 กลุ่ม", "ระบบไม่รับหรือไม่ผูกแบบฟอร์มได้"],
        ["C04", "C กติกา", "P1", "FormTypeLine ชี้แบบฟอร์มไม่มี", "form_type_code ไม่มีใน FormType", "error ชัดเจน ไม่สร้างแท็บค้าง"],
        ["C05", "C กติกา", "P2", "แถวไม่มี code", "มีแถวที่ code ว่าง", "ข้ามแถวว่าง ไม่ทำให้ทั้งไฟล์ล้ม"],
        ["D01", "D ตรวจหน้าจอ", "P1", "เปิดเมนูแหล่งเงิน", "งบประมาณ → การกำหนดค่า → แหล่งเงิน", "รหัสและชื่อตรงไฟล์"],
        ["D02", "D ตรวจหน้าจอ", "P1", "เปิดแบบฟอร์มคำของบ", "หน่วยงานสร้างคำของบ", "เลือกได้แค่ 4 ใบ แท็บรายการตรงประเภท"],
        ["D03", "D ตรวจหน้าจอ", "P1", "สร้างคำของบวัสดุหลังนำเข้า", "เลือกหน่วยงาน แหล่งเงิน กลุ่มวัสดุ", "ค่าที่นำเข้าเลือกได้ ไม่ต้องสร้างมือซ้ำ"],
        ["D04", "D ตรวจหน้าจอ", "P2", "รายการที่ปิดใช้งานไม่โชว์", "นำเข้าแถว active=0", "ไม่ขึ้นให้เลือกตอนสร้างคำขอ"],
        ["E01", "E อัปเดต/แผนปี", "P1", "Import ซ้ำโหมดอัปเดต", "แก้ชื่อแหล่งเงิน แล้ว Import ตาม code", "ชื่อเปลี่ยน จำนวนรายการไม่เพิ่ม"],
        ["E02", "E อัปเดต/แผนปี", "P1", "ผูกแผนประจำปีหลังมี master", "สร้างแผนปี 2569 ใส่หมวดและหน่วยงาน แล้ว validate", "PR เช็คงบเจอวงเงิน"],
        ["E03", "E อัปเดต/แผนปี", "P1", "ห้ามใช้งบก่อนมีแผนที่อนุมัติ", "ลองเช็คงบบน PR ก่อนมีบรรทัดงบ", "ไม่เจอวงเงิน หรืองบไม่พอ"],
        ["E04", "E อัปเดต/แผนปี", "P2", "เก็บ log เป็นหลักฐาน", "คัดลอก log จากวิซาร์ด", "มีวันที่ จำนวนสร้าง/อัปเดต และชื่อไฟล์ต้นทาง"],
    ],
}

PURCHASE = {
    "file": "UAT-นำเข้าข้อมูล-จัดซื้อ.xlsx",
    "head": "UAT นำเข้าข้อมูล ระบบจัดซื้อ",
    "title": "แผน UAT นำเข้าข้อมูล ระบบจัดซื้อ",
    "subtitle": "Vendor / PR / PO จาก Template นำเข้า Vendor / PR / PO",
    "system": "Odoo 18 · Purchase Request / Purchase Order",
    "doc": "UAT-IMP-PUR-001",
    "sample": "Vendor V001–V003 · PR-2569-0001 · PO-2569-0001",
    "roles": ["พัสดุ / จัดซื้อ", "หน่วยงาน", "IT"],
    "howto": (
        "วิธีใช้\n1) ต้องมีสินค้า หน่วยนับ คลัง จากชุดคลังก่อน  2) นำเข้า Vendor ก่อน PR/PO\n"
        "3) PR/PO นำเข้าสถานะ draft แล้วยืนยันในระบบ\n"
        "เมนู: จัดซื้อ → การกำหนดค่า → Template นำเข้า Vendor / PR / PO\n"
        "บัญชีอบรม: uat.procurement   รหัสผ่าน VpkUat@2569"
    ),
    "sum_title": "สรุปผลการนำเข้าข้อมูลระบบจัดซื้อ",
    "groups": [("A Template / สิทธิ์", "A"), ("B ผู้จำหน่าย", "B"), ("C ใบขอซื้อ PR", "C"), ("D ใบสั่งซื้อ PO", "D"), ("E กติกา", "E")],
    "seq": [
        ["1", "ผู้จำหน่าย (Vendor)", "Contacts / จัดซื้อ → ผู้จำหน่าย (Import)", "—", "ref ไม่ซ้ำ supplier_rank = 1"],
        ["2", "บัญชีธนาคารผู้ขาย", "ชีต VendorBank", "Vendor.ref ตรง", "เห็นบัญชีบนการ์ดผู้ขาย"],
        ["3", "ราคาผู้ขาย / MOQ", "ชีต VendorPricelist", "มีสินค้า + ผู้ขาย", "ราคา MOQ lead time ตรงไฟล์"],
        ["4", "ใบขอซื้อ Header", "Import purchase.request สถานะ draft", "หน่วยงาน คลัง มิติงบ สินค้า", "เลข PR ตามไฟล์หรือตาม sequence"],
        ["5", "ใบขอซื้อ Lines", "Import purchase.request.line", "pr_name ตรง Header", "จำนวน หน่วย ราคารวมตรง"],
        ["6", "ใบสั่งซื้อ Header", "Import purchase.order draft", "Vendor.ref และ PR ถ้าอ้าง", "PO ผูกผู้ขายถูก"],
        ["7", "ใบสั่งซื้อ Lines", "Import purchase.order.line", "po_name + product_code", "จำนวน ราคา ภาษีตรง"],
        ["8", "ยืนยันในระบบ (ห้าม import เป็น purchase/done)", "Confirm ที่หน้าจอ", "เช็คงบ/สต็อกพร้อม", "สร้างใบรับของอัตโนมัติ"],
    ],
    "seq_note": "แถวแรกของ Template เป็นไทย แถวถัดไปเป็นชื่อฟิลด์อังกฤษ  PR/PO แนะนำ draft เสมอ เพราะ Confirm จะจองงบและสร้างใบรับสินค้า  partner_ref / pr_name / po_name ต้องตรงทุกชีต",
    "objs": [
        ["Vendor", "ผู้จำหน่าย", "ref", "ref, name, vat, company_type", "นิติบุคคลใช้ชื่อบริษัท"],
        ["VendorBank", "บัญชีธนาคารผู้ขาย", "partner_ref + เลขบัญชี", "partner_ref, acc_number, bank_name", "จับคู่ partner จาก ref"],
        ["VendorPricelist", "ราคา / MOQ", "partner + product", "partner_ref, product_code, price, min_qty, delay", "โมเดล product.supplierinfo"],
        ["PurchaseRequest", "ใบขอซื้อ Header", "name", "requested_by, warehouse_code, budget_*, state", "state = draft"],
        ["PurchaseRequestLine", "ใบขอซื้อ Lines", "pr_name", "product_code, qty, uom, estimated_cost", "estimated_cost = ยอดรวมบรรทัด"],
        ["PurchaseOrder", "ใบสั่งซื้อ Header", "name", "partner_ref, pr_name, warehouse_code, state", "state = draft"],
        ["PurchaseOrderLine", "ใบสั่งซื้อ Lines", "po_name", "product_code, qty, price, taxes", "หน่วยต้องหมวดเดียวกับสินค้า"],
    ],
    "cases": [
        ["A01", "A Template/สิทธิ์", "P1", "ดาวน์โหลด Template จัดซื้อ", "เปิด Template นำเข้า Vendor / PR / PO", "ได้ชีต Vendor, VendorBank, VendorPricelist, PurchaseRequest, PurchaseRequestLine, PurchaseOrder, PurchaseOrderLine"],
        ["A02", "A Template/สิทธิ์", "P1", "หัวคอลัมน์สองแถว", "ดูแถว 1 ไทย และแถว 2 ชื่อฟิลด์อังกฤษ", "ระบบ Import จับแถวฟิลด์อังกฤษได้"],
        ["A03", "A Template/สิทธิ์", "P2", "หน่วยงานนำเข้า Vendor ทั้งไฟล์ไม่ได้", "login uat.dept พยายาม Import ผู้จำหน่ายจำนวนมาก", "ถูกจำกัดสิทธิ์ หรือต้องให้พัสดุเป็นคนนำเข้า"],
        ["B01", "B ผู้จำหน่าย", "P1", "นำเข้า Vendor 3 ราย", "นิติบุคคลจด VAT และบุคคลไม่จด VAT", "มี 3 รายการ เลขผู้เสียภาษีถูกต้อง supplier_rank=1"],
        ["B02", "B ผู้จำหน่าย", "P1", "ค้นด้วยรหัส ref", "ใช้ V001 V002 V003", "เปิดจากจัดซื้อ → ผู้จำหน่าย แล้วค้นเจอ"],
        ["B03", "B ผู้จำหน่าย", "P1", "นำเข้าบัญชีธนาคารผู้ขาย", "VendorBank อ้าง partner_ref", "เห็นเลขบัญชีบนฟอร์มผู้ขาย"],
        ["B04", "B ผู้จำหน่าย", "P1", "นำเข้าราคาผู้ขาย", "VendorPricelist ชี้สินค้าที่มีในคลัง", "แท็บ Purchase ของสินค้ามีผู้ขาย ราคา MOQ lead time"],
        ["B05", "B ผู้จำหน่าย", "P2", "VAT ผิดรูปแบบ", "ใส่ vat สั้นเกิน", "ระบบเตือนหรือไม่รับแถวนั้น"],
        ["C01", "C ใบขอซื้อ", "P1", "นำเข้า PR Header draft", "PR หนึ่งใบ หน่วยงาน คลัง Receipt มิติงบครบ", "ได้ใบสถานะร่าง"],
        ["C02", "C ใบขอซื้อ", "P1", "นำเข้า PR Lines", "บรรทัดสินค้า qty หน่วย estimated_cost", "ยอดหัวใบเท่าผลรวมบรรทัด"],
        ["C03", "C ใบขอซื้อ", "P1", "บรรทัดชี้ PR ที่ไม่มี", "pr_name ผิด", "error แถวนั้น ไม่สร้างบรรทัดกำพร้า"],
        ["C04", "C ใบขอซื้อ", "P1", "สินค้าไม่มีในระบบ", "product_code ที่ยังไม่ได้นำเข้าคลัง", "error ชัดเจน"],
        ["C05", "C ใบขอซื้อ", "P2", "PR เร่งด่วนต้องมีเหตุผล", "urgency=urgent แต่ไม่ใส่เหตุผล", "ระบบกันหรือเตือนตามกติกาโมดูล"],
        ["D01", "D ใบสั่งซื้อ", "P1", "นำเข้า PO Header", "partner_ref และ pr_name ตรง", "PO ร่างผูกผู้ขายและอ้าง PR"],
        ["D02", "D ใบสั่งซื้อ", "P1", "นำเข้า PO Lines", "สินค้า จำนวน ราคา หน่วย", "ยอด PO ตรงไฟล์"],
        ["D03", "D ใบสั่งซื้อ", "P1", "Confirm บนหน้าจอหลังนำเข้า", "อย่า import state=purchase ใช้ปุ่มยืนยัน", "ได้ใบรับของ และงบเป็นผูกพันถ้าเปิดเช็คงบ"],
        ["D04", "D ใบสั่งซื้อ", "P2", "PO ไม่มีผู้ขาย", "partner_ref ว่างหรือไม่มีในระบบ", "นำเข้าไม่ได้"],
        ["E01", "E กติกา", "P1", "ห้ามนำเข้า PO เป็น done", "ลองใส่ state=done", "ไม่สร้างสต็อกข้ามขั้นตอน หรือมีคำเตือนใน Template"],
        ["E02", "E กติกา", "P1", "หน่วยนับคนละหมวดกับสินค้า", "uom คนละ category", "error"],
        ["E03", "E กติกา", "P1", "นำเข้า PR แล้วยังไม่จองงบ", "ตรวจคงเหลือหลัง import draft", "ยังไม่กันเงิน จองเมื่อส่งอนุมัติ"],
        ["E04", "E กติกา", "P2", "นับจำนวนเทียบไฟล์", "นับ Vendor/PR/PO ในระบบ", "ตรงจำนวนแถวไม่ซ้ำในไฟล์ บันทึกในสรุป"],
    ],
}

WAREHOUSE = {
    "file": "UAT-นำเข้าข้อมูล-คลังสินค้า.xlsx",
    "head": "UAT นำเข้าข้อมูล ระบบคลังสินค้า",
    "title": "แผน UAT นำเข้าข้อมูล ระบบคลังสินค้า",
    "subtitle": "หน่วยนับ หมวดสินค้า สินค้า คลัง ตำแหน่ง และยกยอดสต็อก",
    "system": "Odoo 18 · Inventory",
    "doc": "UAT-IMP-STK-001",
    "sample": "คลังยา PHAR · สินค้า 3012815 · ตำแหน่ง A01",
    "roles": ["งานคลัง", "พัสดุ", "IT"],
    "howto": (
        "วิธีใช้\n1) นำเข้าหน่วยนับก่อนหมวดสินค้า และหมวดก่อนสินค้า  2) สร้างคลังก่อนตำแหน่ง\n"
        "3) ยกยอดสต็อกหลังมีสินค้าและตำแหน่ง  4) สินค้าคุมสต็อกต้องเป็น Storable\n"
        "เมนู: สินค้าคงคลัง → การกำหนดค่า → Template Master Data / Import Product Master"
    ),
    "sum_title": "สรุปผลการนำเข้าข้อมูลระบบคลังสินค้า",
    "groups": [("A Template / สิทธิ์", "A"), ("B หน่วยนับและหมวด", "B"), ("C สินค้า", "C"), ("D คลังและตำแหน่ง", "D"), ("E ยกยอดสต็อก", "E")],
    "seq": [
        ["1", "หมวดหน่วยนับ (UoMCategory)", "Template Master Data", "—", "มีหมวดหน่วย น้ำหนัก กล่องแพ็ค"],
        ["2", "หน่วยนับ (UoM)", "ชีต UoM", "แต่ละหมวดมี reference เพียง 1", "หน่วย กล่อง-100 ใช้ได้"],
        ["3", "หมวดหมู่สินค้า", "ชีต ProductCategory", "parent_path เริ่ม All", "path เต็มถูกต้อง"],
        ["4", "สินค้า (Product)", "Inventory → Import Product Master", "หมวด + หน่วยมีแล้ว", "default_code ไม่ซ้ำ tracking ถูก"],
        ["5", "คลังสินค้า (Warehouse)", "ชีต Warehouse", "code สูงสุด 5 ตัว", "ระบบสร้าง Stock location ให้"],
        ["6", "ตำแหน่งจัดเก็บ (Location)", "ชีต Location", "parent_path ชี้คลังที่สร้างแล้ว", "usage=internal สำหรับเก็บของ"],
        ["7", "ยกยอดสต็อก", "Inventory Adjustment / Import stock.quant", "สินค้า + location พร้อม", "On Hand ตรงไฟล์หลัง Validate"],
        ["8", "สินค้า Lot / วันหมดอายุ", "ตรวจ tracking=lot", "ต้องมี lot เมื่อรับของหรือยกยอด", "บังคับ lot ตามที่ตั้ง"],
    ],
    "seq_note": "code คลังสูงสุด 5 ตัวอักษร  สินค้า type=consu และ is_storable=1 จึงเข้าคลังได้  ห้ามยกยอดเข้าคลังที่ยังไม่มี location ภายใน",
    "objs": [
        ["UoMCategory", "หมวดหน่วยนับ", "name", "name", "สร้างก่อน UoM"],
        ["UoM", "หน่วยนับ", "name + category", "uom_type, factor, rounding", "reference ได้หมวดละ 1"],
        ["ProductCategory", "หมวดหมู่สินค้า", "path", "name, parent_path, active", "parent ว่างจะอยู่ใต้ All"],
        ["Product", "สินค้า", "default_code", "name, categ_path, uom, tracking, is_storable", "egp_purchase_name ใช้ตอนซื้อ e-GP"],
        ["Warehouse", "คลังสินค้า", "code", "code, name, reception_steps, delivery_steps", "สร้าง location หลักอัตโนมัติ"],
        ["Location", "ตำแหน่งจัดเก็บ", "complete_name", "name, parent_path, usage, warehouse_code", "internal = เก็บของ"],
        ["stock.quant / inventory", "ยกยอดคงเหลือ", "product + location + lot", "quantity, lot, expiration", "Validate แล้ว On Hand เพิ่ม"],
    ],
    "cases": [
        ["A01", "A Template/สิทธิ์", "P1", "ดาวน์โหลด Template คลัง/สินค้า", "เปิด Template Master Data เลือก UoM Category Product Warehouse", "ได้ชีตครบตามที่เลือก"],
        ["A02", "A Template/สิทธิ์", "P1", "เมนู Import Product Master", "เข้า Inventory หาเมนูนำเข้าสินค้า", "อัปโหลดไฟล์สินค้าได้"],
        ["A03", "A Template/สิทธิ์", "P2", "ผู้ไม่มีสิทธิ์คลัง Import ไม่ได้", "login หน่วยงาน", "ไม่มีเมนูกำหนดค่าคลัง"],
        ["B01", "B หน่วยนับ/หมวด", "P1", "นำเข้าหมวดหน่วยนับ", "สร้างหมวด หน่วย น้ำหนัก กล่องแพ็ค", "มีครบ"],
        ["B02", "B หน่วยนับ/หมวด", "P1", "นำเข้า UoM พร้อม reference", "แต่ละหมวดมี reference เดียว", "แปลงหน่วยได้"],
        ["B03", "B หน่วยนับ/หมวด", "P1", "สอง reference ในหมวดเดียว", "ใส่ uom_type=reference สองแถวหมวดเดียวกัน", "ระบบกัน"],
        ["B04", "B หน่วยนับ/หมวด", "P1", "นำเข้าหมวดสินค้าแบบต้นไม้", "กลุ่มพัสดุ แล้วหมวดย่อยเวชภัณฑ์", "path เต็มถูกต้อง"],
        ["C01", "C สินค้า", "P1", "นำเข้าสินค้าคุมสต็อก", "default_code หมวด หน่วย is_storable=1", "เป็น Storable เข้าคลังได้"],
        ["C02", "C สินค้า", "P1", "สินค้าติดตาม lot + วันหมดอายุ", "tracking=lot use_expiration_date=1", "รับของ/ยกยอดต้องระบุ lot"],
        ["C03", "C สินค้า", "P1", "รหัสสินค้าซ้ำ", "default_code ซ้ำกับของที่มี", "อัปเดตรายการเดิม ไม่สร้างซ้ำ"],
        ["C04", "C สินค้า", "P1", "หมวดหรือหน่วยยังไม่มี", "categ_path / uom ที่ยังไม่ได้นำเข้า", "error แถวนั้น"],
        ["C05", "C สินค้า", "P2", "สินค้าบริการไม่เข้าคลัง", "type=service หรือ is_storable=0", "ไม่มี On Hand"],
        ["D01", "D คลัง/ตำแหน่ง", "P1", "นำเข้าคลังยาและคลังกลาง", "code PHAR / WH ความยาวไม่เกิน 5", "ได้คลังและ Stock location"],
        ["D02", "D คลัง/ตำแหน่ง", "P1", "code คลังเกิน 5 ตัว", "ใส่โค้ดยาวเกิน", "ระบบไม่รับ"],
        ["D03", "D คลัง/ตำแหน่ง", "P1", "นำเข้าตำแหน่งใต้คลัง", "parent_path เช่น WH/Stock", "ตำแหน่งอยู่ถูกคลัง usage=internal"],
        ["D04", "D คลัง/ตำแหน่ง", "P2", "parent_path ผิด", "ชี้ path ที่ไม่มี", "error"],
        ["E01", "E ยกยอด", "P1", "ยกยอดสินค้าไม่มี lot", "ปรับสต็อกเข้าตำแหน่ง แล้ว Validate", "On Hand ตรงจำนวน"],
        ["E02", "E ยกยอด", "P1", "ยกยอดสินค้า lot", "ระบุ lot และวันหมดอายุ", "สต็อกแยก lot ได้"],
        ["E03", "E ยกยอด", "P1", "ยกยอดเข้าตำแหน่งที่ไม่มี", "location ผิด", "นำเข้าไม่ได้"],
        ["E04", "E ยกยอด", "P1", "นับรายการเทียบไฟล์", "นับสินค้าและปริมาณรวม", "ตรงไฟล์ต้นทาง บันทึกในสรุป"],
        ["E05", "E ยกยอด", "P2", "หลังยกยอดจ่ายของได้", "ทำใบจ่าย/โอนจำนวนน้อย", "สต็อกลดตามจำนวน"],
    ],
}

FINANCE = {
    "file": "UAT-นำเข้าข้อมูล-เจ้าหนี้บัญชีการเงิน.xlsx",
    "head": "UAT นำเข้าข้อมูล เจ้าหนี้ บัญชี การเงิน",
    "title": "แผน UAT นำเข้าข้อมูล เจ้าหนี้ บัญชี การเงิน",
    "subtitle": "ผังบัญชี ตั้งค่าบัญชี ตรวจรับ ใบแจ้งหนี้ จ่ายชำระ ยกยอด GL เช็ค WHT",
    "system": "Odoo 18 · Accounting / AP / Treasury",
    "doc": "UAT-IMP-FIN-001",
    "sample": "ผังบัญชีโรงพยาบาล · BILL-2569-0001 · ยกยอดตามวันที่ที่ตกลง",
    "roles": ["การเงิน", "บัญชี", "พัสดุ (ตรวจรับ)", "IT"],
    "howto": (
        "วิธีใช้\n1) นำเข้าผังบัญชีและสมุดรายวันก่อนเอกสาร  2) ต้องมี Vendor/PO จากชุดจัดซื้อถ้าจะโยงบิล\n"
        "3) ใบแจ้งหนี้/จ่ายชำระ/JE นำเข้า draft แล้ว Post ในระบบ\n"
        "เมนู: บัญชี → การกำหนดค่า → Template นำเข้าเจ้าหนี้ / การเงิน / บัญชี"
    ),
    "sum_title": "สรุปผลการนำเข้าข้อมูลเจ้าหนี้ บัญชี การเงิน",
    "groups": [("A Template / สิทธิ์", "A"), ("B ผังบัญชีและตั้งค่า", "B"), ("C ตรวจรับและเจ้าหนี้", "C"), ("D จ่ายชำระ / เช็ค / WHT", "D"), ("E GL / ยกยอด", "E")],
    "seq": [
        ["1", "ผังบัญชี (ChartOfAccounts)", "บัญชี → Template นำเข้า…", "บริษัท ภาษีไทย", "รหัสบัญชีไม่ซ้ำ ประเภทบัญชีถูก"],
        ["2", "สมุดรายวัน / ภาษี / WHT / เงื่อนไขชำระ", "ชีต Journal Tax WithholdingTax PaymentTerm", "ผังบัญชีมีแล้ว", "มี BILL CASH BNK1 MISC และ VAT 7%"],
        ["3", "บัญชีธนาคารบริษัท", "ชีต CompanyBank", "สมุดรายวัน type=bank", "ผูกเลขบัญชีกับสมุดธนาคาร"],
        ["4", "ใบตรวจรับงาน (WA)", "Header + Lines ต้องมี PO ก่อน", "PO ยืนยันแล้ว", "WA ร่างโยง PO"],
        ["5", "ใบแจ้งหนี้ผู้ขาย + บรรทัด", "VendorBill draft", "Vendor + บัญชีค่าใช้จ่าย + ภาษี", "ยอดพร้อม Post"],
        ["6", "ใบลดหนี้ / ใบกำกับภาษีซื้อ", "ถ้ามีในชุดข้อมูล", "บิลต้นทาง", "อ้างอิงบิลเดิมได้"],
        ["7", "จ่ายชำระเจ้าหนี้ / Statement", "หลัง Post บิล", "สมุดธนาคาร/เงินสด", "ลดยอดเจ้าหนี้เมื่อ Post/Reconcile"],
        ["8", "สมุดเช็ค / WHT / JE / ยกยอด", "Checkbook WHTCert JournalEntry OpeningBalance", "บัญชีและสมุดรายวันพร้อม", "ยกยอดวันที่ตรงงวด เดบิต=เครดิต"],
    ],
    "seq_note": "Post บิลในระบบเท่านั้น อย่า import เป็น posted ถ้ายังไม่ตรวจภาษีและงบประมาณ  ยกยอด OpeningBalance เดบิตต้องเท่าเครดิตทุกชุด  ref ใบแจ้งหนี้ห้ามซ้ำคู่ผู้ขาย+เลขที่",
    "objs": [
        ["ChartOfAccounts", "ผังบัญชี", "code", "code, name, account_type, reconcile, wht_account", "เจ้าหนี้ = liability_payable reconcile=1"],
        ["Journal", "สมุดรายวัน", "code", "code, name, type, default_account_code", "code สูงสุด 5 ตัว"],
        ["Tax / WithholdingTax", "VAT / ภาษีหัก ณ ที่จ่าย", "name", "type_tax_use, amount, account_code", "VAT ซื้อ 7% และ WHT"],
        ["PaymentTerm", "เงื่อนไขชำระ", "name", "name + บรรทัดวันครบกำหนด", "Immediate / 30 Days"],
        ["WorkAcceptance", "ใบตรวจรับ", "name", "po_name, date, committee", "ต้องมี PO"],
        ["VendorBill", "ใบแจ้งหนี้ผู้ขาย", "name / ref", "partner_ref, invoice_date, journal_code, po_name", "move_type = in_invoice"],
        ["VendorBillLine", "รายการบิล", "bill_name", "account_code, qty, price, taxes, analytic", "บัญชีค่าใช้จ่ายถูกหมวด"],
        ["VendorPayment", "จ่ายชำระ", "name", "partner, journal, amount, bill_name", "draft แล้ว Post"],
        ["BankStatement", "Statement ธนาคาร", "name", "journal_code, date, amount +/−", "กระทบยอดหลังนำเข้า"],
        ["JournalEntry / OpeningBalance", "รายการทั่วไป / ยกยอด", "name + date", "journal, debit, credit, account_code", "เดบิต = เครดิต"],
        ["WHTCert", "หนังสือรับรอง หัก ณ ที่จ่าย", "number", "partner, date, lines", "หลังมีบิล/จ่าย"],
        ["Checkbook", "สมุดเช็ค", "name", "journal_code, เลขที่เริ่ม–จบ, lines", "สถานะสมุดถูกต้อง"],
    ],
    "cases": [
        ["A01", "A Template/สิทธิ์", "P1", "ดาวน์โหลด Template บัญชี/เจ้าหนี้/การเงิน", "เลือก CoA ตั้งค่า บิล จ่าย ยกยอด", "ได้ชีตผังบัญชี Journal Tax Bill Payment OpeningBalance อย่างน้อย"],
        ["A02", "A Template/สิทธิ์", "P1", "สิทธิ์หน่วยงานนำเข้าบิลไม่ได้", "login uat.dept", "ไม่มีเมนูบัญชีกำหนดค่า"],
        ["A03", "A Template/สิทธิ์", "P2", "ลำดับในหน้าวิซาร์ดอ่านรู้เรื่อง", "อ่านลำดับเตรียมข้อมูลบนฟอร์ม", "มีข้อ CoA ก่อนบิล และ draft ก่อน Post"],
        ["B01", "B ผังบัญชี/ตั้งค่า", "P1", "นำเข้าผังบัญชีชุดโรงพยาบาล", "รหัสไม่ซ้ำ ประเภทบัญชีครบ", "เปิดผังบัญชีแล้วค้นรหัสเจอ จำนวนตรงไฟล์"],
        ["B02", "B ผังบัญชี/ตั้งค่า", "P1", "บัญชีเจ้าหนี้ reconcile", "210101 liability_payable reconcile=1", "ใช้ตั้งเจ้าหนี้ได้"],
        ["B03", "B ผังบัญชี/ตั้งค่า", "P1", "นำเข้าสมุดรายวัน", "BILL purchase, BNK1 bank, CASH cash, MISC general", "สมุดครบและผูกบัญชีเงินสด/ธนาคาร"],
        ["B04", "B ผังบัญชี/ตั้งค่า", "P1", "นำเข้า VAT 7% และ WHT", "Tax type purchase/sale และบัญชี WHT", "เลือกได้ตอนลงบิล"],
        ["B05", "B ผังบัญชี/ตั้งค่า", "P1", "รหัสบัญชีซ้ำ", "code ซ้ำ", "อัปเดตหรือ error ไม่สร้างสองบัญชี"],
        ["B06", "B ผังบัญชี/ตั้งค่า", "P2", "สมุดธนาคารไม่มีบัญชีเงินฝาก", "type=bank แต่ default_account_code ว่าง", "ระบบกัน"],
        ["C01", "C ตรวจรับ/เจ้าหนี้", "P1", "นำเข้า WA จาก PO", "มี PO ยืนยันแล้ว นำเข้า Header+Lines", "ใบตรวจรับร่างอ้าง PO จำนวนไม่เกินที่สั่ง"],
        ["C02", "C ตรวจรับ/เจ้าหนี้", "P1", "นำเข้าใบแจ้งหนี้ draft", "VendorBill + Lines ภาษี บัญชีค่าใช้จ่าย", "ยอดบิลตรงไฟล์ สถานะร่าง"],
        ["C03", "C ตรวจรับ/เจ้าหนี้", "P1", "Post บิลบนหน้าจอ", "ตรวจภาษีแล้วกด Post", "ตั้งเจ้าหนี้ และกระทบงบถ้าผูกงบ"],
        ["C04", "C ตรวจรับ/เจ้าหนี้", "P1", "ref ใบแจ้งหนี้ซ้ำผู้ขาย", "ผู้ขายเดียวกัน เลขที่ใบแจ้งหนี้ซ้ำ", "ระบบเตือนซ้ำ"],
        ["C05", "C ตรวจรับ/เจ้าหนี้", "P1", "บิลไม่มีผู้ขายหรือบัญชี", "partner_ref หรือ account_code ว่าง", "error"],
        ["C06", "C ตรวจรับ/เจ้าหนี้", "P2", "ใบลดหนี้อ้างบิลเดิม", "นำเข้า VendorCreditNote ถ้ามีในชุด", "ลดยอดเจ้าหนี้เมื่อ Post"],
        ["D01", "D จ่ายชำระ", "P1", "นำเข้าจ่ายชำระ draft แล้ว Post", "VendorPayment ชี้บิลที่โพสต์แล้ว", "ยอดเจ้าหนี้ลดตามที่จ่าย"],
        ["D02", "D จ่ายชำระ", "P1", "นำเข้า Bank Statement", "ยอด + เข้า − ออก ตามสมุดธนาคาร", "กระทบยอดรายการที่จับคู่ได้"],
        ["D03", "D จ่ายชำระ", "P1", "นำเข้าสมุดเช็ค", "Checkbook + เลขที่เช็ค", "ช่วงเลขที่ต่อเนื่อง ไม่ซ้ำ"],
        ["D04", "D จ่ายชำระ", "P2", "หนังสือรับรอง WHT", "WHTCert หลังมีบิลหักภาษี", "ยอดภาษีตรงบรรทัดบิล"],
        ["E01", "E GL/ยกยอด", "P1", "นำเข้า Journal Entry เดบิต=เครดิต", "สองบรรทัดสมดุล แล้ว Post", "ยอดบัญชีขยับตามไฟล์"],
        ["E02", "E GL/ยกยอด", "P1", "JE เดบิตไม่เท่าเครดิต", "ตั้งใจทำยอดไม่สมดุล", "Post ไม่ได้ / import ไม่ผ่าน"],
        ["E03", "E GL/ยกยอด", "P1", "นำเข้า Opening Balance", "ชุดยกยอดตามวันที่ที่ตกลง", "ทดลองงบหลังยกยอดสมดุล"],
        ["E04", "E GL/ยกยอด", "P1", "นับบัญชีและเอกสารเทียบไฟล์", "นับผังบัญชี บิลที่โพสต์ รายการยกยอด", "ตรงไฟล์ต้นทาง บันทึกในสรุป"],
        ["E05", "E GL/ยกยอด", "P2", "งวดบัญชีล็อกแล้วนำเข้าไม่ได้", "ลองนำเข้าเอกสารในงวดที่ปิด", "ระบบกันตามงวด"],
    ],
}


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for spec in (BUDGET, PURCHASE, WAREHOUSE, FINANCE):
        save(spec)


if __name__ == "__main__":
    main()
