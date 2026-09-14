#!/usr/bin/env python3
"""Presentation deck for explaining VPK budget workflow to users (UAT)."""
from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = "/opt/odoo18vpk/docs/UAT-งบประมาณ-workflow.pptx"
FONT = "Prompt"

NAVY = RGBColor(0x12, 0x3A, 0x56)
TEAL = RGBColor(0x0F, 0x7A, 0x72)
TEAL_SOFT = RGBColor(0xE5, 0xF5, 0xF3)
AMBER = RGBColor(0xC4, 0x7E, 0x14)
AMBER_SOFT = RGBColor(0xFD, 0xF3, 0xE0)
RED = RGBColor(0xB4, 0x43, 0x3A)
RED_SOFT = RGBColor(0xFB, 0xEC, 0xEA)
GREEN = RGBColor(0x2C, 0x7A, 0x45)
GREEN_SOFT = RGBColor(0xE5, 0xF5, 0xEA)
SLATE = RGBColor(0x2F, 0x3A, 0x48)
MUTED = RGBColor(0x5B, 0x68, 0x78)
LINE = RGBColor(0xD7, 0xDE, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF6, 0xF8, 0xFA)
INK = RGBColor(0x9F, 0xC9, 0xC4)

TOTAL = 14


def set_run(run, size, bold=False, color=SLATE):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", FONT)


def apply_theme_font(prs):
    try:
        master = prs.slide_masters[0]
        theme_part = None
        for rel in master.part.rels.values():
            if "theme" in (rel.reltype or ""):
                theme_part = rel.target_part
                break
        if theme_part is None:
            return
        root = theme_part.element
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for font_el in root.xpath(".//a:majorFont | .//a:minorFont", namespaces=ns):
            for child in font_el:
                tag = child.tag.split("}")[-1]
                if tag in ("latin", "ea", "cs"):
                    child.set("typeface", FONT)
    except Exception:
        return


def fill_text(shape, lines, align=PP_ALIGN.CENTER, anchor="ctr"):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    tf._txBody.bodyPr.set("anchor", anchor)
    for i, item in enumerate(lines):
        text, size, bold, color = item if isinstance(item, tuple) else (item, 16, False, SLATE)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        if i == 0:
            p.text = ""
        run = p.add_run()
        run.text = text
        set_run(run, size, bold, color)


def tb(slide, l, t, w, h, text, size=18, bold=False, color=SLATE, align=PP_ALIGN.LEFT):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return sh


def rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    return sh


def card(slide, l, t, w, h, fill=WHITE, line=LINE, lw=1.25):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(lw)
    return sh


def notes(slide, text):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    tf.text = text
    for p in tf.paragraphs:
        for run in p.runs:
            set_run(run, 14, False, SLATE)


def footer(slide, page):
    rect(slide, Inches(0), Inches(7.22), Inches(13.333), Inches(0.28), NAVY)
    tb(slide, Inches(0.4), Inches(7.23), Inches(9.8), Inches(0.26),
       "โรงพยาบาลวชิระภูเก็ต  ·  อธิบายระบบงบประมาณให้ผู้ใช้", 11, False, WHITE)
    tb(slide, Inches(11.3), Inches(7.23), Inches(1.6), Inches(0.26),
       f"{page} / {TOTAL}", 11, False, WHITE, PP_ALIGN.RIGHT)


def header(slide, kicker, title):
    rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.22), NAVY)
    rect(slide, Inches(0), Inches(0), Inches(0.14), Inches(1.22), TEAL)
    tb(slide, Inches(0.45), Inches(0.14), Inches(12.3), Inches(0.32), kicker, 13, False, INK)
    tb(slide, Inches(0.45), Inches(0.44), Inches(12.3), Inches(0.66), title, 26, True, WHITE)


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), BG)
    return slide


def arrow(slide, x, y):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, Inches(0.26), Inches(0.16))
    sh.fill.solid()
    sh.fill.fore_color.rgb = TEAL
    sh.line.fill.background()


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    apply_theme_font(prs)

    # 1 cover
    s = blank(prs)
    rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    rect(s, Inches(0), Inches(0), Inches(0.16), Inches(7.5), TEAL)
    tb(s, Inches(0.7), Inches(1.55), Inches(12), Inches(0.36),
       "ก่อนเริ่มทดสอบใช้งาน  ·  UAT", 14, False, INK)
    tb(s, Inches(0.7), Inches(2.05), Inches(12), Inches(1.35),
       "ระบบงบประมาณทำงานอย่างไร", 36, True, WHITE)
    tb(s, Inches(0.7), Inches(3.45), Inches(11.4), Inches(0.9),
       "อธิบายภาพรวมให้ผู้ใช้เข้าใจก่อนลงมือกดในระบบ\nตั้งงบ  →  ใช้งบ  →  ดูคงเหลือ",
       18, False, RGBColor(0xD7, 0xE4, 0xEE))
    c1 = card(s, Inches(0.7), Inches(5.15), Inches(3.7), Inches(1.05), TEAL, TEAL, 0)
    fill_text(c1, [("โรงพยาบาลวชิระภูเก็ต", 15, True, WHITE), ("โมดูลงบประมาณใน Odoo", 13, False, WHITE)])
    c2 = card(s, Inches(4.6), Inches(5.15), Inches(3.7), Inches(1.05), RGBColor(0x1C, 0x4E, 0x6E), RGBColor(0x1C, 0x4E, 0x6E), 0)
    fill_text(c2, [("ผู้ฟัง", 15, True, WHITE), ("หน่วยงาน  งานงบ  พัสดุ  การเงิน", 13, False, WHITE)])
    c3 = card(s, Inches(8.5), Inches(5.15), Inches(3.7), Inches(1.05), RGBColor(0x1C, 0x4E, 0x6E), RGBColor(0x1C, 0x4E, 0x6E), 0)
    fill_text(c3, [("ใช้เวลาพูด", 15, True, WHITE), ("ประมาณ 15–20 นาที แล้วค่อยเปิดระบบ", 13, False, WHITE)])
    notes(s, "เปิดด้วยประโยค: วันนี้ยังไม่ให้กดระบบทันที จะเล่าภาพรวมก่อน 15 นาที แล้วค่อยเดินทดสอบตามลำดับ")

    # 2 agenda
    s = blank(prs)
    header(s, "ลำดับการอธิบาย", "วันนี้จะเล่า 5 เรื่องนี้")
    items = [
        ("1", "ภาพรวม", "ระบบมี 2 ช่วง: ตั้งงบ แล้วค่อยใช้งบ"),
        ("2", "ตั้งงบประจำปี", "หน่วยงานขอ งานงบพิจารณา แล้วเข้าแผน"),
        ("3", "ใช้งบ", "ใบขอซื้อ → ใบสั่งซื้อ → ตั้งเจ้าหนี้"),
        ("4", "ตัวเลขคงเหลือ", "ใช้ได้เท่าไร ดูจากสูตรเดียว"),
        ("5", "ลำดับทดสอบ", "6 ข้อ ที่ต้องทำตามลำดับ"),
    ]
    for i, (num, title, desc) in enumerate(items):
        y = Inches(1.5) + i * Inches(1.05)
        card(s, Inches(0.5), y, Inches(12.3), Inches(0.92), WHITE, LINE, 1)
        n = card(s, Inches(0.7), y + Inches(0.2), Inches(0.72), Inches(0.52), TEAL, TEAL, 0)
        fill_text(n, [(num, 20, True, WHITE)])
        tb(s, Inches(1.7), y + Inches(0.14), Inches(10.7), Inches(0.36), title, 20, True, NAVY)
        tb(s, Inches(1.7), y + Inches(0.5), Inches(10.7), Inches(0.32), desc, 15, False, MUTED)
    footer(s, 2)
    notes(s, "บอกผู้ฟังว่าไม่ต้องจำรายละเอียดหน้าจอทุกปุ่ม จำลำดับงานและความหมายของเงินให้ได้ก่อน")

    # 3 one sentence
    s = blank(prs)
    header(s, "จำประโยคนี้ไว้ก่อน", "ยังซื้อของไม่ได้ จนกว่าจะมีงบที่อนุมัติ")
    left = card(s, Inches(0.5), Inches(1.7), Inches(6.0), Inches(4.9), WHITE, TEAL, 2)
    fill_text(
        left,
        [
            ("ช่วงที่ 1", 14, True, TEAL),
            ("ตั้งงบ", 32, True, NAVY),
            ("", 10, False, WHITE),
            ("หน่วยงานขอเงินประจำปี", 16, False, SLATE),
            ("งานงบจัดสรรแล้วอนุมัติ", 16, False, SLATE),
            ("ได้วงเงินในแผน", 16, False, SLATE),
        ],
    )
    right = card(s, Inches(6.8), Inches(1.7), Inches(6.0), Inches(4.9), WHITE, AMBER, 2)
    fill_text(
        right,
        [
            ("ช่วงที่ 2", 14, True, AMBER),
            ("ใช้งบ", 32, True, NAVY),
            ("", 10, False, WHITE),
            ("จะซื้อของ ต้องมีใบขอซื้อ", 16, False, SLATE),
            ("ระบบกันเงินให้อัตโนมัติ", 16, False, SLATE),
            ("ใช้ได้ไม่เกินวงเงินที่จัดสรร", 16, False, SLATE),
        ],
    )
    footer(s, 3)
    notes(s, "ย้ำว่าคนมักเข้าใจผิดว่าของบแล้วซื้อได้เลย จริงๆ ต้องรออนุมัติเข้าแผนก่อน แล้วค่อยเปิดใบขอซื้อ")

    # 4 cycle
    s = blank(prs)
    header(s, "ภาพรวมวงจร", "งานเดินจากซ้ายไปขวา อย่าข้ามขั้น")
    steps = [
        ("1", "ข้อมูลหลัก", "แหล่งเงิน หน่วยงาน\nแบบฟอร์ม", NAVY),
        ("2", "แผนปี", "เปิดกรอบปีงบ", NAVY),
        ("3", "คำของบ", "หน่วยงานกรอกขอ", TEAL),
        ("4", "อนุมัติ", "งานงบจัดสรร", TEAL),
        ("5", "ใช้งบ", "PR → PO → บิล", AMBER),
        ("6", "ดูคงเหลือ", "Dashboard", GREEN),
    ]
    for i, (num, title, body, col) in enumerate(steps):
        x = Inches(0.38) + i * Inches(2.16)
        top = card(s, x, Inches(1.6), Inches(2.02), Inches(0.58), col, col, 0)
        fill_text(top, [(f"{num}  {title}", 14, True, WHITE)])
        b = card(s, x, Inches(2.22), Inches(2.02), Inches(2.35), WHITE, col, 1.5)
        fill_text(b, [(body, 15, False, SLATE)])
        if i < 5:
            arrow(s, x + Inches(1.94), Inches(3.15))
    note = card(s, Inches(0.38), Inches(4.85), Inches(12.55), Inches(2.05), TEAL_SOFT, TEAL, 1.25)
    tb(s, Inches(0.6), Inches(5.05), Inches(12.1), Inches(0.4), "พูดกับผู้ใช้แบบนี้", 16, True, TEAL)
    tb(
        s,
        Inches(0.6),
        Inches(5.5),
        Inches(12.1),
        Inches(1.15),
        "แผนปีเป็นกรอบ หน่วยงานของบเข้าไปในกรอบนี้\nงานงบเป็นคนบอกว่าใช้ได้เท่าไร   การซื้อของต้องมีเอกสารจัดซื้อ ระบบจะกันเงินให้เอง",
        17,
        False,
        SLATE,
    )
    footer(s, 4)
    notes(s, "ชี้ทีละกล่อง ช้าๆ เน้นข้อ 3-5 เพราะเป็นจุดที่ user จะได้กดเอง")

    # 5 four forms
    s = blank(prs)
    header(s, "เลือกใบให้ถูก", "คำของบมีแค่ 4 แบบฟอร์ม")
    forms = [
        ("01", "แบบฟอร์มตั้งงบวัสดุ", "ยา เวชภัณฑ์ วัสดุทั่วไป\nต้องมีรายละเอียดสินค้า", TEAL),
        ("02", "แบบฟอร์มครุภัณฑ์", "เครื่องมือ เครื่องใช้\nคอมพิวเตอร์ รวมใบนี้", NAVY),
        ("03", "แบบฟอร์ม ก่อสร้าง", "งานซ่อม ต่อเติม อาคาร", AMBER),
        ("04", "แบบฟอร์มโครงการ", "ค่าใช้จ่ายโครงการ\nกิจกรรมต่างๆ", GREEN),
    ]
    for i, (num, title, body, col) in enumerate(forms):
        x = Inches(0.4) + i * Inches(3.22)
        c = card(s, x, Inches(1.65), Inches(3.05), Inches(5.15), WHITE, col, 2)
        n = card(s, x + Inches(0.2), Inches(1.9), Inches(0.7), Inches(0.55), col, col, 0)
        fill_text(n, [(num, 16, True, WHITE)])
        tb(s, x + Inches(0.18), Inches(2.65), Inches(2.7), Inches(1.15), title, 20, True, NAVY)
        tb(s, x + Inches(0.18), Inches(4.0), Inches(2.7), Inches(2.2), body, 16, False, MUTED)
    footer(s, 5)
    notes(s, "ย้ำว่าเหลือ 4 ใบ ไม่มีใบต่ำกว่าแสน และครุภัณฑ์คอมใช้ใบครุภัณฑ์ เลือกผิดแบบฟอร์มจะกรอกคนละหน้า")

    # 6 request who
    s = blank(prs)
    header(s, "ช่วงตั้งงบ  ·  หน่วยงาน", "หน่วยงานทำ 3 อย่างนี้")
    acts = [
        ("1", "สร้างคำของบ", "เลือกหน่วยงาน แหล่งเงิน แผนปี และแบบฟอร์ม"),
        ("2", "กรอกรายการ", "ใส่ของที่ต้องการ พร้อมเหตุผล ว่างไม่ได้"),
        ("3", "กดส่งงานงบ", "ส่งแล้วแก้เองไม่ได้ ต้องรองานงบรับเรื่อง"),
    ]
    for i, (num, title, body) in enumerate(acts):
        y = Inches(1.6) + i * Inches(1.7)
        card(s, Inches(0.5), y, Inches(12.3), Inches(1.5), WHITE, LINE, 1)
        n = card(s, Inches(0.75), y + Inches(0.4), Inches(0.85), Inches(0.7), TEAL, TEAL, 0)
        fill_text(n, [(num, 24, True, WHITE)])
        tb(s, Inches(1.9), y + Inches(0.28), Inches(10.5), Inches(0.5), title, 24, True, NAVY)
        tb(s, Inches(1.9), y + Inches(0.82), Inches(10.5), Inches(0.45), body, 17, False, MUTED)
    footer(s, 6)
    notes(s, "ถ้าเป็นงบวัสดุ เน้นว่าต้องมีรายละเอียดสินค้าให้ครบ ประเภทวัสดุต้องตรงกับที่ตั้งไว้")

    # 7 budget office
    s = blank(prs)
    header(s, "ช่วงตั้งงบ  ·  งานงบประมาณ", "งานงบเป็นคนปิดงานให้")
    flow = [
        ("รับเรื่อง", "รับเข้าคิว", TEAL),
        ("พิจารณา", "ดูรายการ ปรับได้", AMBER),
        ("ใส่ยอดจัดสรร", "ต้องครบทุกบรรทัด", AMBER),
        ("อนุมัติ", "เข้าแผนปีทันที", GREEN),
    ]
    for i, (title, sub, col) in enumerate(flow):
        x = Inches(0.45) + i * Inches(3.2)
        c = card(s, x, Inches(1.6), Inches(2.95), Inches(2.4), WHITE, col, 2)
        fill_text(c, [(title, 22, True, col), (sub, 15, False, MUTED)])
        if i < 3:
            arrow(s, x + Inches(2.88), Inches(2.7))
    rej = card(s, Inches(0.45), Inches(4.35), Inches(6.05), Inches(2.45), RED_SOFT, RED, 1.5)
    fill_text(rej, [("ถ้าไม่ผ่าน", 14, True, RED), ("ตีกลับ", 28, True, RED), ("หน่วยงานแก้ แล้วส่งใหม่", 16, False, SLATE)])
    ok = card(s, Inches(6.8), Inches(4.35), Inches(6.05), Inches(2.45), GREEN_SOFT, GREEN, 1.5)
    fill_text(ok, [("ถ้าผ่าน", 14, True, GREEN), ("ได้วงเงินใช้จ่าย", 28, True, GREEN), ("ใช้ได้เท่าที่จัดสรร ไม่ใช่เท่าที่ขอ", 16, False, SLATE)])
    footer(s, 7)
    notes(s, "จุดสำคัญ: อนุมัติได้เมื่อใส่ยอดจัดสรรครบ ยอดนี้คือเพดานใช้จ่ายทั้งปีของรายการนั้น")

    # 8 spend
    s = blank(prs)
    header(s, "ช่วงใช้งบ", "ซื้อของทีละใบ ตามนี้")
    spend = [
        ("ร่าง PR", "ยังไม่กันเงิน", NAVY),
        ("เช็คงบ", "พอไหม", TEAL),
        ("ส่ง PR", "จองงบ", AMBER),
        ("ยืนยัน PO", "ผูกพันงบ", AMBER),
        ("ตั้งเจ้าหนี้", "ใช้จริง", GREEN),
    ]
    for i, (title, sub, col) in enumerate(spend):
        x = Inches(0.4) + i * Inches(2.55)
        c = card(s, x, Inches(1.7), Inches(2.38), Inches(3.15), WHITE, col, 2)
        fill_text(c, [(title, 22, True, col), (sub, 16, False, MUTED)])
        if i < 4:
            arrow(s, x + Inches(2.28), Inches(3.15))
    row = [
        ("เช็คงบไม่ผ่าน", "ส่งอนุมัติไม่ได้ ต้องลดยอดหรือเลือกงบให้ถูก", RED_SOFT, RED),
        ("ส่ง PR แล้ว", "เงินถูกกันไว้ ใบอื่นใช้วงเงินซ้ำไม่ได้", AMBER_SOFT, AMBER),
        ("มีบิลแล้ว", "เงินกลายเป็นใช้จริง คงเหลือลดลงถาวร", GREEN_SOFT, GREEN),
    ]
    for i, (title, body, bg, fg) in enumerate(row):
        x = Inches(0.4) + i * Inches(4.25)
        c = card(s, x, Inches(5.1), Inches(4.08), Inches(1.75), bg, fg, 1.25)
        fill_text(c, [(title, 16, True, fg), (body, 14, False, SLATE)])
    footer(s, 8)
    notes(s, "เปรียบเหมือนจองโต๊ะ: ส่ง PR = จองไว้, PO = นั่งแล้ว, ตั้งบิล = จ่ายแล้ว")

    # 9 formula
    s = blank(prs)
    header(s, "ตัวเลขที่ต้องจำ", "คงเหลือดูจากสูตรนี้สูตรเดียว")
    bar = card(s, Inches(0.4), Inches(1.55), Inches(12.5), Inches(1.2), NAVY, NAVY, 0)
    fill_text(bar, [("คงเหลือ  =  งบจัดสรร  −  ใช้จริง  −  จอง PR  −  ผูกพัน PO", 22, True, WHITE)])
    nums = [
        ("1,000,000", "งบจัดสรร", "ที่อนุมัติในแผน", NAVY),
        ("250,000", "ใช้จริง", "บิลที่โพสต์แล้ว", GREEN),
        ("150,000", "จอง PR", "ส่งแล้วยังไม่ครบ PO", AMBER),
        ("200,000", "ผูกพัน PO", "PO ยังไม่ครบบิล", AMBER),
        ("400,000", "คงเหลือ", "ซื้อเพิ่มได้เท่านั้น", TEAL),
    ]
    for i, (val, label, hint, col) in enumerate(nums):
        x = Inches(0.4) + i * Inches(2.55)
        c = card(s, x, Inches(3.0), Inches(2.4), Inches(2.55), WHITE, col, 1.75)
        fill_text(c, [(val, 20, True, col), (label, 16, True, NAVY), (hint, 13, False, MUTED)])
    footer(s, 9)
    notes(s, "เดินตัวเลขจากซ้ายไปขวา 1,000,000 ลบ 250+150+200 เหลือ 400,000 ถามผู้ฟังว่าซื้อของ 500,000 ได้ไหม คำตอบคือไม่ได้")

    # 10 matching
    s = blank(prs)
    header(s, "ถ้าเช็คงบไม่ผ่าน", "ดูให้ครบ 4 อย่างนี้ก่อน")
    keys = [
        ("หน่วยงาน", "แผนกที่ขอต้องตรงกับงบในแผน"),
        ("แหล่งเงิน", "เงินบำรุง / งบประมาณ / เงินบริจาค"),
        ("หมวดงบ", "วัสดุ ครุภัณฑ์ ก่อสร้าง โครงการ"),
        ("กลุ่มวัสดุ", "ใช้เฉพาะตอนซื้อวัสดุ เช่น ยา เวชภัณฑ์"),
    ]
    for i, (title, body) in enumerate(keys):
        r, c = divmod(i, 2)
        x = Inches(0.45) + c * Inches(6.4)
        y = Inches(1.6) + r * Inches(2.55)
        boxc = card(s, x, y, Inches(6.15), Inches(2.35), WHITE, TEAL, 1.5)
        n = card(s, x + Inches(0.25), y + Inches(0.3), Inches(0.7), Inches(0.7), TEAL, TEAL, 0)
        fill_text(n, [(str(i + 1), 22, True, WHITE)])
        tb(s, x + Inches(1.15), y + Inches(0.35), Inches(4.7), Inches(0.55), title, 24, True, NAVY)
        tb(s, x + Inches(0.3), y + Inches(1.2), Inches(5.55), Inches(0.8), body, 17, False, MUTED)
    footer(s, 10)
    notes(s, "ถ้า UAT ติดที่เช็คงบ ให้พาเปิดใบขอซื้อดู 4 ช่องนี้ทีละช่อง อย่าเพิ่งโทษว่าระบบบั๊ก")

    # 11 roles
    s = blank(prs)
    header(s, "ใครกดอะไร", "จำบทบาทตัวเองให้ชัด")
    roles = [
        ("หน่วยงาน", TEAL, "สร้างคำของบ แล้วส่ง\nสร้างใบขอซื้อ เช็คงบ ส่งอนุมัติ"),
        ("งานงบประมาณ", AMBER, "ตั้งข้อมูลหลักและแผนปี\nรับเรื่อง จัดสรร อนุมัติหรือตีกลับ"),
        ("พัสดุ", NAVY, "ยืนยันใบสั่งซื้อ (PO)\nยอดจองจะกลายเป็นผูกพัน"),
        ("การเงิน", GREEN, "ตั้งเจ้าหนี้จากใบแจ้งหนี้\nยอดกลายเป็นใช้จริง"),
    ]
    for i, (title, col, body) in enumerate(roles):
        x = Inches(0.4) + i * Inches(3.22)
        c = card(s, x, Inches(1.7), Inches(3.05), Inches(5.1), WHITE, col, 2)
        fill_text(c, [(title, 22, True, col), ("", 8, False, WHITE), (body, 16, False, SLATE)])
    footer(s, 11)
    notes(s, "ให้คนในห้องชี้ว่าตัวเองอยู่กล่องไหน กันกดข้ามสิทธิ์ตอนทดสอบ")

    # 12 uat script
    s = blank(prs)
    header(s, "ลำดับทดสอบวันนี้", "ทำทีละข้อ ห้ามข้าม")
    script = [
        ("1", "ตรวจข้อมูลหลัก", "แหล่งเงิน หน่วยงาน แผน แบบฟอร์ม เลือกได้ครบ"),
        ("2", "หน่วยงานของบวัสดุ", "สร้าง แล้วส่งงานงบ"),
        ("3", "งานงบอนุมัติ", "ใส่ยอดจัดสรรครบ แล้วกดอนุมัติ"),
        ("4", "PR เกินงบ", "ต้องขึ้นว่างบไม่พอ และส่งไม่ได้"),
        ("5", "PR ในงบ แล้วส่ง", "เงินถูกจอง คงเหลือลด"),
        ("6", "PO แล้วตั้งเจ้าหนี้", "ยอดย้าย จอง → ผูกพัน → ใช้จริง"),
    ]
    for i, (num, title, body) in enumerate(script):
        r, c = divmod(i, 2)
        x = Inches(0.4) + c * Inches(6.45)
        y = Inches(1.55) + r * Inches(1.75)
        card(s, x, y, Inches(6.25), Inches(1.55), WHITE, LINE, 1)
        n = card(s, x + Inches(0.18), y + Inches(0.42), Inches(0.7), Inches(0.7), NAVY, NAVY, 0)
        fill_text(n, [(num, 22, True, WHITE)])
        tb(s, x + Inches(1.05), y + Inches(0.28), Inches(4.95), Inches(0.5), title, 18, True, NAVY)
        tb(s, x + Inches(1.05), y + Inches(0.82), Inches(4.95), Inches(0.5), body, 14, False, MUTED)
    footer(s, 12)
    notes(s, "บอกว่าข้อ 4 ตั้งใจให้ไม่ผ่าน อย่าตกใจ ข้อ 6 ต้องทำครบถึงจะเห็นสูตรคงเหลือจริง")

    # 13 remember
    s = blank(prs)
    header(s, "ก่อนเปิดระบบ", "จำ 3 ข้อนี้ก็ไปต่อได้")
    recs = [
        ("1", "ของบยังซื้อไม่ได้", "ต้องรออนุมัติเข้าแผนก่อน"),
        ("2", "ใช้ได้เท่าที่จัดสรร", "ไม่ใช่เท่าที่ขอตอนต้นปี"),
        ("3", "ซื้อของต้องมี PR", "ส่งแล้วระบบกันเงินให้เอง"),
    ]
    for i, (num, title, body) in enumerate(recs):
        y = Inches(1.6) + i * Inches(1.7)
        card(s, Inches(0.5), y, Inches(12.3), Inches(1.5), WHITE, LINE, 1)
        n = card(s, Inches(0.75), y + Inches(0.35), Inches(0.9), Inches(0.8), TEAL, TEAL, 0)
        fill_text(n, [(num, 26, True, WHITE)])
        tb(s, Inches(1.95), y + Inches(0.28), Inches(10.4), Inches(0.55), title, 26, True, NAVY)
        tb(s, Inches(1.95), y + Inches(0.88), Inches(10.4), Inches(0.4), body, 16, False, MUTED)
    footer(s, 13)
    notes(s, "ทวน 3 ข้อ แล้วถามว่ามีคำถามไหม ถ้าไม่มี เปิดระบบข้อ 1")

    # 14 close
    s = blank(prs)
    rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    rect(s, Inches(0), Inches(0), Inches(0.16), Inches(7.5), TEAL)
    tb(s, Inches(0.7), Inches(1.9), Inches(12), Inches(0.4), "พร้อมลงมือ", 15, False, INK)
    tb(s, Inches(0.7), Inches(2.4), Inches(12), Inches(1.5),
       "เริ่มที่ข้อมูลหลัก\nยังไม่ต้องเปิดใบซื้อ", 34, True, WHITE)
    tb(s, Inches(0.7), Inches(4.3), Inches(11.8), Inches(1.1),
       "เดินตามข้อ 1 ถึง 6 ทีละขั้น\nถ้าเช็คงบไม่ผ่าน ให้ดูหน่วยงาน แหล่งเงิน หมวดงบ และกลุ่มวัสดุ",
       17, False, RGBColor(0xD7, 0xE4, 0xEE))
    go = card(s, Inches(0.7), Inches(5.7), Inches(5.6), Inches(0.85), TEAL, TEAL, 0)
    fill_text(go, [("เริ่มข้อ 1 : ตรวจข้อมูลหลัก", 18, True, WHITE)])
    notes(s, "ปิดไฟล์นี้ แล้วสลับไปหน้าจอ Odoo ทันที อย่าเปิดเมนูซื้อของก่อนมีแผนที่อนุมัติ")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
