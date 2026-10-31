#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a Thai user guide explaining how the official document module works."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "คู่มือการใช้งาน" / "คู่มือการใช้งาน-ระบบหนังสือราชการ.docx"
LOGO = ROOT / "manual_assets" / "figures" / "moph_logo.png"

NAVY = RGBColor(0x0B, 0x68, 0x48)
SLATE = RGBColor(0x2A, 0x33, 0x3A)
MUTED = RGBColor(0x7D, 0x8A, 0x9E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "0B6848"


def set_run(run, size=12, bold=False, color=SLATE, font_name="Prompt"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font_name)


def add_text(p, text, size=12, bold=False, color=SLATE):
    run = p.add_run(text)
    set_run(run, size, bold, color)
    return run


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, color="D7DEE6"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_narrow_cell_margin(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), "80")
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def cell_text(cell, text, size=11, bold=False, color=SLATE, align="left", fill=None):
    if fill:
        shade_cell(cell, fill)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.clear()
    add_text(p, text, size, bold, color)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    set_narrow_cell_margin(cell)


def make_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        cell_text(table.rows[0].cells[i], header, 11, True, WHITE, "center", HEADER_FILL)
    for r, row in enumerate(rows, 1):
        fill = "F6F8FA" if r % 2 == 0 else "FFFFFF"
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, 11, False, SLATE, "left", fill)
    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


def add_bottom_line(paragraph, color="0B6848", size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), "6")
    el.set(qn("w:color"), color)
    pBdr.append(el)
    pPr.append(pBdr)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, 16 if level == 1 else 13, True, NAVY)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    if level == 1:
        add_bottom_line(p)
    return p


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text, 12)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        add_text(p, item, 12)
        p.paragraph_format.space_after = Pt(3)


def numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(3)
        add_text(p, f"{index}. {item}", 12)


def callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade_cell(cell, "EAF3F1")
    set_cell_border(cell, "0B6848")
    p = cell.paragraphs[0]
    p.clear()
    add_text(p, title, 12, True, NAVY)
    p2 = cell.add_paragraph()
    add_text(p2, text, 12)
    p2.paragraph_format.space_after = Pt(4)
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(header, "โรงพยาบาลวชิระภูเก็ต  ·  คู่มือระบบหนังสือราชการ", 9, False, MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "กลุ่มงานพัสดุ  ·  ใช้ประกอบการอบรมและปฏิบัติงานจริง", 9, False, MUTED)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Cm(2.4))

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(t, "คู่มือสรุปการทำงาน", 12, False, MUTED)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(title, "ระบบหนังสือราชการ", 26, True, NAVY)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(16)
    add_text(sub, "โมดูลออกเอกสารราชการ ลงนามอิเล็กทรอนิกส์ และเชื่อมเลขสารบรรณ", 13)

    callout(
        doc,
        "จำแค่นี้ก่อนก็ใช้ได้",
        "โมดูลนี้ไม่ใช่ตู้เก็บไฟล์ทั่วไป แต่เป็นโรงพิมพ์เอกสารราชการของงานจัดซื้อ "
        "ระบบดึงข้อมูลจากใบขอซื้อ / ใบตรวจรับ / e-GP มาใส่แบบฟอร์ม Word แล้วแปลงเป็น PDF "
        "ส่งให้ผู้มีหน้าที่ลงนามตามลำดับ เมื่อลงนามครบ เอกสารนั้นจึงใช้เดินงานขั้นถัดไปได้",
    )

    heading(doc, "1. โมดูลนี้คืออะไร")
    body(
        doc,
        "งานจัดซื้อจัดจ้างภาครัฐต้องมีหนังสือราชการหลายฉบับ เช่น คำสั่งแต่งตั้งคณะกรรมการ "
        "รายงานขออนุมัติจัดซื้อ ประกาศผู้ชนะ และใบตรวจรับ เดิมเจ้าหน้าที่พิมพ์จากแบบฟอร์ม "
        "แล้วนำไปลงนามบนกระดาษ โมดูลหนังสือราชการทำหน้าที่นี้ในระบบ ERP",
    )
    body(doc, "แต่ละฉบับในระบบจึงมี 3 ส่วนที่ต้องจำคู่กัน")
    numbered(
        doc,
        [
            "ข้อมูลต้นทาง เช่น ใบขอซื้อ ใบสั่งซื้อ ใบรับมอบงาน หรือเอกสาร e-GP",
            "แบบฟอร์ม Word ของประเภทนั้น ซึ่งระบบเติมข้อความ รายชื่อคณะกรรมการ และเลขที่สารบรรณให้",
            "ไฟล์ PDF ที่ส่งเข้ากล่องอนุมัติ เพื่อให้ผู้เกี่ยวข้องเปิดดูและลงนามบนไฟล์",
        ],
    )
    body(
        doc,
        "เข้าเมนูได้ที่ ระบบจัดซื้อจัดจ้าง → หนังสือราชการ → ภาพรวม "
        "หน้านี้เป็นแผงการ์ดตามประเภทเอกสาร กดการ์ดแล้วจะเห็นรายการของประเภทนั้น",
    )

    heading(doc, "2. อยู่ตรงไหนของงานจัดซื้อ")
    body(
        doc,
        "หนังสือราชการไม่ได้ทำงานคนเดียว แต่แทรกอยู่ในสายงานจัดซื้อ "
        "จุดที่สร้างเอกสารไม่ใช่ที่เดียวกันทั้งหมด",
    )
    make_table(
        doc,
        ["ช่วงงาน", "เอกสารที่มักออก", "จุดที่กดสร้าง"],
        [
            [
                "ก่อน/ระหว่างขอซื้อ",
                "บันทึกข้อความอนุมัติขอซื้อ, ขอแต่งตั้งคณะกรรมการกำหนดคุณลักษณะ, แบบแสดงความบริสุทธิ์ใจ",
                "ใบขอซื้อ หรือเมนูหนังสือราชการ",
            ],
            [
                "ขออนุมัติจัดซื้อ",
                "รายงานขออนุมัติจัดซื้อจัดจ้างโดยวิธีเฉพาะเจาะจง, คำสั่งแต่งตั้งคณะกรรมการตรวจรับ",
                "ใบขอซื้อ (ระบบสร้างคู่กันเมื่อส่งอนุมัติ) หรือสร้างเองจากแท็บเอกสาร",
            ],
            [
                "e-Bidding / คัดเลือกผู้ชนะ",
                "รายงานผลการพิจารณา, ประกาศผู้ชนะการเสนอราคา",
                "ชุดเอกสาร e-GP ของใบประกวดราคา",
            ],
            [
                "ตรวจรับพัสดุ",
                "ใบตรวจรับการจัดซื้อ/จัดจ้าง",
                "ใบรับมอบงาน (Work Acceptance)",
            ],
        ],
        col_widths=[4.0, 7.4, 5.8],
    )
    body(
        doc,
        "ดังนั้นถ้าเปิดเมนูหนังสือราชการแล้วยังว่าง ไม่ได้แปลว่าระบบพัง "
        "ส่วนใหญ่เอกสารถูกสร้างจากใบขอซื้อ ใบตรวจรับ หรือ e-GP แล้วค่อยไปโผล่ที่ภาพรวม",
    )

    heading(doc, "3. ประเภทเอกสารในระบบ")
    body(
        doc,
        "ประเภทเอกสารเก็บในเมนู ประเภทเอกสาร ผู้จัดการระบบเพิ่มประเภทใหม่ได้ "
        "แต่ประเภทหลักที่มากับโรงพยาบาลมีดังนี้",
    )
    make_table(
        doc,
        ["ประเภท", "ใช้เมื่อใด", "ใครลงนามโดยสรุป"],
        [
            [
                "คำสั่งแต่งตั้งคณะกรรมการตรวจรับพัสดุ",
                "แต่งตั้งกรรมการตรวจรับของงานซื้อนั้น",
                "เลขาฯ กลั่นกรอง แล้วผู้อำนวยการลงนาม",
            ],
            [
                "ขออนุมัติแต่งตั้งคณะกรรมการกำหนดคุณลักษณะ",
                "แต่งตั้งกรรมการกำหนดรายละเอียด/ราคากลาง",
                "เลขาฯ กลั่นกรอง แล้วผู้อำนวยการลงนาม",
            ],
            [
                "แบบแสดงความบริสุทธิ์ใจ",
                "วงเงินมากกว่า 100,000 บาท",
                "เจ้าหน้าที่ → หัวหน้าเจ้าหน้าที่ → กรรมการ",
            ],
            [
                "รายงานขออนุมัติจัดซื้อโดยวิธีเฉพาะเจาะจง",
                "ขออนุมัติวิธีจัดซื้อและวงเงิน",
                "เจ้าหน้าที่ → หัวหน้าเจ้าหน้าที่ → เลขาฯ → ผู้อำนวยการ",
            ],
            [
                "บันทึกข้อความอนุมัติขอซื้อ",
                "บันทึกข้อความขออนุมัติซื้อ",
                "เจ้าหน้าที่ → หัวหน้าเจ้าหน้าที่ → เลขาฯ → ผู้อำนวยการ",
            ],
            [
                "รายงานผลการพิจารณาจัดซื้อจัดจ้าง",
                "สรุปผลพิจารณาหลังเปิดซอง/คัดเลือก",
                "เจ้าหน้าที่ → หัวหน้าเจ้าหน้าที่ → เลขาฯ → ผู้อำนวยการ",
            ],
            [
                "ประกาศผู้ชนะการเสนอราคา",
                "ประกาศผู้ชนะหลังอนุมัติผล",
                "เลขาฯ กลั่นกรอง แล้วผู้อำนวยการลงนาม",
            ],
            [
                "ใบตรวจรับการจัดซื้อ/จัดจ้าง",
                "ตรวจรับพัสดุตามใบสั่งซื้อ",
                "คณะกรรมการตรวจรับตามรายชื่อในเอกสาร",
            ],
        ],
        col_widths=[5.8, 5.6, 5.8],
    )

    heading(doc, "4. วงจรชีวิตของเอกสารทุกฉบับ")
    body(
        doc,
        "ไม่ว่าจะเป็นคำสั่ง รายงาน หรือใบตรวจรับ ขั้นตอนหลักเหมือนกัน "
        "ต่างกันที่ใครลงนามและเอกสารถูกสร้างจากจุดไหน",
    )
    numbered(
        doc,
        [
            "สร้างเอกสาร จากใบขอซื้อ / ใบรับมอบงาน / e-GP หรือจากเมนูหนังสือราชการ สถานะเริ่มต้นคือ ร่าง",
            "ระบบดึงข้อมูลอัตโนมัติ เช่น หน่วยงาน เรียน รายการจัดซื้อ รายชื่อคณะกรรมการ ผู้ลงนามจากค่าตั้งต้นของโรงพยาบาล",
            "ตรวจแก้ข้อความ ถ้าเป็นคำสั่งหรือรายงานที่ต้องมีกรรมการ กดดึงรายชื่อจากแท็บคณะกรรมการของใบขอซื้อ",
            "ขอเลขจากระบบสารบรรณ เพื่อให้ช่อง “ที่” ในหนังสือมีเลขหนังสือราชการ",
            "พิมพ์ PDF ระบบสร้างไฟล์ Word จากแบบฟอร์ม แล้วแปลงเป็น PDF",
            "ส่งเพื่อลงนาม เอกสารเข้ากล่องอนุมัติ สถานะเป็น รออนุมัติ / รอลงนาม",
            "ผู้มีหน้าที่เปิดไฟล์ลงนามตามลำดับ เลขาฯ ต้องกดตรวจสอบแล้ว ก่อนส่งต่อผู้อำนวยการ",
            "ลงนามครบ สถานะเป็น อนุมัติแล้ว ลายเซ็นถูกประทับบน PDF และงานขั้นถัดไปจึงเดินต่อได้",
        ],
    )
    make_table(
        doc,
        ["สถานะเอกสาร", "ความหมาย"],
        [
            ["ร่าง", "เพิ่งสร้าง ยังไม่พิมพ์ไฟล์หรือยังไม่ส่งลงนาม"],
            ["สร้างไฟล์แล้ว", "มีไฟล์ Word/PDF แล้ว แต่ยังไม่ส่งลงนาม"],
            ["รออนุมัติ", "อยู่ในกล่องอนุมัติ รอผู้ลงนามตามลำดับ"],
            ["อนุมัติแล้ว", "ลงนามครบ ใช้เป็นหลักฐานได้"],
            ["ยกเลิก", "ไม่ใช้ฉบับนี้แล้ว"],
        ],
        col_widths=[4.4, 13.0],
    )
    callout(
        doc,
        "ลำดับที่มักผิด",
        "ต้องมีไฟล์ PDF ก่อนจึงส่งลงนามได้ และต้องขอเลขสารบรรณก่อนพิมพ์ "
        "ถ้าพิมพ์ก่อนขอเลข ช่อง “ที่” ในหนังสืออาจว่าง "
        "ถ้าต้องการแก้ข้อความหลังส่งลงนามแล้ว ต้องตีกลับหรือยกเลิกแล้วสร้างใหม่ ไม่ใช่แก้ไฟล์เดิมระหว่างรอลงนาม",
    )

    heading(doc, "5. สร้างเอกสารจากจุดไหน")
    heading(doc, "จากใบขอซื้อ", 2)
    body(
        doc,
        "เปิดใบขอซื้อแล้วดูปุ่มหนังสือราชการที่หัวเอกสาร หรือเลือกประเภทแล้วสร้างจากแท็บที่เกี่ยวข้อง "
        "ระบบจะผูกเอกสารกับใบขอซื้อนั้น ถ้าเคยสร้างประเภทเดียวกันไว้แล้ว จะเปิดฉบับเดิม ไม่สร้างซ้ำ",
    )
    body(
        doc,
        "เมื่อส่งใบขอซื้อให้อนุมัติ ระบบอาจสร้างชุดเอกสารอนุมัติให้อัตโนมัติ "
        "ได้แก่ คำสั่งแต่งตั้งคณะกรรมการตรวจรับ และรายงานขออนุมัติจัดซื้อโดยวิธีเฉพาะเจาะจง "
        "พร้อมพิมพ์ PDF และตั้งสถานะรออนุมัติ เพื่อให้ผู้อำนวยการลงนามเป็นชุด",
    )
    heading(doc, "จากใบรับมอบงาน / ตรวจรับ", 2)
    body(
        doc,
        "ใบตรวจรับการจัดซื้อ/จัดจ้าง สร้างจากใบรับมอบงาน "
        "ระบบดึงรายชื่อคณะกรรมการตรวจรับและรายการพัสดุมาใส่แบบฟอร์ม "
        "คำสั่งแต่งตั้งและแบบแสดงความบริสุทธิ์ใจถ้าออกจากใบขอซื้อแล้ว ใบรับมอบงานจะชี้ไปฉบับนั้น",
    )
    heading(doc, "จากเอกสาร e-GP", 2)
    body(
        doc,
        "รายงานผลการพิจารณา และประกาศผู้ชนะ สร้างจากแถวเอกสารในชุด e-GP "
        "ระบบสร้างหนังสือราชการคู่กับไฟล์นั้น แล้วส่งเข้ากล่องลงนามแบบเดียวกับเอกสารอื่น",
    )
    heading(doc, "จากเมนูหนังสือราชการ", 2)
    body(
        doc,
        "ใช้เมื่อต้องการดูรายการทั้งโรงพยาบาล หรือสร้างเอกสารที่ไม่ได้ผูกกับใบใดใบหนึ่งในขณะนั้น "
        "กดการ์ดประเภทที่ต้องการ แล้วกด สร้าง "
        "ถ้าเอกสารนั้นควรผูกกับใบขอซื้อ ให้เลือกใบขอซื้อในฟอร์มด้วย จะได้ดึงข้อมูลและขอเลขสารบรรณได้ถูกฉบับ",
    )

    heading(doc, "6. ภาพรวมหนังสือราชการ")
    body(
        doc,
        "หน้าภาพรวมเป็นจุดเข้าหลักหลังปรับเมนู การ์ดหนึ่งใบคือหนึ่งประเภทเอกสาร "
        "ตัวเลขบนการ์ดคือจำนวนฉบับของประเภทนั้น",
    )
    bullets(
        doc,
        [
            "กดการ์ดทั้งใบ เพื่อเปิดรายการเอกสารประเภทนั้น",
            "กดตัวเลข ร่าง / รออนุมัติ / อนุมัติแล้ว เพื่อกรองเฉพาะสถานะนั้น",
            "กดดูเอกสารทั้งหมด เพื่อเปิดรายการข้ามประเภท",
            "ประเภทที่เพิ่มใหม่ในเมนูประเภทเอกสาร จะขึ้นเป็นการ์ดให้อัตโนมัติ",
        ],
    )

    heading(doc, "7. เลขสารบรรณ")
    body(
        doc,
        "หนังสือราชการต้องมีเลขที่ในช่อง “ที่” โมดูลนี้ไม่ป้อนเลขเองจากใจ "
        "แต่ขอเลขจากระบบสารบรรณของโรงพยาบาล",
    )
    bullets(
        doc,
        [
            "กดปุ่ม ขอเลขจากระบบสารบรรณ ที่หัวเอกสารของฉบับนั้น",
            "แต่ละหนังสือราชการที่ผูกกับใบขอซื้อ ต้องขอเลขของตนเอง ห้ามใช้เลขใบขอซื้อไปใส่ทุกฉบับ",
            "ข้อยกเว้น: เอกสารอนุมัติขอซื้อบางประเภทใช้เลขสารบรรณจากใบขอซื้อตามที่ระบบกำหนด",
            "เมื่อขอเลขสำเร็จ จะมีงานสารบรรณคู่กัน กดเปิดงานสารบรรณจากฟอร์มเอกสารได้",
        ],
    )

    heading(doc, "8. การลงนามและบทบาท")
    body(
        doc,
        "การลงนามไม่ได้อยู่ที่เมนูหนังสือราชการอย่างเดียว ผู้ลงนามเปิดงานจากกล่องอนุมัติเอกสาร "
        "ระบบจะเปิดไฟล์ PDF ให้ลงนามบนหน้าจอ ลายเซ็นถูกประทับลงไฟล์",
    )
    make_table(
        doc,
        ["บทบาท", "ทำอะไรในโมดูลนี้"],
        [
            [
                "เจ้าหน้าที่พัสดุ",
                "สร้างเอกสาร ดึงรายชื่อ ขอเลขสารบรรณ พิมพ์ PDF ส่งเพื่อลงนาม",
            ],
            [
                "หัวหน้าเจ้าหน้าที่",
                "ลงนามลำดับต้นของรายงาน/บันทึก/แบบแสดงความบริสุทธิ์ใจ",
            ],
            [
                "เลขาผู้อำนวยการ",
                "กลั่นกรองก่อนถึงผู้อำนวยการ กดตรวจสอบแล้ว จึงส่งผู้อำนวยการ หรือตีกลับ",
            ],
            [
                "ผู้อำนวยการ",
                "ลงนามลำดับสุดท้ายของคำสั่ง รายงาน ประกาศผู้ชนะ",
            ],
            [
                "คณะกรรมการตรวจรับ",
                "ลงนามใบตรวจรับตามหน้าที่ในรายชื่อ",
            ],
        ],
        col_widths=[4.6, 12.8],
    )
    body(
        doc,
        "ปุ่มสำคัญบนฟอร์มเอกสารระหว่างรอลงนาม ได้แก่ เปิดไฟล์เพื่อลงนาม, ตรวจสอบแล้ว, ส่งผู้อำนวยการ, ตีกลับ "
        "เลขาฯ ต้องตรวจสอบก่อน จึงจะส่งต่อผู้อำนวยการได้",
    )

    heading(doc, "9. แบบฟอร์มและการตั้งค่า")
    body(
        doc,
        "ข้อความในไฟล์ไม่ได้พิมพ์จากหน้าจอ HTML แต่มาจากแบบฟอร์ม Word ที่อัปโหลดในเมนู แบบฟอร์มเอกสาร "
        "แต่ละประเภทชี้ไปที่ไฟล์ Word ของตน เมื่อกดพิมพ์ PDF ระบบเติมข้อมูลลงช่องในแบบฟอร์มแล้วนำไปแปลงไฟล์",
    )
    bullets(
        doc,
        [
            "ผู้จัดการระบบแก้ชื่อประเภทได้ แต่รหัสประเภทของระบบห้ามเปลี่ยนหรือลบ ปิดการใช้งานแทน",
            "ผู้ลงนาม คำขึ้นต้นหน่วยงาน และคำว่าเรียน มีค่าตั้งต้นที่ตั้งค่าบริษัท/โรงพยาบาล",
            "ถ้าพิมพ์แล้วข้อความไม่ครบ ให้กลับไปแก้ข้อมูลต้นทางหรือกดสร้างข้อความจากข้อมูล แล้วพิมพ์ใหม่ก่อนส่งลงนาม",
        ],
    )

    heading(doc, "10. สิ่งที่โมดูลนี้ไม่ใช่")
    bullets(
        doc,
        [
            "ไม่ใช่ระบบสารบรรณทั้งโรงพยาบาล สารบรรณมีโมดูลของตนเอง หนังสือราชการมาขอเลขจากตรงนั้น",
            "ไม่ใช่คลังไฟล์แนบทั่วไป ของ e-GP ยังอยู่ที่แท็บเอกสารกระบวนการ e-GP ของใบขอซื้อ",
            "ไม่ใช่ขั้นตอนจัดซื้อเอง การคัดเลือกผู้ขาย เปิด PO และตรวจรับ ยังทำในโมดูลจัดซื้อ/คลัง ตามเดิม",
            "ไม่แทนการอนุมัติใบขอซื้อทั้งหมด ใบขอซื้อมีลำดับอนุมัติของตนเอง ส่วนหนังสือราชการเป็นการลงนามบนไฟล์เอกสาร",
        ],
    )

    heading(doc, "11. ลำดับตรวจเมื่อเอกสารเดินไม่ครบ")
    numbered(
        doc,
        [
            "เอกสารถูกสร้างหรือยัง ดูที่ภาพรวมหรือปุ่มนับหนังสือราชการบนใบขอซื้อ",
            "มีเลขสารบรรณหรือยัง ถ้าช่อง “ที่” ว่าง ให้กดขอเลข",
            "มีไฟล์ PDF หรือยัง ถ้ายัง ให้กดพิมพ์ PDF",
            "ส่งลงนามหรือยัง สถานะต้องเป็นรออนุมัติ จึงจะอยู่ในกล่องผู้อนุมัติ",
            "ค้างที่ใคร ดูตารางผู้ลงนามหรือกล่องอนุมัติ ถ้าค้างที่เลขาฯ ต้องกดตรวจสอบแล้วก่อนส่งผู้อำนวยการ",
            "ลงนามครบหรือยัง เมื่อครบ สถานะเป็นอนุมัติแล้ว จึงใช้เป็นหลักฐานและเดินงานต่อ",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
