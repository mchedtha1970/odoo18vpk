#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Create the UAT result workbook for the budget module."""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "/opt/odoo18vpk/docs/UAT-งบประมาณ-ใบบันทึกผล.xlsx"
FONT = "Prompt"

NAVY = "123A56"
TEAL = "0F7A72"
RED = "B4433A"
GREEN = "2C7A45"
LINE = "D7DEE6"
WHITE = "FFFFFF"
SOFT = "E5F5F3"
P1 = "FDECDC"
P2 = "E8F1F8"
PASS = "E5F5EA"
FAIL = "FBECEA"
SKIP = "F6F8FA"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


def font(size=11, bold=False, color="2F3A48"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def fill(color):
    return PatternFill("solid", fgColor=color)


def style_cell(cell, size=11, bold=False, color="2F3A48", bg=None, align="left", wrap=True):
    cell.font = font(size, bold, color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border = thin
    if bg:
        cell.fill = fill(bg)


def header_row(ws, row, values, bg=NAVY):
    for i, value in enumerate(values, 1):
        cell = ws.cell(row, i, value)
        style_cell(cell, 11, True, WHITE, bg, "center")
    ws.row_dimensions[row].height = 28


def landscape(ws, fit_height=1):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.page_setup.leftMargin = 0.4
    ws.page_setup.rightMargin = 0.4
    ws.page_setup.topMargin = 0.5
    ws.page_setup.bottomMargin = 0.5
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "โรงพยาบาลวชิระภูเก็ต · UAT งบประมาณ"
    ws.oddFooter.left.text = "เอกสารภายใน สำหรับทดสอบระบบ"
    ws.oddFooter.right.text = "หน้า &P / &N"


def build_cover(wb):
    ws = wb.active
    ws.title = "ปกแผน"
    ws.sheet_view.showGridLines = False
    for col, width in enumerate([4, 22, 28, 22, 18, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.merge_cells("B2:F2")
    ws["B2"].value = "โรงพยาบาลวชิระภูเก็ต"
    ws["B2"].font = font(14, True, TEAL)
    ws.merge_cells("B3:F3")
    ws["B3"].value = "ใบบันทึกผล UAT ระบบงบประมาณ"
    ws["B3"].font = font(22, True, NAVY)
    ws.merge_cells("B4:F4")
    ws["B4"].value = "ใช้คู่กับเอกสารแผนการทดสอบ และไฟล์นำเสนอ workflow"
    ws["B4"].font = font(12, False, "5B6878")

    info = [
        ("ระบบ", "Odoo 18 · VPK Budget"),
        ("ฐานข้อมูล", ""),
        ("รอบทดสอบ", "UAT ครั้งที่ 1"),
        ("วันที่เริ่ม", ""),
        ("วันที่จบ", ""),
        ("ผู้ประสาน", ""),
        ("ปีงบที่ทดสอบ", ""),
        ("หน่วยงานตัวอย่าง", ""),
    ]
    for i, (label, value) in enumerate(info):
        row = 6 + i
        label_cell = ws.cell(row, 2, label)
        style_cell(label_cell, 11, True, WHITE, NAVY)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        value_cell = ws.cell(row, 3, value)
        style_cell(value_cell, 11, False, "2F3A48", WHITE)

    ws.merge_cells("B16:F16")
    ws["B16"].value = "เกณฑ์ผ่านรอบนี้: เคส P1 ผ่านทั้งหมด · ไม่มี Blocker ค้าง · ผู้แทน 4 ฝ่ายลงนาม"
    ws["B16"].font = font(12, True, TEAL)

    header_row(ws, 18, ["ลำดับ", "ฝ่าย", "ผู้ทดสอบ", "ลงชื่อ", "วันที่", "ผ่านภาพรวม"])
    for i, role in enumerate(["หน่วยงาน", "งานงบประมาณ", "พัสดุ", "การเงิน"], 1):
        row = 18 + i
        style_cell(ws.cell(row, 1, i), align="center", bg=SOFT)
        style_cell(ws.cell(row, 2, role), bold=True)
        for col in range(3, 7):
            style_cell(ws.cell(row, col, ""))

    ws.merge_cells("B24:F27")
    ws["B24"].value = (
        "วิธีใช้\n"
        "1) กรอกปกนี้ก่อนเริ่ม   2) เดินเคสตามลำดับในชีต “บันทึกผลเคส”\n"
        "3) ถ้าไม่ผ่าน ให้เปิดรายการในชีต “บั๊ก”   4) ดูจำนวนผ่านที่ชีต “สรุป”\n"
        "ห้ามข้ามชุด A และ B ก่อนทำชุด D เพราะขั้นใช้งบต้องมีวงเงินที่อนุมัติแล้ว"
    )
    ws["B24"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B24"].font = font(12)
    ws["B24"].fill = fill(SOFT)
    landscape(ws, 1)


def build_cases(wb):
    ws = wb.create_sheet("บันทึกผลเคส")
    headers = [
        "รหัส", "ชุด", "P", "หัวข้อ", "ขั้นตอนโดยย่อ", "ผลที่ต้องเห็น",
        "ผล", "ผู้ทดสอบ", "วันที่", "หลักฐาน/หมายเหตุ",
    ]
    widths = [10, 16, 6, 30, 50, 48, 12, 16, 14, 28]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    header_row(ws, 1, headers)
    cases = [
        ("A01", "A ข้อมูลหลัก", "P1", "เมนูแบบฟอร์มเหลือ 4 ใบ",
         "เข้าเมนูคำของบหน่วยงาน ดูรายการแบบฟอร์ม",
         "เห็นเฉพาะ ตั้งงบวัสดุ / ครุภัณฑ์ / ก่อสร้าง / โครงการ"),
        ("A02", "A ข้อมูลหลัก", "P1", "เปิดแบบฟอร์มวัสดุ",
         "สร้างคำของบจากแบบฟอร์มตั้งงบวัสดุ",
         "เป็นหน้าวัสดุ มีที่กรอกรายละเอียดสินค้า"),
        ("A03", "A ข้อมูลหลัก", "P1", "เปิด 3 แบบฟอร์มที่เหลือ",
         "เปิดครุภัณฑ์ ก่อสร้าง โครงการ",
         "แต่ละใบคนละหน้าจอ ไม่ปนแท็บวัสดุ"),
        ("A04", "A ข้อมูลหลัก", "P1", "แหล่งเงินเลือกได้",
         "เปิดคำของบ แล้วเปิดช่องแหล่งเงิน",
         "มีแหล่งเงินที่ใช้จริง เช่น เงินบำรุง งบประมาณ UC"),
        ("A05", "A ข้อมูลหลัก", "P1", "แผนประจำปีพร้อมใช้",
         "เลือกแผนงบประมาณประจำปี",
         "มีแผนปีที่ทดสอบ และมีช่วงวันที่"),
        ("A06", "A ข้อมูลหลัก", "P1", "สิทธิ์หน่วยงาน",
         "login user หน่วยงาน แล้วหาเมนูกำหนดค่า master",
         "เข้าแก้ไข master ไม่ได้"),
        ("A07", "A ข้อมูลหลัก", "P2", "สิทธิ์งานงบ",
         "login user งานงบ เข้าเมนูกำหนดค่า",
         "แก้แหล่งเงิน / แบบฟอร์ม / กลุ่มวัสดุได้"),
        ("B01", "B ตั้งงบผ่าน", "P1", "หน่วยงานส่งคำของบวัสดุ",
         "สร้างใบวัสดุ กรอกเหตุผล รายการ รายละเอียดสินค้า ยอดขอ เช่น 10,000 แล้วส่ง",
         "สถานะ Submitted to Budget Unit"),
        ("B02", "B ตั้งงบผ่าน", "P1", "งานงบรับเรื่อง",
         "login งานงบ เปิดใบจาก B01 กดรับเรื่อง",
         "สถานะ Received by Budget Unit"),
        ("B03", "B ตั้งงบผ่าน", "P1", "ใส่ยอดจัดสรร",
         "เริ่มพิจารณา ใส่ยอดจัดสรรครบทุกบรรทัด น้อยกว่ายอดขอ เช่น 8,000",
         "บันทึกได้"),
        ("B04", "B ตั้งงบผ่าน", "P1", "อนุมัติวัสดุเข้าแผน",
         "กดอนุมัติ",
         "สถานะ Approved มีบรรทัดในแผนปี ยอดใช้จ่าย = ยอดจัดสรร"),
        ("B05", "B ตั้งงบผ่าน", "P1", "อนุมัติครุภัณฑ์",
         "ทำ B01–B04 ซ้ำด้วยแบบฟอร์มครุภัณฑ์",
         "เข้าแผนหมวดครุภัณฑ์"),
        ("B06", "B ตั้งงบผ่าน", "P1", "อนุมัติก่อสร้าง",
         "ทำ B01–B04 ซ้ำด้วยแบบฟอร์ม ก่อสร้าง",
         "เข้าแผนหมวดก่อสร้าง"),
        ("B07", "B ตั้งงบผ่าน", "P1", "อนุมัติโครงการ",
         "ทำ B01–B04 ซ้ำด้วยแบบฟอร์มโครงการ",
         "เข้าแผนหมวดโครงการ"),
        ("B08", "B ตั้งงบผ่าน", "P2", "หน่วยงานแก้ใบที่อนุมัติแล้ว",
         "login หน่วยงาน เปิดใบ Approved แล้วพยายามแก้รายการ",
         "แก้รายการหลักไม่ได้"),
        ("C01", "C กติกากันได้", "P1", "ส่งโดยไม่มีเหตุผล",
         "สร้างคำของบวัสดุ ไม่ใส่เหตุผล แล้วส่ง",
         "ส่งไม่ได้ ระบบบังคับเหตุผล"),
        ("C02", "C กติกากันได้", "P1", "วัสดุไม่มีรายละเอียดสินค้า",
         "มีบรรทัดงบ แต่ไม่มีรายละเอียดสินค้า แล้วส่ง",
         "ส่งไม่ได้"),
        ("C03", "C กติกากันได้", "P1", "อนุมัติทั้งที่จัดสรรว่าง",
         "งานงบพยายามอนุมัติทั้งที่ยอดจัดสรรบางบรรทัดว่าง",
         "อนุมัติไม่ได้"),
        ("C04", "C กติกากันได้", "P1", "ตีกลับคำขอ",
         "งานงบตีกลับ",
         "สถานะ Rejected และไม่มีวงเงินเข้าแผน"),
        ("C05", "C กติกากันได้", "P1", "แก้ใบที่ถูกตีกลับแล้วส่งใหม่",
         "ตั้งเป็นร่าง แก้ แล้วส่งใหม่",
         "ส่งได้ เข้าคิวงานงบใหม่"),
        ("C06", "C กติกากันได้", "P2", "ดึงใบอนุมัติกลับเป็นร่าง",
         "อนุมัติแล้วตั้งกลับเป็นร่าง",
         "บรรทัดงบที่ซิงก์เข้าแผนของใบนั้นถูกลบ"),
        ("D01", "D ใช้งบ", "P1", "เช็คงบ PR ในวงเงิน",
         "สร้าง PR วัสดุ มิติตรงแผน ยอด 3,000 แล้วกดเช็คงบ",
         "งบเพียงพอ ยังไม่กันเงิน"),
        ("D02", "D ใช้งบ", "P1", "ส่ง PR แล้วจองงบ",
         "ส่งอนุมัติ PR จาก D01",
         "จอง 3,000 คงเหลือ = จัดสรร − 3,000"),
        ("D03", "D ใช้งบ", "P1", "PR เกินคงเหลือ",
         "สร้าง PR ใหม่ยอดเกินคงเหลือ แล้วเช็คงบ / ส่ง",
         "งบไม่พอ และส่งอนุมัติไม่ได้"),
        ("D04", "D ใช้งบ", "P1", "ยืนยัน PO",
         "สร้างและยืนยัน PO จาก PR ใน D01 ทั้งจำนวน",
         "ยอดย้ายจากจอง PR เป็นผูกพัน PO"),
        ("D05", "D ใช้งบ", "P1", "โพสต์ใบตั้งเจ้าหนี้",
         "ตั้งเจ้าหนี้ตาม PO แล้วโพสต์",
         "ผูกพันลด ใช้จริงเพิ่ม ตามยอดบิล"),
        ("D06", "D ใช้งบ", "P1", "เทียบสูตรคงเหลือ",
         "คำนวณมือ: จัดสรร − ใช้จริง − จอง PR − ผูกพัน PO เทียบหน้าจอ",
         "ตัวเลขตรงกัน"),
        ("D07", "D ใช้งบ", "P2", "ทิ้ง PR ฉบับร่าง",
         "เช็คงบแล้วแต่ยังไม่ส่ง จากนั้นยกเลิก/ไม่ใช้",
         "ไม่กันเงิน"),
        ("D08", "D ใช้งบ", "P2", "ยกเลิก PR ที่จองแล้ว",
         "ยกเลิก PR ที่ส่งไปแล้ว",
         "คืนวงเงินที่จอง"),
        ("E01", "E จับคู่มิติ", "P1", "PR หน่วยงานไม่ตรงแผน",
         "เลือกหน่วยงานคนละแผนกกับงบที่อนุมัติ แล้วเช็คงบ",
         "ไม่เจอวงเงิน หรืองบไม่พอ"),
        ("E02", "E จับคู่มิติ", "P1", "PR แหล่งเงินไม่ตรงแผน",
         "เลือกแหล่งเงินคนละแหล่ง แล้วเช็คงบ",
         "ไม่เจอวงเงิน หรืองบไม่พอ"),
        ("E03", "E จับคู่มิติ", "P1", "PR วัสดุไม่เลือกกลุ่มวัสดุ",
         "สร้าง PR วัสดุเว้นกลุ่มวัสดุ",
         "บังคับกรอก หรือจับคู่ไม่เจอ"),
        ("E04", "E จับคู่มิติ", "P2", "PR ครุภัณฑ์ไม่ต้องมีกลุ่มวัสดุ",
         "สร้าง PR ครุภัณฑ์ ไม่ใส่กลุ่มวัสดุ แล้วเช็คงบ",
         "เช็คงบได้"),
        ("E05", "E จับคู่มิติ", "P2", "ห้ามใช้งบข้ามแหล่งเงิน",
         "มีงบเหลือแหล่ง B แต่ซื้อด้วยแหล่ง A",
         "ใช้เงิน B ไม่ได้"),
        ("E06", "E จับคู่มิติ", "P2", "PO บางส่วนของ PR",
         "ยืนยัน PO น้อยกว่ายอด PR",
         "จอง PR ลดเฉพาะส่วนที่ PO ครอบคลุม"),
        ("F01", "F สรุปผล", "P1", "Dashboard รายปี",
         "เปิดภาพรวมรายจ่ายประจำปี",
         "เห็นจัดสรร ใช้จริง คงเหลือ ของปีที่ทดสอบ"),
        ("F02", "F สรุปผล", "P2", "Dashboard รายแผนก",
         "เปิดรายแผนกของหน่วยงานที่ทดสอบ",
         "ยอดสอดคล้องชุด D"),
        ("F03", "F สรุปผล", "P2", "สรุปคำของบแยกแบบฟอร์ม",
         "เปิดวิซาร์ดสรุป เลือกแบบฟอร์มวัสดุ",
         "รวมเฉพาะใบวัสดุ"),
        ("F04", "F สรุปผล", "P2", "ส่งออกรายงานวัสดุ",
         "พิมพ์หรือส่งออกถ้าใช้ในรอบนี้",
         "เปิดได้ ยอดรวมตรง"),
    ]
    for r, row in enumerate(cases, 2):
        for col, value in enumerate(row, 1):
            bg = P1 if (col == 3 and value == "P1") else P2 if (col == 3 and value == "P2") else None
            align = "center" if col in (1, 3, 7, 8, 9) else "left"
            style_cell(ws.cell(r, col, value), bg=bg, align=align)
        ws.row_dimensions[r].height = 46
    dv = DataValidation(type="list", formula1='"ผ่าน,ไม่ผ่าน,ข้าม"', allow_blank=True)
    dv.error = "เลือก ผ่าน / ไม่ผ่าน / ข้าม"
    dv.errorTitle = "ผลทดสอบ"
    ws.add_data_validation(dv)
    dv.add("G2:G80")
    ws.conditional_formatting.add(
        "G2:G80",
        CellIsRule(operator="equal", formula=['"ผ่าน"'], fill=fill(PASS), font=font(11, True, GREEN)),
    )
    ws.conditional_formatting.add(
        "G2:G80",
        CellIsRule(operator="equal", formula=['"ไม่ผ่าน"'], fill=fill(FAIL), font=font(11, True, RED)),
    )
    ws.conditional_formatting.add(
        "G2:G80",
        CellIsRule(operator="equal", formula=['"ข้าม"'], fill=fill(SKIP), font=font(11, False, "5B6878")),
    )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:J40"
    landscape(ws, 0)
    return len(cases)


def build_bugs(wb):
    ws = wb.create_sheet("บั๊ก")
    headers = [
        "รหัสบั๊ก", "เคสที่พบ", "ระดับ", "สรุปปัญหา", "ขั้นตอนที่เกิด",
        "ผลที่คาด", "ผลที่เกิดจริง", "ผู้พบ", "วันที่", "สถานะแก้",
    ]
    widths = [12, 12, 14, 28, 32, 26, 26, 14, 14, 18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    header_row(ws, 1, headers, RED)
    for row in range(2, 21):
        for col in range(1, 11):
            style_cell(ws.cell(row, col, f"BUG-{row-1:02d}" if col == 1 else ""), align="center" if col in (1, 3, 9, 10) else "left")
    sev = DataValidation(type="list", formula1='"Blocker,Major,Minor,Cosmetic"', allow_blank=True)
    st = DataValidation(type="list", formula1='"เปิด,กำลังแก้,แก้แล้วรอทดสอบซ้ำ,ปิด"', allow_blank=True)
    ws.add_data_validation(sev)
    ws.add_data_validation(st)
    sev.add("C2:C20")
    st.add("J2:J20")
    ws.freeze_panes = "A2"
    landscape(ws, 1)


def build_summary(wb):
    ws = wb.create_sheet("สรุป")
    for col, width in enumerate([6, 22, 12, 12, 12, 12, 12, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.merge_cells("B2:H2")
    ws["B2"].value = "สรุปผลการทดสอบระบบงบประมาณ"
    ws["B2"].font = font(18, True, NAVY)
    header_row(ws, 4, ["", "ชุด", "ทั้งหมด", "ผ่าน", "ไม่ผ่าน", "ข้าม", "ยังไม่ทำ", "% ผ่าน P1"])
    groups = [
        ("A ข้อมูลหลัก", "A"),
        ("B ตั้งงบผ่าน", "B"),
        ("C กติกากันได้", "C"),
        ("D ใช้งบ", "D"),
        ("E จับคู่มิติ", "E"),
        ("F สรุปผล", "F"),
    ]
    for i, (name, prefix) in enumerate(groups):
        row = 5 + i
        style_cell(ws.cell(row, 1, i + 1), align="center", bg=SOFT)
        style_cell(ws.cell(row, 2, name), bold=True)
        formulas = {
            3: f'=COUNTIF(\'บันทึกผลเคส\'!A:A,"{prefix}*")',
            4: f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",\'บันทึกผลเคส\'!G:G,"ผ่าน")',
            5: f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",\'บันทึกผลเคส\'!G:G,"ไม่ผ่าน")',
            6: f'=COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",\'บันทึกผลเคส\'!G:G,"ข้าม")',
            7: f"=C{row}-D{row}-E{row}-F{row}",
            8: (
                f'=IFERROR(COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",'
                f'\'บันทึกผลเคส\'!C:C,"P1",\'บันทึกผลเคส\'!G:G,"ผ่าน")/'
                f'COUNTIFS(\'บันทึกผลเคส\'!A:A,"{prefix}*",\'บันทึกผลเคส\'!C:C,"P1"),0)'
            ),
        }
        for col, formula in formulas.items():
            style_cell(ws.cell(row, col, formula), align="center")
        ws.cell(row, 8).number_format = "0%"

    row = 12
    style_cell(ws.cell(row, 2, "รวมทุกชุด"), 12, True, WHITE, NAVY)
    for col in range(3, 8):
        letter = get_column_letter(col)
        style_cell(ws.cell(row, col, f"=SUM({letter}5:{letter}10)"), 12, True, WHITE, NAVY, "center")
    style_cell(
        ws.cell(
            row,
            8,
            '=IFERROR(COUNTIFS(\'บันทึกผลเคส\'!C:C,"P1",\'บันทึกผลเคส\'!G:G,"ผ่าน")/'
            'COUNTIF(\'บันทึกผลเคส\'!C:C,"P1"),0)',
        ),
        12,
        True,
        WHITE,
        NAVY,
        "center",
    )
    ws.cell(row, 8).number_format = "0%"

    ws.merge_cells("B15:H18")
    ws["B15"].value = (
        "ตัดสินใจ:  ผ่าน UAT และไป Go-Live ได้   /   ผ่านแบบมีเงื่อนไข   /   ต้องทดสอบซ้ำ\n"
        "เงื่อนไขที่ค้าง: .................................................................\n"
        "ลงนามผู้ประสาน: ....................................     วันที่: ................"
    )
    ws["B15"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B15"].font = font(12)
    landscape(ws, 1)


def main():
    wb = Workbook()
    build_cover(wb)
    count = build_cases(wb)
    build_bugs(wb)
    build_summary(wb)
    wb.save(OUT)
    print(count, OUT)


if __name__ == "__main__":
    main()
