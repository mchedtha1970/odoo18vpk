#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""สรุป Data Migration: master data ที่นำเข้าแล้ว แยกตามโมดูล จากฐาน VPK-S1"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = Path("/opt/odoo18vpk/docs/UAT-Data-Migrate")
DBNAME = "VPK-S1"
FONT = "Prompt"
NAVY, TEAL, GREEN, RED = "123A56", "0F7A72", "2C7A45", "B4433A"
LINE, WHITE, SOFT, PASS, FAIL, AMBER = (
    "D7DEE6",
    "FFFFFF",
    "E5F5F3",
    "E5F5EA",
    "FBECEA",
    "F4E4C8",
)
PALE = "F4F7FA"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
TH_MONTHS = "ม.ค. ก.พ. มี.ค. เม.ย. พ.ค. มิ.ย. ก.ค. ส.ค. ก.ย. ต.ค. พ.ย. ธ.ค.".split()

REAL_VENDOR_SQL = """
    p.supplier_rank > 0
    AND p.parent_id IS NULL
    AND p.name NOT ILIKE '%ตัวอย่าง%'
    AND p.name NOT ILIKE '%DEMO%'
    AND p.name NOT ILIKE '%เอกสารทางการค้า%'
    AND COALESCE(p.ref, '') NOT IN (
        'VENDOR-DEMO-001', 'EXT-001', 'EXT-002', 'HTTP-EXT-100',
        'DOC-EXT-200', 'SUP-PRICE-001', 'SUP-PRICE-002', 'VENDOR-RED-CROSS-SAMPLE'
    )
    AND COALESCE(p.vpk_vendor_external_id, '') NOT IN (
        'EXT-001', 'EXT-002', 'HTTP-EXT-100', 'DOC-EXT-200', 'TRCS'
    )
"""

DIRECT_VENDOR_SQL = f"""
    {REAL_VENDOR_SQL}
    AND EXISTS (
        SELECT 1 FROM procurement_vendor_group g
        WHERE g.id = p.procurement_vendor_group_id AND g.code = 'DIRECT'
    )
"""

PURCHASE_VENDOR_SQL = f"""
    {REAL_VENDOR_SQL}
    AND NOT EXISTS (
        SELECT 1 FROM procurement_vendor_group g
        WHERE g.id = p.procurement_vendor_group_id AND g.code = 'DIRECT'
    )
"""

MODULES = [
    ("budget", "งบประมาณ", "vpk_budget · บัญชีวิเคราะห์"),
    ("purchase", "จัดซื้อ", "Purchase · ผู้จำหน่ายจัดซื้อ"),
    ("stock", "คลังสินค้า / สินค้า", "Inventory · Product Master"),
    ("finance", "บัญชี / การเงิน", "Accounting · เจ้าหนี้ตรง · CoA"),
    ("asset", "สินทรัพย์ถาวร", "account.asset · ครุภัณฑ์"),
    ("outstanding", "ข้อมูลคงค้าง", "PO คงค้าง · เจ้าหนี้คงค้าง · TB ปีงบ 2569"),
]


def fnt(size=11, bold=False, color="2F3A48"):
    return Font(name=FONT, size=size, bold=bold, color=color)


def fill(c):
    return PatternFill("solid", fgColor=c)


def sc(cell, size=11, bold=False, color="2F3A48", bg=None, align="left"):
    cell.font = fnt(size, bold, color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = thin
    if bg:
        cell.fill = fill(bg)


def hdr(ws, row, values, bg=NAVY):
    for i, v in enumerate(values, 1):
        sc(ws.cell(row, i, v), 11, True, WHITE, bg, "center")
    ws.row_dimensions[row].height = 28


def thai_dt(dt: datetime) -> str:
    return f"{dt.day} {TH_MONTHS[dt.month - 1]} {dt.year + 543}  {dt.strftime('%H:%M')}"


def dated_xlsx(name: str, dt: datetime | None = None) -> str:
    dt = dt or datetime.now()
    p = Path(name)
    return f"{p.stem}-{dt.year + 543}-{dt.month:02d}-{dt.day:02d}{p.suffix}"


def jname(alias="name"):
    return f"COALESCE({alias}->>'th_TH', {alias}->>'en_US', {alias}::text)"


def page(ws, title):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.leftMargin = ws.page_setup.rightMargin = 0.35
    ws.page_setup.topMargin = ws.page_setup.bottomMargin = 0.45
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = f"โรงพยาบาลวชิระภูเก็ต · {title}"
    ws.oddFooter.left.text = "สรุป Data Migration · Master Data + ข้อมูลคงค้าง"
    ws.oddFooter.right.text = "หน้า &P / &N"


def connect():
    return psycopg2.connect(
        dbname=DBNAME,
        user="odoo18vpk",
        password="odoo18vpk",
        host="/var/run/postgresql",
        connect_timeout=15,
    )


def qcount(cur, sql):
    try:
        cur.execute(sql)
        row = cur.fetchone()
        if row is None:
            return 0
        return int(next(iter(row.values() if isinstance(row, dict) else row)))
    except Exception as exc:
        print(f"COUNT FAIL: {exc}\n  {sql[:120]}")
        cur.connection.rollback()
        return None


def qrows(cur, sql):
    try:
        cur.execute(sql)
        return cur.fetchall()
    except Exception as exc:
        print(f"LIST FAIL: {exc}\n  {sql[:120]}")
        cur.connection.rollback()
        return []


def status_of(n, kind="Master"):
    if n is None:
        return "ไม่มีตาราง / ยังไม่ติดตั้ง"
    if kind == "คงค้าง":
        if n <= 0:
            return "ยังไม่นำเข้า (คงค้าง)"
        return "มีรายการคงค้าง"
    if n <= 0:
        return "ยังไม่นำเข้า"
    return "นำเข้าแล้ว"


def status_bg(st):
    if st == "นำเข้าแล้ว":
        return PASS
    if st == "มีรายการคงค้าง":
        return AMBER
    if "ยังไม่" in st:
        return FAIL
    return AMBER


MASTERS = [
    # ---- งบประมาณ ----
    dict(
        mod="budget",
        name="แหล่งเงิน",
        model="vpk.budget.fund.source",
        key="code",
        how="XML ติดตั้งโมดูล + Wizard นำเข้า Excel (FundSource)",
        source="vpk_budget/data/budget_fund_source_data.xml",
        count="SELECT COUNT(*) FROM vpk_budget_fund_source",
    ),
    dict(
        mod="budget",
        name="กลุ่มงบประมาณ-วัสดุ",
        model="vpk.budget.group",
        key="code",
        how="XML ติดตั้งโมดูล + Wizard นำเข้า Excel (BudgetGroup)",
        source="vpk_budget/data/budget_group_data.xml",
        count="SELECT COUNT(*) FROM vpk_budget_group",
    ),
    dict(
        mod="budget",
        name="ประเภทวัสดุย่อย",
        model="vpk.budget.material.sub.type",
        key="code + กลุ่มงบ",
        how="Wizard นำเข้า Excel (MaterialSubType)",
        source="เมนู งบประมาณ → นำเข้าข้อมูลหลัก",
        count="SELECT COUNT(*) FROM vpk_budget_material_sub_type",
    ),
    dict(
        mod="budget",
        name="ประเภทงบประมาณ",
        model="vpk.budget.type",
        key="code",
        how="XML ติดตั้งโมดูล + Wizard นำเข้า Excel (BudgetType)",
        source="vpk_budget/data/budget_type_data.xml",
        count="SELECT COUNT(*) FROM vpk_budget_type",
    ),
    dict(
        mod="budget",
        name="กลุ่มครุภัณฑ์ (งบประมาณ)",
        model="vpk.budget.asset.category",
        key="code",
        how="XML ติดตั้งโมดูล + Wizard นำเข้า Excel (AssetCategory)",
        source="vpk_budget/data/budget_asset_category_data.xml",
        count="SELECT COUNT(*) FROM vpk_budget_asset_category",
    ),
    dict(
        mod="budget",
        name="แบบฟอร์มคำของบ",
        model="vpk.budget.request.form.type",
        key="code",
        how="XML ติดตั้งโมดูล + Wizard นำเข้า Excel (FormType)",
        source="vpk_budget/data/budget_request_form_type_data.xml",
        count="SELECT COUNT(*) FROM vpk_budget_request_form_type",
    ),
    dict(
        mod="budget",
        name="ประเภทงบในแบบฟอร์ม",
        model="vpk.budget.request.form.type.line",
        key="form + ประเภทงบ",
        how="XML ติดตั้งโมดูล + Wizard นำเข้า Excel (FormTypeLine)",
        source="vpk_budget/data/budget_request_form_type_line_data.xml",
        count="SELECT COUNT(*) FROM vpk_budget_request_form_type_line",
    ),
    dict(
        mod="budget",
        name="แผนวิเคราะห์ / มิติหน่วยงาน",
        model="account.analytic.plan",
        key="name",
        how="แม่แบบ Excel Analytic / สร้างในระบบ",
        source="vpk_master_data_templates · ชีต AnalyticPlan",
        count="SELECT COUNT(*) FROM account_analytic_plan",
    ),
    dict(
        mod="budget",
        name="บัญชีวิเคราะห์ / หน่วยงาน",
        model="account.analytic.account",
        key="code",
        how="แม่แบบ Excel Analytic / สร้างในระบบ",
        source="vpk_master_data_templates · ชีต AnalyticAccount",
        count="SELECT COUNT(*) FROM account_analytic_account",
    ),
    dict(
        mod="budget",
        name="แผนงบประมาณประจำปี",
        model="budget.budget",
        key="name + ช่วงวันที่",
        how="สร้างในระบบงบประมาณ",
        source="เมนู งบประมาณ → แผนงบประมาณ",
        count="SELECT COUNT(*) FROM budget_budget",
    ),
    dict(
        mod="budget",
        name="หมวดงบประมาณ",
        model="account.budget.post",
        key="name",
        how="สร้างอัตโนมัติจากประเภทวัสดุย่อย / ตั้งค่าแผนงบ",
        source="ระบบงบประมาณ (account.budget.post)",
        count="SELECT COUNT(*) FROM account_budget_post",
    ),
    # ---- จัดซื้อ (ไม่รวมเจ้าหนี้ตรง — เจ้าหนี้ตรงอยู่โมดูลบัญชี) ----
    dict(
        mod="purchase",
        name="ผู้จำหน่ายจัดซื้อ (ไม่รวมเจ้าหนี้ตรง)",
        model="res.partner (supplier, ไม่ใช่กลุ่ม DIRECT)",
        key="ref / vat / ชื่อ",
        how="สคริปต์นำเข้า JSON + กลุ่มผู้จำหน่ายจัดซื้อ (ไม่รวมเจ้าหนี้ตรง)",
        source="import/import_vendor_master.py · import_utility_vendors.py",
        count=f"SELECT COUNT(*) FROM res_partner p WHERE {PURCHASE_VENDOR_SQL}",
    ),
    dict(
        mod="purchase",
        name="กลุ่มผู้จำหน่ายจัดซื้อ (ไม่รวมเจ้าหนี้ตรง)",
        model="procurement.vendor.group",
        key="code",
        how="XML โมดูลจัดซื้อ",
        source="vpk_procurement_auto_pr/data/vendor_group_data.xml",
        count="SELECT COUNT(*) FROM procurement_vendor_group WHERE code <> 'DIRECT'",
    ),
    dict(
        mod="purchase",
        name="บัญชีธนาคารผู้ขาย",
        model="res.partner.bank",
        key="partner + เลขบัญชี",
        how="นำเข้าพร้อมผู้จำหน่ายจัดซื้อ",
        source="import/import_vendor_master.py",
        count=f"""
            SELECT COUNT(*) FROM res_partner_bank b
            JOIN res_partner p ON p.id = b.partner_id
            WHERE {PURCHASE_VENDOR_SQL}
        """,
    ),
    dict(
        mod="purchase",
        name="ราคา / MOQ ผู้ขาย",
        model="product.supplierinfo",
        key="ผู้ขาย + สินค้า",
        how="แม่แบบ Excel VendorPrice / สร้างในระบบ",
        source="vpk_master_data_templates · ชีต VendorPrice",
        count=f"""
            SELECT COUNT(*) FROM product_supplierinfo s
            LEFT JOIN res_partner p ON p.id = s.partner_id
            WHERE p.id IS NULL OR ({REAL_VENDOR_SQL})
        """,
    ),
    # ---- คลัง ----
    dict(
        mod="stock",
        name="หมวดหน่วยนับ",
        model="uom.category",
        key="name",
        how="มาตรฐาน Odoo + สร้างเพิ่มตอนนำเข้าสินค้า",
        source="Odoo UoM + สคริปต์นำเข้าสินค้า",
        count="SELECT COUNT(*) FROM uom_category",
    ),
    dict(
        mod="stock",
        name="หน่วยนับ",
        model="uom.uom",
        key="name + หมวด",
        how="มาตรฐาน Odoo + สร้างบรรจุภัณฑ์ตอนนำเข้าสินค้า",
        source="Odoo UoM + สคริปต์นำเข้าสินค้า (กล่อง/แพ็ค)",
        count="SELECT COUNT(*) FROM uom_uom",
    ),
    dict(
        mod="stock",
        name="หมวดหมู่สินค้า",
        model="product.category",
        key="complete_name",
        how="สคริปต์นำเข้าสินค้า / แคตตาล็อกครุภัณฑ์",
        source="สร้างตามไฟล์ต้นทางแต่ละกลุ่มวัสดุ",
        count="SELECT COUNT(*) FROM product_category",
    ),
    dict(
        mod="stock",
        name="สินค้า (Product Master)",
        model="product.template",
        key="default_code",
        how="สคริปต์ import/*.py + Wizard นำเข้าสินค้า",
        source="ดูชีต “สินค้าแยกหมวด / แหล่งนำเข้า”",
        count="SELECT COUNT(*) FROM product_template",
    ),
    dict(
        mod="stock",
        name="คลังสินค้า (เปิดใช้ · รับเข้า 1 ขั้น)",
        model="stock.warehouse",
        key="code",
        how="สคริปต์นำเข้าคลัง + ตั้งค่าคลังยา / ธนาคารเลือด",
        source="import/import_warehouses.py ← vpk_data_warehouse_template.xlsx",
        count="""
            SELECT COUNT(*) FROM stock_warehouse
            WHERE active AND reception_steps = 'one_step'
        """,
    ),
    dict(
        mod="stock",
        name="ตำแหน่งจัดเก็บ (internal · เปิดใช้)",
        model="stock.location",
        key="complete_name",
        how="สร้างอัตโนมัติจากคลัง + สคริปต์คลังยา (bin)",
        source="import/setup_pharmacy_warehouses.py · import_warehouses.py",
        count="""
            SELECT COUNT(*) FROM stock_location loc
            LEFT JOIN stock_warehouse wh ON wh.id = loc.warehouse_id
            WHERE loc.usage = 'internal'
              AND loc.active
              AND (wh.id IS NULL OR (wh.active AND wh.reception_steps = 'one_step'))
        """,
    ),
    # ---- บัญชี ----
    dict(
        mod="finance",
        name="ผังบัญชี",
        model="account.account",
        key="code",
        how="สคริปต์นำเข้า Excel ผังบัญชี รพ.",
        source="import/import_coa_vpk.py ← coa-vpk.xlsx",
        count="SELECT COUNT(*) FROM account_account",
    ),
    dict(
        mod="finance",
        name="สมุดรายวัน",
        model="account.journal",
        key="code",
        how="ตั้งค่าบัญชีมาตรฐาน + แม่แบบ Excel Journal",
        source="Odoo Accounting setup · ชีต Journal",
        count="SELECT COUNT(*) FROM account_journal",
    ),
    dict(
        mod="finance",
        name="ภาษี (VAT)",
        model="account.tax",
        key="name",
        how="ตั้งค่าบัญชีมาตรฐาน / แม่แบบ Excel Tax",
        source="Odoo Accounting setup · ชีต Tax",
        count="SELECT COUNT(*) FROM account_tax",
    ),
    dict(
        mod="finance",
        name="ภาษีหัก ณ ที่จ่าย (WHT)",
        model="account.withholding.tax",
        key="name",
        how="โมดูล l10n_th / ตั้งค่าบัญชี",
        source="l10n_th withholding tax",
        count="SELECT COUNT(*) FROM account_withholding_tax",
    ),
    dict(
        mod="finance",
        name="เงื่อนไขการชำระ",
        model="account.payment.term",
        key="name",
        how="ตั้งค่าบัญชีมาตรฐาน / แม่แบบ Excel PaymentTerm",
        source="Odoo Accounting setup · ชีต PaymentTerm",
        count="SELECT COUNT(*) FROM account_payment_term",
    ),
    dict(
        mod="finance",
        name="เจ้าหนี้ตรง",
        model="res.partner (กลุ่ม DIRECT)",
        key="ref / ชื่อ",
        how="สคริปต์นำเข้าหน่วยงานราชการ / โรงพยาบาล / กองทุน (กลุ่มเจ้าหนี้ตรง)",
        source="import/import_direct_vendors.py · import_gov_hospital_vendors.py",
        count=f"SELECT COUNT(*) FROM res_partner p WHERE {DIRECT_VENDOR_SQL}",
    ),
    dict(
        mod="finance",
        name="กลุ่มเจ้าหนี้ตรง",
        model="procurement.vendor.group (DIRECT)",
        key="code = DIRECT",
        how="XML โมดูลจัดซื้อ ใช้ฝั่งบัญชีเป็นกลุ่มเจ้าหนี้จ่ายตรง",
        source="vpk_procurement_auto_pr/data/vendor_group_data.xml (vendor_group_direct)",
        count="SELECT COUNT(*) FROM procurement_vendor_group WHERE code = 'DIRECT'",
    ),
    dict(
        mod="finance",
        name="บัญชีธนาคารของโรงพยาบาล",
        model="res.partner.bank (บริษัท) + account.journal BK*",
        key="เลขบัญชี / รหัสสมุดรายวัน",
        how="สคริปต์นำเข้าบัญชีกองทุน รพ.",
        source="import/import_company_bank_accounts.py",
        count="""
            SELECT COUNT(*) FROM res_partner_bank pb
            JOIN res_company c ON c.partner_id = pb.partner_id
        """,
    ),
    # ---- สินทรัพย์ ----
    dict(
        mod="asset",
        name="กลุ่มสินทรัพย์",
        model="account.asset.group",
        key="code",
        how="แม่แบบ Excel AssetGroup / สร้างในระบบ",
        source="vpk_master_data_templates · ชีต AssetGroup",
        count="SELECT COUNT(*) FROM account_asset_group",
    ),
    dict(
        mod="asset",
        name="โปรไฟล์สินทรัพย์ถาวร",
        model="account.asset.profile",
        key="name",
        how="สคริปต์นำเข้า Excel ทะเบียนครุภัณฑ์",
        source="import/import_assets.py ← vpk_master_asset_v6.0.1.xlsx (AssetProfile)",
        count="SELECT COUNT(*) FROM account_asset_profile",
    ),
    dict(
        mod="asset",
        name="สถานะย่อยสินทรัพย์",
        model="account.asset.sub.state",
        key="name",
        how="ตั้งค่าโมดูลสินทรัพย์",
        source="account_asset_management",
        count="SELECT COUNT(*) FROM account_asset_sub_state",
    ),
    dict(
        mod="asset",
        name="ชุดสินทรัพย์ (Asset Set)",
        model="account.asset.parent",
        key="code",
        how="แม่แบบ Excel / สร้างในระบบ",
        source="account.asset.parent",
        count="SELECT COUNT(*) FROM account_asset_parent",
    ),
    dict(
        mod="asset",
        name="บัตรสินทรัพย์ / ครุภัณฑ์",
        model="account.asset",
        key="number / code",
        how="สคริปต์นำเข้า Excel ทะเบียนครุภัณฑ์",
        source="import/import_assets.py ← vpk_master_asset_v6.0.1.xlsx (AssetCard)",
        count="SELECT COUNT(*) FROM account_asset",
    ),
    dict(
        mod="asset",
        name="หมวดครุภัณฑ์ซ่อมบำรุง",
        model="maintenance.equipment.category",
        key="name",
        how="โมดูลซ่อมบำรุง / ลิงก์จากสินทรัพย์",
        source="maintenance",
        count="SELECT COUNT(*) FROM maintenance_equipment_category",
    ),
    dict(
        mod="asset",
        name="ครุภัณฑ์ซ่อมบำรุง (Equipment)",
        model="maintenance.equipment",
        key="asset_number / serial / ชื่อ",
        how="สร้างคู่กับบัตรสินทรัพย์ / แม่แบบ Excel Equipment",
        source="import_assets / vpk_master_data_templates",
        count="SELECT COUNT(*) FROM maintenance_equipment",
    ),
    # ---- ข้อมูลคงค้าง ----
    dict(
        mod="outstanding",
        kind="คงค้าง",
        name="PO คงค้าง (ค้างรับ / ค้างตั้งหนี้)",
        model="purchase.order (purchase/done ที่ยังไม่ครบ)",
        key="name",
        how="นำเข้าเอกสาร PO จากระบบเดิม (ยังไม่นำเข้า)",
        source="ดูชีต “คงค้าง”",
        count="""
            SELECT COUNT(*) FROM (
                SELECT po.id
                FROM purchase_order po
                LEFT JOIN purchase_order_line pol ON pol.order_id = po.id
                WHERE po.state IN ('purchase', 'done')
                GROUP BY po.id, po.invoice_status
                HAVING SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                THEN pol.product_qty - COALESCE(pol.qty_received, 0) ELSE 0 END) > 0
                    OR SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                THEN COALESCE(pol.qty_received, 0) - COALESCE(pol.qty_invoiced, 0) ELSE 0 END) > 0
                    OR po.invoice_status IS DISTINCT FROM 'invoiced'
            ) t
        """,
    ),
    dict(
        mod="outstanding",
        kind="คงค้าง",
        name="เจ้าหนี้คงค้าง (บิลที่ยังไม่จ่ายครบ)",
        model="account.move (in_invoice/in_refund posted, residual ≠ 0)",
        key="name / เจ้าหนี้",
        how="นำเข้าบิลค้างจ่ายจากระบบเดิม (ยังไม่นำเข้า)",
        source="ดูชีต “คงค้าง”",
        count="""
            SELECT COUNT(*) FROM account_move
            WHERE move_type IN ('in_invoice', 'in_refund')
              AND state = 'posted'
              AND COALESCE(amount_residual, 0) <> 0
        """,
    ),
    dict(
        mod="outstanding",
        kind="คงค้าง",
        name="งบทดลอง (TB) ปีงบประมาณ 2569",
        model="account.move.line posted · 1 ต.ค. 2568 – 30 ก.ย. 2569",
        key="รหัสบัญชี",
        how="นำเข้า TB / ยอดยกเข้าปีงบ 2569 (ยังไม่มียอดยกเข้าจากระบบเดิม)",
        source="ดูชีต “คงค้าง”",
        count="""
            SELECT COUNT(*) FROM (
                SELECT aml.account_id
                FROM account_move_line aml
                JOIN account_move am ON am.id = aml.move_id
                WHERE am.state = 'posted'
                  AND aml.account_id IS NOT NULL
                  AND COALESCE(aml.display_type, '') NOT IN ('line_section', 'line_note', 'cogs')
                  AND am.date BETWEEN DATE '2025-10-01' AND DATE '2026-09-30'
                GROUP BY aml.account_id
                HAVING SUM(aml.debit) <> 0 OR SUM(aml.credit) <> 0
            ) t
        """,
    ),
]

PRODUCT_SOURCE_HINTS = [
    (
        "ยา / คลังยาใหญ่",
        "01. เวชภัณฑ์ ยา",
        "import/import_drug_warehouse.py",
        "วัสดุ_ยา_คลังยาใหญ่.xlsx",
    ),
    (
        "เวชภัณฑ์มิใช่ยา (ชุดหลัก)",
        "02. เวชภัณฑ์มิใช่ยา",
        "import/import_medical_supplies.py",
        "vpk_data_none_medical.xlsx",
    ),
    (
        "เวชภัณฑ์มิใช่ยา (ตามหน่วยงาน)",
        "02. เวชภัณฑ์มิใช่ยา / 21–29",
        "import/import_department_nondrug_products.py",
        "Item_เวชภัณฑ์หมวดต่างๆ.xlsx",
    ),
    (
        "วัสดุวิทยาศาสตร์",
        "04. วัสดุวิทยาศาสตร์",
        "import/import_scientific_supplies.py",
        "วัสดุวิทยาศาสตร์.xlsx",
    ),
    (
        "วัสดุงานช่าง ก่อสร้าง/ไฟฟ้า/ยานพาหนะ",
        "07. ช่าง",
        "import/import_construction_craftsman_products.py",
        "รายงานคงคลังงานช่างก่อสร้าง_ปี_2569.xlsx",
    ),
    (
        "วัสดุบริโภค / โภชนาการ",
        "08. โภชนาการ",
        "import/import_nutrition_consumables.py",
        "รายการวัสดุบริโภค.xlsx",
    ),
    (
        "งานพิมพ์ / งานบ้านงานครัว / สำนักงาน",
        "16–18 ภายใต้กลุ่มวัสดุ",
        "import/import_print_house_office_supplies.py",
        "vpk_master_data_จัดซื้อ-คลังพัสดุ-งานพิมพ์-งานบ้าน-สำนักงาน.xlsx",
    ),
    (
        "ธนาคารเลือด / แล็บ / พยาธิวิทยา",
        "04 / 19. วัสดุธนาคารเลือด / 20. วัสดุพยาธิวิทยา",
        "import/import_blood_bank_lab_products.py",
        "vpk_master_data_ระบบคลัง_พัสดุ_templates_เพิ่มธนาคารเลือด.xlsx",
    ),
    (
        "แคตตาล็อกครุภัณฑ์ สป.สธ.",
        "ครุภัณฑ์ / ประเภท-กลุ่มเครื่องมือ",
        "import/import_equipment_catalog.py",
        "บัญชีรายการครุภัณฑ์_สำนักงานปลัด-8.xlsx",
    ),
    (
        "Wizard นำเข้าสินค้าทั่วไป",
        "ตาม categ_path ในไฟล์",
        "โมดูล vpk_product_import",
        "แม่แบบ Excel สินค้า (Product)",
    ),
]


def load_masters(cur):
    for ds in MASTERS:
        n = qcount(cur, ds["count"])
        ds["n"] = 0 if n is None else n
        ds["missing"] = n is None
        ds["status"] = status_of(n, ds.get("kind", "Master"))


def widths(ws, vals):
    for i, w in enumerate(vals, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_table_rows(ws, start, items, cols):
    for i, ds in enumerate(items):
        r = start + i
        n = ds["n"]
        st = ds["status"]
        bg = status_bg(st)
        zebra = PALE if i % 2 else WHITE
        values = []
        for c in cols:
            if c == "seq":
                values.append(i + 1)
            elif c == "n":
                values.append(n)
            else:
                values.append(ds.get(c, ""))
        for col, v in enumerate(values, 1):
            cell_bg = bg if cols[col - 1] in ("status",) else zebra
            align = "center" if cols[col - 1] in ("seq", "n", "status") else "left"
            sc(ws.cell(r, col, v), align=align, bg=cell_bg, bold=(cols[col - 1] == "n"))
            if cols[col - 1] == "n":
                ws.cell(r, col).number_format = "#,##0"
        ws.row_dimensions[r].height = 32
    return start + len(items) - 1


def write_cover(wb, snapshot, masters):
    ws = wb.active
    ws.title = "ปก"
    ws.sheet_view.showGridLines = False
    page(ws, "สรุป Data Migration")
    widths(ws, [4, 28, 42, 22, 18, 18])
    ws.merge_cells("B2:F2")
    ws["B2"].value, ws["B2"].font = "โรงพยาบาลวชิระภูเก็ต", fnt(14, True, TEAL)
    ws.merge_cells("B3:F3")
    ws["B3"].value, ws["B3"].font = "สรุป Data Migration · Master Data และข้อมูลคงค้าง", fnt(20, True, NAVY)
    ws.merge_cells("B4:F4")
    ws["B4"].value, ws["B4"].font = (
        "เอกสารภาพรวมว่าแต่ละโมดูลมีข้อมูลหลักอะไรบ้าง ที่นำเข้าแล้วกี่รายการ และมียอดคงค้างอะไรบ้าง"
    ), fnt(12, False, "5B6878")

    masters_only = [d for d in masters if d.get("kind", "Master") != "คงค้าง"]
    outstanding = [d for d in masters if d.get("kind") == "คงค้าง"]
    imported = [d for d in masters_only if d["n"] > 0]
    empty = [d for d in masters_only if d["n"] == 0]
    info = [
        ("ระบบ", "Odoo 18 · โรงพยาบาลวชิระภูเก็ต"),
        ("ฐานข้อมูล", f"{DBNAME}"),
        ("วัน-เวลาที่ดึงข้อมูล", snapshot),
        ("ขอบเขต", "Master Data + ข้อมูลคงค้าง (PO / เจ้าหนี้คงค้าง / TB ปีงบ 2569)"),
        ("เจ้าหนี้ตรง", "อยู่โมดูลบัญชี / การเงิน — ไม่นับในจัดซื้อ"),
        ("ประเภท Master", f"{len(masters_only)} ประเภท"),
        ("นำเข้าแล้ว", f"{len(imported)} ประเภท  รวม {sum(d['n'] for d in imported):,} รายการ"),
        ("ยังไม่มีข้อมูล Master", f"{len(empty)} ประเภท" if empty else "ไม่มี"),
        (
            "ข้อมูลคงค้าง",
            f"{len(outstanding)} ประเภท · รวม {sum(d['n'] for d in outstanding):,} รายการ — ดูชีต คงค้าง",
        ),
    ]
    for i, (k, v) in enumerate(info):
        r = 6 + i
        sc(ws.cell(r, 2, k), 11, True, WHITE, NAVY)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        sc(ws.cell(r, 3, v), bg=SOFT)

    r = 16
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"].value = (
        "วิธีอ่านเอกสาร   "
        "ชีต “สรุปภาพรวม” = รายการทุกโมดูลในหน้าเดียว   ·   "
        "ชีตตามโมดูล = รายละเอียดแหล่งนำเข้า   ·   "
        "ชีต “เจ้าหนี้ตรง” = รายชื่อเจ้าหนี้จ่ายตรง (ฝั่งบัญชี)   ·   "
        "ชีต “คงค้าง” = สรุปจำนวน PO / เจ้าหนี้คงค้าง / TB   ·   "
        "ชีต “รายการคงค้าง” = รายการจริงให้ยืนยัน"
    )
    ws[f"B{r}"].font = fnt(11)
    ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 52

    hdr(ws, 18, ["ลำดับ", "โมดูล", "Addon / ระบบ", "ประเภทข้อมูล", "จำนวนในระบบ", "สถานะ"])
    for i, (key, th, sys) in enumerate(MODULES, 1):
        items = [d for d in masters if d["mod"] == key]
        ok = [d for d in items if d["n"] > 0]
        total = sum(d["n"] for d in items)
        if key == "outstanding":
            st = "มีรายการคงค้าง" if ok else "ยังไม่นำเข้า (คงค้าง)"
        else:
            st = "นำเข้าแล้ว" if ok else "ยังไม่นำเข้า"
        rr = 18 + i
        sc(ws.cell(rr, 1, i), align="center", bg=SOFT)
        sc(ws.cell(rr, 2, th), bold=True)
        sc(ws.cell(rr, 3, sys))
        sc(ws.cell(rr, 4, f"{len(ok)} / {len(items)} ประเภท"), align="center")
        sc(ws.cell(rr, 5, total), align="center", bold=True, bg=status_bg(st))
        ws.cell(rr, 5).number_format = "#,##0"
        sc(ws.cell(rr, 6, st), align="center", bold=True, bg=status_bg(st))
        ws.row_dimensions[rr].height = 26

    note = 26
    ws.merge_cells(f"B{note}:F{note}")
    ws[f"B{note}"].value = (
        "หมายเหตุ  จำนวนเป็น snapshot จากฐานจริง ณ วันที่ดึงข้อมูล "
        "เจ้าหนี้ตรงอยู่โมดูลบัญชี ไม่นับในจัดซื้อ "
        "ไม่นับข้อมูลตัวอย่าง (DEMO / ตัวอย่าง) ของผู้จำหน่าย "
        "คลังนับเฉพาะคลังที่เปิดใช้และรับเข้า 1 ขั้น"
    )
    ws[f"B{note}"].font = fnt(10, False, "5B6878")
    ws.row_dimensions[note].height = 36


def write_overview(wb, snapshot, masters):
    ws = wb.create_sheet("สรุปภาพรวม")
    page(ws, "สรุปภาพรวม Master Data")
    widths(ws, [8, 22, 32, 36, 22, 16, 18, 44, 48])
    ws.merge_cells("A1:I1")
    ws["A1"].value = f"สรุป Master Data และข้อมูลคงค้าง ทุกโมดูล  ·  {snapshot}"
    ws["A1"].font = fnt(14, True, NAVY)
    ws.row_dimensions[1].height = 28
    headers = [
        "ลำดับ",
        "โมดูล",
        "ประเภทข้อมูล",
        "โมเดล Odoo",
        "คีย์จับคู่",
        "จำนวนในระบบ",
        "สถานะ",
        "วิธีนำเข้า",
        "แหล่งข้อมูล / ไฟล์",
    ]
    hdr(ws, 2, headers)
    mod_th = {k: th for k, th, _ in MODULES}
    rows = []
    for ds in masters:
        rows.append(
            {
                **ds,
                "mod_th": mod_th.get(ds["mod"], ds["mod"]),
            }
        )
    cols = ["seq", "mod_th", "name", "model", "key", "n", "status", "how", "source"]
    last = write_table_rows(ws, 3, rows, cols)
    ws.auto_filter.ref = f"A2:I{last}"
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:I{last}"


def write_module_sheet(wb, mod_key, title, masters):
    items = [d for d in masters if d["mod"] == mod_key]
    ws = wb.create_sheet(title)
    page(ws, title)
    widths(ws, [8, 34, 38, 22, 16, 18, 44, 52])
    ws.merge_cells("A1:H1")
    heading = (
        f"โมดูล {title}  ·  ข้อมูลคงค้าง"
        if mod_key == "outstanding"
        else f"โมดูล {title}  ·  Master Data ที่นำเข้าเข้าระบบ"
    )
    ws["A1"].value = heading
    ws["A1"].font = fnt(14, True, NAVY)
    ws.row_dimensions[1].height = 28
    hdr(
        ws,
        2,
        [
            "ลำดับ",
            "ประเภทข้อมูล",
            "โมเดล Odoo",
            "คีย์จับคู่",
            "จำนวนในระบบ",
            "สถานะ",
            "วิธีนำเข้า",
            "แหล่งข้อมูล / ไฟล์",
        ],
    )
    last = write_table_rows(
        ws, 3, items, ["seq", "name", "model", "key", "n", "status", "how", "source"]
    )
    tot = last + 2
    sc(ws.cell(tot, 1, ""), bg=NAVY)
    sc(ws.cell(tot, 2, "รวมรายการในโมดูลนี้"), 11, True, WHITE, NAVY)
    for c in range(3, 5):
        sc(ws.cell(tot, c, ""), bg=NAVY)
    sc(ws.cell(tot, 5, sum(d["n"] for d in items)), 11, True, WHITE, NAVY, "center")
    ws.cell(tot, 5).number_format = "#,##0"
    imported = sum(1 for d in items if d["n"] > 0)
    sc(ws.cell(tot, 6, f"{imported}/{len(items)} ประเภท"), 11, True, WHITE, NAVY, "center")
    sc(ws.cell(tot, 7, ""), bg=NAVY)
    sc(ws.cell(tot, 8, ""), bg=NAVY)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:H{last}"
    return ws


def write_product_sheet(wb, cur):
    ws = wb.create_sheet("สินค้าแยกหมวด")
    page(ws, "สินค้าแยกหมวด")
    widths(ws, [8, 48, 16, 16, 16, 52, 44])
    ws.merge_cells("A1:G1")
    ws["A1"].value = "Product Master แยกตามหมวดหมู่ในระบบ  และแหล่งไฟล์ที่นำเข้า"
    ws["A1"].font = fnt(14, True, NAVY)
    ws.row_dimensions[1].height = 28

    sql = f"""
        SELECT pc.complete_name AS หมวด,
               COUNT(*) AS จำนวน,
               COUNT(*) FILTER (WHERE t.active) AS ใช้งาน,
               COUNT(*) FILTER (WHERE NOT t.active) AS ปิด
        FROM product_template t
        LEFT JOIN product_category pc ON pc.id = t.categ_id
        GROUP BY pc.complete_name
        ORDER BY pc.complete_name NULLS LAST
    """
    rows = qrows(cur, sql)
    hdr(ws, 2, ["ลำดับ", "หมวดหมู่สินค้า", "จำนวน", "ใช้งาน", "ปิด", "กลุ่มที่สัมพันธ์กับไฟล์นำเข้า", "สคริปต์โดยประมาณ"])

    def hint_for(categ):
        text = categ or ""
        mapping = [
            ("เวชภัณฑ์ ยา", "ยา / คลังยาใหญ่", "import_drug_warehouse.py"),
            ("เวชภัณฑ์ยา", "ยา / คลังยาใหญ่", "import_drug_warehouse.py"),
            ("เวชภัณฑ์มิใช่ยา", "เวชภัณฑ์มิใช่ยา", "import_medical_supplies.py / import_department_nondrug_products.py"),
            ("วัสดุวิทยาศาสตร์", "วัสดุวิทยาศาสตร์ / แล็บ", "import_scientific_supplies.py / import_blood_bank_lab_products.py"),
            ("ธนาคารเลือด", "วัสดุธนาคารเลือด", "import_blood_bank_lab_products.py"),
            ("พยาธิ", "วัสดุพยาธิวิทยา", "import_blood_bank_lab_products.py"),
            ("ช่าง", "วัสดุงานช่าง", "import_construction_craftsman_products.py"),
            ("โภชนาการ", "วัสดุบริโภค / โภชนาการ", "import_nutrition_consumables.py"),
            ("งานพิมพ์", "วัสดุงานพิมพ์", "import_print_house_office_supplies.py"),
            ("งานบ้าน", "วัสดุงานบ้านงานครัว", "import_print_house_office_supplies.py"),
            ("สำนักงาน", "วัสดุสำนักงาน", "import_print_house_office_supplies.py"),
            ("ครุภัณฑ์", "แคตตาล็อกครุภัณฑ์ สป.สธ.", "import_equipment_catalog.py"),
        ]
        for key, group, script in mapping:
            if key in text:
                return group, script
        return "หมวดอื่น / สร้างในระบบ", ""

    last = 2
    total = 0
    for i, row in enumerate(rows, 1):
        r = 2 + i
        categ = row[0] if not isinstance(row, dict) else row["หมวด"]
        n = row[1] if not isinstance(row, dict) else row["จำนวน"]
        act = row[2] if not isinstance(row, dict) else row["ใช้งาน"]
        off = row[3] if not isinstance(row, dict) else row["ปิด"]
        group, script = hint_for(categ)
        zebra = PALE if i % 2 else WHITE
        sc(ws.cell(r, 1, i), align="center", bg=zebra)
        sc(ws.cell(r, 2, categ or "(ไม่มีหมวด)"), bg=zebra)
        sc(ws.cell(r, 3, n), align="center", bold=True, bg=zebra)
        ws.cell(r, 3).number_format = "#,##0"
        sc(ws.cell(r, 4, act), align="center", bg=zebra)
        ws.cell(r, 4).number_format = "#,##0"
        sc(ws.cell(r, 5, off), align="center", bg=zebra)
        ws.cell(r, 5).number_format = "#,##0"
        sc(ws.cell(r, 6, group), bg=zebra)
        sc(ws.cell(r, 7, script), bg=zebra)
        total += int(n or 0)
        last = r
        ws.row_dimensions[r].height = 22

    tot = last + 2
    sc(ws.cell(tot, 2, "รวมสินค้าทั้งหมด"), 11, True, WHITE, NAVY)
    sc(ws.cell(tot, 3, total), 11, True, WHITE, NAVY, "center")
    ws.cell(tot, 3).number_format = "#,##0"

    ws.freeze_panes = "A3"
    if last > 2:
        ws.auto_filter.ref = f"A2:G{last}"

    # warehouse list
    r0 = tot + 3
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=5)
    ws.cell(r0, 1).value = "คลังสินค้าที่เปิดใช้ (รับเข้า 1 ขั้น)"
    ws.cell(r0, 1).font = fnt(13, True, TEAL)
    hdr(ws, r0 + 1, ["ลำดับ", "รหัสคลัง", "ชื่อคลัง", "รับเข้า", "จ่ายออก"], bg=TEAL)
    wh_sql = """
        SELECT code, name,
               CASE reception_steps
                 WHEN 'one_step' THEN 'รับเข้า 1 ขั้น' ELSE reception_steps END,
               CASE delivery_steps
                 WHEN 'ship_only' THEN 'จ่ายออก 1 ขั้น'
                 WHEN 'pick_ship' THEN 'จ่ายออก 2 ขั้น'
                 WHEN 'pick_pack_ship' THEN 'จ่ายออก 3 ขั้น'
                 ELSE delivery_steps END
        FROM stock_warehouse
        WHERE active AND reception_steps = 'one_step'
        ORDER BY id
    """
    for i, row in enumerate(qrows(cur, wh_sql), 1):
        rr = r0 + 1 + i
        zebra = PALE if i % 2 else WHITE
        vals = [i, row[0], row[1], row[2], row[3]]
        for c, v in enumerate(vals, 1):
            sc(ws.cell(rr, c, v), align="center" if c != 3 else "left", bg=zebra)


def _write_block_table(ws, start, title, headers, rows, empty_msg, note=None):
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=max(len(headers), 6))
    ws.cell(start, 1).value = title
    ws.cell(start, 1).font = fnt(13, True, TEAL)
    hdr(ws, start + 1, headers, bg=TEAL)
    if not rows:
        r = start + 2
        sc(ws.cell(r, 1, empty_msg), bg=AMBER)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        ws.row_dimensions[r].height = 28
        last = r
    else:
        last = start + 1
        for i, row in enumerate(rows, 1):
            last = start + 1 + i
            zebra = PALE if i % 2 else WHITE
            vals = [i] + list(row)
            for c, v in enumerate(vals, 1):
                if hasattr(v, "isoformat") and not isinstance(v, datetime):
                    v = v.isoformat()
                elif isinstance(v, datetime):
                    v = v.strftime("%Y-%m-%d %H:%M")
                sc(ws.cell(last, c, v), align="center" if c == 1 else "left", bg=zebra)
            ws.row_dimensions[last].height = 22
    if note:
        last += 1
        ws.merge_cells(start_row=last, start_column=1, end_row=last, end_column=len(headers))
        ws.cell(last, 1).value = note
        ws.cell(last, 1).font = fnt(10, False, "5B6878")
    return last + 2


def write_direct_vendors(wb, cur):
    ws = wb.create_sheet("เจ้าหนี้ตรง")
    page(ws, "เจ้าหนี้ตรง")
    widths(ws, [8, 18, 52, 18, 16, 18, 14])
    ws.merge_cells("A1:G1")
    ws["A1"].value = "เจ้าหนี้ตรง (กลุ่ม DIRECT) · อยู่โมดูลบัญชี / การเงิน ไม่ใช่จัดซื้อ"
    ws["A1"].font = fnt(14, True, NAVY)
    ws.row_dimensions[1].height = 28
    sql = f"""
        SELECT COALESCE(p.ref, '') AS รหัส,
               p.name AS ชื่อ,
               COALESCE(p.vat, '') AS เลขผู้เสียภาษี,
               COALESCE(p.city, '') AS อำเภอ,
               COALESCE(s.name, '') AS จังหวัด,
               CASE WHEN p.active THEN 'ใช้งาน' ELSE 'ปิด' END AS สถานะ
        FROM res_partner p
        LEFT JOIN res_country_state s ON s.id = p.state_id
        WHERE {DIRECT_VENDOR_SQL}
        ORDER BY p.name
    """
    rows = qrows(cur, sql)
    hdr(ws, 2, ["ลำดับ", "รหัสอ้างอิง", "ชื่อเจ้าหนี้ตรง", "เลขผู้เสียภาษี", "อำเภอ", "จังหวัด", "สถานะ"])
    last = 2
    for i, row in enumerate(rows, 1):
        r = 2 + i
        zebra = PALE if i % 2 else WHITE
        vals = [i] + list(row)
        for c, v in enumerate(vals, 1):
            sc(ws.cell(r, c, v), align="center" if c in (1, 7) else "left", bg=zebra)
        last = r
        ws.row_dimensions[r].height = 22
    tot = last + 2
    sc(ws.cell(tot, 2, "รวมเจ้าหนี้ตรง"), 11, True, WHITE, NAVY)
    sc(ws.cell(tot, 3, len(rows)), 11, True, WHITE, NAVY, "center")
    ws.cell(tot, 3).number_format = "#,##0"
    ws.freeze_panes = "A3"
    if last > 2:
        ws.auto_filter.ref = f"A2:G{last}"


def write_outstanding(wb, cur):
    ws = wb.create_sheet("รายการคงค้าง")
    page(ws, "รายการคงค้าง")
    widths(ws, [8, 22, 22, 42, 16, 18, 16, 16, 16, 16, 14, 14])
    ws.merge_cells("A1:L1")
    ws["A1"].value = "ข้อมูลคงค้างที่นำเข้าเข้าระบบ · PO คงค้าง / เจ้าหนี้คงค้าง / งบทดลองปีงบ 2569"
    ws["A1"].font = fnt(14, True, NAVY)
    ws.row_dimensions[1].height = 28

    r = 3
    po_sql = """
        SELECT po.name,
               po.date_order,
               COALESCE(rp.name, ''),
               po.state,
               po.invoice_status,
               po.amount_total,
               COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                 THEN pol.product_qty ELSE 0 END), 0),
               COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                 THEN COALESCE(pol.qty_received, 0) ELSE 0 END), 0),
               COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                 THEN COALESCE(pol.qty_invoiced, 0) ELSE 0 END), 0),
               COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                 THEN pol.product_qty - COALESCE(pol.qty_received, 0) ELSE 0 END), 0),
               COALESCE(SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                                 THEN COALESCE(pol.qty_received, 0) - COALESCE(pol.qty_invoiced, 0) ELSE 0 END), 0)
        FROM purchase_order po
        LEFT JOIN res_partner rp ON rp.id = po.partner_id
        LEFT JOIN purchase_order_line pol ON pol.order_id = po.id
        WHERE po.state IN ('purchase', 'done')
        GROUP BY po.id, po.name, po.date_order, rp.name, po.state, po.invoice_status, po.amount_total
        HAVING SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                        THEN pol.product_qty - COALESCE(pol.qty_received, 0) ELSE 0 END) > 0
            OR SUM(CASE WHEN COALESCE(pol.display_type, '') IN ('', 'product')
                        THEN COALESCE(pol.qty_received, 0) - COALESCE(pol.qty_invoiced, 0) ELSE 0 END) > 0
            OR po.invoice_status IS DISTINCT FROM 'invoiced'
        ORDER BY po.date_order NULLS LAST, po.name
    """
    r = _write_block_table(
        ws,
        r,
        "1) PO คงค้าง (ค้างรับ / ค้างตั้งหนี้)",
        [
            "ลำดับ", "เลขที่ PO", "วันที่สั่ง", "ผู้ขาย", "สถานะ", "สถานะใบแจ้งหนี้",
            "มูลค่ารวม", "จำนวนสั่ง", "รับแล้ว", "ตั้งหนี้แล้ว", "ค้างรับ", "ค้างตั้งหนี้",
        ],
        qrows(cur, po_sql),
        "ยังไม่มี PO คงค้างในระบบ — ให้ยืนยันว่ายังไม่นำเข้า",
        "PO คงค้าง = สถานะ purchase/done ที่ยังรับของไม่ครบ หรือยังตั้งหนี้ไม่ครบ",
    )

    bill_sql = """
        SELECT am.name, am.ref, am.invoice_date, am.invoice_date_due,
               COALESCE(rp.name, ''), am.state, am.payment_state,
               am.amount_total, am.amount_residual
        FROM account_move am
        LEFT JOIN res_partner rp ON rp.id = am.partner_id
        WHERE am.move_type IN ('in_invoice', 'in_refund')
          AND am.state = 'posted'
          AND COALESCE(am.amount_residual, 0) <> 0
        ORDER BY am.invoice_date NULLS LAST, am.name
    """
    r = _write_block_table(
        ws,
        r,
        "2) เจ้าหนี้คงค้าง (บิลที่ผ่านรายการแล้วยังไม่จ่ายครบ)",
        [
            "ลำดับ", "เลขที่เอกสาร", "อ้างอิง", "วันที่บิล", "วันครบกำหนด",
            "เจ้าหนี้", "สถานะ", "สถานะชำระ", "มูลค่ารวม", "ยอดคงค้าง",
        ],
        qrows(cur, bill_sql),
        "ยังไม่มีเจ้าหนี้คงค้างในระบบ — ให้ยืนยันว่ายังไม่นำเข้าบิลค้างจ่าย",
        "เจ้าหนี้คงค้าง = ใบแจ้งหนี้/ใบลดหนี้ผู้ขาย posted และ amount_residual ≠ 0",
    )

    tb_sql = f"""
        SELECT COALESCE(aa.code_store->>'1', '') AS รหัส,
               {jname('aa.name')} AS ชื่อบัญชี,
               aa.account_type,
               SUM(aml.debit), SUM(aml.credit), SUM(aml.debit - aml.credit)
        FROM account_move_line aml
        JOIN account_move am ON am.id = aml.move_id
        JOIN account_account aa ON aa.id = aml.account_id
        WHERE am.state = 'posted'
          AND aml.account_id IS NOT NULL
          AND COALESCE(aml.display_type, '') NOT IN ('line_section', 'line_note', 'cogs')
          AND am.date BETWEEN DATE '2025-10-01' AND DATE '2026-09-30'
        GROUP BY aa.id, aa.code_store, aa.name, aa.account_type
        HAVING SUM(aml.debit) <> 0 OR SUM(aml.credit) <> 0
        ORDER BY 1
    """
    _write_block_table(
        ws,
        r,
        "3) งบทดลอง (TB) ปีงบประมาณ 2569  (1 ต.ค. 2568 – 30 ก.ย. 2569)",
        ["ลำดับ", "รหัสบัญชี", "ชื่อบัญชี", "ประเภทบัญชี", "เดบิต", "เครดิต", "ยอดคงเหลือเดบิตสุทธิ"],
        qrows(cur, tb_sql),
        "ยังไม่มีรายการในงบทดลองปีงบ 2569 — ให้ยืนยันว่ายังไม่นำเข้า TB / ยอดยกเข้า",
        "ยังไม่มียอดยกเข้าจากระบบเดิม  รายการที่เห็น (ถ้ามี) มาจากเอกสารที่บันทึกใน Odoo ไม่ใช่ TB ยกยอดที่นำเข้า",
    )


def write_sources(wb):
    ws = wb.create_sheet("ไฟล์ต้นทาง")
    page(ws, "ไฟล์ต้นทาง")
    widths(ws, [8, 22, 40, 48, 56, 36])
    ws.merge_cells("A1:F1")
    ws["A1"].value = "สคริปต์และไฟล์ Excel ที่ใช้ Data Migration (Product / Vendor / CoA / Asset / Warehouse)"
    ws["A1"].font = fnt(14, True, NAVY)
    ws.row_dimensions[1].height = 28
    hdr(ws, 2, ["ลำดับ", "กลุ่มข้อมูล", "หมวดที่สัมพันธ์", "สคริปต์", "ไฟล์ต้นทาง", "โฟลเดอร์"])
    for i, (group, categ, script, xlsx) in enumerate(PRODUCT_SOURCE_HINTS, 1):
        r = 2 + i
        zebra = PALE if i % 2 else WHITE
        sc(ws.cell(r, 1, i), align="center", bg=zebra)
        sc(ws.cell(r, 2, group), bg=zebra)
        sc(ws.cell(r, 3, categ), bg=zebra)
        sc(ws.cell(r, 4, script), bg=zebra)
        sc(ws.cell(r, 5, xlsx), bg=zebra)
        sc(ws.cell(r, 6, "/opt/odoo18vpk/import/"), bg=zebra)
        ws.row_dimensions[r].height = 28

    extra = [
        ("เจ้าหนี้ตรง (บัญชี)", "res.partner กลุ่ม DIRECT", "import/import_direct_vendors.py", "รายชื่อหน่วยงานราชการ / รพ. / กองทุน"),
        ("เจ้าหนี้ตรง (บัญชี)", "res.partner กลุ่ม DIRECT", "import/import_gov_hospital_vendors.py", "รายชื่อจากภาพ / รายการโรงพยาบาล"),
        ("บัญชีธนาคาร รพ.", "res.partner.bank + journal BK*", "import/import_company_bank_accounts.py", "รายการกองทุน / เลขที่บัญชี รพ."),
        ("ผู้จำหน่ายจัดซื้อ", "res.partner (ไม่รวม DIRECT)", "import/import_vendor_master.py", "vendor_import_actions.json"),
        ("สาธารณูปโภค", "res.partner กลุ่ม UTILITY", "import/import_utility_vendors.py", "รายชื่อผู้ให้บริการสาธารณูปโภค"),
        ("คลังสินค้า", "stock.warehouse", "import/import_warehouses.py", "vpk_data_warehouse_template.xlsx"),
        ("คลังยา / ตำแหน่ง bin", "stock.location", "import/setup_pharmacy_warehouses.py", "วัสดุ_ยา_คลังยาใหญ่.xlsx"),
        ("คลังธนาคารเลือด", "stock.warehouse", "import/setup_blood_bank_warehouse.py", "ตั้งค่าคลัง BBK"),
        ("ผังบัญชี", "account.account", "import/import_coa_vpk.py", "coa-vpk.xlsx"),
        ("บัตรสินทรัพย์ / โปรไฟล์", "account.asset", "import/import_assets.py", "vpk_master_asset_v6.0.1.xlsx"),
        ("Master งบประมาณ", "vpk.budget.*", "Wizard ในโมดูล vpk_budget", "ไฟล์ Excel ตามชีต FundSource / BudgetGroup / …"),
        ("แม่แบบ Master รวม", "หลายโมเดล", "โมดูล vpk_master_data_templates", "ดาวน์โหลดแม่แบบจากระบบ แล้ว import กลับ"),
    ]
    start = 3 + len(PRODUCT_SOURCE_HINTS) + 1
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
    ws.cell(start, 1).value = "Master อื่นที่นำเข้าด้วยสคริปต์ / Wizard"
    ws.cell(start, 1).font = fnt(13, True, TEAL)
    hdr(ws, start + 1, ["ลำดับ", "กลุ่มข้อมูล", "โมเดล", "สคริปต์ / ช่องทาง", "ไฟล์ต้นทาง", "หมายเหตุ"], bg=TEAL)
    for i, (group, model, script, xlsx) in enumerate(extra, 1):
        r = start + 1 + i
        zebra = PALE if i % 2 else WHITE
        sc(ws.cell(r, 1, i), align="center", bg=zebra)
        sc(ws.cell(r, 2, group), bg=zebra)
        sc(ws.cell(r, 3, model), bg=zebra)
        sc(ws.cell(r, 4, script), bg=zebra)
        sc(ws.cell(r, 5, xlsx), bg=zebra)
        sc(ws.cell(r, 6, ""), bg=zebra)
        ws.row_dimensions[r].height = 26


def main():
    snapshot = thai_dt(datetime.now())
    conn = connect()
    try:
        with conn.cursor() as cur:
            load_masters(cur)
            wb = Workbook()
            write_cover(wb, snapshot, MASTERS)
            write_overview(wb, snapshot, MASTERS)
            write_module_sheet(wb, "budget", "งบประมาณ", MASTERS)
            write_module_sheet(wb, "purchase", "จัดซื้อ", MASTERS)
            write_module_sheet(wb, "stock", "คลังสินค้า", MASTERS)
            write_product_sheet(wb, cur)
            write_module_sheet(wb, "finance", "บัญชีการเงิน", MASTERS)
            write_direct_vendors(wb, cur)
            write_module_sheet(wb, "asset", "สินทรัพย์ถาวร", MASTERS)
            write_module_sheet(wb, "outstanding", "คงค้าง", MASTERS)
            write_outstanding(wb, cur)
            write_sources(wb)
            OUT_DIR.mkdir(parents=True, exist_ok=True)
            out = OUT_DIR / dated_xlsx("สรุป-Data-Migration-Master-Data.xlsx")
            wb.save(out)
            print(f"Wrote {out}")
            for ds in MASTERS:
                print(f"  {ds['mod']:10} {ds['n']:6,}  {ds['status']:16}  {ds['name']}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
